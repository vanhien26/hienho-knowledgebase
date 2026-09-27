import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


from scripts._lib import load_workbook_safe, save_with_backup, standard_argparser

def main(args=None):
    wb = load_workbook_safe('05_HUBS/financial-hub-roadmap.xlsx')

    # 1. Update Content Plan sheet
    ws_cp = wb['Content Plan']
    updated_cp_count = 0
    for r in range(3, ws_cp.max_row + 1):
        val = ws_cp.cell(row=r, column=7).value
        if val and '/diem-tin-dung/blog/' in str(val):
            new_val = str(val).replace('/diem-tin-dung/blog/', '/tai-chinh/blog/')
            ws_cp.cell(row=r, column=7).value = new_val
            updated_cp_count += 1
    print(f"Updated {updated_cp_count} URLs in Content Plan to /tai-chinh/blog/*")

    # 2. Update Sitemap sheet
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    section_font = Font(name='Arial', size=10.5, bold=True, color='1F4E78')
    thin_side = Side(style='thin', color='64748B')
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    font_bold = Font(name='Arial', size=9.5, bold=True)
    font_regular = Font(name='Arial', size=9.5)
    font_url = Font(name='Arial', size=9.5, color='0563C1', underline='single')

    ws_old = wb['Sitemap']
    refined_rows = []
    stt_counter = 1

    for r in range(4, ws_old.max_row + 1):
        track = ws_old.cell(row=r, column=2).value
        ptype = ws_old.cell(row=r, column=3).value
        pname = ws_old.cell(row=r, column=4).value
        purl = ws_old.cell(row=r, column=5).value
        pkw = ws_old.cell(row=r, column=6).value
        parch = ws_old.cell(row=r, column=7).value

        # DROP /diem-tin-dung/blog
        if purl == '/diem-tin-dung/blog':
            continue

        # Update the blog details URL
        if '/diem-tin-dung/blog/[slug]' in str(purl):
            purl = '/tai-chinh/blog/[slug]'

        refined_rows.append((stt_counter, track, ptype, pname, purl, pkw, parch))
        stt_counter += 1

    del wb['Sitemap']
    ws_s = wb.create_sheet('Sitemap', index=3)

    ws_s.append(["DANH BẠ KIẾN TRÚC THÔNG TIN & SITEMAP CÁC TRANG CẦN TẠO (FINANCIAL HUB)"])
    ws_s.merge_cells("A1:G1")
    ws_s.cell(row=1, column=1).fill = section_fill
    ws_s.cell(row=1, column=1).font = section_font
    ws_s.row_dimensions[1].height = 25.0

    ws_s.append(["Bản đồ định tuyến toàn bộ các trang cần tạo thuộc Financial Hub: Cổng trung tâm, các công cụ tiện ích trực thuộc /tai-chinh/*, hệ thống trang con cặp tiền, danh bạ 34 ngân hàng, các chuyên trang độc lập (/diem-tin-dung, /chung-khoan, /chung-chi-quy) và blog hub tập trung (/tai-chinh/blog*)."])
    ws_s.merge_cells("A2:G2")
    ws_s.cell(row=2, column=1).font = Font(name='Arial', size=9.5, italic=True, color='475569')
    ws_s.row_dimensions[2].height = 20.0

    s_headers = [
        "STT",
        "Cụm Sản Phẩm (Track)",
        "Phân Cấp / Loại Trang",
        "Tên Trang Cần Tạo (H1 / Title)",
        "Định Tuyến URL Chuẩn Hóa (Canonical URL)",
        "Mục Tiêu Từ Khóa & Mục Đích Tạo Trang",
        "Phương Thức Kỹ Thuật (Architecture / CMS)"
    ]
    ws_s.append(s_headers)
    for c_i in range(1, len(s_headers) + 1):
        c = ws_s.cell(row=3, column=c_i)
        c.fill = header_fill; c.font = header_font; c.border = thin_border
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws_s.row_dimensions[3].height = 30.0

    for row_item in refined_rows:
        ws_s.append(list(row_item))
        r = ws_s.max_row
        num_lines = str(row_item[5] or '').count('\n') + 1
        ws_s.row_dimensions[r].height = max(36.0, num_lines * 18.0)

        ws_s.cell(row=r, column=1).alignment = Alignment(horizontal='center', vertical='center')
        ws_s.cell(row=r, column=2).alignment = Alignment(horizontal='center', vertical='center'); ws_s.cell(row=r, column=2).font = font_bold
        ws_s.cell(row=r, column=3).alignment = Alignment(horizontal='center', vertical='center')
        ws_s.cell(row=r, column=4).alignment = Alignment(horizontal='left', vertical='center'); ws_s.cell(row=r, column=4).font = font_bold
        ws_s.cell(row=r, column=5).alignment = Alignment(horizontal='left', vertical='center'); ws_s.cell(row=r, column=5).font = font_url
        ws_s.cell(row=r, column=6).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws_s.cell(row=r, column=7).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

        for c_i in range(1, 8):
            ws_s.cell(row=r, column=c_i).border = thin_border

    col_widths_s = [6, 22, 32, 42, 38, 55, 42]
    for idx, w in enumerate(col_widths_s, start=1):
        ws_s.column_dimensions[get_column_letter(idx)].width = w

    save_with_backup(wb, '05_HUBS/financial-hub-roadmap.xlsx', dry_run=args.dry_run, make_backup=not args.no_backup)
    print(f"Successfully updated Sitemap! Kept {len(refined_rows)} clean rows (1 to {len(refined_rows)}).")

if __name__ == "__main__":
    _ap = standard_argparser("Fix runner - safe with --dry-run")
    main(_ap.parse_args())

