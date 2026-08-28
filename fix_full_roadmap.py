import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

file_path = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/vehicle-hub-roadmap.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb['Roadmap']

# Clear rows from 4 downwards
ws.delete_rows(4, ws.max_row)

# Styles
font_bold = Font(bold=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

fill_p1 = PatternFill("solid", fgColor="D9EAD3")
fill_p2 = PatternFill("solid", fgColor="FFF2CC")
fill_p3 = PatternFill("solid", fgColor="CFE2F3")
fill_block = PatternFill("solid", fgColor="F4CCCC")

data = [
    # PHASE 1
    ['Phase 1: Build-Up & Foundation (Nền tảng hút Traffic & Tiện ích Ưu tiên P1 từ Cell Team)', 'Master Hub Gateway', 'Trang Chủ Master Hub (/tien-ich-giao-thong)', '1. Mặt tiền trung tâm với Hero Widget tra cứu 3-in-1.\n2. Gắn điểm chạm Onelink/QR Code điều hướng In-App.', 'Go-live Sớm'],
    ['None', 'Magnet + Transaction (P1)', 'Cổng Tra Cứu Phạt Nguội (/phat-nguoi)', '1. Trả kết quả Inline 2 nhánh (Có/Không vi phạm).\n2. Block Cross-sell gợi ý nộp phạt In-App & Mua Bảo hiểm.\nVai trò: Phễu thu thập Lead lớn nhất (3.6M search).', 'Check-in Giữa Sprint'],
    ['None', 'Magnet + Transaction (P1)', 'Thu Phí Không Dừng ePass/VETC (/phi-khong-dung, /tram-thu-phi)', '1. Hướng dẫn & luồng liên kết ePass/VETC theo biển số.\n2. Nút W2A nạp ví thu phí không dừng định kỳ.', 'Check-in Giữa Sprint'],
    ['None', 'Giữ Chân & Acquisition (P2)', 'Bảng Giá Xăng & Cây Xăng (/gia-xang, /cay-xang)', '1. Giá xăng update realtime & Bản đồ định vị GPS.\nVai trò: Phễu Acquisition tiếp cận tệp chủ xe hàng ngày (10M+ search).', 'Check-in Đầu Sprint'],
    ['None', 'Giữ Chân & Acquisition (P3)', 'Bản Đồ Trạm Sạc EV (/tram-sac, /tram-sac/vinfast)', '1. Bản đồ trạm sạc xe điện theo tỉnh/thành phố, brand.\nVai trò: Đặt nền tảng cho tệp chủ xe điện EV (1.2M search).', 'Check-in Đầu Sprint'],

    # PHASE 2
    ['Phase 2: Convert & Monetize (Tối ưu Chuyển Đổi & Bán Chéo - Master Plan)', 'Conversion & SEO', 'Tối Ưu Kỹ Thuật SEO, Chuyển Đổi & W2A Handoff', '1. Fix lỗi kỹ thuật SEO (Index CMS, URL trùng lặp).\n2. Onelink giữ context biển số khi quét QR mở App.\n3. Xây UI Consent lưu biển số theo Nghị định 13.', 'Tháng 10/2026'],
    ['None', 'Conversion & Cross-sell', 'Bảo Hiểm Phương Tiện (/bao-hiem-o-to, /bao-hiem-xe-may)', '1. Block cross-sell chuẩn hóa theo loại xe (Ô tô/Xe máy).\n2. Prefill Vehicle Profile (Auto-fill) khi mua BH.\n3. Cải thiện CTR luồng mua bảo hiểm từ Web.', 'Tháng 10/2026'],
    ['None', 'Calculator & Guide', 'Tính Phí Trước Bạ & Nuôi Xe (/lan-banh, /chi-phi-nuoi-xe, /so-sanh-xe)', '1. Công cụ tính phí trước bạ ô tô, giá lăn bánh, so sánh xe.\nVai trò: Lead chất lượng bán chéo Bảo hiểm VCX/TNDS xe mới.', 'Tháng 10/2026'],
    ['None', 'Magnet + Transaction', 'Danh Bạ Cứu Hộ Ô tô 24/7 (/cuu-ho)', '1. Danh sách/bản đồ đơn vị cứu hộ theo vị trí.\nVai trò: Xử lý Intent khẩn cấp (41K search).', 'Tháng 10/2026'],
    ['None', 'Giữ Chân', 'Bản Đồ Bãi Đỗ Xe (/bai-do-xe)', '1. Bản đồ bãi đỗ xe VETC Parking theo vị trí, giá, chỗ trống.\nVai trò: Tiện ích giữ chân (53K search).', 'Tháng 10/2026'],

    # PHASE 3
    ['Phase 3: Scale-Up & Complete Ecosystem (Hoàn thiện 21 URLs)', 'Legal & Rules', 'Tra Cứu Luật Giao Thông (/tra-cuu-muc-phat, /bien-bao-giao-thong, /kinh-nghiem-lai-xe)', '1. Từ điển tra cứu mức phạt theo Nghị định 100/123.\n2. Cẩm nang biển báo giao thông & kinh nghiệm lái xe.', 'Tháng 11/2026'],
    ['None', 'Smart City', 'Camera Giao Thông (/camera-giao-thong)', '1. Tích hợp quan sát luồng live stream camera giao thông.', 'Tháng 11/2026'],
    ['None', 'Maintenance & Tracking', 'Chăm Sóc & Theo Dõi Hành Trình (/bao-duong, /chuyen-di)', '1. Đặt lịch rửa xe, bảo dưỡng xe.\n2. Nhật ký chuyến đi cá nhân.', 'Tháng 12/2026'],

    # BLOCKED
    ['Blocked / Pending (Vướng Phụ Thuộc Đối Tác)', 'Magnet + Transaction (P1)', 'Tra Cứu Hạn Đăng Kiểm (/dang-kiem)', '1. Tra cứu hạn đăng kiểm & đặt lịch kiểm định.\nTRẠNG THÁI: Blocked. API bảo trì từ phía TTĐK. Chờ mở lại.', 'Blocked']
]

for idx, row_data in enumerate(data, start=4):
    for col, val in enumerate(row_data, start=1):
        cell = ws.cell(row=idx, column=col, value=val if val != 'None' else None)
        cell.alignment = align_left if col in [3, 4] else align_center
        cell.border = border
        
        if idx <= 8:
            cell.fill = fill_p1
        elif idx <= 13:
            cell.fill = fill_p2
        elif idx <= 16:
            cell.fill = fill_p3
        else:
            cell.fill = fill_block

ws.merge_cells('A4:A8')
ws.merge_cells('A9:A13')
ws.merge_cells('A14:A16')

wb.save(file_path)
print("Full Roadmap updated successfully.")
