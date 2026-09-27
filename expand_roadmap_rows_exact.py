import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


from scripts._lib import load_workbook_safe, save_with_backup, standard_argparser

def main(args=None):
    wb = load_workbook_safe('05_HUBS/student-hub-roadmap.xlsx')

    font_title = Font(name='Calibri', size=14, bold=True, color='1E293B')
    font_subtitle = Font(name='Calibri', size=11, italic=True, color='64748B')
    font_header = Font(name='Calibri', size=11, bold=True, color='1E293B')
    font_bold = Font(name='Calibri', size=11, bold=True, color='000000')
    font_regular = Font(name='Calibri', size=11, color='000000')

    fill_header = PatternFill('solid', fgColor='F1F5F9')
    fill_completed = PatternFill('solid', fgColor='DCFCE7')
    fill_inprogress = PatternFill('solid', fgColor='FEF08A')
    fill_planned = PatternFill('solid', fgColor='F3F4F6')

    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    border_thin = Side(style='thin', color='CBD5E1')
    border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    if 'Roadmap' in wb.sheetnames:
        ws_roadmap = wb['Roadmap']
        ws_roadmap.delete_rows(1, ws_roadmap.max_row)
        ws_roadmap.views.sheetView[0].showGridLines = True

        ws_roadmap.cell(row=1, column=1, value='LỘ TRÌNH PHÁT TRIỂN SẢN PHẨM STUDENT HUB (STUDENT HUB ROADMAP 2026)').font = font_title
        ws_roadmap.cell(row=2, column=1, value='Quy hoạch chi tiết các dòng hạng mục lớn theo từng Phase triển khai (Foundation Launch T8 -> Scale, Verification Drive & Gamification T9 -> Career Master & pSEO Scale 100+ Q4).').font = font_subtitle

        headers_roadmap = ['Phase', 'Thời Gian', 'Giai Đoạn & Nhóm Dịch Vụ', 'Tên Trang / Utility / Công Việc', 'Mô Tả Sản Phẩm & Quy Hoạch Kỹ Thuật (UX/UI & Specs)', 'Dữ Liệu, API & Phối Hợp Kỹ Thuật', 'Phễu Chuyển Đổi Kinh Doanh (Web-to-App)']
        for c_idx, h in enumerate(headers_roadmap, start=1):
            cell = ws_roadmap.cell(row=4, column=c_idx, value=h)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = border

        detailed_roadmap_rows = [
            # PHASE 1 (Completed)
            ['Phase 1', 'Tháng 8/2026', 'Foundation | Gateway Hub', 'Student Hub Master Gateway (/sinh-vien)', '• Dựng khung HTML/CSS Master Gateway (/sinh-vien) với Hero Banner Student Pass, 5 Content Pillars, Bạn đồng hành MoMo.\n• ĐÃ GO-LIVE TRÊN PRODUCTION.', '- Design Tokens MoBase Web\n- GA4/GTM Tracking Schema\n- Asset Student Pass & Back2School', 'Entry Points Onelink & Dynamic QR dẫn vào Mini App Student Pass In-App; Nút CTA "Xác thực ngay / Nhận đặc quyền".'],
            ['Phase 1', 'Tháng 8/2026', 'Foundation | Pilot Schools', '05 Trang Trường ĐH Pilot (/sinh-vien/[truong-slug])', '• Xây dựng 05 trang trường Pilot trọng điểm (UEH, FTU, TDTU, HCMUT, VLU).\n• Hiển thị thông tin tổng quan, điểm chuẩn 3 năm và học phí tín chỉ.\n• ĐÃ GO-LIVE ON STAGING/DEV.', '- File Content Structure Master Data\n- DB Điểm chuẩn tuyển sinh 3 năm\n- DB Học phí tín chỉ các ngành', 'Anchor links #hoc-phi, #diem-chuan & OneLink eKYC Thẻ Sinh Viên Số custom theo từng trường.'],
            ['Phase 1', 'Tháng 8/2026', 'Foundation | Campus POI', 'Utility Tìm Điểm Thanh Toán Quanh Campus', '• Bản đồ hiển thị cửa hàng F&B, quán ăn, siêu thị quanh campus chấp nhận thanh toán MoMo QR.\n• ĐÃ GO-LIVE ON STAGING/DEV.', '- Merchant POI Database\n- Google Maps Location API', 'Nút CTA "Tìm điểm thanh toán MoMo gần bạn" + Voucher ưu đãi F&B.'],
            ['Phase 1', 'Tháng 8/2026', 'Foundation | Transport Utility', 'Utility Tra Cứu Tuyến Xe Buýt Đến Trường', '• Công cụ tra cứu thông tin các tuyến xe buýt ghé qua từng cơ sở campus đại học.\n• ĐÃ GO-LIVE ON STAGING/DEV.', '- Google Transit API (GTFS data)\n- Dữ liệu tuyến xe buýt Sở GTVT', 'Nút CTA "Nạp tiền vé xe buýt / Mua vé xe buýt MoMo".'],

            # PHASE 2 (In Progress)
            ['Phase 2', 'Tháng 9/2026', 'Scale-out | School Expansion', 'Mở Rộng 07 Trang Trường ĐH Mới (/sinh-vien/[truong-slug])', '• Mở rộng 07 trang trường mới (HUST, NEU, HANU, HAU, USSH, HUIT, UEL -> nâng quy mô lên 12 trường trọng điểm).\n• Tích hợp full 4 JTBD thông tin trường.', '- Master Data 7 Trường Mới\n- DB Điểm chuẩn & Học phí', 'OneLink eKYC Thẻ Sinh Viên Số custom theo từng trường đại học.'],
            ['Phase 2', 'Tháng 9/2026', 'Verification Drive | Student Pass', 'Phễu Drive Verify Student Pass (Custom School Card UI)', '• Tập trung Promote & Drive 100% sinh viên từ Web đến phễu Xác Thực Sinh Viên In-App.\n• Hiển thị giao diện Thẻ Sinh Viên Số Custom Card UI theo tên trường sau eKYC.', '- API Xác thực Thẻ Sinh Viên Số In-App\n- Student Pass Card Rendering Engine', 'Focus 100% phễu Web-to-App CTR kéo sinh viên eKYC Verify Student Pass In-App nhận gói voucher 500K.'],
            ['Phase 2', 'Tháng 9/2026', 'Gamification | Ambassador Showcase', 'Trang Showcase Đại Sứ Sinh Viên MoMo (/sinh-vien/ambassador)', '• Trang Showcase Project Đại sứ sinh viên (5-7 nhóm đại sứ, bảng thành tích 1 năm).\n• Công cụ giao nhiệm vụ tự động và Leaderboard vinh danh đại sứ xuất sắc.', '- Database 50 Đại sứ sinh viên\n- System tích điểm & Leaderboard In-App', 'Nút CTA "Đăng ký trở thành Đại Sứ Sinh Viên MoMo" mở phễu đăng ký In-App.'],
            ['Phase 2', 'Tháng 9/2026', 'Gamification | Realtime O2O Event', 'Module Photobooth Real-time Event Offline Campus (26-27/09)', '• Xây dựng Web Photobooth Real-time cho Offline Event Campus 26-27/09.\n• Sinh viên chụp ảnh tại Booth -> Upload Web -> Nhận ảnh có khung MoMo Student Pass -> Tích lũy footprint nhận quà.', '- Web Realtime Image Storage & Frame Service\n- QR Scanner tại gian hàng Offline Event', 'Quét QR Code nhận voucher 50K & Kích hoạt Thẻ Sinh Viên Số ngay tại sự kiện.'],
            ['Phase 2', 'Tháng 9/2026', 'Gamification | Minigame Rewards', 'Minigame Tích Điểm Đổi Voucher In-App', '• Minigame tương tác trên Web kéo sinh viên làm nhiệm vụ tra cứu thông tin trường nhận lượt quay thưởng voucher.', '- Gamification Mission Engine\n- MoMo Rewards / Voucher Distribution API', 'Kéo 100% sinh viên mở App MoMo nhận voucher ăn uống / xem phim 50K.'],
            ['Phase 2', 'Tháng 9/2026', 'Content Pipeline | 100 Blog Posts', 'Sản Xuất 100 Bài Viết Cẩm Nang Sinh Viên', '• Xuất bản 100 bài viết cẩm nang chất lượng cao: Cẩm nang tân sinh viên, quy chế học vụ, thuê trọ an toàn, quản lý chi tiêu, kỹ năng mềm AI.\n• Vận hành MoSpark GenAI Content Pipeline.', '- Kế hoạch biên tập 100 bài viết\n- Bộ từ khóa SEO phân cấp từ Search Volume', 'Inline Banner điều hướng sử dụng dịch vụ tài chính sinh viên & Student Pass trên App MoMo.'],

            # PHASE 3 (Planned)
            ['Phase 3', 'Quý 4/2026', 'Career Master | Gateway', 'Cổng Việc Làm Sinh Viên Master (/sinh-vien/viec-lam)', '• Ra mắt Cổng Việc Làm Sinh Viên Master (/sinh-vien/viec-lam).\n• Hiển thị danh sách vị trí việc làm part-time và thực tập sinh từ các chuỗi F&B/bán lẻ đối tác MoMo.\n(Search Volume: 125.000 search/tháng)', '- Merchant Jobs Database từ Partnership Team\n- Job Search Engine & Filter Module', 'Nút CTA "Ứng Tuyển Qua Student Pass" điều hướng mở App MoMo.'],
            ['Phase 3', 'Quý 4/2026', 'Career Master | AI Resume', 'Công Cụ AI Resume Builder & ATS Scoring', '• Công cụ tạo CV tự động ngay trên Web.\n• GenAI Pipeline tổng hợp thông tin, xuất CV PDF chuẩn ATS & chấm điểm tương thích JD thời gian thực.', '- MoMo GenAI ATS Resume Engine Service\n- PDF Export & Preview Component', 'Nút CTA "Xuất CV PDF / Ứng Tuyển Nhanh 1-Tap" mở phễu eKYC Student Pass.'],
            ['Phase 3', 'Quý 4/2026', 'Career Master | Location Map', 'Bản Đồ Việc Làm Campus Bán Kính 5km', '• Tích hợp bộ lọc bán kính 500m, 1km, 3km, 5km xung quanh campus đại học.\n• Lọc việc làm part-time gần nơi học/nơi ở của sinh viên.', '- Location Engine (Google Maps API)\n- Merchant Store Address Database', 'Nút CTA "Xem công việc gần campus của bạn" trên bản đồ.'],
            ['Phase 3', 'Quý 4/2026', 'Career Master | Financial Unlock', 'Phễu 1-Tap Apply & Mở Khóa Ví Trả Sau 0% HSSV', '• Sinh viên bấm ứng tuyển 1-Tap qua Student Pass -> Mở App eKYC xác thực sinh viên -> Mở gói voucher HSSV & Ví Trả Sau 0% (2-5 triệu VNĐ).\n• KHÔNG CÓ luồng nhận lương qua Ví MoMo.', '- API xác thực Thẻ Sinh Viên Số In-App\n- Credit PayLater Engine', 'Kích hoạt Thẻ Sinh Viên Số & Mở hạn mức Ví Trả Sau 0% HSSV.'],
            ['Phase 3', 'Quý 4/2026', 'Programmatic pSEO | Scale 100+', 'Hệ Thống pSEO Engine Scale-out 100+ Trang Trường ĐH/CĐ', '• Programmatic pSEO Engine tự động tạo 100+ trang trường ĐH/CĐ toàn quốc theo các Tier (Tier 1 >100K search, Tier 2 50K-100K).\n• Tự động hóa MoSpark GenAI Content Pipeline.', '- Programmatic Content Generation Engine\n- Master Data 100+ Trường ĐH/CĐ toàn quốc', 'OneLink eKYC Thẻ Sinh Viên Số cá nhân hóa theo từng trường đại học.'],
            ['Phase 3', 'Quý 4/2026', 'Sub-pages Ecosystem | Housing', 'Sub-page Nhà Trọ & KTX An Toàn (/nha-tro)', '• Sub-page Nhà trọ & KTX (/nha-tro) kết hợp bản đồ an toàn quanh trường & công cụ AI Bill Splitter (chia tiền trọ).', '- Data danh bạ nhà trọ sạch từ Apify Google Maps\n- AI Bill Splitter Component', 'Nút CTA "Tìm Phòng Trọ An Toàn" + "Tạo QR Chia Tiền Nhà Trọ".'],
            ['Phase 3', 'Quý 4/2026', 'Sub-pages Ecosystem | UGC Review', 'Sub-page Review UGC Chi Tiết (/review)', '• Sub-page Review UGC Chi tiết (/review) vinh danh Top Reviewers và tổng hợp đánh giá trường học 5 sao.', '- UGC Review & Rating Engine In-App Sync', 'Nút CTA "Viết Review Trường Nhận Xu MoMo".'],
            ['Phase 3', 'Quý 4/2026', 'Sub-pages Ecosystem | Workshop', 'Sub-page Webinar & Event Workshop (/workshop)', '• Sub-page Webinar & Workshop (/workshop) đăng ký tham gia sự kiện kỹ năng mềm, AI và định hướng nghề nghiệp sinh viên.', '- Event Management Service\n- Webinar Registration Engine', 'Nút CTA "Đăng Ký Tham Gia Workshop Kỹ Năng".']
        ]

        for r_offset, r_data in enumerate(detailed_roadmap_rows, start=5):
            for c_idx, val in enumerate(r_data, start=1):
                cell = ws_roadmap.cell(row=r_offset, column=c_idx, value=val)
                cell.font = font_bold if c_idx in [1, 2, 4] else font_regular
                cell.border = border
                cell.alignment = align_center if c_idx in [1, 2] else align_left
                if 'Phase 1' in r_data[0]:
                    cell.fill = fill_completed
                elif 'Phase 2' in r_data[0]:
                    cell.fill = fill_inprogress
                else:
                    cell.fill = fill_planned

        ws_roadmap.column_dimensions['A'].width = 16
        ws_roadmap.column_dimensions['B'].width = 24
        ws_roadmap.column_dimensions['C'].width = 34
        ws_roadmap.column_dimensions['D'].width = 45
        ws_roadmap.column_dimensions['E'].width = 68
        ws_roadmap.column_dimensions['F'].width = 50
        ws_roadmap.column_dimensions['G'].width = 50

    save_with_backup(wb, '05_HUBS/student-hub-roadmap.xlsx', dry_run=args.dry_run, make_backup=not args.no_backup)
    print('Successfully expanded Roadmap sheet into 18 detailed activity rows across 3 Phases!')

if __name__ == "__main__":
    _ap = standard_argparser("Fix runner - safe with --dry-run")
    main(_ap.parse_args())

