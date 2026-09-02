import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook('05_HUBS/financial-hub-roadmap.xlsx')

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
section_font = Font(name='Arial', size=10.5, bold=True, color='1F4E78')
thin_side = Side(style='thin', color='64748B')
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

p0_font = Font(name='Arial', size=9.5, bold=True, color='B91C1C')
p1_font = Font(name='Arial', size=9.5, bold=True, color='1D4ED8')
p2_font = Font(name='Arial', size=9.5, bold=False, color='475569')

font_bold = Font(name='Arial', size=9.5, bold=True)
font_bold_blue = Font(name='Arial', size=9.5, bold=True, color='002060')
font_regular = Font(name='Arial', size=9.5)

# Refined Consolidated Financial Markets (No InsurTech, No Source column, Full Stock/Ticker data)
refined_markets = [
    # STT, Market, Pillar, Volume, Priority/Position
    (1, "Giá Vàng", "Đầu Tư & Tích Sản", 78879270, "Big Bet (P1)"),
    (2, "Chứng Khoán & Mã Cổ Phiếu", "Đầu Tư & Tích Sản", 12543260, "Big Bet (P1)"),
    (3, "Tỷ Giá", "Tiền Tệ & Ngoại Hối", 8516910, "Big Bet (P1)"),
    (4, "Ngân Hàng", "Ngân Hàng & Thẻ", 6629420, "Big Bet / Programmatic (P0)"),
    (5, "Mã Số Thuế", "Thu Nhập & Thuế", 3903480, "Big Bet (P1)"),
    (6, "Lãi Suất", "Tiết Kiệm & Lãi Suất", 1246620, "Big Bet (P1)"),
    (7, "Thuế (TNCN)", "Thu Nhập & Thuế", 1159450, "Top Quick Win (P0)"),
    (8, "Ngoại Tệ", "Tiền Tệ & Ngoại Hối", 959340, "Standard Feature (P2)"),
    (9, "Tiền Số & Crypto", "Tài Sản Số", 573920, "Low Priority (P2)"),
    (10, "Bảo Hiểm Xã Hội", "Thu Nhập & An Sinh Xã Hội", 441060, "Top Quick Win (P0)"),
    (11, "Thẻ Tín Dụng", "Ngân Hàng & Thẻ", 394590, "Big Bet (P1)"),
    (12, "Vay Tín Chấp (Vay Nhanh)", "Tín Dụng & Vay Vốn", 389820, "Big Bet (P1)"),
    (13, "CIC (Điểm Tín Dụng)", "Tín Dụng & Sức Khỏe Tín Dụng", 339000, "Top Quick Win #1 (P0)"),
    (14, "Thẻ Visa", "Ngân Hàng & Thẻ", 243110, "Core Driver (P1)"),
    (15, "Gửi Tiết Kiệm", "Tiết Kiệm & Lãi Suất", 230110, "High-Impact Quick Win (P0)"),
    (16, "Tính Lương", "Thu Nhập & Thuế", 190690, "Top Quick Win (P0)"),
    (17, "Vay Thế Chấp (Mua Nhà)", "Tín Dụng & Vay Vốn", 125960, "Core Driver (P1)"),
    (18, "Trả Góp (Ví Trả Sau)", "Tín Dụng & Vay Vốn", 116390, "Top Quick Win (P0)"),
    (19, "Trái Phiếu", "Đầu Tư & Tích Sản", 79580, "Standard Feature (P2)"),
    (20, "Thẻ Ghi Nợ (ATM)", "Ngân Hàng & Thẻ", 73290, "Top Quick Win (P0)"),
    (21, "Nợ Xấu", "Tín Dụng & Sức Khỏe Tín Dụng", 59890, "Top Quick Win (P0)"),
    (22, "Napas", "Ngân Hàng & Thẻ", 51100, "Top Quick Win (P0)"),
    (23, "Chứng Chỉ Quỹ", "Đầu Tư & Tích Sản", 25790, "Top Quick Win (P0)"),
    (24, "Mastercard", "Ngân Hàng & Thẻ", 18630, "Core Driver (P1)")
]

total_refined_vol = sum(m[3] for m in refined_markets)

if 'Consolidated Financial Markets' in wb.sheetnames:
    del wb['Consolidated Financial Markets']

ws_cm = wb.create_sheet('Consolidated Financial Markets', index=3)
ws_cm.append(["TỔNG HỢP TOÀN BỘ CÁC THỊ TRƯỜNG DỊCH VỤ TÀI CHÍNH FINANCIAL HUB (KHÔNG BAO GỒM BẢO HIỂM)"])
ws_cm.merge_cells("A1:F1")
ws_cm.cell(row=1, column=1).fill = section_fill
ws_cm.cell(row=1, column=1).font = section_font
ws_cm.row_dimensions[1].height = 24.0

cm_headers = ["STT", "Thị Trường / Dịch Vụ", "Trụ Cột Nghiệp Vụ", "Search Volume / Tháng", "Tỷ Trọng Ngành (%)", "Định Vị & Mức Độ Ưu Tiên"]
ws_cm.append(cm_headers)
for c_i in range(1, len(cm_headers) + 1):
    c = ws_cm.cell(row=2, column=c_i)
    c.fill = header_fill; c.font = header_font; c.border = thin_border
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws_cm.row_dimensions[2].height = 26.0

for stt, name, pillar, vol, prio_str in refined_markets:
    pct = f"{(vol / total_refined_vol)*100:.2f}%"
    ws_cm.append([stt, name, pillar, vol, pct, prio_str])
    r = ws_cm.max_row
    ws_cm.row_dimensions[r].height = 22.0
    ws_cm.cell(row=r, column=1).alignment = Alignment(horizontal='center', vertical='center')
    ws_cm.cell(row=r, column=2).alignment = Alignment(horizontal='left', vertical='center'); ws_cm.cell(row=r, column=2).font = font_bold
    ws_cm.cell(row=r, column=3).alignment = Alignment(horizontal='left', vertical='center')
    ws_cm.cell(row=r, column=4).alignment = Alignment(horizontal='right', vertical='center'); ws_cm.cell(row=r, column=4).font = font_bold_blue; ws_cm.cell(row=r, column=4).number_format = '#,##0'
    ws_cm.cell(row=r, column=5).alignment = Alignment(horizontal='center', vertical='center')
    ws_cm.cell(row=r, column=6).alignment = Alignment(horizontal='center', vertical='center')

    if "P0" in prio_str: ws_cm.cell(row=r, column=6).font = p0_font
    elif "P1" in prio_str: ws_cm.cell(row=r, column=6).font = p1_font
    else: ws_cm.cell(row=r, column=6).font = p2_font

    for c_i in range(1, 7): ws_cm.cell(row=r, column=c_i).border = thin_border

# Grand Total Row
ws_cm.append(["TỔNG CỘNG FINANCIAL MASTER HUB", "", "", total_refined_vol, "100.00%", "24 Thị Trường Tài Chính Hoàn Chỉnh"])
r_tot = ws_cm.max_row
ws_cm.merge_cells(f"A{r_tot}:C{r_tot}")
ws_cm.row_dimensions[r_tot].height = 24.0
ws_cm.cell(row=r_tot, column=1).fill = section_fill; ws_cm.cell(row=r_tot, column=1).font = section_font
ws_cm.cell(row=r_tot, column=1).alignment = Alignment(horizontal='center', vertical='center')
for c_i in range(1, 7):
    c = ws_cm.cell(row=r_tot, column=c_i)
    c.fill = section_fill; c.font = font_bold; c.border = thin_border
ws_cm.cell(row=r_tot, column=4).number_format = '#,##0'

ws_cm.column_dimensions['A'].width = 8
ws_cm.column_dimensions['B'].width = 32
ws_cm.column_dimensions['C'].width = 30
ws_cm.column_dimensions['D'].width = 24
ws_cm.column_dimensions['E'].width = 20
ws_cm.column_dimensions['G'].width = 30
ws_cm.column_dimensions['F'].width = 30

wb.save('05_HUBS/financial-hub-roadmap.xlsx')
print(f"Successfully updated Consolidated Financial Markets: Removed InsurTech, removed Source column, full Stock volume: {total_refined_vol:,d}!")
