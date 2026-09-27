import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


from scripts._lib import load_workbook_safe, save_with_backup, standard_argparser

def main(args=None):
    wb = load_workbook_safe('05_HUBS/financial-hub-roadmap.xlsx')

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    section_font = Font(name='Arial', size=10.5, bold=True, color='1F4E78')
    thin_side = Side(style='thin', color='64748B')
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    p0_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    p1_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
    p2_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    status_done_fill = PatternFill(start_color="1E7E34", end_color="1E7E34", fill_type="solid")
    status_done_font = Font(name='Arial', size=9.5, bold=True, color='FFFFFF')
    status_prog_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    status_prog_font = Font(name='Arial', size=9.5, bold=True, color='B9770E')
    status_plan_fill = PatternFill(start_color="1B4F72", end_color="1B4F72", fill_type="solid")
    status_plan_font = Font(name='Arial', size=9.5, bold=True, color='FFFFFF')

    font_bold = Font(name='Arial', size=9.5, bold=True)
    font_regular = Font(name='Arial', size=9.5)
    font_url = Font(name='Arial', size=9.5, color='0563C1', underline='single')

    if 'Sitemap' in wb.sheetnames:
        del wb['Sitemap']

    # Insert Sitemap right after Utilities (index 3)
    ws_s = wb.create_sheet('Sitemap', index=3)

    ws_s.append(["DANH BẠ KIẾN TRÚC THÔNG TIN & SITEMAP CÁC TRANG FINANCIAL MASTER HUB"])
    ws_s.merge_cells("A1:H1")
    ws_s.cell(row=1, column=1).fill = section_fill
    ws_s.cell(row=1, column=1).font = section_font
    ws_s.row_dimensions[1].height = 25.0

    ws_s.append(["Quy hoạch toàn diện hệ thống URL Kênh Web: Master Gateway, Standalone Utilities, Programmatic Cặp Tiền, Programmatic 34 Ngân Hàng, và Blog Hubs. Mọi URL blog bắt buộc có tiền tố /blog/."])
    ws_s.merge_cells("A2:H2")
    ws_s.cell(row=2, column=1).font = Font(name='Arial', size=9.5, italic=True, color='475569')
    ws_s.row_dimensions[2].height = 20.0

    s_headers = [
        "STT",
        "Cụm Sản Phẩm (Track)",
        "Phân Cấp / Loại Trang",
        "Tên Trang (Page Title / H1)",
        "URL Chuẩn Hóa (Canonical Slug)",
        "Từ Khóa Mục Tiêu & Nhu Cầu Tìm Kiếm",
        "Phương Thức Kỹ Thuật (Architecture)",
        "Trạng Thái"
    ]
    ws_s.append(s_headers)
    for c_i in range(1, len(s_headers) + 1):
        c = ws_s.cell(row=3, column=c_i)
        c.fill = header_fill; c.font = header_font; c.border = thin_border
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws_s.row_dimensions[3].height = 30.0

    sitemap_items = [
        # 1. CỔNG TRUNG TÂM
        (1, "Master Gateway", "Cổng Trung Tâm (Master Hub)", "Financial Master Hub MoMo", "/tai-chinh",
         "tài chính momo, dịch vụ tài chính momo (Cổng điều hướng 5 trụ cột)", "Static Core Hub (Next.js SSR/SSG)", "Done (Live)"),

        (2, "Master Gateway", "Trang Chuyển Hướng (301 Redirect)", "Trung Tâm Tài Chính (Cũ)", "/trung-tam-tai-chinh",
         "Bảo toàn thẩm quyền SEO và backlink cũ sang /tai-chinh", "Nginx / Cloudflare 301 Permanent Redirect", "Done (Live)"),

        (3, "Sức Khỏe Tín Dụng", "Chuyên Trang Độc Lập (Vertical Hub)", "Điểm Tín Dụng & Tra Cứu CIC", "/diem-tin-dung",
         "điểm tín dụng, tra cứu cic, kiểm tra nợ xấu (339000 search/tháng)", "Interactive SPA + Pre-scoring Lead Capture", "Done (Live)"),

        # 2. CÔNG CỤ TIỆN ÍCH ĐỘC LẬP (STANDALONE UTILITIES)
        (4, "Đầu Tư & Vàng", "Công Cụ Tiện Ích (Standalone Tool)", "Bảng Giá Vàng SJC, 9999 Mới Nhất", "/gia-vang",
         "giá vàng hôm nay, giá vàng sjc, giá vàng 9999 (78879270 search/tháng)", "Realtime Feed API + Client-side Converter", "Done (Live)"),

        (5, "Thu Nhập & Thuế", "Công Cụ Tiện Ích (Standalone Tool)", "Tính Lương Gross Sang Net 2026", "/tinh-luong",
         "tính lương gross net, cách tính lương net (190690 search/tháng)", "Pure Client-side JS Calculator (Zero Latency)", "Done (Live)"),

        (6, "Tiết Kiệm & Lãi Suất", "Công Cụ Tiện Ích (Standalone Tool)", "Bảng So Sánh Lãi Suất Tiền Gửi 30+ Ngân Hàng", "/lai-suat",
         "lãi suất ngân hàng, lãi suất tiết kiệm cao nhất (1246620 search/tháng)", "Data Feed Sheet + Dynamic Matrix Filter", "In Progress (Tháng 9)"),

        (7, "Thu Nhập & Chi Tiêu", "Công Cụ Tiện Ích (Standalone Tool)", "Phân Bổ Lương 50/30/20 & Hũ Tài Chính", "/tai-chinh/phan-bo-luong",
         "phân bổ lương, quy tắc 50 30 20, 6 chiếc hũ tài chính (50000 search/tháng)", "Client-side Budgeting Simulator", "In Progress (Tháng 9)"),

        (8, "Tiền Tệ & Ngoại Hối", "Cổng Danh Mục Master", "Cổng Tỷ Giá Ngoại Tệ 20+ Ngân Hàng", "/tai-chinh/ty-gia",
         "tỷ giá ngoại tệ, bảng tỷ giá ngân hàng, tỷ giá vietcombank (8516910 search/tháng)", "Master Exchange Rate Hub + Currency Selector", "In Progress (Tháng 9)"),

        (9, "An Sinh Xã Hội", "Công Cụ Tiện Ích (Standalone Tool)", "Công Cụ Ước Tính Rút Bảo Hiểm Xã Hội 1 Lần", "/tai-chinh/bao-hiem-xa-hoi",
         "rút bảo hiểm xã hội 1 lần, cách tính bhxh 1 lần (441060 search/tháng)", "Client-side BHXH Formula Engine", "Done (Live)"),

        (10, "An Sinh Xã Hội", "Công Cụ Tiện Ích (Standalone Tool)", "Dự Toán Mức Hưởng Lương Hưu Tuổi Già", "/tai-chinh/luong-huu",
         "cách tính lương hưu, tuổi nghỉ hưu, tỷ lệ hưởng lương hưu (150000 search/tháng)", "Retirement Pension Estimator", "Done (Live)"),

        (11, "Tín Dụng & Tiêu Dùng", "Trang Sản Phẩm & Simulator", "Ví Trả Sau (PayLater)", "/vi-tra-sau",
         "ví trả sau momo, cách dùng ví trả sau (116390 search/tháng)", "Standalone Product Landing + Fee Calculator", "Done (Live)"),

        (12, "Tín Dụng & Cho Vay", "Trang Sản Phẩm & Simulator", "Vay Nhanh Tiêu Dùng Tín Chấp", "/vay-nhanh",
         "vay nhanh momo, vay tiền online, vay tín chấp (389820 search/tháng)", "Standalone Product Landing + Loan Schedule Estimator", "Done (Live)"),

        # 3. PROGRAMMATIC CÁC TRANG CẶP TIỀN (CURRENCY PAIRS)
        (13, "Tiền Tệ & Ngoại Hối", "Trang Con Cặp Tiền (Programmatic Subpage)", "Tỷ Giá USD/VND: Đổi Đô La Mỹ Sang Tiền Việt", "/tai-chinh/ty-gia/usd-vnd",
         "tỷ giá usd, usd to vnd, 1 usd bằng bao nhiêu tiền việt (3200000 search/tháng)", "Programmatic Dynamic Route + Realtime Converter", "In Progress (Tháng 9)"),

        (14, "Tiền Tệ & Ngoại Hối", "Trang Con Cặp Tiền (Programmatic Subpage)", "Tỷ Giá JPY/VND: Đổi Yên Nhật Sang Tiền Việt", "/tai-chinh/ty-gia/jpy-vnd",
         "tỷ giá yên nhật, jpy to vnd, yên nhật hôm nay (1850000 search/tháng)", "Programmatic Dynamic Route + Remittance Hook", "In Progress (Tháng 9)"),

        (15, "Tiền Tệ & Ngoại Hối", "Trang Con Cặp Tiền (Programmatic Subpage)", "Tỷ Giá EUR/VND: Đổi Đồng Euro Sang Tiền Việt", "/tai-chinh/ty-gia/eur-vnd",
         "tỷ giá euro, eur to vnd, 1 euro bằng bao nhiêu vnd (850000 search/tháng)", "Programmatic Dynamic Route + FX Matrix", "In Progress (Tháng 9)"),

        (16, "Tiền Tệ & Ngoại Hối", "Trang Con Cặp Tiền (Programmatic Subpage)", "Tỷ Giá KRW/VND: Đổi Won Hàn Quốc Sang Tiền Việt", "/tai-chinh/ty-gia/krw-vnd",
         "tỷ giá won, krw to vnd, 1 won bằng bao nhiêu vnd (720000 search/tháng)", "Programmatic Dynamic Route + Remittance Hook", "In Progress (Tháng 9)"),

        (17, "Tiền Tệ & Ngoại Hối", "Trang Con Cặp Tiền (Programmatic Subpage)", "Tỷ Giá CNY/VND: Đổi Nhân Dân Tệ Sang Tiền Việt", "/tai-chinh/ty-gia/cny-vnd",
         "tỷ giá nhân dân tệ, cny to vnd, đổi tiền trung quốc (680000 search/tháng)", "Programmatic Dynamic Route + FX Matrix", "Planned"),

        (18, "Tiền Tệ & Ngoại Hối", "Trang Con Cặp Tiền (Programmatic Subpage)", "Tỷ Giá GBP/VND: Đổi Bảng Anh Sang Tiền Việt", "/tai-chinh/ty-gia/gbp-vnd",
         "tỷ giá bảng anh, gbp to vnd, bảng anh hôm nay (410000 search/tháng)", "Programmatic Dynamic Route + FX Matrix", "Planned"),

        (19, "Tiền Tệ & Ngoại Hối", "Trang Con Cặp Tiền (Programmatic Subpage)", "Tỷ Giá AUD/VND: Đổi Đô La Úc Sang Tiền Việt", "/tai-chinh/ty-gia/aud-vnd",
         "tỷ giá aud, aud to vnd, đô la úc hôm nay (350000 search/tháng)", "Programmatic Dynamic Route + FX Matrix", "Planned"),

        (20, "Tiền Tệ & Ngoại Hối", "Trang Con Cặp Tiền (Programmatic Subpage)", "Tỷ Giá SGD/VND: Đổi Đô La Singapore Sang Tiền Việt", "/tai-chinh/ty-gia/sgd-vnd",
         "tỷ giá sgd, sgd to vnd, đô singapore (220000 search/tháng)", "Programmatic Dynamic Route + FX Matrix", "Planned"),

        (21, "Tiền Tệ & Ngoại Hối", "Trang Con Cặp Tiền (Programmatic Subpage)", "Tỷ Giá THB/VND: Đổi Baht Thái Sang Tiền Việt", "/tai-chinh/ty-gia/thb-vnd",
         "tỷ giá baht thái, thb to vnd, đổi tiền thái lan (190000 search/tháng)", "Programmatic Dynamic Route + Travel Hook", "Planned"),

        # 4. PROGRAMMATIC 34 NGÂN HÀNG (BANK HUB & SPOKES)
        (22, "Ngân Hàng & Thẻ", "Cổng Danh Mục Master", "Danh Bạ 34 Ngân Hàng Việt Nam (Lãi Suất & Hotline)", "/tai-chinh/ngan-hang",
         "danh sách ngân hàng việt nam, số tổng đài ngân hàng (21063400 search/tháng)", "Programmatic Hub Grid + Search Filter", "Planned"),

        (23, "Ngân Hàng & Thẻ", "Trang Chi Tiết Ngân Hàng (Bank Spoke)", "Thông Tin Ngân Hàng VietinBank", "/tai-chinh/ngan-hang/vietinbank",
         "vietinbank, ngân hàng vietinbank, tổng đài vietinbank (1786700 search/tháng)", "Programmatic Dynamic Route ([bank-slug])", "Planned"),

        (24, "Ngân Hàng & Thẻ", "Trang Chi Tiết Ngân Hàng (Bank Spoke)", "Thông Tin Ngân Hàng Vietcombank", "/tai-chinh/ngan-hang/vietcombank",
         "vietcombank, vcb, hotline vietcombank, tỷ giá vietcombank (1770850 search/tháng)", "Programmatic Dynamic Route ([bank-slug])", "Planned"),

        (25, "Ngân Hàng & Thẻ", "Trang Chi Tiết Ngân Hàng (Bank Spoke)", "Thông Tin Ngân Hàng Agribank", "/tai-chinh/ngan-hang/agribank",
         "agribank, ngân hàng nông nghiệp, lãi suất agribank (1628790 search/tháng)", "Programmatic Dynamic Route ([bank-slug])", "Planned"),

        (26, "Ngân Hàng & Thẻ", "Trang Chi Tiết Ngân Hàng (Bank Spoke)", "Thông Tin Ngân Hàng BIDV", "/tai-chinh/ngan-hang/bidv",
         "bidv, ngân hàng bidv, smartbanking bidv (1483320 search/tháng)", "Programmatic Dynamic Route ([bank-slug])", "Planned"),

        (27, "Ngân Hàng & Thẻ", "Trang Chi Tiết Ngân Hàng (Bank Spoke)", "Thông Tin Ngân Hàng MBBank", "/tai-chinh/ngan-hang/mbbank",
         "mbbank, mb bank, ngân hàng quân đội (1478140 search/tháng)", "Programmatic Dynamic Route ([bank-slug])", "Planned"),

        (28, "Ngân Hàng & Thẻ", "Trang Chi Tiết Ngân Hàng (Bank Spoke)", "Thông Tin Ngân Hàng Techcombank", "/tai-chinh/ngan-hang/techcombank",
         "techcombank, tcb, ngân hàng techcombank (1247920 search/tháng)", "Programmatic Dynamic Route ([bank-slug])", "Planned"),

        (29, "Ngân Hàng & Thẻ", "Trang Chi Tiết Ngân Hàng (Bank Spoke)", "Thông Tin Ngân Hàng ACB", "/tai-chinh/ngan-hang/acb",
         "acb, ngân hàng á châu, hotline acb (1152430 search/tháng)", "Programmatic Dynamic Route ([bank-slug])", "Planned"),

        (30, "Ngân Hàng & Thẻ", "Trang Chi Tiết Ngân Hàng (Bank Spoke)", "Thông Tin Ngân Hàng VPBank", "/tai-chinh/ngan-hang/vpbank",
         "vpbank, neo vpbank, tiết kiệm vpbank (1014380 search/tháng)", "Programmatic Dynamic Route + TKO Partner Hook", "Planned"),

        (31, "Ngân Hàng & Thẻ", "Trang Chi Tiết Ngân Hàng (Bank Spoke)", "Thông Tin Ngân Hàng Bản Việt (BVBank)", "/tai-chinh/ngan-hang/bvbank",
         "bvbank, ngân hàng bản việt, tiết kiệm bản việt momo (105000 search/tháng)", "Programmatic Dynamic Route + Strategic Partner Hook", "Planned"),

        (32, "Ngân Hàng & Thẻ", "Hệ Thống 25 Ngân Hàng Còn Lại", "Trang Chi Tiết 25 Ngân Hàng Khác (TPB, VIB, STB...)", "/tai-chinh/ngan-hang/[bank-slug]",
         "25 ngân hàng thương mại cổ phần theo danh sách sheet Banks", "Programmatic Multi-tenant Template Engine", "Planned"),

        # 5. CHƯƠNG TRÌNH ĐẦU TƯ & TÍCH SẢN
        (33, "Đầu Tư & Chứng Chỉ Quỹ", "Trang Chương Trình Trọng Tâm", "Chương Trình Thực Tập Sinh Đầu Tư", "/tai-chinh/thuc-tap-sinh-dau-tu",
         "thực tập sinh đầu tư, học đầu tư chứng khoán, tích sản cho giới trẻ", "Campaign Landing Page + Gamification Missions", "Planned (Tháng 10)"),

        (34, "Đầu Tư & Chứng Chỉ Quỹ", "Trang Cổng Dịch Vụ", "Cổng Hướng Dẫn Mở Tài Khoản Chứng Khoán CVX", "/tai-chinh/chung-khoan",
         "chứng khoán momo, mở tài khoản chứng khoán vietcap cvx (12648090 search/tháng)", "Vertical Hub + Vietcap Onboarding Flow", "Planned (Tháng 10)"),

        (35, "Đầu Tư & Chứng Chỉ Quỹ", "Trang Cổng Dịch Vụ", "Cổng Thông Tin Chứng Chỉ Quỹ (SIP)", "/tai-chinh/chung-chi-quy",
         "chứng chỉ quỹ, quỹ mở dragon capital, vinacapital, ssiam (25790 search/tháng)", "Vertical Hub + Fund Comparison Matrix", "Planned (Tháng 10)"),

        # 6. BLOG HUBS & NỘI DUNG CHUYÊN ĐỀ
        (36, "Content Strategy", "Trang Danh Mục Blog (Blog Listing)", "Cẩm Nang Tài Chính Cá Nhân & Đầu Tư", "/tai-chinh/blog",
         "cẩm nang tài chính, kiến thức tài chính cá nhân, quản lý dòng tiền", "CollectionPage + Category Filters (5 Pillars)", "In Progress"),

        (37, "Sức Khỏe Tín Dụng", "Trang Danh Mục Blog (Blog Listing)", "Cẩm Nang Điểm Tín Dụng & Nợ Xấu CIC", "/diem-tin-dung/blog",
         "hướng dẫn xóa nợ xấu, cách tăng điểm tín dụng, cảnh báo lừa đảo vay", "CollectionPage + Credit Scoring Guides", "In Progress"),

        (38, "Content Strategy", "Kho 50 Bài Viết Chi Tiết (Blog Details)", "50 Bài Viết Chuyên Đề Chuẩn SEO (Theo Sheet Content Plan)", "/tai-chinh/blog/[slug] & /diem-tin-dung/blog/[slug]",
         "Toàn bộ 50 từ khóa phân bổ trên 5 trụ cột theo kế hoạch Tháng 9", "MoSpark GenAI Content Architecture + Schema Article", "In Progress (50 Articles)")
    ]

    for row_item in sitemap_items:
        ws_s.append(list(row_item))
        r = ws_s.max_row
        num_lines = str(row_item[5] or '').count('\n') + 1
        ws_s.row_dimensions[r].height = max(38.0, num_lines * 18.0)

        ws_s.cell(row=r, column=1).alignment = Alignment(horizontal='center', vertical='center')
        ws_s.cell(row=r, column=2).alignment = Alignment(horizontal='center', vertical='center'); ws_s.cell(row=r, column=2).font = font_bold
        ws_s.cell(row=r, column=3).alignment = Alignment(horizontal='center', vertical='center')
        ws_s.cell(row=r, column=4).alignment = Alignment(horizontal='left', vertical='center'); ws_s.cell(row=r, column=4).font = font_bold
        ws_s.cell(row=r, column=5).alignment = Alignment(horizontal='left', vertical='center'); ws_s.cell(row=r, column=5).font = font_url
        ws_s.cell(row=r, column=6).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws_s.cell(row=r, column=7).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws_s.cell(row=r, column=8).alignment = Alignment(horizontal='center', vertical='center')

        time_or_status = str(row_item[7] or '')
        c_st = ws_s.cell(row=r, column=8)
        if "Done" in time_or_status: c_st.fill = status_done_fill; c_st.font = status_done_font
        elif "In Progress" in time_or_status: c_st.fill = status_prog_fill; c_st.font = status_prog_font
        elif "Planned" in time_or_status: c_st.fill = status_plan_fill; c_st.font = status_plan_font

        for c_i in range(1, 9):
            ws_s.cell(row=r, column=c_i).border = thin_border

    col_widths_s = [6, 22, 28, 38, 42, 55, 36, 18]
    for idx, w in enumerate(col_widths_s, start=1):
        ws_s.column_dimensions[get_column_letter(idx)].width = w

    save_with_backup(wb, '05_HUBS/financial-hub-roadmap.xlsx', dry_run=args.dry_run, make_backup=not args.no_backup)
    print(f"Successfully created 'Sitemap' sheet with {len(sitemap_items)} categorized architecture entries in 05_HUBS/financial-hub-roadmap.xlsx!")

if __name__ == "__main__":
    _ap = standard_argparser("Fix runner - safe with --dry-run")
    main(_ap.parse_args())

