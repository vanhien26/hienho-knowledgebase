import openpyxl

wb = openpyxl.load_workbook('05_HUBS/financial-hub-roadmap.xlsx')
ws = wb['Roadmap']

# Update row for Dự Án 7: Bộ Quy Đổi Ngoại Tệ & Tỷ Giá
for r in range(3, ws.max_row + 1):
    val_name = str(ws.cell(row=r, column=4).value or '')
    if 'Tỷ Giá' in val_name or 'Dự Án 7' in val_name or 'Ngoại Tệ' in val_name:
        ws.cell(row=r, column=4).value = "Dự Án 7: Bộ Quy Đổi Tỷ Giá & Hub Các Cặp Tiền Tệ"
        ws.cell(row=r, column=5).value = (
            "• Xây dựng Trang Cổng Tỷ Giá Master (/ty-gia) tích hợp bảng so sánh realtime 30+ ngân hàng.\n"
            "• Triển khai Kiến trúc Programmatic Subpages theo từng Cặp Tiền Tệ (/ty-gia/[pair-slug]) cho 20+ ngoại tệ chủ lực (USD/VND, JPY/VND, EUR/VND, KRW/VND, CNY/VND, GBP/VND, AUD/VND, SGD/VND, THB/VND...).\n"
            "• Mỗi trang con cặp tiền tích hợp: Bộ quy đổi 2 chiều tức thì, Biểu đồ biến động lịch sử 7-30 ngày của riêng cặp tiền, Bảng so sánh tỷ giá Mua/Bán giữa các ngân hàng thương mại và Cẩm nang đổi tiền/du lịch.\n"
            "• Biểu đồ theo dõi biến động tỷ giá ngân hàng trung ương và thị trường tự do."
        )
        ws.cell(row=r, column=6).value = (
            "CTA theo ngữ cảnh cặp tiền: 'Mở Thẻ Quốc Tế Visa/Mastercard MoMo' miễn phí chuyển đổi ngoại tệ chi tiêu nước ngoài / 'Chuyển Tiền Quốc Tế' in-app."
        )

wb.save('05_HUBS/financial-hub-roadmap.xlsx')
print("Successfully updated Currency Pair Programmatic Architecture in 05_HUBS/financial-hub-roadmap.xlsx!")

# Update BRD as well
with open('05_HUBS/financial-hub-brd.md', 'r', encoding='utf-8') as f:
    brd = f.read()

# Enhance Section 5.1 & Section 10 for Currency Pairs Hub
old_s10_str = "Dự Án 7: Bộ Quy Đổi Ngoại Tệ & Tỷ Giá"
if old_s10_str in brd:
    brd = brd.replace(
        "Bộ công cụ quy đổi tiền tệ tức thì giữa 20+ ngoại tệ (USD, JPY, EUR...) và bảng tỷ giá so sánh giữa các ngân hàng thương mại lớn.",
        "Trang Cổng Master /ty-gia và Hệ thống Programmatic Subpages theo từng cặp tiền (/ty-gia/[pair-slug]: USD/VND, JPY/VND, EUR/VND, KRW/VND, CNY/VND, GBP/VND...) kèm biểu đồ biến động lịch sử và so sánh tỷ giá đa ngân hàng."
    )
    with open('05_HUBS/financial-hub-brd.md', 'w', encoding='utf-8') as f:
        f.write(brd)
    print("Successfully updated BRD with Currency Pair Programmatic Architecture!")
