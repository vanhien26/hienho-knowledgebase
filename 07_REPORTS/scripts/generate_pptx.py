#!/usr/bin/env python3
"""Generate PLG MoSpark Proposal PPTX — Clean layout, MoMo branding, correct slide order."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ── MoMo Brand Colors ──
C_PRIMARY = RGBColor(0xA5, 0x00, 0x64)
C_DARK = RGBColor(0x8B, 0x00, 0x55)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_GRAY50 = RGBColor(0xF9, 0xFA, 0xFB)
C_GRAY200 = RGBColor(0xE5, 0xE7, 0xEB)
C_GRAY500 = RGBColor(0x6B, 0x72, 0x80)
C_GRAY700 = RGBColor(0x37, 0x41, 0x51)
C_GRAY900 = RGBColor(0x11, 0x18, 0x27)
C_GREEN = RGBColor(0x10, 0xB9, 0x81)
C_BLUE = RGBColor(0x3B, 0x82, 0xF6)
C_PINK_LIGHT = RGBColor(0xFF, 0xCC, 0xE5)
C_PINK_BG = RGBColor(0xFD, 0xF2, 0xF7)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

FONT = 'Calibri'
M_LEFT = Inches(1.0)
M_TOP = Inches(0.5)
CONTENT_W = Inches(11.3)

# ── Helpers ──

def solid_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def txt(slide, l, t, w, h, text, sz=14, bold=False, color=C_GRAY700, align=PP_ALIGN.LEFT):
    """Simple single-paragraph textbox."""
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(sz)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = FONT
    p.alignment = align
    return tb

def rich_text(slide, l, t, w, h, runs_list, align=PP_ALIGN.LEFT):
    """Textbox with multiple styled runs in one paragraph."""
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    for text, sz, bold, color in runs_list:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(sz)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = FONT
    return tb

def bullet_box(slide, l, t, w, items, sz=12):
    """Multi-paragraph bullet list."""
    tb = slide.shapes.add_textbox(l, t, w, Inches(5))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        p.space_before = Pt(2)
        # dot
        r = p.add_run()
        r.text = "●  "
        r.font.size = Pt(7)
        r.font.color.rgb = C_PRIMARY
        r.font.name = FONT
        if isinstance(item, tuple):
            bold_part, rest_part = item
            if bold_part:
                r2 = p.add_run()
                r2.text = bold_part
                r2.font.size = Pt(sz)
                r2.font.bold = True
                r2.font.color.rgb = C_GRAY900
                r2.font.name = FONT
            r3 = p.add_run()
            r3.text = rest_part
            r3.font.size = Pt(sz)
            r3.font.color.rgb = C_GRAY700
            r3.font.name = FONT
        else:
            r2 = p.add_run()
            r2.text = str(item)
            r2.font.size = Pt(sz)
            r2.font.color.rgb = C_GRAY700
            r2.font.name = FONT
    return tb

def badge(slide, l, t, text):
    """Slide category badge."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, Inches(3.5), Inches(0.32))
    shape.fill.solid()
    shape.fill.fore_color.rgb = C_PRIMARY
    shape.line.fill.background()
    tf = shape.text_frame
    tf.margin_left = Pt(8)
    tf.margin_right = Pt(8)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.font.name = FONT
    p.alignment = PP_ALIGN.CENTER

def section_heading(slide, l, t, w, text):
    """Section title with left accent bar."""
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t + Inches(0.04), Inches(0.05), Inches(0.26))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_PRIMARY
    bar.line.fill.background()
    txt(slide, l + Inches(0.14), t, w, Inches(0.35), text, sz=15, bold=True, color=C_GRAY900)

def divider_line(slide, l, t, w):
    """Thin horizontal line."""
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = C_GRAY200
    line.line.fill.background()


# ═══════════════════════════════════════
# SLIDE 0: COVER
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
solid_bg(s, C_DARK)

txt(s, Inches(0), Inches(1.5), prs.slide_width, Inches(0.4),
    "MOMO  ·  GROWTH PLATFORM DIVISION", sz=13, bold=True, color=C_PINK_LIGHT, align=PP_ALIGN.CENTER)

txt(s, Inches(1.5), Inches(2.5), Inches(10.3), Inches(1.4),
    "Đề Xuất Chiến Lược PLG\nTrên Nền Tảng MoSpark", sz=36, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)

txt(s, Inches(2), Inches(4.1), Inches(9.3), Inches(0.5),
    "Pilot: Chiến dịch Tra cứu Phạt Nguội", sz=17, color=C_PINK_LIGHT, align=PP_ALIGN.CENTER)

line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.0), Inches(5.0), Inches(1.3), Inches(0.03))
line.fill.solid()
line.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
line.line.fill.background()

txt(s, Inches(2), Inches(5.5), Inches(9.3), Inches(0.4),
    "Tài liệu đọc hiểu dành cho C-Level và các Head of BUs  ·  Tháng 7/2026",
    sz=11, color=RGBColor(0xBB, 0x88, 0xA0), align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════
# SLIDE 1: Web Platform & MoSpark Roles
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
solid_bg(s, C_WHITE)

badge(s, M_LEFT, Inches(0.5), "SLIDE 01  —  ĐỊNH VỊ")
txt(s, M_LEFT, Inches(1.0), CONTENT_W, Inches(0.55),
    "Web Platform & MoSpark — Vai Trò Nền Tảng", sz=24, bold=True, color=C_GRAY900)
txt(s, M_LEFT, Inches(1.55), CONTENT_W, Inches(0.3),
    "Vai trò chiến lược của momo.vn ngoài App và hạ tầng MoSpark", sz=12, color=C_GRAY500)

divider_line(s, M_LEFT, Inches(1.9), CONTENT_W)

# Left Column (Web Platform)
section_heading(s, M_LEFT, Inches(2.1), Inches(5.2), "Web Platform (momo.vn) — Kênh ngoài App")
bullet_box(s, Inches(1.2), Inches(2.55), Inches(5.0), [
    ("Financial Authority: ", "Xây dựng momo.vn thành nguồn thông tin tin cậy hàng đầu ngành tài chính và thanh toán, được cả người dùng lẫn AI Search Engine công nhận."),
    ("Growth & MAU: ", "Thu hút traffic tự nhiên ngoài App, tối ưu phễu Web-to-App để mang về người dùng mới và tăng MAU."),
    ("Product-Led Growth: ", "Phát triển công cụ tương tác (Tra cứu, Calculator...) giải quyết nhu cầu thực tế, tạo lực kéo tự nhiên."),
], sz=11)

# Right Column (MoSpark)
section_heading(s, Inches(7.0), Inches(2.1), Inches(5.2), "MoSpark — Hệ thống tự động hóa Web")
bullet_box(s, Inches(7.2), Inches(2.55), Inches(5.0), [
    ("Platform-as-a-Service: ", "Cung cấp hạ tầng và Quality Gate để các Cell Team tự triển khai Mini Web mà không phụ thuộc Dev."),
    ("GenAI Content: ", "Tự động sản xuất nội dung quy mô lớn với quy trình kiểm duyệt chất lượng và pháp lý tài chính chặt chẽ."),
    ("SEO Inventory: ", "Hệ thống dữ liệu nghiên cứu thị trường, theo dõi và đo lường thị phần tìm kiếm (Share of Voice)."),
], sz=11)


# ═══════════════════════════════════════
# SLIDE 2: Executive Summary
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
solid_bg(s, C_WHITE)

badge(s, M_LEFT, Inches(0.5), "SLIDE 02  —  EXECUTIVE SUMMARY")
txt(s, M_LEFT, Inches(1.0), CONTENT_W, Inches(0.55),
    "Tăng Trưởng Dẫn Dắt Bởi Sản Phẩm (PLG) Trên Nền Tảng MoSpark", sz=24, bold=True, color=C_GRAY900)
txt(s, M_LEFT, Inches(1.55), CONTENT_W, Inches(0.3),
    "Tóm tắt cốt lõi", sz=12, color=C_GRAY500)

divider_line(s, M_LEFT, Inches(1.9), CONTENT_W)

bullet_box(s, M_LEFT, Inches(2.1), CONTENT_W, [
    ("Dịch chuyển chiến lược: ", 'Chuyển đổi mô hình Web MoMo.vn từ "Sản xuất bài viết SEO" sang "Cung cấp Tiện ích tương tác (PLG Utilities)" để giải quyết trực tiếp nhu cầu thực tế của người dùng.'),
    ("Kim chỉ nam: ", "Sử dụng hệ thống dữ liệu SEO Inventory để định lượng thị phần tìm kiếm (Share of Voice) và định hướng đầu tư Use Case có nhu cầu thực tế cao."),
    ("Dự án Pilot: ", "Chiến dịch Phạt Nguội chứng minh tính khả thi và hiệu năng bứt phá của mô hình PLG thông qua 4 bước:"),
], sz=13)

# Flow as simple table
tbl_shape = s.shapes.add_table(1, 7, Inches(1.5), Inches(4.8), Inches(10.3), Inches(0.7))
tbl = tbl_shape.table
labels = ["SEO Inventory\n(Định vị)", "→", "Vibe Code\n(Phát triển)", "→", "PLG Project\n(Quản trị)", "→", "Performance\n(Hiệu quả)"]
for i, label in enumerate(labels):
    cell = tbl.cell(0, i)
    cell.text = label
    p = cell.text_frame.paragraphs[0]
    p.font.name = FONT
    p.alignment = PP_ALIGN.CENTER
    if label == "→":
        tbl.columns[i].width = Inches(0.5)
        p.font.size = Pt(18)
        p.font.color.rgb = C_GRAY500
        p.font.bold = False
        cell.fill.background()
    else:
        tbl.columns[i].width = Inches(2.3)
        p.font.size = Pt(11)
        p.font.bold = True
        is_last = (i == 6)
        if is_last:
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_PRIMARY
            p.font.color.rgb = C_WHITE
        else:
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_WHITE
            p.font.color.rgb = C_PRIMARY
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE


# ═══════════════════════════════════════
# SLIDE 3: Bối cảnh
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
solid_bg(s, C_WHITE)

badge(s, M_LEFT, Inches(0.5), "SLIDE 03  —  BỐI CẢNH")
txt(s, M_LEFT, Inches(1.0), CONTENT_W, Inches(0.55),
    "Sự Dịch Chuyển Của Search & Nút Thắt Web Platform", sz=24, bold=True, color=C_GRAY900)

# Left section
section_heading(s, M_LEFT, Inches(1.8), Inches(5), "Sự dịch chuyển hành vi tìm kiếm")
bullet_box(s, Inches(1.2), Inches(2.2), Inches(5.2), [
    ("", "Lượng tìm kiếm truyền thống trên Google sụt giảm ~20% YoY, nhường chỗ cho các nền tảng AI Search (ChatGPT, Perplexity)."),
    ("", "Người dùng có xu hướng tìm kiếm giải pháp tức thời thay vì đọc nội dung dài."),
], sz=12)

# Right section
section_heading(s, Inches(7.0), Inches(1.8), Inches(5), "Nút thắt của Web MoMo hiện tại")
bullet_box(s, Inches(7.2), Inches(2.2), Inches(5.2), [
    ("", "Lưu lượng truy cập lớn ngoài App (Non-MoMo Users) chưa được tối ưu hóa chuyển đổi Web-to-App do thiếu công cụ giữ chân."),
    ("", "Người dùng tìm đến trang Blog, đọc thông tin và rời đi mà không phát sinh hành vi cài đặt App hay liên kết tài khoản."),
], sz=12)

divider_line(s, M_LEFT, Inches(4.3), CONTENT_W)

# Highlight
highlight_bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, M_LEFT, Inches(4.5), CONTENT_W, Inches(1.0))
highlight_bg.fill.solid()
highlight_bg.fill.fore_color.rgb = C_PINK_BG
highlight_bg.line.fill.background()

bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, M_LEFT, Inches(4.5), Inches(0.05), Inches(1.0))
bar.fill.solid()
bar.fill.fore_color.rgb = C_PRIMARY
bar.line.fill.background()

rich_text(s, Inches(1.3), Inches(4.65), Inches(10.7), Inches(0.7), [
    ("Nhận định chiến lược: ", 13, True, C_PRIMARY),
    ("MoMo.vn cần dịch chuyển từ mô hình Blog SEO sang Nền tảng Công cụ Tiện ích để giữ vai trò là nguồn dữ liệu tin cậy được các AI Search Engine trích dẫn.", 13, False, C_GRAY700),
])


# ═══════════════════════════════════════
# SLIDE 4: PLG & Co-Investment
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
solid_bg(s, C_WHITE)

badge(s, M_LEFT, Inches(0.5), "SLIDE 04  —  CHIẾN LƯỢC")
txt(s, M_LEFT, Inches(1.0), CONTENT_W, Inches(0.55),
    "Chuyển Dịch Sang Chiến Lược PLG & Co-Investment", sz=24, bold=True, color=C_GRAY900)

section_heading(s, M_LEFT, Inches(1.75), Inches(10), "Triết lý PLG (Product-Led Growth) trên Web")
txt(s, Inches(1.2), Inches(2.15), CONTENT_W, Inches(0.45),
    "Lấy các công cụ tương tác tiện ích làm hạt nhân để giải quyết nhu cầu thực tế. Content bài viết chỉ đóng vai trò tối ưu hóa khả năng hiển thị.",
    sz=12, color=C_GRAY700)

# Flow table
tbl_shape = s.shapes.add_table(1, 7, Inches(1.2), Inches(2.8), Inches(10.9), Inches(0.6))
tbl = tbl_shape.table
labels = ["Tìm kiếm", "→", "Trải nghiệm Tool Web Lite", "→", "Kích hoạt Nhu cầu", "→", "Chuyển đổi mở App (W2A)"]
for i, label in enumerate(labels):
    cell = tbl.cell(0, i)
    cell.text = label
    p = cell.text_frame.paragraphs[0]
    p.font.name = FONT
    p.alignment = PP_ALIGN.CENTER
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if label == "→":
        tbl.columns[i].width = Inches(0.4)
        p.font.size = Pt(16)
        p.font.color.rgb = C_GRAY500
        cell.fill.background()
    else:
        tbl.columns[i].width = Inches(2.5)
        p.font.size = Pt(11)
        p.font.bold = True
        if i == 6:
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_PRIMARY
            p.font.color.rgb = C_WHITE
        else:
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_WHITE
            p.font.color.rgb = C_PRIMARY

divider_line(s, M_LEFT, Inches(3.65), CONTENT_W)

section_heading(s, M_LEFT, Inches(3.85), Inches(10), "Hợp tác đồng đầu tư giữa các BU (Cross-BU Co-Investment)")
txt(s, Inches(1.2), Inches(4.25), CONTENT_W, Inches(0.4),
    "Phạt Nguội đóng vai trò phễu hút traffic đại chúng chất lượng cao để phân phối chéo Lead cho các BU hưởng lợi:",
    sz=12, color=C_GRAY700)

# BU table
tbl_shape = s.shapes.add_table(3, 3, M_LEFT, Inches(4.85), CONTENT_W, Inches(2.0))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(1.8)
tbl.columns[1].width = Inches(5.5)
tbl.columns[2].width = Inches(4.0)

headers = ["BU", "Vai trò", "Mức đóng góp"]
for i, h in enumerate(headers):
    cell = tbl.cell(0, i)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = C_PRIMARY
    p = cell.text_frame.paragraphs[0]
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.font.name = FONT

rows_data = [
    ["BU Bảo Hiểm", "Tích hợp bán chéo Bảo hiểm TNDS và thân vỏ ngay tại màn hình kết quả xe không vi phạm", "Tài trợ 35-40% chi phí tiếp thị và vận hành"],
    ["BU Tài Chính", "Đề xuất giải ngân đóng phạt nhanh qua Ví Trả Sau hoặc Vay Nhanh cho lỗi phạt nặng ≥1,000,000đ", "Tài trợ 30% chi phí tiếp thị"],
]
for r, row in enumerate(rows_data):
    for c, val in enumerate(row):
        cell = tbl.cell(r + 1, c)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_WHITE
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(11)
        p.font.name = FONT
        p.font.color.rgb = C_GRAY700
        if c == 0:
            p.font.bold = True
            p.font.color.rgb = C_GRAY900


# ═══════════════════════════════════════
# SLIDE 5: MoSpark & SEO Inventory
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
solid_bg(s, C_WHITE)

badge(s, M_LEFT, Inches(0.5), "SLIDE 05  —  NỀN TẢNG")
txt(s, M_LEFT, Inches(1.0), CONTENT_W, Inches(0.55),
    "MoSpark & SEO Inventory — Hệ Sinh Thái Tự Động Hóa PLG", sz=24, bold=True, color=C_GRAY900)
txt(s, M_LEFT, Inches(1.55), CONTENT_W, Inches(0.3),
    "Quy trình phát triển và vận hành PLG được tự động hóa thông qua 3 thành phần cốt lõi:", sz=12, color=C_GRAY500)

# 3 components as table
tbl_shape = s.shapes.add_table(4, 3, M_LEFT, Inches(2.2), CONTENT_W, Inches(4.5))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(3.77)
tbl.columns[1].width = Inches(3.77)
tbl.columns[2].width = Inches(3.76)

components = [
    ("THÀNH PHẦN 1", "SEO Inventory", "Định vị nhu cầu",
     "Khai thác dữ liệu tìm kiếm thị trường để định lượng dung lượng và định hướng Use Case đầu tư, đảm bảo mọi quyết định phát triển công cụ đều dựa trên nhu cầu thực tế."),
    ("THÀNH PHẦN 2", "PLG Project", "Quy hoạch & Quản trị chiến dịch",
     "Bộ não điều phối cấu trúc website. Cung cấp hệ quản trị nội dung (CMS) linh hoạt trong việc xây dựng, cập nhật trang và cơ chế phân quyền vai trò rõ ràng giữa Platform Team (quản trị kỹ thuật & khung chuẩn) và Cell Teams (tự chủ biên tập nội dung)."),
    ("THÀNH PHẦN 3", "GenAI Content", "Tự động hóa nội dung",
     "Hệ thống sản xuất nội dung quy mô lớn, cho phép tích hợp đa dạng các AI Model để tạo lập nội dung chuẩn SEO và thiết kế hình ảnh chuẩn Brand Guideline, đi kèm quy trình kiểm duyệt chất lượng nghiêm ngặt."),
]

for c, comp in enumerate(components):
    cell = tbl.cell(0, c)
    cell.fill.solid()
    cell.fill.fore_color.rgb = [C_PRIMARY, RGBColor(0xE9, 0x1E, 0x8C), C_BLUE][c]
    tbl.rows[0].height = Inches(0.08)

    cell = tbl.cell(1, c)
    cell.text = comp[0]
    cell.fill.solid()
    cell.fill.fore_color.rgb = C_GRAY50
    p = cell.text_frame.paragraphs[0]
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY
    p.font.name = FONT
    tbl.rows[1].height = Inches(0.35)

    cell = tbl.cell(2, c)
    cell.fill.solid()
    cell.fill.fore_color.rgb = C_GRAY50
    tf = cell.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = comp[1]
    p1.font.size = Pt(18)
    p1.font.bold = True
    p1.font.color.rgb = C_GRAY900
    p1.font.name = FONT
    p2 = tf.add_paragraph()
    p2.text = comp[2]
    p2.font.size = Pt(11)
    p2.font.color.rgb = C_GRAY500
    p2.font.name = FONT
    p2.space_before = Pt(4)
    tbl.rows[2].height = Inches(0.9)

    cell = tbl.cell(3, c)
    cell.text = comp[3]
    cell.fill.solid()
    cell.fill.fore_color.rgb = C_GRAY50
    p = cell.text_frame.paragraphs[0]
    p.font.size = Pt(12)
    p.font.color.rgb = C_GRAY700
    p.font.name = FONT
    cell.text_frame.word_wrap = True
    tbl.rows[3].height = Inches(2.5)
    cell.vertical_anchor = MSO_ANCHOR.TOP


# ═══════════════════════════════════════
# SLIDE 6: SEO Inventory — Market Potential
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
solid_bg(s, C_WHITE)

badge(s, M_LEFT, Inches(0.5), "SLIDE 06  —  CASE STUDY: PHẠT NGUỘI")
txt(s, M_LEFT, Inches(1.0), CONTENT_W, Inches(0.55),
    "1. SEO Inventory — Tiềm Năng Thị Trường", sz=24, bold=True, color=C_GRAY900)
txt(s, M_LEFT, Inches(1.55), CONTENT_W, Inches(0.3),
    "Nhu cầu thị trường & Quy mô tiếp cận (Market Potential)", sz=12, color=C_GRAY500)

# Stat table
tbl_shape = s.shapes.add_table(2, 2, M_LEFT, Inches(2.1), Inches(8), Inches(1.2))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(4)
tbl.columns[1].width = Inches(4)

stats = [("~3.56M", "lượt search/tháng — Volume lớn nhất nhóm Dịch vụ công"),
         ("~356K", "organic sessions/tháng — Mục tiêu 10% SoV")]

for c, (val, label) in enumerate(stats):
    cell_v = tbl.cell(0, c)
    cell_v.text = val
    cell_v.fill.solid()
    cell_v.fill.fore_color.rgb = C_GRAY50
    p = cell_v.text_frame.paragraphs[0]
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY
    p.font.name = FONT
    p.alignment = PP_ALIGN.CENTER
    tbl.rows[0].height = Inches(0.65)

    cell_l = tbl.cell(1, c)
    cell_l.text = label
    cell_l.fill.solid()
    cell_l.fill.fore_color.rgb = C_GRAY50
    p = cell_l.text_frame.paragraphs[0]
    p.font.size = Pt(10)
    p.font.color.rgb = C_GRAY500
    p.font.name = FONT
    p.alignment = PP_ALIGN.CENTER
    tbl.rows[1].height = Inches(0.45)

divider_line(s, M_LEFT, Inches(3.5), CONTENT_W)

bullet_box(s, M_LEFT, Inches(3.7), CONTENT_W, [
    ("Quy mô tìm kiếm lớn: ", "Đạt ~3.56M lượt search/tháng. Đây là cụm từ khóa có volume lớn nhất thuộc nhóm Dịch vụ công và Giao thông."),
    ("Nhu cầu thường trực: ", "Quy định pháp luật tăng mức phạt vi phạm giao thông, thúc đẩy nhu cầu tra cứu thường trực của tài xế."),
    ("Tệp người dùng giá trị: ", "Hướng đến nhóm sở hữu phương tiện (ô tô, xe máy) có khả năng chi trả cao — đối tượng mục tiêu của các dịch vụ bảo hiểm và tài chính tiêu dùng."),
    ("Mục tiêu chiếm lĩnh: ", "Đón đầu và chuyển đổi 10% thị phần tìm kiếm (Share of Voice), tương đương ~356,000 lượt truy cập tự nhiên/tháng."),
], sz=12)


# ═══════════════════════════════════════
# SLIDE 7: Vibe Code
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
solid_bg(s, C_WHITE)

badge(s, M_LEFT, Inches(0.5), "SLIDE 07  —  CASE STUDY: PHẠT NGUỘI")
txt(s, M_LEFT, Inches(1.0), CONTENT_W, Inches(0.55),
    "2. Vibe Code — Tốc Độ & Phát Triển", sz=24, bold=True, color=C_GRAY900)
txt(s, M_LEFT, Inches(1.55), CONTENT_W, Inches(0.3),
    "Phương thức phát triển siêu tốc với sự hỗ trợ của AI", sz=12, color=C_GRAY500)

section_heading(s, M_LEFT, Inches(2.1), Inches(5), "Rút ngắn thời gian ra mắt (Time-to-Market)")
bullet_box(s, Inches(1.2), Inches(2.5), Inches(5.2), [
    ("Chu kỳ phát triển tối giản: ", "Vibe Code hỗ trợ bởi AI giúp rút ngắn chu kỳ từ 1 tháng xuống còn 1 tuần để Go-live toàn bộ hệ thống (Trang chủ, trang ngách Ô tô, Xe máy, Xe máy điện và Blog)."),
    ("Tự chủ vận hành: ", "PM và Editor có thể tự cấu hình SEO/GEO On-page trực tiếp mà không phụ thuộc đội ngũ lập trình Web."),
], sz=12)

section_heading(s, Inches(7.0), Inches(2.1), Inches(5), "Trải nghiệm Mobile-First")
bullet_box(s, Inches(7.2), Inches(2.5), Inches(5.2), [
    ("", "Tối ưu tốc độ tải trang nhanh dưới 2 giây."),
    ("", "Giao diện tương tác, nhập biển số và trả kết quả mượt mà tương đương trải nghiệm in-app."),
], sz=12)

divider_line(s, M_LEFT, Inches(4.7), CONTENT_W)

# Key stats table
tbl_shape = s.shapes.add_table(2, 2, M_LEFT, Inches(5.0), Inches(8), Inches(1.2))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(4)
tbl.columns[1].width = Inches(4)

for c, (val, label, clr) in enumerate([("1 tuần", "Go-live toàn hệ thống (thay vì 1 tháng)", C_GREEN),
                                         ("< 2 giây", "Thời gian tải trang Mobile-First", C_BLUE)]):
    cell_v = tbl.cell(0, c)
    cell_v.text = val
    cell_v.fill.solid()
    cell_v.fill.fore_color.rgb = C_GRAY50
    p = cell_v.text_frame.paragraphs[0]
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = clr
    p.font.name = FONT
    p.alignment = PP_ALIGN.CENTER
    tbl.rows[0].height = Inches(0.65)

    cell_l = tbl.cell(1, c)
    cell_l.text = label
    cell_l.fill.solid()
    cell_l.fill.fore_color.rgb = C_GRAY50
    p = cell_l.text_frame.paragraphs[0]
    p.font.size = Pt(10)
    p.font.color.rgb = C_GRAY500
    p.font.name = FONT
    p.alignment = PP_ALIGN.CENTER
    tbl.rows[1].height = Inches(0.45)


# ═══════════════════════════════════════
# SLIDE 8: PLG Project
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
solid_bg(s, C_WHITE)

badge(s, M_LEFT, Inches(0.5), "SLIDE 08  —  CASE STUDY: PHẠT NGUỘI")
txt(s, M_LEFT, Inches(1.0), CONTENT_W, Inches(0.55),
    "3. PLG Project — Quản Trị Chiến Dịch", sz=24, bold=True, color=C_GRAY900)
txt(s, M_LEFT, Inches(1.55), CONTENT_W, Inches(0.3),
    "Phân tách mục tiêu và cơ chế vận hành tự động hóa", sz=12, color=C_GRAY500)

section_heading(s, M_LEFT, Inches(2.1), Inches(10), "Phân tách mục tiêu quản trị")

tbl_shape = s.shapes.add_table(3, 3, M_LEFT, Inches(2.55), CONTENT_W, Inches(1.3))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(2.0)
tbl.columns[1].width = Inches(2.5)
tbl.columns[2].width = Inches(6.8)

headers = ["Quản lý", "Dự án", "Mô tả"]
for i, h in enumerate(headers):
    cell = tbl.cell(0, i)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = C_PRIMARY
    p = cell.text_frame.paragraphs[0]
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.font.name = FONT

proj_rows = [
    ["BU vận hành", "Dự án Use Case", "Triển khai công cụ tiện ích (Phạt Nguội, eSIM) nhằm thu hút traffic tìm kiếm và chuyển đổi Web-to-App."],
    ["Platform vận hành", "Dự án Merchant Page", "Tăng Local SEO và thúc đẩy thanh toán Ví Trả Sau tại cửa hàng đối tác."],
]
for r, row in enumerate(proj_rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r + 1, c)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_WHITE
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(11)
        p.font.name = FONT
        p.font.color.rgb = C_GRAY700
        if c <= 1:
            p.font.bold = True
            p.font.color.rgb = C_GRAY900

divider_line(s, M_LEFT, Inches(4.15), CONTENT_W)

section_heading(s, M_LEFT, Inches(4.35), Inches(10), "Cơ chế vận hành tự động hóa & Quản trị")
bullet_box(s, M_LEFT, Inches(4.75), CONTENT_W, [
    ("Content Plan & Direction: ", "Quy hoạch từ khóa phân cấp (Topic → Cluster → Keyword) làm định hướng sản xuất nội dung có chủ đích, ngăn chặn chồng chéo từ khóa."),
    ("Quản trị CMS linh hoạt & Phân quyền rõ ràng: ", "Cho phép xây dựng, cập nhật nội dung tức thời mà không cần deploy lại hệ thống. Phân quyền rõ ràng giữa Platform Team (kiểm soát Technical Gate & Template chuẩn) và Cell Teams (tự chủ quản trị, biên tập nội dung)."),
    ("Tích hợp đa dạng AI Model (GenAI Content): ", "Kết nối linh hoạt nhiều mô hình AI khác nhau để tự động hóa sản xuất bài viết chuẩn SEO và thiết kế hình ảnh chuẩn Brand Guideline."),
    ("Cổng kiểm duyệt chất lượng & pháp lý (Quality Gate): ", "Đóng gói Prompt riêng biệt để chặn nội dung vi phạm pháp lý hoặc các từ khóa nhạy cảm về mặt tài chính."),
    ("Quy trình phê duyệt 2 lớp (Human-in-the-loop): ", "AI sinh dàn ý → PM phê duyệt → AI sinh bài chi tiết để xuất bản."),
], sz=11)


# ═══════════════════════════════════════
# SLIDE 9: Performance
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
solid_bg(s, C_WHITE)

badge(s, M_LEFT, Inches(0.5), "SLIDE 09  —  CASE STUDY: PHẠT NGUỘI")
txt(s, M_LEFT, Inches(1.0), CONTENT_W, Inches(0.55),
    "4. SEO/GEO Performance — Hiệu Quả Thực Tế", sz=24, bold=True, color=C_GRAY900)
txt(s, M_LEFT, Inches(1.55), CONTENT_W, Inches(0.3),
    "Số liệu tháng 6/2026", sz=12, color=C_GRAY500)

# KPI stats table
tbl_shape = s.shapes.add_table(3, 3, M_LEFT, Inches(2.0), CONTENT_W, Inches(1.4))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(3.77)
tbl.columns[1].width = Inches(3.77)
tbl.columns[2].width = Inches(3.76)

kpis = [("267.3K", "+17.1% MoM", "Views"),
        ("78.16%", "", "Tỷ lệ điền tra cứu"),
        ("14,886", "+47.4% MoM", "Lượt đăng nhập App từ Web")]

for c, (val, delta, label) in enumerate(kpis):
    cell_v = tbl.cell(0, c)
    cell_v.text = val
    cell_v.fill.solid()
    cell_v.fill.fore_color.rgb = C_GRAY50
    p = cell_v.text_frame.paragraphs[0]
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY
    p.font.name = FONT
    p.alignment = PP_ALIGN.CENTER
    tbl.rows[0].height = Inches(0.6)

    cell_d = tbl.cell(1, c)
    cell_d.text = delta if delta else "—"
    cell_d.fill.solid()
    cell_d.fill.fore_color.rgb = C_GRAY50
    p = cell_d.text_frame.paragraphs[0]
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_GREEN if delta else C_GRAY200
    p.font.name = FONT
    p.alignment = PP_ALIGN.CENTER
    tbl.rows[1].height = Inches(0.3)

    cell_l = tbl.cell(2, c)
    cell_l.text = label
    cell_l.fill.solid()
    cell_l.fill.fore_color.rgb = C_GRAY50
    p = cell_l.text_frame.paragraphs[0]
    p.font.size = Pt(10)
    p.font.color.rgb = C_GRAY500
    p.font.name = FONT
    p.alignment = PP_ALIGN.CENTER
    tbl.rows[2].height = Inches(0.4)

# Ecosystem section
section_heading(s, M_LEFT, Inches(3.65), Inches(10), "Kênh dẫn dắt Login App & Tăng trưởng MAU (Ecosystem Entry Point)")
bullet_box(s, M_LEFT, Inches(4.1), CONTENT_W, [
    ("Kênh mang lại Login App hiệu quả: ", "Chứng minh Web là phễu gom và điều hướng người dùng ngoài App vào App MoMo tối ưu."),
    ("Tăng trưởng MAU cho các dịch vụ khác trong hệ sinh thái: ", "Ghi nhận 9,825 người dùng đăng nhập thành công phát sinh giao dịch chéo (Ví Trả Sau, Bảo hiểm...), chiếm 66.0% tổng lượng đăng nhập từ Web."),
], sz=12)

divider_line(s, M_LEFT, Inches(5.4), CONTENT_W)

# GEO stats
section_heading(s, M_LEFT, Inches(5.6), Inches(10), "Hiệu quả hiển thị trên AI Search (GEO)")

tbl_shape = s.shapes.add_table(2, 3, M_LEFT, Inches(6.0), CONTENT_W, Inches(1.0))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(3.77)
tbl.columns[1].width = Inches(3.77)
tbl.columns[2].width = Inches(3.76)

geo_stats = [("1,920", "searches/tháng — Branded Search"),
             ("697", "lượt trích dẫn — GEO Citations"),
             ("7.84%", "SoA trên ChatGPT/Perplexity")]

for c, (val, label) in enumerate(geo_stats):
    cell_v = tbl.cell(0, c)
    cell_v.text = val
    cell_v.fill.solid()
    cell_v.fill.fore_color.rgb = C_GRAY50
    p = cell_v.text_frame.paragraphs[0]
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY
    p.font.name = FONT
    p.alignment = PP_ALIGN.CENTER
    tbl.rows[0].height = Inches(0.5)

    cell_l = tbl.cell(1, c)
    cell_l.text = label
    cell_l.fill.solid()
    cell_l.fill.fore_color.rgb = C_GRAY50
    p = cell_l.text_frame.paragraphs[0]
    p.font.size = Pt(10)
    p.font.color.rgb = C_GRAY500
    p.font.name = FONT
    p.alignment = PP_ALIGN.CENTER
    tbl.rows[1].height = Inches(0.4)


# ═══════════════════════════════════════
# SLIDE 10: Quy trình & Điều kiện
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
solid_bg(s, C_WHITE)

badge(s, M_LEFT, Inches(0.5), "SLIDE 10  —  VẬN HÀNH")
txt(s, M_LEFT, Inches(1.0), CONTENT_W, Inches(0.55),
    "Quy Trình Vận Hành & Điều Kiện Tiên Quyết", sz=24, bold=True, color=C_GRAY900)
txt(s, M_LEFT, Inches(1.55), CONTENT_W, Inches(0.4),
    "Để bắt đầu một dự án Product Growth mới trên nền tảng MoSpark, Cell Team và Web Platform cần đáp ứng các điều kiện cốt lõi sau",
    sz=12, color=C_GRAY500)

section_heading(s, M_LEFT, Inches(2.15), Inches(10), "Hai điều kiện tiên quyết áp dụng Use Case")
bullet_box(s, M_LEFT, Inches(2.55), CONTENT_W, [
    ("Sản phẩm phù hợp tinh thần Product-Led Growth: ", "Cell Team sở hữu sản phẩm hoặc công cụ tiện ích tương tác giải quyết trực tiếp nhu cầu (JTBD) của người dùng trên Web (thay vì bài viết thông tin đơn thuần)."),
    ("Cam kết đầu tư tiếp thị (SEO/GEO & Off-Page): ", "Cell Team sẵn sàng đầu tư ngân sách và nguồn lực tiếp thị tương ứng để tạo lực đẩy traffic ban đầu và xây dựng Domain Authority."),
], sz=12)

divider_line(s, M_LEFT, Inches(3.8), CONTENT_W)

# Left section
section_heading(s, M_LEFT, Inches(4.0), Inches(5), "Dữ liệu đầu vào cần chuẩn bị")
bullet_box(s, Inches(1.2), Inches(4.4), Inches(5.2), [
    ("Tài liệu nghiệp vụ (Markdown): ", "Đặc tả nghiệp vụ, luật pháp, USP và thông điệp sản phẩm."),
    ("Danh sách nghiên cứu từ khóa (CSV): ", "Danh mục từ khóa phân cấp kèm volume search thị trường."),
    ("Prompt & API Keys riêng biệt: ", "Cấu hình Prompt phù hợp ngành hàng và gán API key riêng."),
], sz=11)

# Right section
section_heading(s, Inches(7.0), Inches(4.0), Inches(5), "Quy trình phê duyệt & Kích hoạt")
bullet_box(s, Inches(7.2), Inches(4.4), Inches(5.2), [
    ("Technical Readiness: ", "Core API của Cell Team hoàn tất kết nối kỹ thuật cho luồng Web-to-App."),
    ("Quyết định phê duyệt: ", "Đánh giá mức độ sẵn sàng kỹ thuật và chất lượng bởi Web Product Lead để kích hoạt dự án trên hệ thống MoSpark."),
], sz=11)


# ═══════════════════════════════════════
# SLIDE 11: North Star Metrics
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
solid_bg(s, C_WHITE)

badge(s, M_LEFT, Inches(0.5), "SLIDE 11  —  ĐO LƯỜNG")
txt(s, M_LEFT, Inches(1.0), CONTENT_W, Inches(0.55),
    "Kỳ Vọng & Chỉ Số Đo Lường Cốt Lõi", sz=24, bold=True, color=C_GRAY900)
txt(s, M_LEFT, Inches(1.55), CONTENT_W, Inches(0.3),
    "North Star Metrics — H2/2026", sz=12, color=C_GRAY500)

tbl_shape = s.shapes.add_table(5, 3, M_LEFT, Inches(2.2), CONTENT_W, Inches(4.0))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(2.8)
tbl.columns[1].width = Inches(3.5)
tbl.columns[2].width = Inches(5.0)

headers = ["Phân Loại", "Chỉ Số Cốt Lõi", "Kỳ Vọng Đạt Được (H2/2026)"]
for i, h in enumerate(headers):
    cell = tbl.cell(0, i)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = C_PRIMARY
    p = cell.text_frame.paragraphs[0]
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.font.name = FONT

rows_data = [
    ["North Star (Web)", "Web-to-App Conversion (W2A CVR)", "Đạt tỷ lệ trung bình >10% (hướng tới target 12.5%)"],
    ["North Star (SEO)", "Share of Voice (SoV)", "Đạt 10% SoV (tương đương ~356,000 organic sessions/tháng)"],
    ["Business (Acquisition)", "Monthly Engagement Users (MEU)", "Đo lường bằng lượng xe lưu định theo dõi cảnh báo trên Web"],
    ["Business (Ecosystem)", "Cross-BU Conversion (Insurance & VTS)", "Tỷ lệ User mở Ví Trả Sau hoặc mua bảo hiểm từ phễu Phạt Nguội"],
]
for r, row in enumerate(rows_data):
    for c, val in enumerate(row):
        cell = tbl.cell(r + 1, c)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_WHITE if r % 2 == 0 else C_GRAY50
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(11)
        p.font.name = FONT
        p.font.color.rgb = C_GRAY700
        if c == 1:
            p.font.bold = True
            p.font.color.rgb = C_GRAY900


# ═══════════════════════════════════════
# SLIDE 12: Thank You
# ═══════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
solid_bg(s, C_DARK)

txt(s, Inches(0), Inches(2.3), prs.slide_width, Inches(0.8),
    "Thank You", sz=46, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)

line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.0), Inches(3.4), Inches(1.3), Inches(0.03))
line.fill.solid()
line.fill.fore_color.rgb = C_WHITE
line.line.fill.background()

txt(s, Inches(2), Inches(3.8), Inches(9.3), Inches(0.4),
    "Growth Platform Division (GPD) — Hỗ trợ phát triển và vận hành hệ thống MoSpark",
    sz=14, color=C_PINK_LIGHT, align=PP_ALIGN.CENTER)
txt(s, Inches(2), Inches(4.3), Inches(9.3), Inches(0.4),
    "Chúc dự án Pilot Phạt Nguội mang lại tăng trưởng bứt phá trong H2/2026",
    sz=14, color=C_PINK_LIGHT, align=PP_ALIGN.CENTER)
txt(s, Inches(2), Inches(5.1), Inches(9.3), Inches(0.4),
    "Liên hệ hỗ trợ: Văn Hiến (Web Product Lead)",
    sz=12, color=RGBColor(0xBB, 0x88, 0xA0), align=PP_ALIGN.CENTER)


# ── Save ──
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'plg_mospark_proposal.pptx')
prs.save(output_path)
print(f"✅ PPTX saved: {output_path}")
