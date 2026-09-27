import sys
import openpyxl

def merge_volume(volume_file, mapping_file='garage-keyword-mapping-index.xlsx', output_file='partner-garages-with-volume.xlsx'):
    print(f"Reading volume file: {volume_file}...")
    wb_vol = openpyxl.load_workbook(volume_file)
    ws_vol = wb_vol.active

    # Build dict: keyword -> search_volume
    kw_volume = {}
    for r in range(2, ws_vol.max_row + 1):
        kw = str(ws_vol.cell(r, 1).value or '').strip()
        vol = ws_vol.cell(r, 2).value or 0
        if kw:
            try:
                kw_volume[kw] = int(vol)
            except:
                pass

    print(f"Loaded search volumes for {len(kw_volume)} keywords.")

    # Read mapping index
    print(f"Reading mapping index: {mapping_file}...")
    wb_map = openpyxl.load_workbook(mapping_file)
    ws_map = wb_map['Mapping Index']

    # Map volume to Garage ID
    garage_vol_sum = {}
    for r in range(3, ws_map.max_row + 1):
        gid = str(ws_map.cell(r, 1).value or '').strip()
        kw = str(ws_map.cell(r, 2).value or '').strip()
        if gid and kw in kw_volume:
            garage_vol_sum[gid] = garage_vol_sum.get(gid, 0) + kw_volume[kw]

    print(f"Successfully mapped volume to {len(garage_vol_sum)} Garage Entities!")

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python3 merge_search_volume.py <path_to_scraped_volume_file.xlsx>")
    else:
        merge_volume(sys.argv[1])
