import openpyxl
from openpyxl.styles import Alignment, Border, Side, PatternFill, Font

file_path = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/vehicle-hub-roadmap.xlsx'
wb = openpyxl.load_workbook(file_path)

sheet_name = 'Content Structure'
if sheet_name in wb.sheetnames:
    del wb[sheet_name]
ws = wb.create_sheet(sheet_name, index=3) # Insert after Roadmap (index might vary, but 3 is safe)

# Define Styles
font_header = Font(bold=True, color="FFFFFF")
fill_header = PatternFill("solid", fgColor="4C1130") # MoMo pink/dark red
align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

headers = ['Nhóm Tính Năng', 'Tên Trang & Đường dẫn (URL)', '1. Khu vực Hero (Top Section)', '2. Khối Tiện Ích Lõi (Core Utilities)', '3. Khối Bán Chéo & CTA (Cross-sell)', '4. Nội dung SEO (Bottom Section)']
ws.append(headers)

for col_num in range(1, 7):
    cell = ws.cell(row=1, column=col_num)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = border

data = [
    # Master Hub
    ['Gateway', 'Trang Chủ Master Hub\n(/tien-ich-giao-thong)', 'Banner chính & Hero Widget 3-in-1 (Tab tra cứu Phạt Nguội / Giá Xăng / ePass).', '6 Khối Icon Menu chức năng chính. Khối thẻ xe cá nhân (Vehicle Profile).', 'Sticky Button Mobile (Onelink mở App), Dynamic QR Code cho Desktop.', 'Đoạn giới thiệu tổng quan hệ sinh thái, câu hỏi thường gặp (FAQ Schema).'],
    
    # Phạt nguội
    ['Transaction', 'Cổng Tra Cứu Phạt Nguội\n(/phat-nguoi, /phat-nguoi/[tinh])', 'Thanh tìm kiếm biển số cỡ lớn (Focus state). Nút Quét Camera AI (nếu có).', 'Khối trả kết quả Inline: Có lỗi (hiển thị chi tiết lỗi, camera) / Không lỗi (chúc mừng).', 'Có lỗi: Nút CTA Nộp phạt In-App.\nKhông lỗi: Banner Cross-sell giảm giá Mua Bảo hiểm Ô tô/Xe máy.', 'Trích dẫn lỗi theo Nghị định 100/123. Bảng mã lỗi phổ biến theo từng tỉnh/thành.'],
    
    # BOT
    ['Transaction', 'Thu Phí Không Dừng\n(/phi-khong-dung, /tram-thu-phi)', 'Công cụ tra cứu nhanh biển số xem tình trạng tài khoản ePass/VETC.', 'Bảng tra cứu cước phí BOT theo tuyến đường và loại xe (Biểu phí động).', 'Nút W2A CTA: Liên kết tài khoản & Nạp tiền thu phí không dừng.', 'Hướng dẫn dán thẻ, thủ tục đăng ký ePass/VETC.'],
    
    # Xăng/Trạm Sạc
    ['Location & Maps', 'Giá Xăng & Cây Xăng\n(/gia-xang, /cay-xang)', 'Bảng giá xăng dầu (RON 95, E5, Diesel) cập nhật realtime theo kỳ điều hành.', 'Bản đồ GPS định vị cây xăng gần nhất. Widget tính tiền đổ đầy bình xe.', 'Nút W2A CTA: "Thanh toán MoMo tại trạm" kèm hiển thị Voucher xăng.', 'Phân tích biến động giá xăng, lịch sử giá 6 tháng.'],
    ['Location & Maps', 'Trạm Sạc EV & Bãi Đỗ\n(/tram-sac, /bai-do-xe)', 'Thanh tìm kiếm địa điểm / Tỉnh thành hiện tại.', 'Bản đồ Google Maps nhúng. Bộ lọc (Loại trạm sạc: AC/DC, Hãng: VinFast/V-Green).', 'Nút CTA: Gọi điện hotline bãi đỗ / Chỉ đường Google Maps.', 'Danh sách trạm sạc/bãi đỗ theo khu vực (Text list để SEO Local).'],
    
    # Đăng kiểm
    ['Utilities', 'Tra Cứu Hạn Đăng Kiểm\n(/dang-kiem)', 'Input tra cứu Biển số & Số tem giấy chứng nhận.', 'Kết quả chu kỳ đăng kiểm, ngày hết hạn. Khối đếm ngược (Countdown).', 'Nút W2A: Đặt lịch hẹn Trung Tâm Đăng Kiểm In-App.', 'Bảng chu kỳ kiểm định các loại xe, thủ tục hồ sơ đăng kiểm.'],
    
    # Cứu hộ
    ['Emergency', 'Cứu Hộ Ô Tô\n(/cuu-ho)', 'Banner SOS & Thanh tìm kiếm Tỉnh/Thành/Quận/Huyện.', 'Danh bạ các đối tác cứu hộ gần nhất (Vá lốp, cẩu xe, kích bình).', 'Nút Bấm Gọi Nhanh (Click-to-call Hotline) 24/7.', 'Bảng giá dịch vụ cứu hộ tham khảo, cẩm nang xử lý sự cố giữa đường.'],
    
    # Bảo hiểm
    ['Monetize', 'Bảo Hiểm Phương Tiện\n(/bao-hiem-o-to)', 'Công cụ so sánh giá nhanh 9 hãng bảo hiểm theo Loại Xe & Năm SX.', 'Form Prefill (Auto-fill) thông tin chủ xe & Bảng so sánh quyền lợi.', 'Nút CTA Mua Ngay, Thanh toán Onelink 1 chạm.', 'Giải thích thuật ngữ bảo hiểm TNDS, Thân vỏ, mức miễn thường.'],
    
    # Tools
    ['Calculator', 'Công Cụ Tài Chính Xe\n(/lan-banh, /chi-phi-nuoi-xe)', 'Dropdown chọn Hãng Xe, Dòng Xe, Tỉnh thành đăng ký.', 'Bảng tính tự động chi tiết: Giá niêm yết, Trước bạ (10/12%), Phí ra biển, Phí đường bộ.', 'Banner: Vay mua xe (In-App) hoặc Gói Bảo hiểm cho xe mới.', 'Quy định thuế trước bạ ô tô cập nhật mới nhất.'],
    
    # Content
    ['Content', 'Blog & Luật Giao Thông\n(/tra-cuu-muc-phat...)', 'Tiêu đề bài viết & Mục lục Table of Contents (TOC).', 'Nội dung chi tiết bài viết (Rich Text, Hình ảnh, Bảng biểu pháp lý).', 'Inline Banner (nằm giữa bài): Gợi ý tra cứu phạt nguội, mua bảo hiểm.', 'Các bài viết liên quan (Related Posts), Tag từ khóa (Internal Linking).']
]

for idx, row_data in enumerate(data, start=2):
    for col, val in enumerate(row_data, start=1):
        cell = ws.cell(row=idx, column=col, value=val)
        cell.alignment = align_left if col > 1 else align_center
        cell.border = border

# Adjust column widths
col_widths = {'A': 18, 'B': 30, 'C': 35, 'D': 40, 'E': 35, 'F': 40}
for col, width in col_widths.items():
    ws.column_dimensions[col].width = width

wb.save(file_path)
print("Content Structure sheet created successfully.")
