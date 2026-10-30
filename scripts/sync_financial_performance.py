#!/usr/bin/env python3
"""
Sync Financial Performance from Master Tracking Sheet to Financial Hub Roadmap.
Source: Google Sheet 1JaaYv2MnDy-dfBdC3P-GMB6SFgjez2VowdjX9PjmPYU (Sheet: Financial)
Destination:
  1. Google Sheet 1_AG4otSTXH7ontpPL0zjylYgVfxOozvoSc5PziUPM6A (Sheet: Performances)
  2. Local Excel 05_HUBS/financial/financial-hub-roadmap.xlsx (Sheet: Performance)
"""

import json
import urllib.parse
import urllib.request
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent
EXCEL_PATH = BASE_DIR / "05_HUBS" / "financial" / "financial-hub-roadmap.xlsx"
CREDS_FILE = BASE_DIR / "07_REPORTS" / "data" / "credentials" / "ga4-credentials.json"
TOKEN_FILE = BASE_DIR / "07_REPORTS" / "data" / "credentials" / "gdrive-token.json"

MASTER_SPREADSHEET_ID = "1JaaYv2MnDy-dfBdC3P-GMB6SFgjez2VowdjX9PjmPYU"
TARGET_SPREADSHEET_ID = "1_AG4otSTXH7ontpPL0zjylYgVfxOozvoSc5PziUPM6A"


def get_access_token():
    with open(CREDS_FILE, "r", encoding="utf-8") as f:
        cdata = json.load(f)
    client_info = cdata.get("installed") or cdata.get("web")

    with open(TOKEN_FILE, "r", encoding="utf-8") as f:
        tdata = json.load(f)

    data = urllib.parse.urlencode({
        "client_id": client_info["client_id"],
        "client_secret": client_info["client_secret"],
        "refresh_token": tdata["refresh_token"],
        "grant_type": "refresh_token"
    }).encode("utf-8")

    req = urllib.request.Request("https://oauth2.googleapis.com/token", data=data)
    with urllib.request.urlopen(req) as resp:
        token_data = json.loads(resp.read().decode("utf-8"))
    return token_data["access_token"]


def fetch_master_financial(token):
    print("1. Fetching Master Financial sheet data and format...")
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{MASTER_SPREADSHEET_ID}?ranges=Financial!A1:S150&includeGridData=true"
    req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
    with urllib.request.urlopen(req) as resp:
        res = json.loads(resp.read().decode("utf-8"))

    sheet = res["sheets"][0]
    col_metadata = sheet.get("data", [{}])[0].get("columnMetadata", [])
    row_data = sheet.get("data", [{}])[0].get("rowData", [])

    values_matrix = []
    clean_format_rows = []

    for r in row_data:
        row_vals = []
        fmt_cells = []
        cells = r.get("values", [])
        for c in cells:
            val = c.get("formattedValue", "")
            row_vals.append(val)
            user_fmt = c.get("userEnteredFormat")
            eff_fmt = c.get("effectiveFormat")
            # Prefer userEnteredFormat, fallback to effectiveFormat
            fmt_to_use = user_fmt if user_fmt else eff_fmt
            if fmt_to_use:
                # Remove read-only or unsupported keys
                clean_fmt = {}
                for k in ["backgroundColor", "textFormat", "horizontalAlignment", "verticalAlignment", "wrapStrategy", "padding"]:
                    if k in fmt_to_use:
                        clean_fmt[k] = fmt_to_use[k]
                fmt_cells.append({"userEnteredFormat": clean_fmt})
            else:
                fmt_cells.append({})
        values_matrix.append(row_vals)
        clean_format_rows.append({"values": fmt_cells})

    # Trim empty trailing rows if any
    while values_matrix and not any(cell.strip() for cell in values_matrix[-1]):
        values_matrix.pop()
        clean_format_rows.pop()

    print(f"Fetched {len(values_matrix)} rows across {len(values_matrix[0]) if values_matrix else 0} columns.")
    return values_matrix, clean_format_rows, col_metadata


def update_target_google_sheet(token, values_matrix, clean_format_rows, col_metadata):
    print("2. Updating Google Sheet 'Performances' in Financial Roadmap...")

    # Find sheetId of Performances
    url_meta = f"https://sheets.googleapis.com/v4/spreadsheets/{TARGET_SPREADSHEET_ID}?fields=sheets(properties)"
    req_meta = urllib.request.Request(url_meta, headers={"Authorization": f"Bearer {token}"})
    with urllib.request.urlopen(req_meta) as resp:
        meta = json.loads(resp.read().decode("utf-8"))

    target_sheet_id = None
    for s in meta["sheets"]:
        p = s["properties"]
        if p["title"] == "Performances":
            target_sheet_id = p["sheetId"]
            break

    if target_sheet_id is None:
        raise ValueError("Sheet 'Performances' not found in target spreadsheet.")

    print(f"Target Sheet ID for Performances: {target_sheet_id}")

    # A. Clear existing values
    clear_url = f"https://sheets.googleapis.com/v4/spreadsheets/{TARGET_SPREADSHEET_ID}/values/Performances!A1:Z500:clear"
    req = urllib.request.Request(
        clear_url,
        data=b"{}",
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
    )
    with urllib.request.urlopen(req) as resp:
        print("Cleared existing values on Performances.")

    # B. Write values
    num_rows = len(values_matrix)
    write_url = f"https://sheets.googleapis.com/v4/spreadsheets/{TARGET_SPREADSHEET_ID}/values/Performances!A1:S{num_rows}?valueInputOption=USER_ENTERED"
    body = json.dumps({"values": values_matrix}).encode("utf-8")
    req = urllib.request.Request(
        write_url,
        data=body,
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="PUT"
    )
    with urllib.request.urlopen(req) as resp:
        res = json.loads(resp.read().decode("utf-8"))
        print(f"Performances values written: {res.get('updatedCells')} cells across {num_rows} rows.")

    # C. BatchUpdate styling
    requests = []

    # Update sheet grid properties
    requests.append({
        "updateSheetProperties": {
            "properties": {
                "sheetId": target_sheet_id,
                "gridProperties": {
                    "rowCount": max(200, num_rows + 20),
                    "columnCount": 26,
                    "frozenRowCount": 1,
                    "frozenColumnCount": 2,
                    "hideGridlines": True
                }
            },
            "fields": "gridProperties(rowCount,columnCount,frozenRowCount,frozenColumnCount,hideGridlines)"
        }
    })

    # Update column widths
    for c_idx in range(19):
        width = 100
        if c_idx < len(col_metadata):
            width = col_metadata[c_idx].get("pixelSize", 100)
        requests.append({
            "updateDimensionProperties": {
                "range": {
                    "sheetId": target_sheet_id,
                    "dimension": "COLUMNS",
                    "startIndex": c_idx,
                    "endIndex": c_idx + 1
                },
                "properties": {"pixelSize": width},
                "fields": "pixelSize"
            }
        })

    # Update cell formats
    requests.append({
        "updateCells": {
            "range": {
                "sheetId": target_sheet_id,
                "startRowIndex": 0,
                "endRowIndex": num_rows,
                "startColumnIndex": 0,
                "endColumnIndex": 19
            },
            "rows": clean_format_rows,
            "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment,verticalAlignment,wrapStrategy,padding)"
        }
    })

    url_batch = f"https://sheets.googleapis.com/v4/spreadsheets/{TARGET_SPREADSHEET_ID}:batchUpdate"
    body_batch = json.dumps({"requests": requests}).encode("utf-8")
    req = urllib.request.Request(
        url_batch,
        data=body_batch,
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="POST"
    )
    with urllib.request.urlopen(req) as resp:
        batch_res = json.loads(resp.read().decode("utf-8"))
        print(f"Performances styling batchUpdate successful! Replies: {len(batch_res.get('replies', []))}")


def update_local_excel_performance(values_matrix, clean_format_rows):
    print("3. Updating local Excel sheet 'Performance'...")
    wb = openpyxl.load_workbook(EXCEL_PATH)

    if "Performance" in wb.sheetnames:
        del wb["Performance"]
    ws = wb.create_sheet("Performance", index=1)

    # Styles
    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="073763", end_color="073763", fill_type="solid")

    font_default = Font(name="Arial", size=9.5, color="1E293B")
    font_bold = Font(name="Arial", size=9.5, bold=True, color="1E293B")

    border_thin = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    for r_idx, row_vals in enumerate(values_matrix, 1):
        for c_idx, val_str in enumerate(row_vals, 1):
            cell = ws.cell(r_idx, c_idx)
            cell.border = border_thin

            if r_idx == 1:
                cell.value = val_str
                cell.font = font_header
                cell.fill = fill_header
                cell.alignment = Alignment(horizontal="center" if c_idx > 2 else "left", vertical="center")
            else:
                # Format cell based on data
                val_clean = val_str.strip()
                if not val_clean:
                    cell.value = ""
                elif val_clean.endswith("%"):
                    try:
                        p_num = float(val_clean.replace("%", "").replace(",", "")) / 100.0
                        cell.value = p_num
                        cell.number_format = "0.0%"
                    except ValueError:
                        cell.value = val_clean
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    # Try parsing as integer / float
                    num_clean = val_clean.replace(",", "")
                    try:
                        num_val = int(num_clean)
                        cell.value = num_val
                        cell.number_format = "#,##0"
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    except ValueError:
                        try:
                            num_val = float(num_clean)
                            cell.value = num_val
                            cell.number_format = "#,##0.0"
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                        except ValueError:
                            cell.value = val_clean
                            align_h = "center" if c_idx == 2 and ("Update to" in val_clean or val_clean.isdigit()) else "left"
                            cell.alignment = Alignment(horizontal=align_h, vertical="center")

                # Font bold detection: Parent rows or Totals
                col_a = str(row_vals[0] if len(row_vals) > 0 else "").strip()
                col_b = str(row_vals[1] if len(row_vals) > 1 else "").strip()
                is_parent = bool(col_a and not col_b) or ("Actual Traffic" in col_a) or ("Target" in col_a) or ("Net change" in col_a)
                if is_parent:
                    cell.font = font_bold
                    if r_idx in [2, 3, 4, 5, 6]:
                        cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
                    elif r_idx % 7 == 1:
                        cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
                else:
                    cell.font = font_default

        ws.row_dimensions[r_idx].height = 20

    # Row 1 height
    ws.row_dimensions[1].height = 26

    # Column widths
    col_widths = {
        1: 30,  # Project
        2: 15,  # Channel / Update note
        3: 13, 4: 13, 5: 13, 6: 13, 7: 13, 8: 13, 9: 13,  # MTD (May - Oct + VS. SPLM)
        10: 12, # Full Month separator
        11: 13, 12: 13, 13: 13, 14: 13, 15: 13, 16: 13, 17: 13, 18: 13, 19: 13 # Full month
    }
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze panes: Row 1 frozen, Cols A & B frozen
    ws.freeze_panes = "C2"

    wb.save(EXCEL_PATH)
    print("Local Excel sheet 'Performance' updated and saved successfully!")


def main():
    token = get_access_token()
    values_matrix, clean_format_rows, col_metadata = fetch_master_financial(token)
    update_target_google_sheet(token, values_matrix, clean_format_rows, col_metadata)
    update_local_excel_performance(values_matrix, clean_format_rows)
    print("\nFinancial Performance synchronization completed successfully across all destinations!")


if __name__ == "__main__":
    main()
