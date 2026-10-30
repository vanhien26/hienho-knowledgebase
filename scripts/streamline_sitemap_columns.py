#!/usr/bin/env python3
"""
Streamline Sitemap Sheet:
1. Remove column 'Tên Trang Cần Tạo (H1 / Title)'.
2. Shorten column name 'Phân Cấp / Loại Trang' -> 'Loại Trang' and simplify category values.
3. Update both local Excel (05_HUBS/financial/financial-hub-roadmap.xlsx) and Google Sheet SSOT (1_AG4otSTXH7ontpPL0zjylYgVfxOozvoSc5PziUPM6A).
"""

import json
import urllib.parse
import urllib.request
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from rebuild_sitemap_by_roadmap import SITEMAP_ITEMS

BASE_DIR = Path(__file__).resolve().parent.parent
EXCEL_PATH = BASE_DIR / "05_HUBS" / "financial" / "financial-hub-roadmap.xlsx"
CREDS_FILE = BASE_DIR / "07_REPORTS" / "data" / "credentials" / "ga4-credentials.json"
TOKEN_FILE = BASE_DIR / "07_REPORTS" / "data" / "credentials" / "gdrive-token.json"
SPREADSHEET_ID = "1_AG4otSTXH7ontpPL0zjylYgVfxOozvoSc5PziUPM6A"

CATEGORY_MAPPING = {
    "Trang Độc Lập Root (Không Đổi URL)": "Độc Lập Root",
    "Cổng Trung Tâm (Master Hub)": "Cổng Master Hub",
    "Cổng Danh Mục Tiện Ích (Tools Hub)": "Cổng Tiện Ích",
    "Công Cụ Tiện Ích Under /tai-chinh": "Công Cụ Tiện Ích",
    "Cổng Danh Mục Under /tai-chinh": "Cổng Danh Mục",
    "Trang Con Chi Tiết (Programmatic Spoke)": "Template Chi Tiết",
    "Trang Con Cặp Tiền (Programmatic Subpage)": "Sub-page Cặp Tiền",
    "Trang Dịch Vụ Under /tai-chinh": "Trang Dịch Vụ",
    "Trang Con Quỹ Mở (Programmatic CCQ)": "Sub-page Quỹ Mở",
    "Trang Danh Mục Blog (Blog Listing)": "Danh Mục Blog",
    "Kho Bài Viết Chi Tiết (Blog Details)": "Kho Bài Viết",
    "Trang Con Cổ Phiếu (Programmatic VN30)": "Sub-page Cổ Phiếu VN30",
    "Cổng Danh Mục Master": "Cổng Danh Bạ",
    "Trang Con Ngân Hàng (Programmatic Bank)": "Sub-page Ngân Hàng"
}

STREAMLINED_HEADER = [
    "STT",
    "Giai Đoạn (Phase)",
    "Thời Gian",
    "Loại Trang",
    "Định Tuyến URL Chuẩn Hóa (Canonical URL)",
    "Trạng Thái",
    "Hiện Trạng & Ghi Chú Kỹ Thuật"
]


def build_streamlined_rows():
    rows = []
    for idx, item in enumerate(SITEMAP_ITEMS, 1):
        phase = item[0]
        time = item[1]
        cat = CATEGORY_MAPPING.get(item[2], item[2])
        # item[3] is Title -> OMITTED
        url = item[4]
        status = item[5]
        note = item[6]

        # Ensure blog pages reflect Done
        if url in ["/tai-chinh/blog", "/tai-chinh/blog/[slug]"]:
            status = "Done"
            if url == "/tai-chinh/blog":
                note = "Đã hoàn thành Go-live. Cổng danh mục bài viết cẩm nang kiến thức tài chính cá nhân MoMo."
            else:
                note = "Đã hoàn thành Go-live. Hệ thống 100 bài viết chi tiết chuẩn SEO & GEO, duy trì xuất bản đều đặn."

        rows.append([idx, phase, time, cat, url, status, note])
    return rows


def get_access_token():
    with open(CREDS_FILE, "r", encoding="utf-8") as f:
        cdata = json.load(f)
    client_info = cdata.get("installed") or cdata.get("web")

    with open(TOKEN_FILE, "r", encoding="utf-8") as f:
        tdata = json.load(f)

    data = urllib.parse.urlencode({
        "client_id": client_info["client_id"],
        "client_secret": client_info["client_secret"],
        "refresh_token": tdata["refresh_token"],
        "grant_type": "refresh_token"
    }).encode("utf-8")

    req = urllib.request.Request("https://oauth2.googleapis.com/token", data=data, method="POST")
    with urllib.request.urlopen(req) as resp:
        res = json.loads(resp.read().decode("utf-8"))
        return res["access_token"]


def batch_update_spreadsheet(token, requests):
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}:batchUpdate"
    body = json.dumps({"requests": requests}).encode("utf-8")
    req = urllib.request.Request(
        url,
        data=body,
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="POST"
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode("utf-8"))


def clear_range(token, range_name):
    encoded_range = urllib.parse.quote(range_name)
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}/values/{encoded_range}:clear"
    req = urllib.request.Request(
        url,
        data=b"{}",
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="POST"
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode("utf-8"))


def update_values(token, range_name, values):
    encoded_range = urllib.parse.quote(range_name)
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}/values/{encoded_range}?valueInputOption=USER_ENTERED"
    body = json.dumps({"values": values}).encode("utf-8")
    req = urllib.request.Request(
        url,
        data=body,
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="PUT"
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode("utf-8"))


def update_local_excel():
    print(f"1. Updating local Excel: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH)

    if "Sitemap" in wb.sheetnames:
        del wb["Sitemap"]
    ws_s = wb.create_sheet("Sitemap", index=3)

    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

    font_default = Font(name="Arial", size=9.5, color="1E293B")
    font_bold = Font(name="Arial", size=9.5, bold=True, color="1E293B")

    border_thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    fill_done = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    font_done = Font(name="Arial", size=9.5, bold=True, color="15803D")

    fill_progress = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    font_progress = Font(name="Arial", size=9.5, bold=True, color="B45309")

    fill_pending = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")
    font_pending = Font(name="Arial", size=9.5, bold=True, color="BE123C")

    fill_planned = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
    font_planned = Font(name="Arial", size=9.5, bold=True, color="4338CA")

    fill_not_started = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    font_not_started = Font(name="Arial", size=9.5, bold=True, color="475569")

    # Header Row
    for c_idx, h in enumerate(STREAMLINED_HEADER, 1):
        cell = ws_s.cell(1, c_idx, h)
        cell.font = font_header
        cell.fill = fill_header
        cell.border = border_thin
        align_h = "center" if c_idx in [1, 2, 3, 6] else "left"
        cell.alignment = Alignment(horizontal=align_h, vertical="center")
    ws_s.row_dimensions[1].height = 26

    # Data Rows
    rows = build_streamlined_rows()
    for r_idx, row_vals in enumerate(rows, 2):
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_s.cell(r_idx, c_idx, val)
            cell.border = border_thin

            align_h = "center" if c_idx in [1, 2, 3, 6] else "left"
            wrap = True if c_idx == 7 else False
            cell.alignment = Alignment(horizontal=align_h, vertical="center", wrap_text=wrap)

            if c_idx == 5:  # Canonical URL
                cell.font = font_bold
            elif c_idx == 6:  # Status
                if val == "Done":
                    cell.fill = fill_done
                    cell.font = font_done
                elif val == "In Progress":
                    cell.fill = fill_progress
                    cell.font = font_progress
                elif val == "Pending":
                    cell.fill = fill_pending
                    cell.font = font_pending
                elif val == "Chưa triển khai":
                    cell.fill = fill_not_started
                    cell.font = font_not_started
                else:  # Planned
                    cell.fill = fill_planned
                    cell.font = font_planned
            else:
                cell.font = font_default

        ws_s.row_dimensions[r_idx].height = 24

    # Column widths (7 columns)
    col_widths = {1: 8, 2: 15, 3: 16, 4: 25, 5: 38, 6: 18, 7: 58}
    for col_idx, width in col_widths.items():
        ws_s.column_dimensions[get_column_letter(col_idx)].width = width

    ws_s.freeze_panes = "A2"
    wb.save(EXCEL_PATH)
    print("Local Excel Sitemap sheet updated and saved successfully!")


def update_google_sheets():
    print("2. Updating Google Sheets Tab: Sitemap...")
    token = get_access_token()

    url_meta = f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}?fields=sheets(properties)"
    req_meta = urllib.request.Request(url_meta, headers={"Authorization": f"Bearer {token}"})
    with urllib.request.urlopen(req_meta) as resp:
        meta = json.loads(resp.read().decode("utf-8"))

    sheet_map = {s["properties"]["title"]: s["properties"]["sheetId"] for s in meta["sheets"]}
    sitemap_id = sheet_map.get("Sitemap")

    # Clear wide range to ensure old column H is wiped
    clear_range(token, "Sitemap!A1:Z200")

    rows = build_streamlined_rows()
    full_table = [STREAMLINED_HEADER] + rows
    num_rows = len(full_table)
    num_cols = len(STREAMLINED_HEADER)

    update_values(token, f"Sitemap!A1:G{num_rows}", full_table)

    reqs_s = []
    # Grid properties
    reqs_s.append({
        "updateSheetProperties": {
            "properties": {
                "sheetId": sitemap_id,
                "gridProperties": {
                    "frozenRowCount": 1,
                    "rowCount": max(150, num_rows + 10),
                    "columnCount": 10
                }
            },
            "fields": "gridProperties(frozenRowCount,rowCount,columnCount)"
        }
    })

    # Column widths: STT (60), Phase (110), Time (110), Type (180), URL (300), Status (130), Note (520)
    col_widths = [60, 110, 110, 180, 300, 130, 520]
    for c_idx, w in enumerate(col_widths):
        reqs_s.append({
            "updateDimensionProperties": {
                "range": {
                    "sheetId": sitemap_id,
                    "dimension": "COLUMNS",
                    "startIndex": c_idx,
                    "endIndex": c_idx + 1
                },
                "properties": {"pixelSize": w},
                "fields": "pixelSize"
            }
        })

    # Header Row 1 formatting
    reqs_s.append({
        "repeatCell": {
            "range": {
                "sheetId": sitemap_id,
                "startRowIndex": 0,
                "endRowIndex": 1,
                "startColumnIndex": 0,
                "endColumnIndex": num_cols
            },
            "cell": {
                "userEnteredFormat": {
                    "backgroundColor": {"red": 0.12, "green": 0.31, "blue": 0.47},
                    "horizontalAlignment": "LEFT",
                    "verticalAlignment": "MIDDLE",
                    "textFormat": {
                        "fontFamily": "Arial",
                        "fontSize": 10,
                        "bold": True,
                        "foregroundColor": {"red": 1.0, "green": 1.0, "blue": 1.0}
                    }
                }
            },
            "fields": "userEnteredFormat(backgroundColor,horizontalAlignment,verticalAlignment,textFormat)"
        }
    })
    for c_idx in [0, 1, 2, 5]:
        reqs_s.append({
            "repeatCell": {
                "range": {
                    "sheetId": sitemap_id,
                    "startRowIndex": 0,
                    "endRowIndex": 1,
                    "startColumnIndex": c_idx,
                    "endColumnIndex": c_idx + 1
                },
                "cell": {"userEnteredFormat": {"horizontalAlignment": "CENTER"}},
                "fields": "userEnteredFormat.horizontalAlignment"
            }
        })

    # Data rows basic formatting
    reqs_s.append({
        "repeatCell": {
            "range": {
                "sheetId": sitemap_id,
                "startRowIndex": 1,
                "endRowIndex": num_rows,
                "startColumnIndex": 0,
                "endColumnIndex": num_cols
            },
            "cell": {
                "userEnteredFormat": {
                    "verticalAlignment": "MIDDLE",
                    "textFormat": {
                        "fontFamily": "Arial",
                        "fontSize": 10,
                        "foregroundColor": {"red": 0.12, "green": 0.16, "blue": 0.23}
                    }
                }
            },
            "fields": "userEnteredFormat(verticalAlignment,textFormat)"
        }
    })

    # Zebra striping (Col A to E, and Col G)
    for r_idx in range(1, num_rows):
        bg = {"red": 0.97, "green": 0.98, "blue": 0.99} if r_idx % 2 == 1 else {"red": 1.0, "green": 1.0, "blue": 1.0}
        reqs_s.append({
            "repeatCell": {
                "range": {
                    "sheetId": sitemap_id,
                    "startRowIndex": r_idx,
                    "endRowIndex": r_idx + 1,
                    "startColumnIndex": 0,
                    "endColumnIndex": 5
                },
                "cell": {"userEnteredFormat": {"backgroundColor": bg}},
                "fields": "userEnteredFormat.backgroundColor"
            }
        })
        reqs_s.append({
            "repeatCell": {
                "range": {
                    "sheetId": sitemap_id,
                    "startRowIndex": r_idx,
                    "endRowIndex": r_idx + 1,
                    "startColumnIndex": 6,
                    "endColumnIndex": 7
                },
                "cell": {"userEnteredFormat": {"backgroundColor": bg}},
                "fields": "userEnteredFormat.backgroundColor"
            }
        })

    # Alignment
    for c_idx in [0, 1, 2]:
        reqs_s.append({
            "repeatCell": {
                "range": {
                    "sheetId": sitemap_id,
                    "startRowIndex": 1,
                    "endRowIndex": num_rows,
                    "startColumnIndex": c_idx,
                    "endColumnIndex": c_idx + 1
                },
                "cell": {"userEnteredFormat": {"horizontalAlignment": "CENTER"}},
                "fields": "userEnteredFormat.horizontalAlignment"
            }
        })
    # URL column bold
    reqs_s.append({
        "repeatCell": {
            "range": {
                "sheetId": sitemap_id,
                "startRowIndex": 1,
                "endRowIndex": num_rows,
                "startColumnIndex": 4,
                "endColumnIndex": 5
            },
            "cell": {"userEnteredFormat": {"textFormat": {"bold": True}}},
            "fields": "userEnteredFormat.textFormat.bold"
        }
    })
    # Note column wrap
    reqs_s.append({
        "repeatCell": {
            "range": {
                "sheetId": sitemap_id,
                "startRowIndex": 1,
                "endRowIndex": num_rows,
                "startColumnIndex": 6,
                "endColumnIndex": 7
            },
            "cell": {"userEnteredFormat": {"wrapStrategy": "WRAP"}},
            "fields": "userEnteredFormat.wrapStrategy"
        }
    })

    # Status Badges
    for r_idx in range(1, num_rows):
        st = full_table[r_idx][5]
        if st == "Done":
            bg = {"red": 0.86, "green": 0.98, "blue": 0.90}
            fg = {"red": 0.08, "green": 0.50, "blue": 0.24}
        elif st == "In Progress":
            bg = {"red": 0.99, "green": 0.95, "blue": 0.78}
            fg = {"red": 0.71, "green": 0.33, "blue": 0.04}
        elif st == "Pending":
            bg = {"red": 1.0, "green": 0.89, "blue": 0.90}
            fg = {"red": 0.75, "green": 0.07, "blue": 0.24}
        elif st == "Chưa triển khai":
            bg = {"red": 0.945, "green": 0.96, "blue": 0.976}
            fg = {"red": 0.28, "green": 0.33, "blue": 0.41}
        else:  # Planned
            bg = {"red": 0.88, "green": 0.91, "blue": 1.0}
            fg = {"red": 0.26, "green": 0.22, "blue": 0.79}

        reqs_s.append({
            "repeatCell": {
                "range": {
                    "sheetId": sitemap_id,
                    "startRowIndex": r_idx,
                    "endRowIndex": r_idx + 1,
                    "startColumnIndex": 5,
                    "endColumnIndex": 6
                },
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": bg,
                        "horizontalAlignment": "CENTER",
                        "verticalAlignment": "MIDDLE",
                        "textFormat": {
                            "fontFamily": "Arial",
                            "fontSize": 10,
                            "bold": True,
                            "foregroundColor": fg
                        }
                    }
                },
                "fields": "userEnteredFormat(backgroundColor,horizontalAlignment,verticalAlignment,textFormat)"
            }
        })

    # Borders
    border_style = {"style": "SOLID", "width": 1, "color": {"red": 0.85, "green": 0.88, "blue": 0.92}}
    reqs_s.append({
        "updateBorders": {
            "range": {
                "sheetId": sitemap_id,
                "startRowIndex": 0,
                "endRowIndex": num_rows,
                "startColumnIndex": 0,
                "endColumnIndex": num_cols
            },
            "top": border_style, "bottom": border_style, "left": border_style, "right": border_style,
            "innerHorizontal": border_style, "innerVertical": border_style
        }
    })

    batch_update_spreadsheet(token, reqs_s)
    print("Google Sheet Tab: Sitemap streamlined and formatted successfully!")


def main():
    update_local_excel()
    update_google_sheets()
    print("\nAll Sitemap streamlining tasks completed successfully!")


if __name__ == "__main__":
    main()
