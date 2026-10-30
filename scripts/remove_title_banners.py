#!/usr/bin/env python3
"""
Remove table title banner rows from Financial Hub Roadmap workbook.
Applies to both local Excel and Google Sheets:
- Sheet: Markets / Financial Markets (Removes "TỔNG HỢP TOÀN BỘ CÁC THỊ TRƯỜNG DỊCH VỤ TÀI CHÍNH...")
- Sheet: Roadmap / Product Roadmap (Removes "LỘ TRÌNH TRIỂN KHAI TOÀN DIỆN CÁC TÍNH NĂNG...")
- Sheet: Sitemap (Removes "DANH BẠ KIẾN TRÚC THÔNG TIN..." and subtitle banner)
Result: All sheets start directly at Row 1 with Column Headers and frozenRowCount = 1.
"""

import csv
import json
import urllib.parse
import urllib.request
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent
EXCEL_PATH = BASE_DIR / "05_HUBS" / "financial" / "financial-hub-roadmap.xlsx"
CREDS_FILE = BASE_DIR / "07_REPORTS" / "data" / "credentials" / "ga4-credentials.json"
TOKEN_FILE = BASE_DIR / "07_REPORTS" / "data" / "credentials" / "gdrive-token.json"
SPREADSHEET_ID = "1_AG4otSTXH7ontpPL0zjylYgVfxOozvoSc5PziUPM6A"


def get_access_token():
    with open(CREDS_FILE, "r", encoding="utf-8") as f:
        cdata = json.load(f)
    client_info = cdata.get("installed") or cdata.get("web")

    with open(TOKEN_FILE, "r", encoding="utf-8") as f:
        tdata = json.load(f)

    data = urllib.parse.urlencode({
        "client_id": client_info["client_id"],
        "client_secret": client_info["client_secret"],
        "refresh_token": tdata["refresh_token"],
        "grant_type": "refresh_token"
    }).encode("utf-8")

    req = urllib.request.Request("https://oauth2.googleapis.com/token", data=data)
    with urllib.request.urlopen(req) as resp:
        token_data = json.loads(resp.read().decode("utf-8"))
    return token_data["access_token"]


def batch_update_spreadsheet(token, requests):
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}:batchUpdate"
    body = json.dumps({"requests": requests}).encode("utf-8")
    req = urllib.request.Request(
        url,
        data=body,
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="POST"
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode("utf-8"))


def clear_range(token, range_name):
    encoded_range = urllib.parse.quote(range_name)
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}/values/{encoded_range}:clear"
    req = urllib.request.Request(
        url,
        data=b"{}",
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode("utf-8"))


def update_values(token, range_name, values):
    encoded_range = urllib.parse.quote(range_name)
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}/values/{encoded_range}?valueInputOption=USER_ENTERED"
    body = json.dumps({"values": values}).encode("utf-8")
    req = urllib.request.Request(
        url,
        data=body,
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="PUT"
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode("utf-8"))


def process_local_excel():
    print("1. Processing local Excel workbook...")
    wb = openpyxl.load_workbook(EXCEL_PATH)

    # Common styles
    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

    font_data = Font(name="Arial", size=10, color="1E293B")
    font_bold = Font(name="Arial", size=10, bold=True, color="1E293B")

    font_total = Font(name="Arial", size=10, bold=True, color="1F4E78")
    fill_total = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    fill_even = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    border_thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    border_total = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="medium", color="1F4E78"),
        bottom=Side(style="double", color="1F4E78")
    )

    # -------------------------------------------------------------
    # A. Sheet: Markets
    # -------------------------------------------------------------
    ws_m = wb["Markets"]
    r1_val = str(ws_m.cell(1, 1).value or "")
    if "TỔNG HỢP TOÀN BỘ CÁC THỊ TRƯỜNG" in r1_val:
        print("  - Removing Row 1 title banner from Sheet Markets...")
        for m in list(ws_m.merged_cells.ranges):
            ws_m.unmerge_cells(str(m))
        ws_m.delete_rows(1, 1)

    # Re-apply header styles to new Row 1
    for col_idx in range(1, ws_m.max_column + 1):
        cell = ws_m.cell(1, col_idx)
        cell.font = font_header
        cell.fill = fill_header
        align_h = "center" if col_idx in [1, 4, 5, 6] else "left"
        cell.alignment = Alignment(horizontal=align_h, vertical="center")
        cell.border = border_thin
    ws_m.row_dimensions[1].height = 26

    # Re-merge Total Row at bottom (Row 27)
    total_row = ws_m.max_row
    ws_m.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=3)
    ws_m.cell(total_row, 1).alignment = Alignment(horizontal="center", vertical="center")
    for col_idx in range(1, ws_m.max_column + 1):
        cell = ws_m.cell(total_row, col_idx)
        cell.font = font_total
        cell.fill = fill_total
        cell.border = border_total
    ws_m.row_dimensions[total_row].height = 26
    ws_m.freeze_panes = "A2"

    # -------------------------------------------------------------
    # B. Sheet: Roadmap
    # -------------------------------------------------------------
    ws_r = wb["Roadmap"]
    r1_val_r = str(ws_r.cell(1, 1).value or "")
    if "LỘ TRÌNH TRIỂN KHAI TOÀN DIỆN" in r1_val_r:
        print("  - Removing Row 1 title banner from Sheet Roadmap...")
        for m in list(ws_r.merged_cells.ranges):
            ws_r.unmerge_cells(str(m))
        ws_r.delete_rows(1, 1)

    for col_idx in range(1, ws_r.max_column + 1):
        cell = ws_r.cell(1, col_idx)
        cell.font = font_header
        cell.fill = fill_header
        align_h = "center" if col_idx in [1, 2, 6] else "left"
        cell.alignment = Alignment(horizontal=align_h, vertical="center")
        cell.border = border_thin
    ws_r.row_dimensions[1].height = 26
    ws_r.freeze_panes = "A2"

    # -------------------------------------------------------------
    # C. Sheet: Sitemap
    # -------------------------------------------------------------
    ws_s = wb["Sitemap"]
    r1_val_s = str(ws_s.cell(1, 1).value or "")
    if "DANH BẠ KIẾN TRÚC THÔNG TIN" in r1_val_s:
        print("  - Removing Row 1 & 2 title banners from Sheet Sitemap...")
        for m in list(ws_s.merged_cells.ranges):
            ws_s.unmerge_cells(str(m))
        ws_s.delete_rows(1, 2)

    for col_idx in range(1, ws_s.max_column + 1):
        cell = ws_s.cell(1, col_idx)
        cell.font = font_header
        cell.fill = fill_header
        align_h = "center" if col_idx == 1 else "left"
        cell.alignment = Alignment(horizontal=align_h, vertical="center")
        cell.border = border_thin
    ws_s.row_dimensions[1].height = 26
    ws_s.freeze_panes = "A2"

    wb.save(EXCEL_PATH)
    print("Local Excel updated and saved successfully!")


def process_google_sheets():
    print("2. Processing Google Sheets...")
    token = get_access_token()

    # Get sheet metadata
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}?fields=sheets(properties)"
    req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
    with urllib.request.urlopen(req) as resp:
        meta = json.loads(resp.read().decode("utf-8"))

    sheet_map = {}
    for s in meta["sheets"]:
        p = s["properties"]
        sheet_map[p["title"]] = p["sheetId"]

    markets_id = sheet_map.get("Financial Markets")
    roadmap_id = sheet_map.get("Product Roadmap")
    sitemap_id = sheet_map.get("Sitemap")

    print(f"Sheet IDs: Financial Markets={markets_id}, Product Roadmap={roadmap_id}, Sitemap={sitemap_id}")

    # Read updated rows from local Excel
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    # -------------------------------------------------------------
    # 1. Update Sheet: Financial Markets
    # -------------------------------------------------------------
    print("Updating Google Sheet: Financial Markets...")
    ws_m = wb["Markets"]
    markets_rows = []
    for r in range(1, ws_m.max_row + 1):
        row_vals = [ws_m.cell(r, c).value for c in range(1, ws_m.max_column + 1)]
        row_clean = []
        for v in row_vals:
            if v is None:
                row_clean.append("")
            elif isinstance(v, (int, float)):
                row_clean.append(f"{v:,}" if isinstance(v, int) else f"{v:.2f}%")
            else:
                row_clean.append(str(v))
        markets_rows.append(row_clean)

    # Clear and unmerge
    clear_range(token, "Financial Markets!A1:G60")
    batch_update_spreadsheet(token, [{
        "unmergeCells": {
            "range": {
                "sheetId": markets_id,
                "startRowIndex": 0,
                "endRowIndex": 60,
                "startColumnIndex": 0,
                "endColumnIndex": 7
            }
        }
    }])

    # Write new values (Row 1 = headers, Rows 2-26 = data, Row 27 = Total)
    res_m = update_values(token, f"Financial Markets!A1:G{len(markets_rows)}", markets_rows)
    print(f"Financial Markets updated: {res_m.get('updatedCells')} cells across {len(markets_rows)} rows.")

    # Styling for Financial Markets
    reqs_m = []
    # Freeze 1 row
    reqs_m.append({
        "updateSheetProperties": {
            "properties": {
                "sheetId": markets_id,
                "gridProperties": {
                    "frozenRowCount": 1
                }
            },
            "fields": "gridProperties.frozenRowCount"
        }
    })
    # Column widths
    col_widths_m = [65, 260, 210, 150, 180, 110, 520]
    for c_idx, width in enumerate(col_widths_m):
        reqs_m.append({
            "updateDimensionProperties": {
                "range": {
                    "sheetId": markets_id,
                    "dimension": "COLUMNS",
                    "startIndex": c_idx,
                    "endIndex": c_idx + 1
                },
                "properties": {
                    "pixelSize": width
                },
                "fields": "pixelSize"
            }
        })
    # Header Row 1 (index 0)
    reqs_m.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_id,
                "startRowIndex": 0,
                "endRowIndex": 1,
                "startColumnIndex": 0,
                "endColumnIndex": 7
            },
            "cell": {
                "userEnteredFormat": {
                    "backgroundColor": {"red": 0.12, "green": 0.31, "blue": 0.47},
                    "horizontalAlignment": "LEFT",
                    "verticalAlignment": "MIDDLE",
                    "textFormat": {
                        "fontFamily": "Arial",
                        "fontSize": 10,
                        "bold": True,
                        "foregroundColor": {"red": 1.0, "green": 1.0, "blue": 1.0}
                    }
                }
            },
            "fields": "userEnteredFormat(backgroundColor,horizontalAlignment,verticalAlignment,textFormat)"
        }
    })
    for c_idx in [0, 3, 4, 5]:
        reqs_m.append({
            "repeatCell": {
                "range": {
                    "sheetId": markets_id,
                    "startRowIndex": 0,
                    "endRowIndex": 1,
                    "startColumnIndex": c_idx,
                    "endColumnIndex": c_idx + 1
                },
                "cell": {
                    "userEnteredFormat": {
                        "horizontalAlignment": "CENTER"
                    }
                },
                "fields": "userEnteredFormat.horizontalAlignment"
            }
        })
    # Data Rows (index 1 to len-1)
    tot_idx_m = len(markets_rows) - 1
    reqs_m.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_id,
                "startRowIndex": 1,
                "endRowIndex": tot_idx_m,
                "startColumnIndex": 0,
                "endColumnIndex": 7
            },
            "cell": {
                "userEnteredFormat": {
                    "verticalAlignment": "MIDDLE",
                    "textFormat": {
                        "fontFamily": "Arial",
                        "fontSize": 10,
                        "foregroundColor": {"red": 0.12, "green": 0.16, "blue": 0.23}
                    }
                }
            },
            "fields": "userEnteredFormat(verticalAlignment,textFormat)"
        }
    })
    # Zebra striping
    for r_idx in range(1, tot_idx_m):
        bg = {"red": 0.97, "green": 0.98, "blue": 0.99} if r_idx % 2 == 1 else {"red": 1.0, "green": 1.0, "blue": 1.0}
        reqs_m.append({
            "repeatCell": {
                "range": {
                    "sheetId": markets_id,
                    "startRowIndex": r_idx,
                    "endRowIndex": r_idx + 1,
                    "startColumnIndex": 0,
                    "endColumnIndex": 7
                },
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": bg
                    }
                },
                "fields": "userEnteredFormat.backgroundColor"
            }
        })
    # Alignments
    reqs_m.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_id,
                "startRowIndex": 1,
                "endRowIndex": tot_idx_m,
                "startColumnIndex": 0,
                "endColumnIndex": 1
            },
            "cell": {"userEnteredFormat": {"horizontalAlignment": "CENTER"}},
            "fields": "userEnteredFormat.horizontalAlignment"
        }
    })
    reqs_m.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_id,
                "startRowIndex": 1,
                "endRowIndex": tot_idx_m,
                "startColumnIndex": 1,
                "endColumnIndex": 2
            },
            "cell": {"userEnteredFormat": {"textFormat": {"bold": True}}},
            "fields": "userEnteredFormat.textFormat.bold"
        }
    })
    reqs_m.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_id,
                "startRowIndex": 1,
                "endRowIndex": tot_idx_m,
                "startColumnIndex": 3,
                "endColumnIndex": 4
            },
            "cell": {"userEnteredFormat": {"horizontalAlignment": "CENTER"}},
            "fields": "userEnteredFormat.horizontalAlignment"
        }
    })
    reqs_m.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_id,
                "startRowIndex": 1,
                "endRowIndex": tot_idx_m,
                "startColumnIndex": 4,
                "endColumnIndex": 5
            },
            "cell": {"userEnteredFormat": {"horizontalAlignment": "RIGHT", "textFormat": {"bold": True}}},
            "fields": "userEnteredFormat(horizontalAlignment,textFormat.bold)"
        }
    })
    reqs_m.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_id,
                "startRowIndex": 1,
                "endRowIndex": tot_idx_m,
                "startColumnIndex": 5,
                "endColumnIndex": 6
            },
            "cell": {"userEnteredFormat": {"horizontalAlignment": "CENTER"}},
            "fields": "userEnteredFormat.horizontalAlignment"
        }
    })
    reqs_m.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_id,
                "startRowIndex": 1,
                "endRowIndex": tot_idx_m,
                "startColumnIndex": 6,
                "endColumnIndex": 7
            },
            "cell": {"userEnteredFormat": {"wrapStrategy": "WRAP"}},
            "fields": "userEnteredFormat.wrapStrategy"
        }
    })
    # Total row styling (Row 27 / index tot_idx_m)
    reqs_m.append({
        "mergeCells": {
            "range": {
                "sheetId": markets_id,
                "startRowIndex": tot_idx_m,
                "endRowIndex": tot_idx_m + 1,
                "startColumnIndex": 0,
                "endColumnIndex": 3
            },
            "mergeType": "MERGE_ALL"
        }
    })
    reqs_m.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_id,
                "startRowIndex": tot_idx_m,
                "endRowIndex": tot_idx_m + 1,
                "startColumnIndex": 0,
                "endColumnIndex": 7
            },
            "cell": {
                "userEnteredFormat": {
                    "backgroundColor": {"red": 0.85, "green": 0.88, "blue": 0.95},
                    "verticalAlignment": "MIDDLE",
                    "textFormat": {
                        "fontFamily": "Arial",
                        "fontSize": 10,
                        "bold": True,
                        "foregroundColor": {"red": 0.12, "green": 0.31, "blue": 0.47}
                    }
                }
            },
            "fields": "userEnteredFormat(backgroundColor,verticalAlignment,textFormat)"
        }
    })
    reqs_m.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_id,
                "startRowIndex": tot_idx_m,
                "endRowIndex": tot_idx_m + 1,
                "startColumnIndex": 0,
                "endColumnIndex": 3
            },
            "cell": {"userEnteredFormat": {"horizontalAlignment": "CENTER"}},
            "fields": "userEnteredFormat.horizontalAlignment"
        }
    })
    reqs_m.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_id,
                "startRowIndex": tot_idx_m,
                "endRowIndex": tot_idx_m + 1,
                "startColumnIndex": 3,
                "endColumnIndex": 4
            },
            "cell": {"userEnteredFormat": {"horizontalAlignment": "CENTER"}},
            "fields": "userEnteredFormat.horizontalAlignment"
        }
    })
    reqs_m.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_id,
                "startRowIndex": tot_idx_m,
                "endRowIndex": tot_idx_m + 1,
                "startColumnIndex": 4,
                "endColumnIndex": 5
            },
            "cell": {"userEnteredFormat": {"horizontalAlignment": "RIGHT"}},
            "fields": "userEnteredFormat.horizontalAlignment"
        }
    })
    reqs_m.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_id,
                "startRowIndex": tot_idx_m,
                "endRowIndex": tot_idx_m + 1,
                "startColumnIndex": 5,
                "endColumnIndex": 6
            },
            "cell": {"userEnteredFormat": {"horizontalAlignment": "CENTER"}},
            "fields": "userEnteredFormat.horizontalAlignment"
        }
    })
    # Borders
    border_style = {"style": "SOLID", "width": 1, "color": {"red": 0.85, "green": 0.88, "blue": 0.92}}
    reqs_m.append({
        "updateBorders": {
            "range": {
                "sheetId": markets_id,
                "startRowIndex": 0,
                "endRowIndex": tot_idx_m + 1,
                "startColumnIndex": 0,
                "endColumnIndex": 7
            },
            "top": border_style, "bottom": border_style, "left": border_style, "right": border_style,
            "innerHorizontal": border_style, "innerVertical": border_style
        }
    })
    batch_update_spreadsheet(token, reqs_m)
    print("Financial Markets styling completed!")

    # -------------------------------------------------------------
    # 2. Update Sheet: Product Roadmap
    # -------------------------------------------------------------
    print("Updating Google Sheet: Product Roadmap...")
    ws_r = wb["Roadmap"]
    roadmap_rows = []
    for r in range(1, ws_r.max_row + 1):
        row_vals = [ws_r.cell(r, c).value for c in range(1, ws_r.max_column + 1)]
        roadmap_rows.append([str(v) if v is not None else "" for v in row_vals])

    clear_range(token, "Product Roadmap!A1:F50")
    batch_update_spreadsheet(token, [{
        "unmergeCells": {
            "range": {
                "sheetId": roadmap_id,
                "startRowIndex": 0,
                "endRowIndex": 50,
                "startColumnIndex": 0,
                "endColumnIndex": 6
            }
        }
    }])

    res_r = update_values(token, f"Product Roadmap!A1:F{len(roadmap_rows)}", roadmap_rows)
    print(f"Product Roadmap updated: {res_r.get('updatedCells')} cells across {len(roadmap_rows)} rows.")

    reqs_r = []
    # Freeze 1 row
    reqs_r.append({
        "updateSheetProperties": {
            "properties": {
                "sheetId": roadmap_id,
                "gridProperties": {
                    "frozenRowCount": 1
                }
            },
            "fields": "gridProperties.frozenRowCount"
        }
    })
    col_widths_r = [130, 110, 190, 260, 540, 110]
    for c_idx, width in enumerate(col_widths_r):
        reqs_r.append({
            "updateDimensionProperties": {
                "range": {
                    "sheetId": roadmap_id,
                    "dimension": "COLUMNS",
                    "startIndex": c_idx,
                    "endIndex": c_idx + 1
                },
                "properties": {"pixelSize": width},
                "fields": "pixelSize"
            }
        })
    # Header Row 1 (index 0)
    reqs_r.append({
        "repeatCell": {
            "range": {
                "sheetId": roadmap_id,
                "startRowIndex": 0,
                "endRowIndex": 1,
                "startColumnIndex": 0,
                "endColumnIndex": 6
            },
            "cell": {
                "userEnteredFormat": {
                    "backgroundColor": {"red": 0.12, "green": 0.31, "blue": 0.47},
                    "horizontalAlignment": "LEFT",
                    "verticalAlignment": "MIDDLE",
                    "textFormat": {
                        "fontFamily": "Arial",
                        "fontSize": 10,
                        "bold": True,
                        "foregroundColor": {"red": 1.0, "green": 1.0, "blue": 1.0}
                    }
                }
            },
            "fields": "userEnteredFormat(backgroundColor,horizontalAlignment,verticalAlignment,textFormat)"
        }
    })
    for c_idx in [0, 1, 5]:
        reqs_r.append({
            "repeatCell": {
                "range": {
                    "sheetId": roadmap_id,
                    "startRowIndex": 0,
                    "endRowIndex": 1,
                    "startColumnIndex": c_idx,
                    "endColumnIndex": c_idx + 1
                },
                "cell": {"userEnteredFormat": {"horizontalAlignment": "CENTER"}},
                "fields": "userEnteredFormat.horizontalAlignment"
            }
        })
    # Data rows
    reqs_r.append({
        "repeatCell": {
            "range": {
                "sheetId": roadmap_id,
                "startRowIndex": 1,
                "endRowIndex": len(roadmap_rows),
                "startColumnIndex": 0,
                "endColumnIndex": 6
            },
            "cell": {
                "userEnteredFormat": {
                    "verticalAlignment": "MIDDLE",
                    "textFormat": {
                        "fontFamily": "Arial",
                        "fontSize": 10,
                        "foregroundColor": {"red": 0.12, "green": 0.16, "blue": 0.23}
                    }
                }
            },
            "fields": "userEnteredFormat(verticalAlignment,textFormat)"
        }
    })
    for r_idx in range(1, len(roadmap_rows)):
        bg = {"red": 0.97, "green": 0.98, "blue": 0.99} if r_idx % 2 == 1 else {"red": 1.0, "green": 1.0, "blue": 1.0}
        reqs_r.append({
            "repeatCell": {
                "range": {
                    "sheetId": roadmap_id,
                    "startRowIndex": r_idx,
                    "endRowIndex": r_idx + 1,
                    "startColumnIndex": 0,
                    "endColumnIndex": 5
                },
                "cell": {"userEnteredFormat": {"backgroundColor": bg}},
                "fields": "userEnteredFormat.backgroundColor"
            }
        })
    reqs_r.append({
        "repeatCell": {
            "range": {
                "sheetId": roadmap_id,
                "startRowIndex": 1,
                "endRowIndex": len(roadmap_rows),
                "startColumnIndex": 4,
                "endColumnIndex": 5
            },
            "cell": {"userEnteredFormat": {"wrapStrategy": "WRAP"}},
            "fields": "userEnteredFormat.wrapStrategy"
        }
    })
    reqs_r.append({
        "repeatCell": {
            "range": {
                "sheetId": roadmap_id,
                "startRowIndex": 1,
                "endRowIndex": len(roadmap_rows),
                "startColumnIndex": 3,
                "endColumnIndex": 4
            },
            "cell": {"userEnteredFormat": {"textFormat": {"bold": True}}},
            "fields": "userEnteredFormat.textFormat.bold"
        }
    })
    for c_idx in [0, 1]:
        reqs_r.append({
            "repeatCell": {
                "range": {
                    "sheetId": roadmap_id,
                    "startRowIndex": 1,
                    "endRowIndex": len(roadmap_rows),
                    "startColumnIndex": c_idx,
                    "endColumnIndex": c_idx + 1
                },
                "cell": {"userEnteredFormat": {"horizontalAlignment": "CENTER"}},
                "fields": "userEnteredFormat.horizontalAlignment"
            }
        })
    # Status badges
    for r_idx in range(1, len(roadmap_rows)):
        status_val = roadmap_rows[r_idx][5] if len(roadmap_rows[r_idx]) > 5 else ""
        if status_val == "Done":
            bg = {"red": 0.86, "green": 0.98, "blue": 0.90}
            fg = {"red": 0.08, "green": 0.50, "blue": 0.24}
        elif status_val == "In Progress":
            bg = {"red": 0.99, "green": 0.95, "blue": 0.78}
            fg = {"red": 0.71, "green": 0.33, "blue": 0.04}
        else:
            bg = {"red": 0.88, "green": 0.91, "blue": 1.0}
            fg = {"red": 0.26, "green": 0.22, "blue": 0.79}

        reqs_r.append({
            "repeatCell": {
                "range": {
                    "sheetId": roadmap_id,
                    "startRowIndex": r_idx,
                    "endRowIndex": r_idx + 1,
                    "startColumnIndex": 5,
                    "endColumnIndex": 6
                },
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": bg,
                        "horizontalAlignment": "CENTER",
                        "verticalAlignment": "MIDDLE",
                        "textFormat": {
                            "fontFamily": "Arial",
                            "fontSize": 10,
                            "bold": True,
                            "foregroundColor": fg
                        }
                    }
                },
                "fields": "userEnteredFormat(backgroundColor,horizontalAlignment,verticalAlignment,textFormat)"
            }
        })
    reqs_r.append({
        "updateBorders": {
            "range": {
                "sheetId": roadmap_id,
                "startRowIndex": 0,
                "endRowIndex": len(roadmap_rows),
                "startColumnIndex": 0,
                "endColumnIndex": 6
            },
            "top": border_style, "bottom": border_style, "left": border_style, "right": border_style,
            "innerHorizontal": border_style, "innerVertical": border_style
        }
    })
    batch_update_spreadsheet(token, reqs_r)
    print("Product Roadmap styling completed!")

    # -------------------------------------------------------------
    # 3. Update Sheet: Sitemap
    # -------------------------------------------------------------
    print("Updating Google Sheet: Sitemap...")
    ws_s = wb["Sitemap"]
    sitemap_rows = []
    for r in range(1, ws_s.max_row + 1):
        row_vals = [ws_s.cell(r, c).value for c in range(1, ws_s.max_column + 1)]
        sitemap_rows.append([str(v) if v is not None else "" for v in row_vals])

    clear_range(token, "Sitemap!A1:D150")
    batch_update_spreadsheet(token, [{
        "unmergeCells": {
            "range": {
                "sheetId": sitemap_id,
                "startRowIndex": 0,
                "endRowIndex": 150,
                "startColumnIndex": 0,
                "endColumnIndex": 4
            }
        }
    }])

    res_s = update_values(token, f"Sitemap!A1:D{len(sitemap_rows)}", sitemap_rows)
    print(f"Sitemap updated: {res_s.get('updatedCells')} cells across {len(sitemap_rows)} rows.")

    reqs_s = []
    # Freeze 1 row
    reqs_s.append({
        "updateSheetProperties": {
            "properties": {
                "sheetId": sitemap_id,
                "gridProperties": {
                    "frozenRowCount": 1
                }
            },
            "fields": "gridProperties.frozenRowCount"
        }
    })
    col_widths_s = [65, 280, 460, 340]
    for c_idx, width in enumerate(col_widths_s):
        reqs_s.append({
            "updateDimensionProperties": {
                "range": {
                    "sheetId": sitemap_id,
                    "dimension": "COLUMNS",
                    "startIndex": c_idx,
                    "endIndex": c_idx + 1
                },
                "properties": {"pixelSize": width},
                "fields": "pixelSize"
            }
        })
    # Header Row 1 (index 0)
    reqs_s.append({
        "repeatCell": {
            "range": {
                "sheetId": sitemap_id,
                "startRowIndex": 0,
                "endRowIndex": 1,
                "startColumnIndex": 0,
                "endColumnIndex": 4
            },
            "cell": {
                "userEnteredFormat": {
                    "backgroundColor": {"red": 0.12, "green": 0.31, "blue": 0.47},
                    "horizontalAlignment": "LEFT",
                    "verticalAlignment": "MIDDLE",
                    "textFormat": {
                        "fontFamily": "Arial",
                        "fontSize": 10,
                        "bold": True,
                        "foregroundColor": {"red": 1.0, "green": 1.0, "blue": 1.0}
                    }
                }
            },
            "fields": "userEnteredFormat(backgroundColor,horizontalAlignment,verticalAlignment,textFormat)"
        }
    })
    reqs_s.append({
        "repeatCell": {
            "range": {
                "sheetId": sitemap_id,
                "startRowIndex": 0,
                "endRowIndex": 1,
                "startColumnIndex": 0,
                "endColumnIndex": 1
            },
            "cell": {"userEnteredFormat": {"horizontalAlignment": "CENTER"}},
            "fields": "userEnteredFormat.horizontalAlignment"
        }
    })
    # Data rows
    reqs_s.append({
        "repeatCell": {
            "range": {
                "sheetId": sitemap_id,
                "startRowIndex": 1,
                "endRowIndex": len(sitemap_rows),
                "startColumnIndex": 0,
                "endColumnIndex": 4
            },
            "cell": {
                "userEnteredFormat": {
                    "verticalAlignment": "MIDDLE",
                    "textFormat": {
                        "fontFamily": "Arial",
                        "fontSize": 10,
                        "foregroundColor": {"red": 0.12, "green": 0.16, "blue": 0.23}
                    }
                }
            },
            "fields": "userEnteredFormat(verticalAlignment,textFormat)"
        }
    })
    for r_idx in range(1, len(sitemap_rows)):
        bg = {"red": 0.97, "green": 0.98, "blue": 0.99} if r_idx % 2 == 1 else {"red": 1.0, "green": 1.0, "blue": 1.0}
        reqs_s.append({
            "repeatCell": {
                "range": {
                    "sheetId": sitemap_id,
                    "startRowIndex": r_idx,
                    "endRowIndex": r_idx + 1,
                    "startColumnIndex": 0,
                    "endColumnIndex": 4
                },
                "cell": {"userEnteredFormat": {"backgroundColor": bg}},
                "fields": "userEnteredFormat.backgroundColor"
            }
        })
    reqs_s.append({
        "repeatCell": {
            "range": {
                "sheetId": sitemap_id,
                "startRowIndex": 1,
                "endRowIndex": len(sitemap_rows),
                "startColumnIndex": 0,
                "endColumnIndex": 1
            },
            "cell": {"userEnteredFormat": {"horizontalAlignment": "CENTER"}},
            "fields": "userEnteredFormat.horizontalAlignment"
        }
    })
    reqs_s.append({
        "repeatCell": {
            "range": {
                "sheetId": sitemap_id,
                "startRowIndex": 1,
                "endRowIndex": len(sitemap_rows),
                "startColumnIndex": 1,
                "endColumnIndex": 2
            },
            "cell": {"userEnteredFormat": {"textFormat": {"bold": True}}},
            "fields": "userEnteredFormat.textFormat.bold"
        }
    })
    reqs_s.append({
        "repeatCell": {
            "range": {
                "sheetId": sitemap_id,
                "startRowIndex": 1,
                "endRowIndex": len(sitemap_rows),
                "startColumnIndex": 3,
                "endColumnIndex": 4
            },
            "cell": {
                "userEnteredFormat": {
                    "textFormat": {
                        "fontFamily": "Consolas",
                        "fontSize": 10,
                        "foregroundColor": {"red": 0.016, "green": 0.47, "blue": 0.34}
                    }
                }
            },
            "fields": "userEnteredFormat.textFormat"
        }
    })
    reqs_s.append({
        "updateBorders": {
            "range": {
                "sheetId": sitemap_id,
                "startRowIndex": 0,
                "endRowIndex": len(sitemap_rows),
                "startColumnIndex": 0,
                "endColumnIndex": 4
            },
            "top": border_style, "bottom": border_style, "left": border_style, "right": border_style,
            "innerHorizontal": border_style, "innerVertical": border_style
        }
    })
    batch_update_spreadsheet(token, reqs_s)
    print("Sitemap styling completed!")


def main():
    process_local_excel()
    process_google_sheets()
    print("\nAll title banner rows removed and sheets updated successfully!")


if __name__ == "__main__":
    main()
