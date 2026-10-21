#!/usr/bin/env python3
import os
import sys
import glob
import argparse
import openpyxl
import json
from datetime import datetime
from collections import defaultdict

DOWNLOADS_DIR = "/Users/hienhv/Downloads"
BASE_DIR = "/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base"
SSOT_EXCEL = os.path.join(BASE_DIR, "07_REPORTS/web_performance_tracking.xlsx")

_parser = argparse.ArgumentParser(description="Sync Web Performance Tracking")
_parser.add_argument("--file", help="Duong dan file xlsx tracking (mac dinh: tim moi nhat trong Downloads)")
_parser.add_argument("--no-push", action="store_true", help="Parse va ghi SSOT, khong push sang hubs-momo")
_args, _unknown = _parser.parse_known_args()

# Quy chuẩn đặt tên dự án / Sub-pages cho Cinema Hub trong Web Performance Tracking
CINEMA_PROJECT_NAME_MAP = {
    "Cinema_Cineplex Page": "Trang Rạp phim",
    "Cinema_Cineplex": "Trang Rạp phim",
    "Cinema_Page": "Trang phim chi tiết",
    "Cinema_Home Page": "Cinema_Home Page (Trang chủ Cinema)",
    "Cinema_home Page": "Cinema_home Page (Trang chủ Cinema)",
    "Cinema_News": "Tin tức",
    "Cinema_Blog": "Cinema_Blog",
    "Cinema_Landing Page": "Cinema_Landing Page",
}

def find_latest_tracking_file():
    if _args.file:
        if not os.path.exists(_args.file):
            raise FileNotFoundError(f"File khong ton tai: {_args.file}")
        return _args.file
    pattern = os.path.join(DOWNLOADS_DIR, "*[Ww]eb*[Pp]erformance*[Tt]racking*.xlsx")
    files = glob.glob(pattern)
    if not files:
        files = glob.glob(os.path.join(DOWNLOADS_DIR, "*[Tt]racking*.xlsx"))
    
    if not files:
        raise FileNotFoundError(f"Không tìm thấy file 'Web Performance Tracking' nào trong thư mục {DOWNLOADS_DIR}")
    
    files.sort(key=lambda x: os.path.getmtime(x), reverse=True)
    return files[0]

def parse_and_sync():
    latest_file = find_latest_tracking_file()
    print(f"File mới nhất phát hiện: {latest_file}")
    print(f"Thời gian cập nhật file: {datetime.fromtimestamp(os.path.getmtime(latest_file)).strftime('%Y-%m-%d %H:%M:%S')}")

    wb = openpyxl.load_workbook(latest_file, data_only=True)
    
    # Check days elapsed from Vehicle or Overall Performance cell (1,1) or (1,2)
    days_elapsed = 4 # Default for Oct 4d
    if 'Vehicle' in wb.sheetnames:
        c_val = wb['Vehicle'].cell(1, 2).value
        if c_val and str(c_val).startswith("Update to:"):
            try:
                days_elapsed = int(str(c_val).replace("Update to:", "").strip())
            except:
                pass
    elif 'Overall Performance' in wb.sheetnames:
        c_val = wb['Overall Performance'].cell(1, 1).value
        if c_val and str(c_val).startswith("Update to:"):
            try:
                days_elapsed = int(str(c_val).replace("Update to:", "").strip())
            except:
                pass

    ws_pivot = wb['MTD Pivot']
    oct_col_idx = 9 # Col 9 is 2026-10-01 MTD!
    
    # Read exact Grand Total from Pivot table
    grand_total_oct_4d = 558009
    for r in range(ws_pivot.max_row, 1, -1):
        c1 = ws_pivot.cell(r, 1).value
        if c1 and 'grand total' in str(c1).lower():
            v = ws_pivot.cell(r, oct_col_idx).value
            if v:
                grand_total_oct_4d = int(v)
                break

    totals_by_hub = defaultdict(int)
    totals_by_channel = defaultdict(int)
    cinema_subpages = defaultdict(lambda: {"total": 0, "channels": defaultdict(int)})
    
    for r in range(3, ws_pivot.max_row):
        proj = ws_pivot.cell(r, 1).value
        if proj and 'grand total' in str(proj).lower():
            continue
        group_hub = ws_pivot.cell(r, 2).value or 'Others'
        group_channel = ws_pivot.cell(r, 3).value or 'Others'
        pv = ws_pivot.cell(r, oct_col_idx).value or 0
        if proj and pv:
            try:
                pv = int(pv)
                totals_by_hub[group_hub] += pv
                totals_by_channel[group_channel] += pv
                if 'Cinema Hub' in str(group_hub):
                    p_name = CINEMA_PROJECT_NAME_MAP.get(str(proj), str(proj))
                    cinema_subpages[p_name]["total"] += pv
                    cinema_subpages[p_name]["channels"][group_channel] += pv
            except:
                pass

    total_oct_4d = grand_total_oct_4d
    target_oct_pv = 3665607 # Target T10
    
    hubs_oct_4d = defaultdict(int)
    for group_hub, pv in totals_by_hub.items():
        if 'New User' in str(group_hub):
            hubs_oct_4d['new_user'] += pv
        elif 'Cinema Hub' in str(group_hub):
            hubs_oct_4d['cinema'] += pv
        elif 'Financial Hub' in str(group_hub):
            hubs_oct_4d['financial'] += pv
        elif any(k in str(group_hub) for k in ['Phạt Nguội', 'Vehicle', 'Bảo Hiểm Ô Tô', 'Bảo Hiểm Xe Máy', 'Phí Không Dừng', 'Tiện Ích Giao Thông']):
            hubs_oct_4d['vehicle'] += pv
        elif 'Student Hub' in str(group_hub):
            hubs_oct_4d['student'] += pv

    daily_pace = total_oct_4d / days_elapsed
    forecast_31d = daily_pace * 31

    import shutil
    shutil.copy2(latest_file, SSOT_EXCEL)
    print(f"Đã lưu bản SSOT sạch vào {SSOT_EXCEL}")

    # Cập nhật chuẩn hóa label dự án Cinema trong sheet 'Cinema' của SSOT Excel
    try:
        wb_ssot = openpyxl.load_workbook(SSOT_EXCEL)
        if 'Cinema' in wb_ssot.sheetnames:
            ws_c = wb_ssot['Cinema']
            ws_c.cell(9, 1).value = 'Trang Rạp phim'
            ws_c.cell(16, 1).value = 'Trang phim chi tiết'
            ws_c.cell(30, 1).value = 'Cinema_Home Page (Trang chủ Cinema)'
            ws_c.cell(37, 1).value = 'Tin tức'
            wb_ssot.save(SSOT_EXCEL)
            print(f"Đã cập nhật chuẩn hóa tên sub-pages Cinema trong sheet 'Cinema' ({SSOT_EXCEL})")
    except Exception as e:
        print(f"Lưu ý: Không thể cập nhật label sheet Cinema: {e}")

    summary = {
        "file_source": latest_file,
        "read_timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "report_type": f"MTD {days_elapsed} Ngày (01/10 - {days_elapsed:02d}/10/2026 - Lũy Kế {days_elapsed} Ngày Tháng 10)",
        "days_elapsed": days_elapsed,
        "days_in_month": 31,
        "target_oct_pv": target_oct_pv,
        "actual_mtd_pv": total_oct_4d,
        "progress_pct": round((total_oct_4d / target_oct_pv) * 100, 2),
        "daily_pace": round(daily_pace, 2),
        "forecast_31d": round(forecast_31d, 2),
        "forecast_pct": round((forecast_31d / target_oct_pv) * 100, 2),
        "channels": {
            "organic": {"pv": totals_by_channel.get('Organic', 0), "pct": round(totals_by_channel.get('Organic', 0) / total_oct_4d * 100, 2)},
            "direct": {"pv": totals_by_channel.get('Direct', 0), "pct": round(totals_by_channel.get('Direct', 0) / total_oct_4d * 100, 2)},
            "referral": {"pv": totals_by_channel.get('Referral', 0), "pct": round(totals_by_channel.get('Referral', 0) / total_oct_4d * 100, 2)},
            "paid": {"pv": totals_by_channel.get('Paid', 0), "pct": round(totals_by_channel.get('Paid', 0) / total_oct_4d * 100, 2)},
            "others": {"pv": totals_by_channel.get('Others', 0), "pct": round(totals_by_channel.get('Others', 0) / total_oct_4d * 100, 2)}
        },
        "hubs": {
            "new_user": {"pv": hubs_oct_4d['new_user'], "target": 300000, "pct": round(hubs_oct_4d['new_user']/300000*100, 2), "forecast": round(hubs_oct_4d['new_user']/days_elapsed*31, 2)},
            "cinema": {"pv": hubs_oct_4d['cinema'], "target": 1117378, "pct": round(hubs_oct_4d['cinema']/1117378*100, 2), "forecast": round(hubs_oct_4d['cinema']/days_elapsed*31, 2)},
            "financial": {"pv": hubs_oct_4d['financial'], "target": 450000, "pct": round(hubs_oct_4d['financial']/450000*100, 2), "forecast": round(hubs_oct_4d['financial']/days_elapsed*31, 2)},
            "vehicle": {"pv": hubs_oct_4d['vehicle'], "target": 100000, "pct": round(hubs_oct_4d['vehicle']/100000*100, 2), "forecast": round(hubs_oct_4d['vehicle']/days_elapsed*31, 2)},
            "student": {"pv": hubs_oct_4d['student'], "target": 100000, "pct": round(hubs_oct_4d['student']/100000*100, 2), "forecast": round(hubs_oct_4d['student']/days_elapsed*31, 2)}
        },
        "cinema_breakdown": {
            name: {
                "pv": data["total"],
                "pct_hub": round((data["total"] / hubs_oct_4d['cinema'] * 100) if hubs_oct_4d['cinema'] else 0, 2),
                "channels": dict(data["channels"])
            }
            for name, data in sorted(cinema_subpages.items(), key=lambda x: x[1]["total"], reverse=True)
        }
    }

    # Load Technical SEO: Crawl Stats & Backlink Profile
    crawl_stats_file = os.path.join(BASE_DIR, "07_REPORTS/technical_seo/crawl_stats/gsc_crawl_stats_latest.json")
    if os.path.exists(crawl_stats_file):
        with open(crawl_stats_file, 'r', encoding='utf-8') as f:
            summary["crawl_stats"] = json.load(f)

    backlink_file = os.path.join(BASE_DIR, "07_REPORTS/technical_seo/backlinks/backlink_profile_latest.json")
    if os.path.exists(backlink_file):
        with open(backlink_file, 'r', encoding='utf-8') as f:
            summary["backlink_profile"] = json.load(f)

    # Write snapshot JSON (pure tracking)
    json_path = os.path.join(BASE_DIR, f"07_REPORTS/data/monthly_10_2026_mtd_{days_elapsed}d.json")
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
    print(f"Đã tạo JSON snapshot tracking tại {json_path}")

    # Write consolidated 3-pillar snapshot
    consolidated_path = os.path.join(BASE_DIR, "07_REPORTS/data/monthly_10_2026_consolidated.json")
    with open(consolidated_path, 'w', encoding='utf-8') as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
    print(f"Đã tạo JSON snapshot hợp nhất (Consolidated 3 Pillars) tại {consolidated_path}")

    # Write lean summary.json for hubs-momo & low-token usage (~800 bytes)
    summary_hubs = {
        "lastFull": {
            "month": "2026-09",
            "pv": 3952859,
            "prevPv": 3465363,
            "target": 3509044
        },
        "running": {
            "month": "2026-10",
            "daysElapsed": days_elapsed,
            "daysInMonth": 31,
            "report": f"monthly_10_2026_mtd_{days_elapsed}d.json",
            "pv": total_oct_4d,
            "target": target_oct_pv,
            "dailyPace": round(daily_pace, 2),
            "forecast31d": round(forecast_31d, 2),
            "forecastPct": round((forecast_31d / target_oct_pv) * 100, 2)
        },
        "channels": {
            "organic": totals_by_channel.get('Organic', 0),
            "direct": totals_by_channel.get('Direct', 0),
            "referral": totals_by_channel.get('Referral', 0),
            "paid": totals_by_channel.get('Paid', 0),
            "others": totals_by_channel.get('Others', 0)
        },
        "hubs": [
            {"id": "cinema", "name": "Cinema Hub", "pv": hubs_oct_4d['cinema'], "target": 1117378},
            {"id": "financial", "name": "Financial Hub", "pv": hubs_oct_4d['financial'], "target": 450000},
            {"id": "new_user", "name": "New User Hub", "pv": hubs_oct_4d['new_user'], "target": 300000},
            {"id": "vehicle", "name": "Vehicle Hub", "pv": hubs_oct_4d['vehicle'], "target": 100000},
            {"id": "student", "name": "Student Hub", "pv": hubs_oct_4d['student'], "target": 100000}
        ]
    }
    summary_path = os.path.join(BASE_DIR, "07_REPORTS/data/summary.json")
    with open(summary_path, 'w', encoding='utf-8') as f:
        json.dump(summary_hubs, f, ensure_ascii=False, indent=2)
    print(f"Đã tạo summary.json siêu nhẹ cho hubs-momo tại {summary_path}")

    # Auto push to hubs-momo if available locally (skip khi --no-push)
    # Lưu ý: Đã ngưng luồng tự động đẩy tasks sang Hubs (/bang-tin); chỉ đẩy số liệu /dashboard
    if not _args.no_push:
        hubs_repo = "/Users/hienhv/HienHv/Klaus/Web-Idea/react-web-momo"
        if os.path.exists(hubs_repo):
            build_script = os.path.join(hubs_repo, "scripts/build-dashboard-data.py")
            if os.path.exists(build_script):
                import subprocess
                p_res = subprocess.run([sys.executable, build_script, "--base", BASE_DIR], capture_output=True, text=True)
                if p_res.returncode == 0:
                    print(f"Đã tự động đẩy dữ liệu sang hubs-momo (/dashboard) thành công! (Luồng tasks đã ngưng)")
                else:
                    print(f"Cảnh báo khi sync sang hubs-momo (/dashboard): {p_res.stderr.strip()}")
    else:
        print("[no-push] Bo qua push sang hubs-momo (dieu phoi boi intake.py).")

    return summary

if __name__ == "__main__":
    res = parse_and_sync()
    print(f"\n=== KẾT QUẢ TRÍCH XUẤT TỰ ĐỘNG THÁNG 10 (MTD {res['days_elapsed']} NGÀY - HỢP NHẤT 3 TRỤ CỘT) ===")
    print(f"- Actual MTD: {res['actual_mtd_pv']:,} / Target: {res['target_oct_pv']:,} ({res['progress_pct']}%)")
    print(f"- Run-rate Forecast 31d: {res['forecast_31d']:,.0f} ({res['forecast_pct']}%)")
    print(f"- Crawl Stats Status: Loaded (Requests T10: {res.get('crawl_stats', {}).get('monthly_run_rate', {}).get('2026-10', {}).get('requests_per_day', 'N/A')} req/ngày)")
    print(f"- Backlink Profile Status: Loaded (Total Referring Sites: {res.get('backlink_profile', {}).get('total_referring_domains', 'N/A')}, Disavow Candidates: {res.get('backlink_profile', {}).get('toxic_audit_summary', {}).get('total_disavow_candidates', 'N/A')})")
    if "cinema_breakdown" in res:
        print(f"\n--- Phân rã Cinema Sub-pages MTD {res['days_elapsed']}d (Quy chuẩn Tracking mới) ---")
        for p_name, p_data in res["cinema_breakdown"].items():
            print(f"  * {p_name}: {p_data['pv']:,} PVs ({p_data['pct_hub']}%)")
