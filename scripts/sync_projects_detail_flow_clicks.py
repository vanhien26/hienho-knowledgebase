#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
SYNC PROJECTS DETAIL FLOW CLICKS TO GOOGLE SHEETS
Bổ sung bảng phân tích 'Click to app by Project (Web-to-App OneLink Clicks)'
vào sheet 'Projects Detail Flow' trên Google Sheets SSOT và file Excel local.
"""

import os
import sys
import json
import urllib.request
import urllib.parse
from pathlib import Path
from collections import defaultdict
import openpyxl

BASE_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE_DIR))

from scripts.sync_to_google_sheet_api import get_sheets_token, TARGET_SHEET_ID
from scripts.sync_ga4_to_hubs import get_access_token, run_ga4_report
import scripts.lib.url_project_resolver as upr

SHEET_TITLE = "Projects Detail Flow"
SHEET_ID = 316527624

def get_clicks_by_project(token, start_date, end_date):
    """Kéo số click onelink theo landing page từ GA4 và map sang Project Name."""
    rules = upr.load_rules()
    inferred = upr.load_inferred_rules(rules)
    
    res = run_ga4_report(token, {
        "dateRanges": [{"startDate": start_date, "endDate": end_date}],
        "dimensions": [{"name": "landingPage"}],
        "dimensionFilter": {
            "andGroup": {
                "expressions": [
                    {"filter": {"fieldName": "eventName", "inListFilter": {"values": ["click_to_app", "click_outbound_link"]}}},
                    {
                        "orGroup": {
                            "expressions": [
                                {"filter": {"fieldName": "customEvent:click_url", "stringFilter": {"matchType": "CONTAINS", "value": "onelink"}}},
                                {"filter": {"fieldName": "linkUrl", "stringFilter": {"matchType": "CONTAINS", "value": "onelink"}}}
                            ]
                        }
                    }
                ]
            }
        },
        "metrics": [{"name": "eventCount"}],
        "limit": 10000
    })
    
    clicks = defaultdict(int)
    for r in res.get("rows", []):
        lp = r["dimensionValues"][0]["value"]
        cnt = int(r["metricValues"][0]["value"])
        p = upr.resolve_primary("https://www.momo.vn" + lp, rules, inferred)
        clicks[p] += cnt
    return clicks

def build_click_table(token):
    print("[1/4] Kéo dữ liệu Onelink Click từ GA4 Data API...")
    oct_mtd = get_clicks_by_project(token, "2026-10-01", "2026-10-26")
    sep_mtd = get_clicks_by_project(token, "2026-09-01", "2026-09-26")
    sep_full = get_clicks_by_project(token, "2026-09-01", "2026-09-30")
    
    print(f"  ✓ Oct MTD (26d): {sum(oct_mtd.values()):,} clicks")
    print(f"  ✓ Sep MTD (26d): {sum(sep_mtd.values()):,} clicks")
    print(f"  ✓ Sep Full(30d): {sum(sep_full.values()):,} clicks")

    # Đọc danh sách project gốc từ hàng 8-72 của sheet
    print("[2/4] Khớp cấu trúc danh mục Dự án từ Projects Detail Flow...")
    range_name = urllib.parse.quote(f"{SHEET_TITLE}!A8:A72")
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{TARGET_SHEET_ID}/values/{range_name}"
    req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
    with urllib.request.urlopen(req) as resp:
        data = json.loads(resp.read().decode())
    sheet_projects = [r[0] if r else "" for r in data.get("values", [])]

    # Map tên dự án không phân biệt hoa thường
    norm_oct = {k.lower(): v for k, v in oct_mtd.items()}
    norm_sep_m = {k.lower(): v for k, v in sep_mtd.items()}
    norm_sep_f = {k.lower(): v for k, v in sep_full.items()}

    rows = []
    # Row 111: Tiêu đề khối
    r111 = ["Click to app by Project (Web-to-App OneLink Clicks)"] + [""] * 19
    rows.append(r111)

    # Row 112: Header các cột
    r112 = [
        "Projects", "May 26", "Jun 26", "Jul 26", "Aug 26", "Sep 26 (MTD)",
        "Oct 26 (MTD)", "Nov 26", "Dec 26", "VS. SPLM", "Oct CVR %",
        "May 26", "Jun 26", "Jul 26", "Aug 26", "Sep 26 (Full)",
        "Oct 26", "Nov 26", "Dec 26", "Sep CVR %"
    ]
    rows.append(r112)

    # Row 113: Dòng Total Click to app
    r113 = [
        "Total Click to app",
        "", "", "", "",
        "=SUM(F114, F121:F123, F124, F139, F144:F168, F169, F172:F178)",
        "=SUM(G114, G121:G123, G124, G139, G144:G168, G169, G172:G178)",
        "", "",
        "=G113-F113",
        "=IF(G3>0, G113/G3, 0)",
        "", "", "", "",
        "=SUM(P114, P121:P123, P124, P139, P144:P168, P169, P172:P178)",
        "", "", "",
        "=IF(P3>0, P113/P3, 0)"
    ]
    rows.append(r113)

    # Hub mapping definition theo row index để tránh trùng tên (như Student Hub ở cả dòng 63 và 65)
    hub_row_defs = {
        8: ("SUM(F115:F120)", "SUM(G115:G120)", "SUM(P115:P120)"),
        18: ("SUM(F125:F138)", "SUM(G125:G138)", "SUM(P125:P138)"),
        33: ("SUM(F140:F143)", "SUM(G140:G143)", "SUM(P140:P143)"),
        63: ("SUM(F170:F171)", "SUM(G170:G171)", "SUM(P170:P171)")
    }

    # Rows 114 to 178
    for i, proj in enumerate(sheet_projects, start=8):
        nr = i + 106 # new row index
        if i in hub_row_defs:
            f_sep_m, f_oct_m, f_sep_f = hub_row_defs[i]
            row_data = [
                proj,
                "", "", "", "",
                f"={f_sep_m}",
                f"={f_oct_m}",
                "", "",
                f"=G{nr}-F{nr}",
                f"=IF(G{i}>0, G{nr}/G{i}, 0)",
                "", "", "", "",
                f"={f_sep_f}",
                "", "", "",
                f"=IF(P{i}>0, P{nr}/P{i}, 0)"
            ]
        else:
            c_oct = oct_mtd.get(proj, norm_oct.get(proj.lower(), 0))
            c_sep_m = sep_mtd.get(proj, norm_sep_m.get(proj.lower(), 0))
            c_sep_f = sep_full.get(proj, norm_sep_f.get(proj.lower(), 0))
            row_data = [
                proj,
                "", "", "", "",
                c_sep_m,
                c_oct,
                "", "",
                f"=G{nr}-F{nr}",
                f"=IF(G{i}>0, G{nr}/G{i}, 0)",
                "", "", "", "",
                c_sep_f,
                "", "", "",
                f"=IF(P{i}>0, P{nr}/P{i}, 0)"
            ]
        rows.append(row_data)

    return rows

def sync_google_sheet_clicks():
    token = get_sheets_token()
    rows = build_click_table(token)

    print(f"[3/4] Đang ghi {len(rows)} hàng vào Google Sheets ({SHEET_TITLE}!A111:T178)...")
    range_target = urllib.parse.quote(f"{SHEET_TITLE}!A111:T178")
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{TARGET_SHEET_ID}/values/{range_target}?valueInputOption=USER_ENTERED"
    body = {
        "range": f"{SHEET_TITLE}!A111:T178",
        "majorDimension": "ROWS",
        "values": rows
    }
    req = urllib.request.Request(
        url,
        data=json.dumps(body).encode("utf-8"),
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="PUT"
    )
    with urllib.request.urlopen(req) as resp:
        res = json.loads(resp.read().decode())
    print(f"  ✓ Đã cập nhật {res.get('updatedCells', 0)} ô thành công!")

    # Format styling qua batchUpdate
    print("[4/4] Áp dụng định dạng (Number format, Bold, Colors) cho bảng mới...")
    format_requests = [
        # Format Clicks là số có phân tách ngàn (#,##0) cho các cột B:J và L:S (hàng 113-178, 0-indexed: 112 to 178)
        {
            "repeatCell": {
                "range": {
                    "sheetId": SHEET_ID,
                    "startRowIndex": 112,
                    "endRowIndex": 178,
                    "startColumnIndex": 1,
                    "endColumnIndex": 10
                },
                "cell": {
                    "userEnteredFormat": {
                        "numberFormat": {"type": "NUMBER", "pattern": "#,##0"}
                    }
                },
                "fields": "userEnteredFormat.numberFormat"
            }
        },
        {
            "repeatCell": {
                "range": {
                    "sheetId": SHEET_ID,
                    "startRowIndex": 112,
                    "endRowIndex": 178,
                    "startColumnIndex": 11,
                    "endColumnIndex": 19
                },
                "cell": {
                    "userEnteredFormat": {
                        "numberFormat": {"type": "NUMBER", "pattern": "#,##0"}
                    }
                },
                "fields": "userEnteredFormat.numberFormat"
            }
        },
        # Format CVR % là phần trăm (0.0%) cho cột K (col 10) và cột T (col 19)
        {
            "repeatCell": {
                "range": {
                    "sheetId": SHEET_ID,
                    "startRowIndex": 112,
                    "endRowIndex": 178,
                    "startColumnIndex": 10,
                    "endColumnIndex": 11
                },
                "cell": {
                    "userEnteredFormat": {
                        "numberFormat": {"type": "PERCENT", "pattern": "0.0%"}
                    }
                },
                "fields": "userEnteredFormat.numberFormat"
            }
        },
        {
            "repeatCell": {
                "range": {
                    "sheetId": SHEET_ID,
                    "startRowIndex": 112,
                    "endRowIndex": 178,
                    "startColumnIndex": 19,
                    "endColumnIndex": 20
                },
                "cell": {
                    "userEnteredFormat": {
                        "numberFormat": {"type": "PERCENT", "pattern": "0.0%"}
                    }
                },
                "fields": "userEnteredFormat.numberFormat"
            }
        },
        # Định dạng Tiêu đề Section (Row 111 - index 110)
        {
            "repeatCell": {
                "range": {
                    "sheetId": SHEET_ID,
                    "startRowIndex": 110,
                    "endRowIndex": 111,
                    "startColumnIndex": 0,
                    "endColumnIndex": 20
                },
                "cell": {
                    "userEnteredFormat": {
                        "textFormat": {"bold": True, "italic": True, "fontSize": 11}
                    }
                },
                "fields": "userEnteredFormat.textFormat"
            }
        },
        # Định dạng Header Cột (Row 112 - index 111)
        {
            "repeatCell": {
                "range": {
                    "sheetId": SHEET_ID,
                    "startRowIndex": 111,
                    "endRowIndex": 112,
                    "startColumnIndex": 0,
                    "endColumnIndex": 20
                },
                "cell": {
                    "userEnteredFormat": {
                        "textFormat": {"bold": True, "fontSize": 10},
                        "backgroundColor": {"red": 0.9, "green": 0.93, "blue": 0.98}
                    }
                },
                "fields": "userEnteredFormat(textFormat,backgroundColor)"
            }
        },
        # Định dạng Dòng Total (Row 113 - index 112)
        {
            "repeatCell": {
                "range": {
                    "sheetId": SHEET_ID,
                    "startRowIndex": 112,
                    "endRowIndex": 113,
                    "startColumnIndex": 0,
                    "endColumnIndex": 20
                },
                "cell": {
                    "userEnteredFormat": {
                        "textFormat": {"bold": True, "fontSize": 10},
                        "backgroundColor": {"red": 0.95, "green": 0.95, "blue": 0.95}
                    }
                },
                "fields": "userEnteredFormat(textFormat,backgroundColor)"
            }
        }
    ]

    # Thêm định dạng bold cho các dòng Hubs (Row 114, 124, 139, 169)
    for hr_idx in [113, 123, 138, 168]: # 0-indexed for 114, 124, 139, 169
        format_requests.append({
            "repeatCell": {
                "range": {
                    "sheetId": SHEET_ID,
                    "startRowIndex": hr_idx,
                    "endRowIndex": hr_idx + 1,
                    "startColumnIndex": 0,
                    "endColumnIndex": 20
                },
                "cell": {
                    "userEnteredFormat": {
                        "textFormat": {"bold": True}
                    }
                },
                "fields": "userEnteredFormat.textFormat"
            }
        })

    batch_url = f"https://sheets.googleapis.com/v4/spreadsheets/{TARGET_SHEET_ID}:batchUpdate"
    batch_req = urllib.request.Request(
        batch_url,
        data=json.dumps({"requests": format_requests}).encode("utf-8"),
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="POST"
    )
    with urllib.request.urlopen(batch_req) as resp:
        json.loads(resp.read().decode())
    print("  ✓ Đã hoàn tất định dạng bảng đẹp mắt chuẩn Product Lead!")

    # Cập nhật bản sao Excel Local
    sync_local_excel(rows)

def sync_local_excel(rows):
    excel_path = BASE_DIR / "07_REPORTS" / "web_performance_tracking.xlsx"
    if not excel_path.exists():
        return
    print(f"\n[Local Excel] Đang đồng bộ vào {excel_path.name}...")
    wb = openpyxl.load_workbook(str(excel_path))
    if SHEET_TITLE in wb.sheetnames:
        ws = wb[SHEET_TITLE]
        for r_idx, row_vals in enumerate(rows, start=111):
            for c_idx, val in enumerate(row_vals, start=1):
                ws.cell(row=r_idx, column=c_idx, value=val)
        wb.save(str(excel_path))
        print("  ✓ Đã cập nhật xong file Excel local!")

if __name__ == "__main__":
    sync_google_sheet_clicks()
