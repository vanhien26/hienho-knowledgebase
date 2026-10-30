#!/usr/bin/env python3
"""
Update Financial Hub Roadmap and eliminate 'SIP' across the entire project.
1. Move Stock stream to Month 11/2026 with status 'Pending' (waiting for Cell Team feedback).
2. Remove 'SIP' across all files, sheets, and documents (replace with 'Chứng Chỉ Quỹ' or 'tích sản định kỳ').
3. Restructure sheet Roadmap with complete 21 items (including Phase 3 & 4 for Nov-Dec targets).
4. Update local Excel 05_HUBS/financial/financial-hub-roadmap.xlsx.
5. Update Google Sheet [WP] Financial Hub | Web Product (1_AG4otSTXH7ontpPL0zjylYgVfxOozvoSc5PziUPM6A).
6. Update 05_HUBS/financial/financial-hub-brd.md, PRD docs.
"""

import json
import re
import urllib.parse
import urllib.request
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent
EXCEL_PATH = BASE_DIR / "05_HUBS" / "financial" / "financial-hub-roadmap.xlsx"
BRD_PATH = BASE_DIR / "05_HUBS" / "financial" / "financial-hub-brd.md"
PRD_CK_1 = BASE_DIR / "05_HUBS" / "financial" / "chung-khoan-profile-prd.md"
PRD_CK_2 = BASE_DIR / "09_PRD" / "chung-khoan-profile-prd.md"
PRD_AI = BASE_DIR / "09_PRD" / "momo-ai-proposal-engine-prd.md"

CREDS_FILE = BASE_DIR / "07_REPORTS" / "data" / "credentials" / "ga4-credentials.json"
TOKEN_FILE = BASE_DIR / "07_REPORTS" / "data" / "credentials" / "gdrive-token.json"
SPREADSHEET_ID = "1_AG4otSTXH7ontpPL0zjylYgVfxOozvoSc5PziUPM6A"


SIP_REPLACEMENTS = [
    ("Chứng Chỉ Quỹ & Quỹ Mở SIP", "Chứng Chỉ Quỹ & Quỹ Mở"),
    ("Chứng Chỉ Quỹ & Máy Tính Tích Sản SIP", "Chứng Chỉ Quỹ & Máy Tính Tích Sản"),
    ("Chứng Chỉ Quỹ SIP", "Chứng Chỉ Quỹ"),
    ("Chứng Chỉ Quỹ (SIP)", "Chứng Chỉ Quỹ"),
    ("quỹ mở (SIP)", "Chứng Chỉ Quỹ"),
    ("Quỹ Mở SIP", "Quỹ Mở"),
    ("SIP Quỹ Mở", "Quỹ Mở"),
    ("SIP Engine", "Đầu Tư Định Kỳ"),
    ("Máy Tính Tích Sản SIP", "Máy Tính Tích Sản"),
    ("máy tính tích sản SIP", "máy tính tích sản Chứng Chỉ Quỹ"),
    ("tool tích sản SIP", "công cụ đầu tư Chứng Chỉ Quỹ"),
    ("Tích sản hưu trí định kỳ SIP từ 10k", "Tích sản hưu trí định kỳ từ 10k"),
    ("Tích sản SIP từ 10.000đ/ngày", "Tích sản Chứng Chỉ Quỹ từ 10.000đ/ngày"),
    ("Tích sản SIP từ 10.000đ", "Tích sản Chứng Chỉ Quỹ từ 10.000đ"),
    ("tích sản SIP từ 10.000đ/ngày", "tích sản Chứng Chỉ Quỹ từ 10.000đ/ngày"),
    ("Giả Lập Tích Sản SIP", "Giả Lập Tích Sản"),
    ("tích sản SIP", "tích sản Chứng Chỉ Quỹ"),
    ("Tích Sản SIP", "Tích Sản Chứng Chỉ Quỹ"),
    ("định kỳ SIP từ 10.000đ/ngày", "định kỳ từ 10.000đ/ngày"),
    ("định kỳ SIP", "định kỳ"),
    ("tích lũy định kỳ SIP", "tích lũy Chứng Chỉ Quỹ định kỳ"),
    ("Tích Sản FIRE & Giả Lập SIP", "Tích Sản FIRE & Đầu Tư Định Kỳ"),
    ("Bộ Giả Lập Lãi Kép SIP", "Bộ Giả Lập Lãi Kép Chứng Chỉ Quỹ"),
    ("Giả Lập SIP", "Giả Lập Chứng Chỉ Quỹ"),
    ("Bộ Giả Lập SIP", "Bộ Giả Lập Chứng Chỉ Quỹ"),
    ("Chiến Lược Tích Sản SIP Là Gì?", "Chiến Lược Tích Sản Chứng Chỉ Quỹ Là Gì?"),
    ("Bắt Đầu Tích Sản SIP Với Quỹ Mở", "Bắt Đầu Đầu Tư Định Kỳ Quỹ Mở"),
    ("bắt đầu tích sản SIP", "bắt đầu tích sản Chứng Chỉ Quỹ"),
    ("Bắt đầu tích sản SIP", "Bắt đầu tích sản Chứng Chỉ Quỹ"),
    ("Bắt tay vào tích sản SIP", "Bắt tay vào tích sản Chứng Chỉ Quỹ"),
    (" (SIP)", ""),
    (" SIP", " Chứng Chỉ Quỹ"),
    ("SIP", "Chứng Chỉ Quỹ")
]


def clean_sip_text(text: str) -> str:
    res = text
    for old, new in SIP_REPLACEMENTS:
        res = res.replace(old, new)
    return res


def get_access_token():
    with open(CREDS_FILE, "r", encoding="utf-8") as f:
        cdata = json.load(f)
    client_info = cdata.get("installed") or cdata.get("web")

    with open(TOKEN_FILE, "r", encoding="utf-8") as f:
        tdata = json.load(f)

    data = urllib.parse.urlencode({
        "client_id": client_info["client_id"],
        "client_secret": client_info["client_secret"],
        "refresh_token": tdata["refresh_token"],
        "grant_type": "refresh_token"
    }).encode("utf-8")

    req = urllib.request.Request("https://oauth2.googleapis.com/token", data=data, method="POST")
    with urllib.request.urlopen(req) as resp:
        res = json.loads(resp.read().decode("utf-8"))
        return res["access_token"]


def batch_update_spreadsheet(token, requests):
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}:batchUpdate"
    body = json.dumps({"requests": requests}).encode("utf-8")
    req = urllib.request.Request(
        url,
        data=body,
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="POST"
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode("utf-8"))


def clear_range(token, range_name):
    encoded_range = urllib.parse.quote(range_name)
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}/values/{encoded_range}:clear"
    req = urllib.request.Request(
        url,
        data=b"{}",
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="POST"
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode("utf-8"))


def update_values(token, range_name, values):
    encoded_range = urllib.parse.quote(range_name)
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}/values/{encoded_range}?valueInputOption=USER_ENTERED"
    body = json.dumps({"values": values}).encode("utf-8")
    req = urllib.request.Request(
        url,
        data=body,
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="PUT"
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode("utf-8"))


# Complete New Roadmap (21 items + 1 header = 22 rows)
NEW_ROADMAP_DATA = [
    ["Giai Đoạn (Phase)", "Thời Gian", "Cụm Sản Phẩm (Track)", "Tên Tính Năng / Sản Phẩm", "Mô Tả Công Cụ & Sub-Pages Trong Sitemap", "Trạng Thái"],
    # PHASE 0
    ["Phase 0", "Tháng 7/2026", "Tín Dụng & Tiêu Dùng", "Ví Trả Sau (/vi-tra-sau)",
     "Tool tính trả góp Ví Trả Sau. Nhập giá trị món hàng và kỳ hạn (3 - 12 tháng) để xem số tiền thanh toán mỗi tháng, phí chuyển đổi và so sánh phương án miễn lãi 45 ngày. Giữ nguyên URL root độc lập.", "Done"],
    ["Phase 0", "Tháng 7/2026", "Tín Dụng & Cho Vay", "Vay Nhanh (/vay-nhanh)",
     "Tool tính lịch trả nợ vay tiêu dùng. Nhập số tiền vay (5 - 100 triệu) và kỳ hạn (6 - 48 tháng) để tính số tiền trả hàng tháng theo dư nợ giảm dần, chi tiết tiền gốc, tiền lãi và tổng lãi thực tế. Giữ nguyên URL root độc lập.", "Done"],
    # PHASE 1 (Tháng 8/2026)
    ["Phase 1", "Tháng 8/2026", "Master Gateway", "Financial Hub (/tai-chinh)",
     "Cổng trung tâm điều hướng dịch vụ tài chính Kênh Web theo 4 trụ cột (Tín Dụng, Tích Lũy, Đầu Tư, Chi Tiêu). Đã triển khai Mega Menu Header responsive toàn cục.", "Done"],
    ["Phase 1", "Tháng 8/2026", "Sức Khỏe Tín Dụng", "Điểm Tín Dụng (/diem-tin-dung)",
     "Tool tự đánh giá điểm tín dụng CIC (thang điểm 150 - 750), xếp hạng nhóm nợ 1-5 và nhận gợi ý cải thiện hồ sơ vay. Giữ nguyên URL root độc lập.", "Done"],
    ["Phase 1", "Tháng 8/2026", "Đầu Tư & Vàng", "Bảng Giá Vàng (/tai-chinh/gia-vang)",
     "Bảng tra cứu giá vàng trực tuyến (SJC, DOJI, PNJ, nhẫn trơn 9999) theo thời gian thực kèm biểu đồ lịch sử và tool quy đổi chỉ, lượng ra tiền mặt.", "Done"],
    ["Phase 1", "Tháng 8/2026", "Thu Nhập & Lao Động", "Tính Lương Gross - Net (/tai-chinh/tinh-luong)",
     "Tool tính lương Net từ lương Gross (tự động khấu trừ 10.5% bảo hiểm bắt buộc và thuế TNCN lũy tiến 7 bậc 2026 trên cùng một giao diện, kèm tính năng chuyển đổi ngược lại).", "Done"],
    ["Phase 1", "Tháng 8/2026", "Hưu Trí & Lao Động", "Dự Toán Lương Hưu & BHXH (/tai-chinh/luong-huu)",
     "Tool dự toán lương hưu và tuổi nghỉ hưu theo lộ trình luật mới, tích hợp tiện ích tính tiền trợ cấp một lần (tích hợp trực tiếp trong trang Lương Hưu, không tạo URL bảo hiểm độc lập).", "Done"],
    # PHASE 1 (Tháng 9/2026)
    ["Phase 1", "Tháng 9/2026", "Tiền Tệ & Ngoại Hối", "Cổng Tỷ Giá Master (/tai-chinh/ty-gia)",
     "Đã Go-live: Cổng tra cứu tỷ giá thị trường và tool quy đổi ngoại tệ (30+ đồng tiền từ 20+ ngân hàng).", "Done"],
    ["Phase 1", "Tháng 9/2026", "Tiền Tệ & Ngoại Hối", "5 Sub-pages Tỷ Giá Cặp Tiền",
     "Đã Go-live 5 cặp tiền tệ phổ biến nhất: USD/VND, EUR/VND, JPY/VND, GBP/VND, CNY/VND kèm tool đổi tiền trực tiếp, bảng tỷ giá đa ngân hàng và biểu đồ biến động.", "Done"],
    ["Phase 1", "Tháng 9/2026", "Quản Lý Chi Tiêu", "Phân Bổ Lương 50/30/20 (/tai-chinh/phan-bo-luong)",
     "Đã Go-live: Tool lập ngân sách chi tiêu và chia hũ tiền lương theo quy tắc 50/30/20 và mô hình 6 Hũ tài chính, kết nối Hũ Chi Tiêu in-app.", "Done"],
    ["Phase 1", "Tháng 9/2026", "Tiết Kiệm & Lãi Suất", "Bảng Lãi Suất 40 Ngân Hàng (/tai-chinh/lai-suat)",
     "Đã Go-live: Bảng tra cứu lãi suất 40 ngân hàng toàn cảnh thị trường và công cụ tính lãi gửi tiết kiệm theo số tiền, tự động hiển thị ngân hàng có lãi suất tối ưu nhất.", "Done"],
    # PHASE 2 (Tháng 10/2026)
    ["Phase 2", "Tháng 10/2026", "Thu Nhập & Thuế", "Thuế Thu Nhập Cá Nhân (/tai-chinh/thue-thu-nhap-ca-nhan)",
     "Bổ sung mới (P0): Tool tính thuế TNCN, tra cứu mức giảm trừ gia cảnh 2026, giải đáp biểu thuế lũy tiến từng phần; chuẩn hóa URL không viết tắt. Kỹ sư Frontend đang hoàn thiện giao diện trang.", "In Progress"],
    ["Phase 2", "Tháng 10/2026", "Kiến Trúc & Tái Định Tuyến URL", "Di Chuyển & 301 Redirect 7 Trang Standalone",
     "Lập bảng mapping và rule Edge CDN 301 Redirect cho 7 trang standalone về /tai-chinh/* (/nhan-tien-quoc-te, /quy-nhom, /thanh-toan-the-tin-dung, /quan-ly-chi-tieu, /chung-khoan, /chung-chi-quy, /tui-sinh-loi). Cam kết giữ nguyên URL root CreditTech.", "In Progress"],
    ["Phase 2", "Tháng 10/2026", "Tăng Trưởng Nội Dung", "Cỗ Máy Content Engine Always-on",
     "Bổ sung mới (P1): Duy trì nhịp sản xuất và xuất bản 5 bài viết/ngày chuẩn E-E-A-T vào Blog Hub (/tai-chinh/blog/*), xây dựng chân rết thâu tóm từ khóa ngách và nuôi dưỡng traffic tự nhiên.", "In Progress"],
    ["Phase 2", "Tháng 10/2026", "Đầu Tư & Tích Sản", "Chứng Chỉ Quỹ (/tai-chinh/chung-chi-quy)",
     "Đang review API dữ liệu giá NAV và tỷ suất sinh lời của 30 quỹ mở từ BU CVS; xây dựng bảng giá tra cứu và công cụ đầu tư tích sản định kỳ từ 10.000đ/ngày. Tích hợp cơ chế chuyển đổi chéo trực tiếp từ các trang mã cổ phiếu.", "In Progress"],
    # PHASE 3 (Tháng 11/2026)
    ["Phase 3", "Tháng 11/2026", "Đầu Tư & Cổ Phiếu", "Cổng Chứng Khoán Master & 30 Profiles VN30",
     "Xây dựng Cổng Master /tai-chinh/chung-khoan và 30 trang Dynamic Ticker Profile cho rổ VN30 (/tai-chinh/chung-khoan/[ticker]). Tích hợp máy tính Stock DCA Backtester và Thước đo Sức khỏe MoMo Stock Health Score 1-100 phục vụ tích sản cổ phiếu lô lẻ từ 10.000đ. Đang pending phản hồi nghiệp vụ từ Cell Team (FinHub BU / CVS), dời mốc Go-live sang Tháng 11/2026.", "Pending"],
    ["Phase 3", "Tháng 11/2026", "Tiết Kiệm & Lãi Suất", "Gửi Tiết Kiệm Online (/tai-chinh/gui-tiet-kiem)",
     "Hoàn tất thẩm định API lãi suất tiền gửi ưu đãi đối tác (+0.2% Bản Việt, VPBank) và tích hợp cổng mở sổ online, so sánh gói gửi tiết kiệm sinh lời.", "Planned"],
    ["Phase 3", "Tháng 11/2026", "Dòng Tiền & Tiêu Dùng", "Quản Lý Dòng Tiền & Quỹ Khẩn Cấp (/tai-chinh/dong-tien)",
     "Tool tính dòng tiền tự do (Free Cash Flow), quy mô quỹ dự phòng rủi ro cần có (3 - 6 tháng chi phí) và số tháng tài chính cầm cự nếu mất nguồn thu. Định tuyến trang /tai-chinh/dong-tien trong Sitemap.", "Planned"],
    ["Phase 3", "Tháng 11/2026", "Tích Sản & Hưu Trí", "Đầu Tư & Tích Sản FIRE (/tai-chinh/tu-do-tai-chinh)",
     "Tool lập kế hoạch tự do tài chính và nghỉ hưu sớm. Tạo trang độc lập /tai-chinh/tu-do-tai-chinh trong Sitemap theo quy tắc rút 4% an toàn và lộ trình tích sản hàng tháng.", "Planned"],
    ["Phase 3", "Tháng 11/2026", "Ngân Hàng & Đối Tác", "pSEO Cổng Danh Bạ 34 Ngân Hàng (/tai-chinh/ngan-hang/*)",
     "Đòn bẩy tăng trưởng 700K Traffic: Triển khai Programmatic SEO cho Cổng Danh bạ và 34 Sub-page ngân hàng chi tiết (/tai-chinh/ngan-hang/[bank-slug]), đánh chiếm dung lượng tìm kiếm 19.49M volume/tháng từ mạng lưới chi nhánh, hotline và mã ngân hàng.", "Planned"],
    # PHASE 4 (Tháng 12/2026)
    ["Phase 4", "Tháng 12/2026", "Chuyển Đổi Nâng Cao", "Trợ Lý MoMo AI Suggestion Engine",
     "Đòn bẩy tăng trưởng 1M Traffic: Triển khai engine gợi ý thông minh tính toán chênh lệch lợi suất thực (Net Yield Gap) nhúng trực tiếp vào các tiện ích tính toán, tự động gợi ý giải pháp tài chính MoMo tối ưu theo mức tiền người dùng nhập vào.", "Planned"]
]


def generate_brd_section_10_html() -> str:
    html = [
        "## 10. LỘ TRÌNH TRIỂN KHAI TOÀN DIỆN CÁC DỰ ÁN FINANCIAL HUB (THỰC TẾ TRIỂN KHAI)\n",
        '<table style="width:100%; border-collapse:collapse; border:1.5px solid #64748b; margin:1em 0; font-size:0.85em;">',
        "  <thead>",
        '    <tr style="background-color:#1e3a5f; color:#ffffff;">',
        '      <th style="border:1.5px solid #64748b; padding:6px; text-align:center; font-weight:700;">Giai Đoạn (Phase)</th>',
        '      <th style="border:1.5px solid #64748b; padding:6px; text-align:center; font-weight:700;">Thời Gian</th>',
        '      <th style="border:1.5px solid #64748b; padding:6px; text-align:left; font-weight:700;">Cụm Sản Phẩm (Track)</th>',
        '      <th style="border:1.5px solid #64748b; padding:6px; text-align:left; font-weight:700;">Tên Tính Năng / Sản Phẩm</th>',
        '      <th style="border:1.5px solid #64748b; padding:6px; text-align:left; font-weight:700;">Chi Tiết Triển Khai & Nhiệm Vụ Kỹ Thuật</th>',
        '      <th style="border:1.5px solid #64748b; padding:6px; text-align:center; font-weight:700;">Trạng Thái</th>',
        "    </tr>",
        "  </thead>",
        "  <tbody>"
    ]

    for item in NEW_ROADMAP_DATA[1:]:
        phase, time, track, name, desc, status = item

        if status == "Done":
            st_style = "color:#15803d; background-color:#dcfce7; font-weight:700; text-align:center;"
            row_bg = "background-color:#f0fdf4;"
        elif status == "In Progress":
            st_style = "color:#b45309; background-color:#fef9c3; font-weight:700; text-align:center;"
            row_bg = "background-color:#fefce8;"
        elif status == "Pending":
            st_style = "color:#be123c; background-color:#ffe4e6; font-weight:700; text-align:center;"
            row_bg = "background-color:#fff1f2;"
        else:
            st_style = "color:#4338ca; background-color:#e0e7ff; font-weight:700; text-align:center;"
            row_bg = "background-color:#f8fafc;"

        html.append(f'    <tr style="{row_bg}">')
        html.append(f'      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700;"><strong>{phase}</strong></td>')
        html.append(f'      <td style="border:1px solid #94a3b8; padding:6px; text-align:center;">{time}</td>')
        html.append(f'      <td style="border:1px solid #94a3b8; padding:6px; font-weight:600;">{track}</td>')
        html.append(f'      <td style="border:1px solid #94a3b8; padding:6px; font-weight:700;">{name}</td>')
        html.append(f'      <td style="border:1px solid #94a3b8; padding:6px;">{desc}</td>')
        html.append(f'      <td style="border:1px solid #94a3b8; padding:6px; {st_style}">{status}</td>')
        html.append("    </tr>")

    html.append("  </tbody>")
    html.append("</table>\n")
    return "\n".join(html)


def update_brd_file():
    print("1. Updating 05_HUBS/financial/financial-hub-brd.md...")
    content = BRD_PATH.read_text(encoding="utf-8")

    # Replace Section 10 table
    s10_start = content.find("## 10. LỘ TRÌNH")
    if s10_start != -1:
        tbl_end = content.find("</table>", s10_start)
        if tbl_end != -1:
            new_s10 = generate_brd_section_10_html()
            content = content[:s10_start] + new_s10 + content[tbl_end + 8:]

    # Clean SIP across BRD
    content = clean_sip_text(content)

    BRD_PATH.write_text(content, encoding="utf-8")
    rem_sip = BRD_PATH.read_text(encoding="utf-8").count("SIP")
    print(f"BRD updated successfully! Remaining SIP count: {rem_sip}")


def update_prd_files():
    print("2. Updating PRD markdown files...")
    for p in [PRD_CK_1, PRD_CK_2, PRD_AI]:
        if p.exists():
            text = p.read_text(encoding="utf-8")
            # For momo-ai PRD, keep deeplink param mutual_fund_sip / micro_investing_sip intact if needed, or clean prose
            cleaned = clean_sip_text(text)
            p.write_text(cleaned, encoding="utf-8")
            print(f"Updated {p.name}! Remaining SIP count: {p.read_text(encoding='utf-8').count('SIP')}")


def update_local_excel():
    print(f"3. Updating local Excel workbook: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH)

    # Clean SIP across all sheets in workbook
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                val = ws.cell(r, c).value
                if val and isinstance(val, str) and "SIP" in val:
                    ws.cell(r, c, clean_sip_text(val))

    # Rebuild Sheet: Roadmap
    if "Roadmap" in wb.sheetnames:
        del wb["Roadmap"]
    ws_r = wb.create_sheet("Roadmap", index=2)

    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

    font_default = Font(name="Arial", size=9.5, color="1E293B")
    font_bold = Font(name="Arial", size=9.5, bold=True, color="1E293B")

    border_thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    fill_done = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    font_done = Font(name="Arial", size=9.5, bold=True, color="15803D")

    fill_progress = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    font_progress = Font(name="Arial", size=9.5, bold=True, color="B45309")

    fill_pending = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")
    font_pending = Font(name="Arial", size=9.5, bold=True, color="BE123C")

    fill_planned = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
    font_planned = Font(name="Arial", size=9.5, bold=True, color="4338CA")

    for r_idx, row_vals in enumerate(NEW_ROADMAP_DATA, 1):
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_r.cell(r_idx, c_idx, val)
            cell.border = border_thin

            if r_idx == 1:
                cell.font = font_header
                cell.fill = fill_header
                align_h = "center" if c_idx in [1, 2, 6] else "left"
                cell.alignment = Alignment(horizontal=align_h, vertical="center")
            else:
                align_h = "center" if c_idx in [1, 2, 6] else "left"
                wrap = True if c_idx == 5 else False
                cell.alignment = Alignment(horizontal=align_h, vertical="center", wrap_text=wrap)

                if c_idx in [3, 4]:
                    cell.font = font_bold
                elif c_idx == 6:
                    if val == "Done":
                        cell.fill = fill_done
                        cell.font = font_done
                    elif val == "In Progress":
                        cell.fill = fill_progress
                        cell.font = font_progress
                    elif val == "Pending":
                        cell.fill = fill_pending
                        cell.font = font_pending
                    else:
                        cell.fill = fill_planned
                        cell.font = font_planned
                else:
                    cell.font = font_default

        ws_r.row_dimensions[r_idx].height = 24 if r_idx > 1 else 26

    # Column widths
    col_widths = {1: 15, 2: 15, 3: 26, 4: 36, 5: 68, 6: 15}
    for col_idx, width in col_widths.items():
        ws_r.column_dimensions[get_column_letter(col_idx)].width = width

    ws_r.freeze_panes = "A2"

    wb.save(EXCEL_PATH)
    print("Local Excel workbook updated and saved successfully!")


def update_google_sheets():
    print("4. Updating Google Sheets...")
    token = get_access_token()

    url_meta = f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}?fields=sheets(properties)"
    req_meta = urllib.request.Request(url_meta, headers={"Authorization": f"Bearer {token}"})
    with urllib.request.urlopen(req_meta) as resp:
        meta = json.loads(resp.read().decode("utf-8"))

    sheet_map = {s["properties"]["title"]: s["properties"]["sheetId"] for s in meta["sheets"]}

    roadmap_id = sheet_map.get("Product Roadmap")
    markets_id = sheet_map.get("Financial Markets")
    sitemap_id = sheet_map.get("Sitemap")

    print(f"Sheet IDs: Product Roadmap={roadmap_id}, Financial Markets={markets_id}, Sitemap={sitemap_id}")

    # -------------------------------------------------------------
    # A. Update Tab: Product Roadmap
    # -------------------------------------------------------------
    print("Updating Google Sheet Tab: Product Roadmap...")
    clear_range(token, "Product Roadmap!A1:F50")
    num_rows = len(NEW_ROADMAP_DATA)
    update_values(token, f"Product Roadmap!A1:F{num_rows}", NEW_ROADMAP_DATA)

    reqs_r = []
    # Freeze 1 row
    reqs_r.append({
        "updateSheetProperties": {
            "properties": {
                "sheetId": roadmap_id,
                "gridProperties": {
                    "frozenRowCount": 1,
                    "rowCount": max(50, num_rows + 5),
                    "columnCount": 10
                }
            },
            "fields": "gridProperties(frozenRowCount,rowCount,columnCount)"
        }
    })

    # Column widths: Phase (120), Time (120), Track (190), Name (260), Desc (560), Status (110)
    widths = [120, 120, 190, 260, 560, 110]
    for c_idx, w in enumerate(widths):
        reqs_r.append({
            "updateDimensionProperties": {
                "range": {
                    "sheetId": roadmap_id,
                    "dimension": "COLUMNS",
                    "startIndex": c_idx,
                    "endIndex": c_idx + 1
                },
                "properties": {"pixelSize": w},
                "fields": "pixelSize"
            }
        })

    # Header Row 1 (index 0)
    reqs_r.append({
        "repeatCell": {
            "range": {
                "sheetId": roadmap_id,
                "startRowIndex": 0,
                "endRowIndex": 1,
                "startColumnIndex": 0,
                "endColumnIndex": 6
            },
            "cell": {
                "userEnteredFormat": {
                    "backgroundColor": {"red": 0.12, "green": 0.31, "blue": 0.47},
                    "horizontalAlignment": "LEFT",
                    "verticalAlignment": "MIDDLE",
                    "textFormat": {
                        "fontFamily": "Arial",
                        "fontSize": 10,
                        "bold": True,
                        "foregroundColor": {"red": 1.0, "green": 1.0, "blue": 1.0}
                    }
                }
            },
            "fields": "userEnteredFormat(backgroundColor,horizontalAlignment,verticalAlignment,textFormat)"
        }
    })
    for c_idx in [0, 1, 5]:
        reqs_r.append({
            "repeatCell": {
                "range": {
                    "sheetId": roadmap_id,
                    "startRowIndex": 0,
                    "endRowIndex": 1,
                    "startColumnIndex": c_idx,
                    "endColumnIndex": c_idx + 1
                },
                "cell": {"userEnteredFormat": {"horizontalAlignment": "CENTER"}},
                "fields": "userEnteredFormat.horizontalAlignment"
            }
        })

    # Data rows (index 1 to num_rows-1)
    reqs_r.append({
        "repeatCell": {
            "range": {
                "sheetId": roadmap_id,
                "startRowIndex": 1,
                "endRowIndex": num_rows,
                "startColumnIndex": 0,
                "endColumnIndex": 6
            },
            "cell": {
                "userEnteredFormat": {
                    "verticalAlignment": "MIDDLE",
                    "textFormat": {
                        "fontFamily": "Arial",
                        "fontSize": 10,
                        "foregroundColor": {"red": 0.12, "green": 0.16, "blue": 0.23}
                    }
                }
            },
            "fields": "userEnteredFormat(verticalAlignment,textFormat)"
        }
    })

    # Zebra striping
    for r_idx in range(1, num_rows):
        bg = {"red": 0.97, "green": 0.98, "blue": 0.99} if r_idx % 2 == 1 else {"red": 1.0, "green": 1.0, "blue": 1.0}
        reqs_r.append({
            "repeatCell": {
                "range": {
                    "sheetId": roadmap_id,
                    "startRowIndex": r_idx,
                    "endRowIndex": r_idx + 1,
                    "startColumnIndex": 0,
                    "endColumnIndex": 5
                },
                "cell": {"userEnteredFormat": {"backgroundColor": bg}},
                "fields": "userEnteredFormat.backgroundColor"
            }
        })

    # Alignments
    for c_idx in [0, 1]:
        reqs_r.append({
            "repeatCell": {
                "range": {
                    "sheetId": roadmap_id,
                    "startRowIndex": 1,
                    "endRowIndex": num_rows,
                    "startColumnIndex": c_idx,
                    "endColumnIndex": c_idx + 1
                },
                "cell": {"userEnteredFormat": {"horizontalAlignment": "CENTER"}},
                "fields": "userEnteredFormat.horizontalAlignment"
            }
        })
    reqs_r.append({
        "repeatCell": {
            "range": {
                "sheetId": roadmap_id,
                "startRowIndex": 1,
                "endRowIndex": num_rows,
                "startColumnIndex": 3,
                "endColumnIndex": 4
            },
            "cell": {"userEnteredFormat": {"textFormat": {"bold": True}}},
            "fields": "userEnteredFormat.textFormat.bold"
        }
    })
    reqs_r.append({
        "repeatCell": {
            "range": {
                "sheetId": roadmap_id,
                "startRowIndex": 1,
                "endRowIndex": num_rows,
                "startColumnIndex": 4,
                "endColumnIndex": 5
            },
            "cell": {"userEnteredFormat": {"wrapStrategy": "WRAP"}},
            "fields": "userEnteredFormat.wrapStrategy"
        }
    })

    # Status badges styling
    for r_idx in range(1, num_rows):
        st = NEW_ROADMAP_DATA[r_idx][5]
        if st == "Done":
            bg = {"red": 0.86, "green": 0.98, "blue": 0.90}  # Emerald 100
            fg = {"red": 0.08, "green": 0.50, "blue": 0.24}  # Emerald 700
        elif st == "In Progress":
            bg = {"red": 0.99, "green": 0.95, "blue": 0.78}  # Amber 100
            fg = {"red": 0.71, "green": 0.33, "blue": 0.04}  # Amber 700
        elif st == "Pending":
            bg = {"red": 1.0, "green": 0.89, "blue": 0.90}   # Rose 100
            fg = {"red": 0.75, "green": 0.07, "blue": 0.24}  # Rose 700
        else:  # Planned
            bg = {"red": 0.88, "green": 0.91, "blue": 1.0}   # Indigo 100
            fg = {"red": 0.26, "green": 0.22, "blue": 0.79}  # Indigo 700

        reqs_r.append({
            "repeatCell": {
                "range": {
                    "sheetId": roadmap_id,
                    "startRowIndex": r_idx,
                    "endRowIndex": r_idx + 1,
                    "startColumnIndex": 5,
                    "endColumnIndex": 6
                },
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": bg,
                        "horizontalAlignment": "CENTER",
                        "verticalAlignment": "MIDDLE",
                        "textFormat": {
                            "fontFamily": "Arial",
                            "fontSize": 10,
                            "bold": True,
                            "foregroundColor": fg
                        }
                    }
                },
                "fields": "userEnteredFormat(backgroundColor,horizontalAlignment,verticalAlignment,textFormat)"
            }
        })

    # Borders
    border_style = {"style": "SOLID", "width": 1, "color": {"red": 0.85, "green": 0.88, "blue": 0.92}}
    reqs_r.append({
        "updateBorders": {
            "range": {
                "sheetId": roadmap_id,
                "startRowIndex": 0,
                "endRowIndex": num_rows,
                "startColumnIndex": 0,
                "endColumnIndex": 6
            },
            "top": border_style, "bottom": border_style, "left": border_style, "right": border_style,
            "innerHorizontal": border_style, "innerVertical": border_style
        }
    })
    batch_update_spreadsheet(token, reqs_r)
    print("Google Sheet Tab: Product Roadmap updated successfully!")

    # -------------------------------------------------------------
    # B. Update Tab: Financial Markets (Remove SIP)
    # -------------------------------------------------------------
    print("Updating Google Sheet Tab: Financial Markets (removing SIP)...")
    update_values(token, "Financial Markets!B25", [["Chứng Chỉ Quỹ & Quỹ Mở"]])
    update_values(token, "Financial Markets!G25", [["Cổng 30 quỹ mở, bảng giá NAV thời gian thực và máy tính tích sản định kỳ từ 10.000đ/ngày"]])

    # -------------------------------------------------------------
    # C. Update Tab: Sitemap (Remove SIP)
    # -------------------------------------------------------------
    print("Updating Google Sheet Tab: Sitemap (removing SIP)...")
    update_values(token, "Sitemap!C57", [["Máy Tính Giả Lập Tích Sản Định Kỳ Quỹ Mở Từ 10.000đ"]])

    # -------------------------------------------------------------
    # D. Final Safety Check: Global Find & Replace for SIP
    # -------------------------------------------------------------
    reqs_s = [{
        "findReplace": {
            "find": "SIP",
            "replacement": "Chứng Chỉ Quỹ",
            "matchCase": True,
            "allSheets": True
        }
    }]
    batch_update_spreadsheet(token, reqs_s)
    print("Global find/replace for 'SIP' executed on Google Sheets!")


def main():
    update_brd_file()
    update_prd_files()
    update_local_excel()
    update_google_sheets()
    print("\nAll updates completed successfully!")


if __name__ == "__main__":
    main()
