import os
import openpyxl
import xlsxwriter

def create_single_sheet_project_name_excel():
    overview_path = '/Users/hienhv/Downloads/[Web MoMo] Overview Web 2026.xlsx'
    out_dir = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/07_REPORTS/sheets'
    os.makedirs(out_dir, exist_ok=True)
    out_xlsx = os.path.join(out_dir, 'project_name.xlsx')
    old_xlsx = os.path.join(out_dir, 'cms_projects_audit_and_restructuring_report.xlsx')

    # Remove old file if exists
    if os.path.exists(old_xlsx):
        os.remove(old_xlsx)
        print(f"Removed old file: {old_xlsx}")

    wb_ref = openpyxl.load_workbook(overview_path, data_only=True)

    all_items = []
    seen = set()

    # 1. Hubs sheet
    ws_hubs = wb_ref['Hubs']
    for r in range(2, ws_hubs.max_row + 1):
        div = ws_hubs.cell(r, 1).value
        uc = ws_hubs.cell(r, 2).value
        prod = ws_hubs.cell(r, 3).value
        srv = ws_hubs.cell(r, 4).value
        url = ws_hubs.cell(r, 5).value
        pt = ws_hubs.cell(r, 6).value
        if div and prod:
            key = (str(div).strip(), str(uc).strip(), str(prod).strip(), str(srv).strip(), str(url).strip())
            if key not in seen:
                seen.add(key)
                all_items.append({
                    'division': str(div).strip(),
                    'use_case': str(uc).strip(),
                    'product': str(prod).strip(),
                    'service': str(srv).strip(),
                    'url': str(url).strip(),
                    'page_type': str(pt).strip() if pt else 'Master Hub',
                    'note': 'Master Hub Platform'
                })

    # 2. Mini Web sheet
    ws_mini = wb_ref['Mini Web']
    for r in range(2, ws_mini.max_row + 1):
        div = ws_mini.cell(r, 1).value
        uc = ws_mini.cell(r, 2).value
        prod = ws_mini.cell(r, 3).value
        srv = ws_mini.cell(r, 4).value
        url = ws_mini.cell(r, 5).value
        pt = ws_mini.cell(r, 6).value
        note = ws_mini.cell(r, 7).value
        if div and prod:
            key = (str(div).strip(), str(uc).strip(), str(prod).strip(), str(srv).strip(), str(url).strip())
            if key not in seen:
                seen.add(key)
                all_items.append({
                    'division': str(div).strip(),
                    'use_case': str(uc).strip() if uc else '',
                    'product': str(prod).strip(),
                    'service': str(srv).strip() if srv else '',
                    'url': str(url).strip(),
                    'page_type': str(pt).strip() if pt else 'Advanced MiniWeb',
                    'note': str(note).strip() if note else ''
                })

    # Sort items by Division and Use Case
    div_order = {'GPD': 1, 'MDS': 2, 'FS': 3, 'PS': 4, 'BMC': 5, 'CX': 6}
    all_items.sort(key=lambda x: (div_order.get(x['division'], 99), x['use_case'], x['product']))

    # Create Workbook with ONLY 1 SHEET
    wb = xlsxwriter.Workbook(out_xlsx)
    ws = wb.add_worksheet('Project_Name')

    # Styles
    f_title = wb.add_format({'bold': True, 'font_size': 13, 'font_color': '#FFFFFF', 'bg_color': '#A50064', 'align': 'left', 'valign': 'vcenter', 'font_name': 'Arial'})
    f_subtitle = wb.add_format({'italic': True, 'font_size': 9, 'font_color': '#475569', 'font_name': 'Arial'})
    f_tbl_header = wb.add_format({'bold': True, 'font_size': 9, 'font_color': '#1E293B', 'bg_color': '#E2E8F0', 'align': 'center', 'valign': 'vcenter', 'border': 1, 'font_name': 'Arial'})
    f_cell_left = wb.add_format({'font_size': 9, 'font_color': '#0F172A', 'align': 'left', 'valign': 'vcenter', 'border': 1, 'font_name': 'Arial'})
    f_cell_center = wb.add_format({'font_size': 9, 'font_color': '#0F172A', 'align': 'center', 'valign': 'vcenter', 'border': 1, 'font_name': 'Arial'})
    f_cell_url = wb.add_format({'font_size': 9, 'font_color': '#0284C7', 'align': 'left', 'valign': 'vcenter', 'border': 1, 'font_name': 'Arial'})
    
    # Zone styles
    f_zone_perf = wb.add_format({'font_size': 9, 'font_color': '#166534', 'bg_color': '#DCFCE7', 'bold': True, 'align': 'center', 'valign': 'vcenter', 'border': 1, 'font_name': 'Arial'})
    f_zone_trans = wb.add_format({'font_size': 9, 'font_color': '#1E40AF', 'bg_color': '#DBEAFE', 'bold': True, 'align': 'center', 'valign': 'vcenter', 'border': 1, 'font_name': 'Arial'})
    f_zone_incub = wb.add_format({'font_size': 9, 'font_color': '#854D0E', 'bg_color': '#FEF9C3', 'bold': True, 'align': 'center', 'valign': 'vcenter', 'border': 1, 'font_name': 'Arial'})
    f_zone_prod = wb.add_format({'font_size': 9, 'font_color': '#475569', 'bg_color': '#F1F5F9', 'align': 'center', 'valign': 'vcenter', 'border': 1, 'font_name': 'Arial'})

    # Column widths
    ws.set_column('A:A', 6)   # STT
    ws.set_column('B:B', 12)  # Division
    ws.set_column('C:C', 22)  # Use Case
    ws.set_column('D:D', 32)  # Project Name
    ws.set_column('E:E', 32)  # Service Name
    ws.set_column('F:F', 45)  # URL
    ws.set_column('G:G', 20)  # Page Type
    ws.set_column('H:H', 24)  # 4-Zone
    ws.set_column('I:I', 30)  # Note

    # Title & Metadata
    ws.merge_range('A1:I1', 'DANH MỤC PROJECT THEO USE CASE WEB MOMO NĂM 2026 (SSOT)', f_title)
    ws.write('A2', 'Tài liệu chuẩn hóa tên dự án theo từng Use Case & Sản phẩm của MoMo | Nguồn SSOT: [Web MoMo] Overview Web 2026.xlsx', f_subtitle)

    # Headers
    headers = [
        'STT', 'Division', 'Use Case', 'Project Name (Tên Chuẩn)',
        'Service Name (Dịch Vụ)', 'URL', 'Loại Trang (Page Type)',
        'Phân Vùng 4-Zone', 'Ghi Chú Kỹ Thuật (Note)'
    ]
    for col_idx, h in enumerate(headers):
        ws.write(3, col_idx, h, f_tbl_header)

    for row_idx, item in enumerate(all_items):
        r_num = 4 + row_idx
        ws.write(r_num, 0, row_idx + 1, f_cell_center)
        ws.write(r_num, 1, item['division'], f_cell_center)
        ws.write(r_num, 2, item['use_case'], f_cell_left)
        ws.write(r_num, 3, item['product'], f_cell_left)
        ws.write(r_num, 4, item['service'], f_cell_left)
        ws.write(r_num, 5, item['url'], f_cell_url)
        ws.write(r_num, 6, item['page_type'], f_cell_center)

        # Determine Zone
        div = item['division']
        uc = item['use_case']
        prod = item['product']
        
        if div == 'MDS' and uc in ['Cinema', 'Flight', 'Bus', 'Hourly Hour', 'OTA']:
            zone = 'Performance Zone'
            z_fmt = f_zone_perf
        elif div == 'FS' or prod in ['PS - Phạt Nguội', 'UTI - Vehicle']:
            zone = 'Transformation Zone'
            z_fmt = f_zone_trans
        elif prod in ['GPD - Sinh Viên', 'GPD - New User', 'UTI - Lottery', 'PS - Lottery']:
            zone = 'Incubator Zone'
            z_fmt = f_zone_incub
        else:
            zone = 'Productivity & Self-serve'
            z_fmt = f_zone_prod

        ws.write(r_num, 7, zone, z_fmt)
        ws.write(r_num, 8, item['note'], f_cell_left)

    ws.autofilter(3, 0, 3 + len(all_items), len(headers) - 1)
    ws.freeze_panes(4, 4)

    wb.close()
    print(f"Successfully generated clean 1-sheet workbook at: {out_xlsx}")

if __name__ == '__main__':
    create_single_sheet_project_name_excel()
