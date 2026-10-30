#!/usr/bin/env python3
"""
Update Financial Markets in both local Excel and Google Sheets.
1. Reads all 107 CSV inventory files in 05_HUBS/inventory/Financial/
2. Computes deduplicated keywords & search volume for 25 granular markets
3. Updates local Excel 05_HUBS/financial/financial-hub-roadmap.xlsx (Sheet: Markets)
4. Updates Google Sheet [WP] Financial Hub | Web Product (Sheet: Financial Markets)
"""

import csv
import json
import urllib.parse
import urllib.request
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent
INV_DIR = BASE_DIR / "05_HUBS" / "inventory" / "Financial"
EXCEL_PATH = BASE_DIR / "05_HUBS" / "financial" / "financial-hub-roadmap.xlsx"
CREDS_FILE = BASE_DIR / "07_REPORTS" / "data" / "credentials" / "ga4-credentials.json"
TOKEN_FILE = BASE_DIR / "07_REPORTS" / "data" / "credentials" / "gdrive-token.json"
SPREADSHEET_ID = "1_AG4otSTXH7ontpPL0zjylYgVfxOozvoSc5PziUPM6A"


def clean_val(v):
    if v is None:
        return ""
    v = str(v).strip()
    if v.startswith('="') and v.endswith('"'):
        v = v[2:-1]
    return v.strip()


def get_access_token():
    with open(CREDS_FILE, "r", encoding="utf-8") as f:
        cdata = json.load(f)
    client_info = cdata.get("installed") or cdata.get("web")

    with open(TOKEN_FILE, "r", encoding="utf-8") as f:
        tdata = json.load(f)

    data = urllib.parse.urlencode({
        "client_id": client_info["client_id"],
        "client_secret": client_info["client_secret"],
        "refresh_token": tdata["refresh_token"],
        "grant_type": "refresh_token"
    }).encode("utf-8")

    req = urllib.request.Request("https://oauth2.googleapis.com/token", data=data)
    with urllib.request.urlopen(req) as resp:
        token_data = json.loads(resp.read().decode("utf-8"))
    return token_data["access_token"]


def batch_update_spreadsheet(token, requests):
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}:batchUpdate"
    body = json.dumps({"requests": requests}).encode("utf-8")
    req = urllib.request.Request(
        url,
        data=body,
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="POST"
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode("utf-8"))


def load_inventory_data():
    files_stats = {}
    csv_files = sorted(INV_DIR.glob("*.csv"))
    print(f"Reading {len(csv_files)} inventory CSV files...")

    for f in csv_files:
        kws = {}
        with open(f, mode="r", encoding="utf-8-sig", errors="ignore") as fp:
            reader = csv.reader(fp)
            header = next(reader, None)
            if not header:
                continue
            vol_idx, kw_idx = 1, 0
            for i, h in enumerate(header):
                cleaned_h = clean_val(h).lower()
                if "volume" in cleaned_h or "search volume" in cleaned_h:
                    vol_idx = i
                if "keyword" in cleaned_h:
                    kw_idx = i
            for row in reader:
                if not row or len(row) <= max(kw_idx, vol_idx):
                    continue
                kw = clean_val(row[kw_idx]).lower()
                v_str = clean_val(row[vol_idx]).replace(",", "").replace(".", "")
                try:
                    v = int(v_str) if v_str else 0
                except ValueError:
                    v = 0
                kws[kw] = max(kws.get(kw, 0), v)
        files_stats[f.stem.strip()] = kws

    return files_stats


def compute_market_stats(files_stats):
    bank_files = [
        "ncb", "viet-bank", "Techcombank", "vietcombank", "vcb", "nam-a-bank", "vpbank",
        "mbbank", "acb", "vrb", "BIDV", "agribank", "tpbank", "MSB", "viettinbank",
        "sacombank", "shb", "ab-bank", "ocb", "Shinhan Bank", "bac", "ban-viet",
        "eximbank", "vib", "hd-bank", "pvcom bank", "donga bank", "seabank",
        "bac-a-bank", "LPBank", "LPBank ", "bao-viet-bank", "uob", "scb", "kien-long-bank",
        "Indovina Bank", "ocean-bank", "saigonbank"
    ]

    GRANULAR_MARKETS = [
        ("Giá Vàng", "Đầu Tư & Tích Sản", ["gia-vang", "vang"],
         "Bảng giá vàng SJC, nhẫn trơn 9999, DOJI, PNJ và công cụ quy đổi số lượng chỉ sang tiền mặt"),
        ("Hệ Thống 34+ Ngân Hàng", "Ngân Hàng & Đối Tác", bank_files,
         "Danh bạ mã ngân hàng, hotline CSKH, mạng lưới chi nhánh và biểu phí 34 ngân hàng Việt Nam"),
        ("Ngân Hàng Số & Chuyển Tiền", "Ngân Hàng & Số Hóa", ["ngan-hang", "internet-banking", "chuyen-khoan", "thanh-toan-qr", "vi-dien-tu"],
         "Cổng thông tin ngân hàng số, Internet Banking, chuyển khoản nhanh 24/7 và mã thanh toán QR"),
        ("Tỷ Giá Ngoại Tệ", "Tiền Tệ & Ngoại Hối", ["ty-gia", "doi-tien"],
         "Bảng tỷ giá trung tâm và 30+ ngoại tệ ngân hàng, công cụ quy đổi ngoại tệ trực tuyến"),
        ("Mã Số Thuế", "Thu Nhập & Thuế", ["ma-so-thue"],
         "Tra cứu mã số thuế cá nhân, mã số thuế doanh nghiệp và người phụ thuộc giảm trừ gia cảnh"),
        ("Chứng Khoán & Mã Cổ Phiếu", "Đầu Tư & Tích Sản", ["chung-khoan", "co-phieu", "ma-chung-khoan", "co-tuc", "chia-co-tuc", "phai-sinh", "ky-quy"],
         "Cổng Cổ phiếu VN30, tra cứu mã chứng khoán, tỷ lệ cổ tức, lịch chia cổ tức, giao dịch phái sinh và ký quỹ"),
        ("Cho Vay & Vay Tín Chấp", "Tín Dụng & Cho Vay", ["vay", "lai-vay", "vay-tin-chap", "tin-dung", "thau-chi", "du-no"],
         "Công cụ tính lịch trả nợ vay tiêu dùng tín chấp theo dư nợ giảm dần, chi tiết gốc và lãi"),
        ("Thuế TNCN & Giảm Trừ", "Thu Nhập & Thuế", ["thue", "thue-tncn", "giam-tru-gia-canh"],
         "Bảng biểu thuế thu nhập cá nhân lũy tiến 7 bậc 2026, mức giảm trừ gia cảnh và quyết toán thuế"),
        ("Lãi Suất Thị Trường", "Tiết Kiệm & Lãi Suất", ["lai-suat"],
         "Bảng tra cứu lãi suất 40 ngân hàng toàn cảnh thị trường và lãi suất điều hành Ngân hàng Nhà nước"),
        ("Ngoại Tệ & Kiều Hối", "Tiền Tệ & Ngoại Hối", ["ngoai-te", "kieu-hoi"],
         "Thị trường tiền tệ quốc tế, dòng kiều hối và thủ tục nhận tiền quốc tế Western Union/MoneyGram"),
        ("Thẻ Tín Dụng", "Ngân Hàng & Thẻ", ["the-tin-dung"],
         "Mở thẻ tín dụng online, biểu phí thường niên, hạn mức chi tiêu và chính sách ưu đãi hoàn tiền"),
        ("Tiền Số & Crypto", "Tài Sản Số & Xu Hướng", ["crypto", "tien-so", "tien-ao", "tien-dien-tu"],
         "Kiến thức thị trường tiền kỹ thuật số, công nghệ chuỗi khối Blockchain và nguyên tắc an toàn"),
        ("Đầu Tư & Quản Lý Tài Sản", "Đầu Tư & Tích Sản", ["dau-tu", "tai-san", "tieu-san", "von-luu-dong"],
         "Mô hình phân bổ danh mục đầu tư, quản trị tài sản cá nhân và quản lý vòng quay vốn lưu động"),
        ("Thẻ Visa & Thẻ Quốc Tế", "Ngân Hàng & Thẻ", ["the-visa", "master-card"],
         "Thẻ thanh toán quốc tế Visa, Mastercard, phí quẹt thẻ ngoại tệ và công nghệ bảo mật 3D Secure"),
        ("Gửi Tiết Kiệm & Lãi Kép", "Tiết Kiệm & Lãi Suất", ["gui-tiet-kiem", "tiet-kiem", "lai-kep", "chung-chi-tien-gui", "tich-luy"],
         "Cổng gửi tiết kiệm online đối tác MoMo, máy tính lãi gửi có kỳ hạn và lãi kép sinh lời dài hạn"),
        ("Tính Lương Gross - Net", "Thu Nhập & Lao Động", ["tinh-luong", "thu-nhap"],
         "Tool tính lương Net từ lương Gross (tự động khấu trừ 10.5% bảo hiểm bắt buộc và thuế TNCN 2026)"),
        ("Vay Thế Chấp (Mua Nhà)", "Tín Dụng & Cho Vay", ["vay-the-chap"],
         "Công cụ tính toán gói vay thế chấp mua nhà, mua ô tô ngân hàng kèm biên độ lãi suất sau ưu đãi"),
        ("Điểm Tín Dụng & CIC", "Tín Dụng & Sức Khỏe Tín Dụng", ["diem-tin-dung", "no"],
         "Tool tự đánh giá điểm tín dụng CIC (150 - 750), phân loại nhóm nợ và hướng dẫn tối ưu hồ sơ vay"),
        ("Trả Góp & Ví Trả Sau", "Tín Dụng & Tiêu Dùng", ["tra-gop", "vi-tra-sau", "mua-truoc-tra-sau"],
         "Công cụ tính phí chuyển đổi trả góp 0%, kỳ hạn thanh toán Ví Trả Sau và giải pháp Mua Trước Trả Sau"),
        ("Thẻ Ghi Nợ ATM & Napas", "Ngân Hàng & Thẻ", ["the-ghi-no", "napas", "the-tra-truoc"],
         "Thẻ ghi nợ nội địa ATM, cổng thanh toán quốc gia Napas và giải pháp thẻ trả trước liên kết"),
        ("Kinh Tế Vĩ Mô (GDP, Lạm Phát)", "Kinh Tế Vĩ Mô", ["gdp", "lam-phat", "giam-phat"],
         "Báo cáo chỉ số tăng trưởng GDP, tỷ lệ lạm phát, CPI và các chỉ số kinh tế vĩ mô định kỳ"),
        ("Trái Phiếu Doanh Nghiệp & CP", "Đầu Tư & Tích Sản", ["trai-phieu"],
         "Cổng tra cứu thị trường trái phiếu doanh nghiệp, trái phiếu chính phủ và tỷ suất lợi nhuận đáo hạn"),
        ("Nợ Xấu & Xóa Nợ", "Tín Dụng & Sức Khỏe Tín Dụng", ["no-xau"],
         "Quy chuẩn phân loại 5 nhóm nợ CIC, thời gian lưu vết lịch sử tín dụng và quy trình xóa nợ xấu"),
        ("Chứng Chỉ Quỹ & Quỹ Mở SIP", "Đầu Tư & Tích Sản", ["chung-chi-quy", "quy-etf", "quy-mo", "uy-thac"],
         "Cổng 30 quỹ mở, bảng giá NAV thời gian thực và máy tính tích sản định kỳ SIP từ 10.000đ/ngày"),
        ("Quản Lý Dòng Tiền & FIRE", "Quản Lý Tài Chính Cá Nhân", ["dong-tien", "ke-hoach-tai-chinh", "tu-do-tai-chinh", "quy-du-phong", "quy-huu-tri"],
         "Kế hoạch tự do tài chính FIRE (quy tắc 4%), thiết lập quỹ dự phòng khẩn cấp và lộ trình hưu trí")
    ]

    all_hub_kws = {}
    market_rows = []

    for name, pillar, flist, desc in GRANULAR_MARKETS:
        m_kws = {}
        for f in flist:
            k_dict = files_stats.get(f.strip(), {})
            for k, v in k_dict.items():
                m_kws[k] = max(m_kws.get(k, 0), v)
                all_hub_kws[k] = max(all_hub_kws.get(k, 0), v)
        m_vol = sum(m_kws.values())
        market_rows.append({
            "name": name,
            "pillar": pillar,
            "kw_count": len(m_kws),
            "volume": m_vol,
            "desc": desc
        })

    total_hub_vol = sum(all_hub_kws.values())
    total_hub_kw = len(all_hub_kws)

    # Sort descending by volume
    market_rows.sort(key=lambda x: x["volume"], reverse=True)

    for r in market_rows:
        r["share"] = (r["volume"] / total_hub_vol * 100) if total_hub_vol > 0 else 0

    return market_rows, total_hub_kw, total_hub_vol


def update_local_excel(market_rows, total_hub_kw, total_hub_vol):
    print("Updating local Excel: " + str(EXCEL_PATH))
    wb = openpyxl.load_workbook(EXCEL_PATH)

    if "Markets" in wb.sheetnames:
        del wb["Markets"]
    ws = wb.create_sheet("Markets", index=8)

    # Styling definitions
    font_title = Font(name="Arial", size=11, bold=True, color="1F4E78")
    fill_title = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

    font_data = Font(name="Arial", size=10, color="1E293B")
    font_bold = Font(name="Arial", size=10, bold=True, color="1E293B")
    font_total = Font(name="Arial", size=10, bold=True, color="1F4E78")
    fill_total = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    fill_even = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    border_thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    border_total = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="medium", color="1F4E78"),
        bottom=Side(style="double", color="1F4E78")
    )

    # Row 1: Merged Title
    ws.merge_cells("A1:G1")
    ws["A1"] = "TỔNG HỢP TOÀN BỘ CÁC THỊ TRƯỜNG DỊCH VỤ TÀI CHÍNH FINANCIAL HUB (DEDUPLICATED SSOT)"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_title
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Row 2: Headers
    headers = [
        "STT",
        "Thị Trường / Dịch Vụ",
        "Trụ Cột Nghiệp Vụ",
        "Số Lượng Từ Khóa",
        "Search Volume / Tháng",
        "Tỷ Trọng (%)",
        "Mô Tả Tiện Ích & Định Hướng Web Product"
    ]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(2, col_idx, h)
        cell.font = font_header
        cell.fill = fill_header
        align_h = "center" if col_idx in [1, 4, 5, 6] else "left"
        cell.alignment = Alignment(horizontal=align_h, vertical="center")
        cell.border = border_thin
    ws.row_dimensions[2].height = 26

    # Rows 3 to 27: Data rows
    for r_idx, r in enumerate(market_rows, 3):
        fill_curr = fill_even if r_idx % 2 == 0 else fill_odd

        # Col A: STT
        c1 = ws.cell(r_idx, 1, r_idx - 2)
        c1.alignment = Alignment(horizontal="center", vertical="center")

        # Col B: Tên thị trường
        c2 = ws.cell(r_idx, 2, r["name"])
        c2.font = font_bold
        c2.alignment = Alignment(horizontal="left", vertical="center")

        # Col C: Trụ cột
        c3 = ws.cell(r_idx, 3, r["pillar"])
        c3.alignment = Alignment(horizontal="left", vertical="center")

        # Col D: Keywords
        c4 = ws.cell(r_idx, 4, r["kw_count"])
        c4.number_format = "#,##0"
        c4.alignment = Alignment(horizontal="center", vertical="center")

        # Col E: Search Volume
        c5 = ws.cell(r_idx, 5, r["volume"])
        c5.font = font_bold
        c5.number_format = "#,##0"
        c5.alignment = Alignment(horizontal="right", vertical="center")

        # Col F: Tỷ trọng
        c6 = ws.cell(r_idx, 6, f"{r['share']:.2f}%")
        c6.alignment = Alignment(horizontal="center", vertical="center")

        # Col G: Mô tả
        c7 = ws.cell(r_idx, 7, r["desc"])
        c7.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        for col_idx in range(1, 8):
            cell = ws.cell(r_idx, col_idx)
            cell.border = border_thin
            cell.fill = fill_curr
            if col_idx not in [2, 5]:
                cell.font = font_data
        ws.row_dimensions[r_idx].height = 22

    # Row 28: Total Row
    tot_row = len(market_rows) + 3
    c1 = ws.cell(tot_row, 1, "TỔNG CỘNG FINANCIAL MASTER HUB (DEDUPLICATED)")
    ws.merge_cells(start_row=tot_row, start_column=1, end_row=tot_row, end_column=3)
    c1.alignment = Alignment(horizontal="center", vertical="center")

    c4 = ws.cell(tot_row, 4, total_hub_kw)
    c4.number_format = "#,##0"
    c4.alignment = Alignment(horizontal="center", vertical="center")

    c5 = ws.cell(tot_row, 5, total_hub_vol)
    c5.number_format = "#,##0"
    c5.alignment = Alignment(horizontal="right", vertical="center")

    c6 = ws.cell(tot_row, 6, "100.00%")
    c6.alignment = Alignment(horizontal="center", vertical="center")

    c7 = ws.cell(tot_row, 7, "Tổng hợp 107 tệp Inventory phân tích; toàn bộ từ khóa và Volume đã được khử trùng lặp chuẩn xác")
    c7.alignment = Alignment(horizontal="left", vertical="center")

    for col_idx in range(1, 8):
        cell = ws.cell(tot_row, col_idx)
        cell.font = font_total
        cell.fill = fill_total
        cell.border = border_total
    ws.row_dimensions[tot_row].height = 26

    # Column widths
    col_widths = {1: 8, 2: 32, 3: 26, 4: 18, 5: 24, 6: 15, 7: 68}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze top 2 rows
    ws.freeze_panes = "A3"

    wb.save(EXCEL_PATH)
    print("Local Excel updated successfully!")


def update_google_sheets(market_rows, total_hub_kw, total_hub_vol):
    print("Updating Google Sheets...")
    token = get_access_token()

    # 1. Get sheet ID for Financial Markets
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}?fields=sheets(properties)"
    req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
    with urllib.request.urlopen(req) as resp:
        meta = json.loads(resp.read().decode("utf-8"))

    markets_sheet_id = None
    for s in meta["sheets"]:
        p = s["properties"]
        if p["title"] == "Financial Markets":
            markets_sheet_id = p["sheetId"]
            break

    if markets_sheet_id is None:
        raise ValueError("Sheet 'Financial Markets' not found in spreadsheet.")

    print(f"Target Sheet: Financial Markets (ID: {markets_sheet_id})")

    # 2. Build rows data
    sheet_data = [
        ["TỔNG HỢP TOÀN BỘ CÁC THỊ TRƯỜNG DỊCH VỤ TÀI CHÍNH FINANCIAL HUB (DEDUPLICATED SSOT)", "", "", "", "", "", ""],
        ["STT", "Thị Trường / Dịch Vụ", "Trụ Cột Nghiệp Vụ", "Số Lượng Từ Khóa", "Search Volume / Tháng", "Tỷ Trọng (%)", "Mô Tả Tiện Ích & Định Hướng Web Product"]
    ]

    for idx, r in enumerate(market_rows, 1):
        sheet_data.append([
            str(idx),
            r["name"],
            r["pillar"],
            f"{r['kw_count']:,}",
            f"{r['volume']:,}",
            f"{r['share']:.2f}%",
            r["desc"]
        ])

    tot_row_data = [
        "TỔNG CỘNG FINANCIAL MASTER HUB (DEDUPLICATED)",
        "",
        "",
        f"{total_hub_kw:,}",
        f"{total_hub_vol:,}",
        "100.00%",
        "Tổng hợp 107 tệp Inventory phân tích; toàn bộ từ khóa và Volume đã được khử trùng lặp chuẩn xác"
    ]
    sheet_data.append(tot_row_data)

    # 3. Clear existing values and unmerge old ranges first
    clear_url = f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}/values/Financial%20Markets!A1:G60:clear"
    req = urllib.request.Request(
        clear_url,
        data=b"{}",
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
    )
    with urllib.request.urlopen(req) as resp:
        print("Cleared existing values on Financial Markets.")

    # Unmerge any previous merges before writing to prevent cell swallowing
    batch_update_spreadsheet(token, [{
        "unmergeCells": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": 0,
                "endRowIndex": len(sheet_data) + 10,
                "startColumnIndex": 0,
                "endColumnIndex": 7
            }
        }
    }])

    # 4. Write new values
    write_url = f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}/values/Financial%20Markets!A1:G{len(sheet_data)}?valueInputOption=USER_ENTERED"
    body = json.dumps({"values": sheet_data}).encode("utf-8")
    req = urllib.request.Request(
        write_url,
        data=body,
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="PUT"
    )
    with urllib.request.urlopen(req) as resp:
        res = json.loads(resp.read().decode("utf-8"))
        print(f"Financial Markets values written: {res.get('updatedCells')} cells updated across {len(sheet_data)} rows.")

    # 5. Apply visual formatting via batchUpdate
    print("Applying styling to Financial Markets sheet...")
    requests = []

    # Freeze top 2 rows
    requests.append({
        "updateSheetProperties": {
            "properties": {
                "sheetId": markets_sheet_id,
                "gridProperties": {
                    "frozenRowCount": 2
                }
            },
            "fields": "gridProperties.frozenRowCount"
        }
    })

    # Column widths: Col A: 65, Col B: 260, Col C: 210, Col D: 150, Col E: 180, Col F: 110, Col G: 520
    col_widths = [65, 260, 210, 150, 180, 110, 520]
    for c_idx, width in enumerate(col_widths):
        requests.append({
            "updateDimensionProperties": {
                "range": {
                    "sheetId": markets_sheet_id,
                    "dimension": "COLUMNS",
                    "startIndex": c_idx,
                    "endIndex": c_idx + 1
                },
                "properties": {
                    "pixelSize": width
                },
                "fields": "pixelSize"
            }
        })

    # Row 1: Merged Title A1:G1
    requests.append({
        "mergeCells": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": 0,
                "endRowIndex": 1,
                "startColumnIndex": 0,
                "endColumnIndex": 7
            },
            "mergeType": "MERGE_ALL"
        }
    })
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": 0,
                "endRowIndex": 1,
                "startColumnIndex": 0,
                "endColumnIndex": 7
            },
            "cell": {
                "userEnteredFormat": {
                    "backgroundColor": {"red": 0.85, "green": 0.88, "blue": 0.95},
                    "horizontalAlignment": "CENTER",
                    "verticalAlignment": "MIDDLE",
                    "textFormat": {
                        "fontFamily": "Arial",
                        "fontSize": 11,
                        "bold": True,
                        "foregroundColor": {"red": 0.12, "green": 0.31, "blue": 0.47}
                    }
                }
            },
            "fields": "userEnteredFormat(backgroundColor,horizontalAlignment,verticalAlignment,textFormat)"
        }
    })

    # Row 2: Header Row A2:G2
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": 1,
                "endRowIndex": 2,
                "startColumnIndex": 0,
                "endColumnIndex": 7
            },
            "cell": {
                "userEnteredFormat": {
                    "backgroundColor": {"red": 0.12, "green": 0.31, "blue": 0.47},
                    "horizontalAlignment": "LEFT",
                    "verticalAlignment": "MIDDLE",
                    "textFormat": {
                        "fontFamily": "Arial",
                        "fontSize": 10,
                        "bold": True,
                        "foregroundColor": {"red": 1.0, "green": 1.0, "blue": 1.0}
                    }
                }
            },
            "fields": "userEnteredFormat(backgroundColor,horizontalAlignment,verticalAlignment,textFormat)"
        }
    })
    for c_idx in [0, 3, 4, 5]:
        requests.append({
            "repeatCell": {
                "range": {
                    "sheetId": markets_sheet_id,
                    "startRowIndex": 1,
                    "endRowIndex": 2,
                    "startColumnIndex": c_idx,
                    "endColumnIndex": c_idx + 1
                },
                "cell": {
                    "userEnteredFormat": {
                        "horizontalAlignment": "CENTER"
                    }
                },
                "fields": "userEnteredFormat.horizontalAlignment"
            }
        })

    # Rows 3 to 27: Data rows
    total_data_rows = len(market_rows)
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": 2,
                "endRowIndex": 2 + total_data_rows,
                "startColumnIndex": 0,
                "endColumnIndex": 7
            },
            "cell": {
                "userEnteredFormat": {
                    "verticalAlignment": "MIDDLE",
                    "textFormat": {
                        "fontFamily": "Arial",
                        "fontSize": 10,
                        "foregroundColor": {"red": 0.12, "green": 0.16, "blue": 0.23}
                    }
                }
            },
            "fields": "userEnteredFormat(verticalAlignment,textFormat)"
        }
    })

    # Zebra striping for data rows
    for r_idx in range(2, 2 + total_data_rows):
        bg = {"red": 0.97, "green": 0.98, "blue": 0.99} if r_idx % 2 == 1 else {"red": 1.0, "green": 1.0, "blue": 1.0}
        requests.append({
            "repeatCell": {
                "range": {
                    "sheetId": markets_sheet_id,
                    "startRowIndex": r_idx,
                    "endRowIndex": r_idx + 1,
                    "startColumnIndex": 0,
                    "endColumnIndex": 7
                },
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": bg
                    }
                },
                "fields": "userEnteredFormat.backgroundColor"
            }
        })

    # Specific column alignments & styling for data rows
    # Col A: Center
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": 2,
                "endRowIndex": 2 + total_data_rows,
                "startColumnIndex": 0,
                "endColumnIndex": 1
            },
            "cell": {
                "userEnteredFormat": {
                    "horizontalAlignment": "CENTER"
                }
            },
            "fields": "userEnteredFormat.horizontalAlignment"
        }
    })
    # Col B: Bold
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": 2,
                "endRowIndex": 2 + total_data_rows,
                "startColumnIndex": 1,
                "endColumnIndex": 2
            },
            "cell": {
                "userEnteredFormat": {
                    "textFormat": {
                        "bold": True
                    }
                }
            },
            "fields": "userEnteredFormat.textFormat.bold"
        }
    })
    # Col D: Center
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": 2,
                "endRowIndex": 2 + total_data_rows,
                "startColumnIndex": 3,
                "endColumnIndex": 4
            },
            "cell": {
                "userEnteredFormat": {
                    "horizontalAlignment": "CENTER"
                }
            },
            "fields": "userEnteredFormat.horizontalAlignment"
        }
    })
    # Col E: Right & Bold
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": 2,
                "endRowIndex": 2 + total_data_rows,
                "startColumnIndex": 4,
                "endColumnIndex": 5
            },
            "cell": {
                "userEnteredFormat": {
                    "horizontalAlignment": "RIGHT",
                    "textFormat": {
                        "bold": True
                    }
                }
            },
            "fields": "userEnteredFormat(horizontalAlignment,textFormat.bold)"
        }
    })
    # Col F: Center
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": 2,
                "endRowIndex": 2 + total_data_rows,
                "startColumnIndex": 5,
                "endColumnIndex": 6
            },
            "cell": {
                "userEnteredFormat": {
                    "horizontalAlignment": "CENTER"
                }
            },
            "fields": "userEnteredFormat.horizontalAlignment"
        }
    })
    # Col G: Wrap
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": 2,
                "endRowIndex": 2 + total_data_rows,
                "startColumnIndex": 6,
                "endColumnIndex": 7
            },
            "cell": {
                "userEnteredFormat": {
                    "wrapStrategy": "WRAP"
                }
            },
            "fields": "userEnteredFormat.wrapStrategy"
        }
    })

    # Row 28: Total Row styling
    tot_row_idx = 2 + total_data_rows
    # Merge A-C for Total Label
    requests.append({
        "mergeCells": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": tot_row_idx,
                "endRowIndex": tot_row_idx + 1,
                "startColumnIndex": 0,
                "endColumnIndex": 3
            },
            "mergeType": "MERGE_ALL"
        }
    })
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": tot_row_idx,
                "endRowIndex": tot_row_idx + 1,
                "startColumnIndex": 0,
                "endColumnIndex": 7
            },
            "cell": {
                "userEnteredFormat": {
                    "backgroundColor": {"red": 0.85, "green": 0.88, "blue": 0.95},
                    "verticalAlignment": "MIDDLE",
                    "textFormat": {
                        "fontFamily": "Arial",
                        "fontSize": 10,
                        "bold": True,
                        "foregroundColor": {"red": 0.12, "green": 0.31, "blue": 0.47}
                    }
                }
            },
            "fields": "userEnteredFormat(backgroundColor,verticalAlignment,textFormat)"
        }
    })
    # Center Total Label
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": tot_row_idx,
                "endRowIndex": tot_row_idx + 1,
                "startColumnIndex": 0,
                "endColumnIndex": 3
            },
            "cell": {
                "userEnteredFormat": {
                    "horizontalAlignment": "CENTER"
                }
            },
            "fields": "userEnteredFormat.horizontalAlignment"
        }
    })
    # Center Total Keywords
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": tot_row_idx,
                "endRowIndex": tot_row_idx + 1,
                "startColumnIndex": 3,
                "endColumnIndex": 4
            },
            "cell": {
                "userEnteredFormat": {
                    "horizontalAlignment": "CENTER"
                }
            },
            "fields": "userEnteredFormat.horizontalAlignment"
        }
    })
    # Right align Total Volume
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": tot_row_idx,
                "endRowIndex": tot_row_idx + 1,
                "startColumnIndex": 4,
                "endColumnIndex": 5
            },
            "cell": {
                "userEnteredFormat": {
                    "horizontalAlignment": "RIGHT"
                }
            },
            "fields": "userEnteredFormat.horizontalAlignment"
        }
    })
    # Center Total %
    requests.append({
        "repeatCell": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": tot_row_idx,
                "endRowIndex": tot_row_idx + 1,
                "startColumnIndex": 5,
                "endColumnIndex": 6
            },
            "cell": {
                "userEnteredFormat": {
                    "horizontalAlignment": "CENTER"
                }
            },
            "fields": "userEnteredFormat.horizontalAlignment"
        }
    })

    # Borders for whole table
    border_style = {
        "style": "SOLID",
        "width": 1,
        "color": {"red": 0.85, "green": 0.88, "blue": 0.92}
    }
    requests.append({
        "updateBorders": {
            "range": {
                "sheetId": markets_sheet_id,
                "startRowIndex": 1,
                "endRowIndex": tot_row_idx + 1,
                "startColumnIndex": 0,
                "endColumnIndex": 7
            },
            "top": border_style,
            "bottom": border_style,
            "left": border_style,
            "right": border_style,
            "innerHorizontal": border_style,
            "innerVertical": border_style
        }
    })

    # Execute batchUpdate
    url_batch = f"https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}:batchUpdate"
    body_batch = json.dumps({"requests": requests}).encode("utf-8")
    req = urllib.request.Request(
        url_batch,
        data=body_batch,
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        method="POST"
    )
    with urllib.request.urlopen(req) as resp:
        batch_res = json.loads(resp.read().decode("utf-8"))
        print(f"batchUpdate successful! Total replies: {len(batch_res.get('replies', []))}")


def main():
    files_stats = load_inventory_data()
    market_rows, total_hub_kw, total_hub_vol = compute_market_stats(files_stats)

    print(f"\nCalculation complete:")
    print(f"- Markets count: {len(market_rows)}")
    print(f"- Total Unique Keywords: {total_hub_kw:,}")
    print(f"- Total Unique Search Volume: {total_hub_vol:,}\n")

    update_local_excel(market_rows, total_hub_kw, total_hub_vol)
    update_google_sheets(market_rows, total_hub_kw, total_hub_vol)
    print("\nAll synchronizations finished successfully!")


if __name__ == "__main__":
    main()
