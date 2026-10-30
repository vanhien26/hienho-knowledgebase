#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GOOGLE DRIVE / SHEETS SSOT CONNECTOR: [WP] Web Cinema | Roadmap
Kết nối trực tiếp Cloud-to-Local và Google Drive Desktop:
- Spreadsheet ID: 1YT2sdU-O4smoy6TUKBic6PTFBxUXNUVkd_S-tmjiplg
- Google Sheets URL: https://docs.google.com/spreadsheets/d/1YT2sdU-O4smoy6TUKBic6PTFBxUXNUVkd_S-tmjiplg/edit?gid=113654251#gid=113654251
- Google Drive Desktop Path: /Users/hienhv/Library/CloudStorage/GoogleDrive-hien.ho@mservice.com.vn/My Drive/WEB GROWTH 2026/#1 - Use Case/[MDS] - Movies/[WP] Web Cinema | Roadmap.gsheet
- Local SSOT File: 05_HUBS/cinema/cinema_hub_roadmap_drive_sync.xlsx
"""

import os
import sys
import json
import argparse
import urllib.request
import urllib.parse
from pathlib import Path
import openpyxl

BASE_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE_DIR / "scripts"))
from sync_to_google_sheet_api import get_sheets_token

TARGET_SPREADSHEET_ID = "1YT2sdU-O4smoy6TUKBic6PTFBxUXNUVkd_S-tmjiplg"
GDRIVE_DESKTOP_PATH = Path("/Users/hienhv/Library/CloudStorage/GoogleDrive-hien.ho@mservice.com.vn/My Drive/WEB GROWTH 2026/#1 - Use Case/[MDS] - Movies/[WP] Web Cinema | Roadmap.gsheet")
LOCAL_EXCEL_SSOT = BASE_DIR / "05_HUBS" / "cinema" / "cinema_hub_roadmap_drive_sync.xlsx"
LOCAL_DATA_DIR = BASE_DIR / "05_HUBS" / "cinema" / "data"

def get_spreadsheet_meta(token=None):
    if not token:
        token = get_sheets_token()
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{TARGET_SPREADSHEET_ID}"
    req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode())

def pull_sheet_values(sheet_title, token=None):
    if not token:
        token = get_sheets_token()
    encoded_title = urllib.parse.quote(sheet_title)
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{TARGET_SPREADSHEET_ID}/values/{encoded_title}"
    req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"})
    with urllib.request.urlopen(req) as resp:
        res = json.loads(resp.read().decode())
        return res.get("values", [])

def update_sheet_values(range_name, values, token=None):
    """Cập nhật dữ liệu vào Google Sheet qua Sheets API."""
    if not token:
        token = get_sheets_token()
    encoded_range = urllib.parse.quote(range_name)
    url = f"https://sheets.googleapis.com/v4/spreadsheets/{TARGET_SPREADSHEET_ID}/values/{encoded_range}?valueInputOption=USER_ENTERED"
    payload = json.dumps({"range": range_name, "values": values}).encode("utf-8")
    req = urllib.request.Request(url, data=payload, headers={
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/json"
    }, method="PUT")
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode())

def pull_all_to_local():
    """Tải toàn bộ các sheet về file Excel và JSON local."""
    token = get_sheets_token()
    meta = get_spreadsheet_meta(token)
    title = meta.get("properties", {}).get("title", "[WP] Web Cinema | Roadmap")
    sheets = [s["properties"]["title"] for s in meta.get("sheets", [])]

    print(f"Bắt đầu đồng bộ từ Google Sheets: '{title}' (ID: {TARGET_SPREADSHEET_ID})")
    LOCAL_DATA_DIR.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Xóa sheet rỗng mặc định

    manifest = {
        "spreadsheet_id": TARGET_SPREADSHEET_ID,
        "spreadsheet_title": title,
        "gdrive_desktop_path": str(GDRIVE_DESKTOP_PATH) if GDRIVE_DESKTOP_PATH.exists() else None,
        "sheets": {}
    }

    for s_name in sheets:
        print(f"  - Đang kéo dữ liệu sheet: {s_name}...")
        rows = pull_sheet_values(s_name, token)
        manifest["sheets"][s_name] = {
            "row_count": len(rows),
            "col_count": max((len(r) for r in rows), default=0)
        }
        
        # Ghi vào Excel
        ws = wb.create_sheet(title=s_name[:31]) # Excel sheet title max 31 chars
        for r_idx, row in enumerate(rows, 1):
            for c_idx, val in enumerate(row, 1):
                ws.cell(row=r_idx, column=c_idx, value=val)

        # Lưu JSON cho các sheet quan trọng
        if s_name in ["Roadmap", "WP - Minigame", "Web Performances", "Opportunity & Priority"]:
            safe_name = s_name.lower().replace(" ", "_").replace("-", "_").replace("!", "")
            json_file = LOCAL_DATA_DIR / f"{safe_name}.json"
            with open(json_file, "w", encoding="utf-8") as f:
                json.dump(rows, f, ensure_ascii=False, indent=2)

    wb.save(LOCAL_EXCEL_SSOT)
    manifest_file = LOCAL_DATA_DIR / "cinema_roadmap_manifest.json"
    with open(manifest_file, "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2)

    print(f"\n[THÀNH CÔNG] Đã đồng bộ toàn bộ {len(sheets)} sheet:")
    for s_name, info in manifest["sheets"].items():
        print(f"  * {s_name}: {info['row_count']} dòng x {info['col_count']} cột")
    print(f"File Excel SSOT cục bộ: {LOCAL_EXCEL_SSOT}")
    if GDRIVE_DESKTOP_PATH.exists():
        print(f"Google Drive Desktop: Đã liên kết ({GDRIVE_DESKTOP_PATH})")
    return manifest

def main():
    parser = argparse.ArgumentParser(description="Đồng bộ Google Drive / Sheets [WP] Web Cinema | Roadmap")
    parser.add_argument("--pull", action="store_true", default=True, help="Kéo toàn bộ dữ liệu từ cloud về máy")
    parser.add_argument("--status", action="store_true", help="Kiểm tra trạng thái kết nối Google Drive Desktop và Cloud API")
    args = parser.parse_args()

    if args.status:
        token = get_sheets_token()
        meta = get_spreadsheet_meta(token)
        print("Trạng thái kết nối Google Sheets API: HOẠT ĐỘNG (200 OK)")
        print(f"Tiêu đề: {meta.get('properties', {}).get('title')}")
        print(f"ID: {TARGET_SPREADSHEET_ID}")
        print(f"Google Drive Desktop local: {'TỒN TẠI' if GDRIVE_DESKTOP_PATH.exists() else 'CHƯA TÌM THẤY'}")
        if GDRIVE_DESKTOP_PATH.exists():
            print(f"Đường dẫn: {GDRIVE_DESKTOP_PATH}")
        return

    pull_all_to_local()

if __name__ == "__main__":
    main()
