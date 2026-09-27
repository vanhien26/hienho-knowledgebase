"""Shared helpers for repo automation scripts (single source for Excel styles + safe save).

Usage:
    from scripts._lib import load_workbook_safe, save_with_backup, EXCEL_STYLES, repo_root
    wb = load_workbook_safe("05_HUBS/financial-hub-roadmap.xlsx")
    ... edit ...
    save_with_backup(wb, "05_HUBS/financial-hub-roadmap.xlsx")  # auto .bak, use --dry-run to skip write
"""
from __future__ import annotations

import argparse
import datetime
import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def repo_root() -> Path:
    return Path(__file__).resolve().parents[1]


def resolve_path(p: str | Path) -> Path:
    path = Path(p)
    if path.is_absolute():
        return path
    return repo_root() / path


def load_workbook_safe(path: str | Path):
    full = resolve_path(path)
    if not full.exists():
        raise FileNotFoundError(f"Workbook not found: {full}")
    return load_workbook(str(full))


def backup_path(path: str | Path) -> Path:
    full = resolve_path(path)
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    return full.with_suffix(f".bak-{stamp}{full.suffix}")


def save_with_backup(wb, path: str | Path, dry_run: bool = False, make_backup: bool = True) -> Path | None:
    full = resolve_path(path)
    if dry_run:
        print(f"[dry-run] skip save: {full}")
        return None
    if make_backup and full.exists():
        bak = backup_path(full)
        shutil.copy2(full, bak)
        print(f"backup: {bak.name}")
    wb.save(str(full))
    print(f"saved: {full}")
    return full


def standard_argparser(desc: str) -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(description=desc)
    ap.add_argument("--dry-run", action="store_true", help="Parse and validate without writing files")
    ap.add_argument("--no-backup", action="store_true", help="Skip .bak backup before save")
    return ap


_THIN = Side(style="thin", color="64748B")

EXCEL_STYLES = {
    "header_fill": PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid"),
    "header_font": Font(name="Arial", size=10, bold=True, color="FFFFFF"),
    "section_fill": PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"),
    "section_font": Font(name="Arial", size=10.5, bold=True, color="1F4E78"),
    "font_bold": Font(name="Arial", size=9.5, bold=True),
    "font_regular": Font(name="Arial", size=9.5),
    "font_url": Font(name="Arial", size=9.5, color="0563C1", underline="single"),
    "thin_border": Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN),
    "align_left": Alignment(horizontal="left", vertical="center", wrap_text=True),
    "align_center": Alignment(horizontal="center", vertical="center", wrap_text=True),
}
