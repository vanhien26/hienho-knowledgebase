import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

wb = openpyxl.load_workbook('05_HUBS/financial-hub-roadmap.xlsx')
ws_rd = wb['Roadmap']

status_done_fill = PatternFill(start_color="1E7E34", end_color="1E7E34", fill_type="solid")
status_done_font = Font(name='Arial', size=9.5, bold=True, color='FFFFFF')

for r in range(3, ws_rd.max_row + 1):
    feat_name = str(ws_rd.cell(row=r, column=4).value or '').strip()
    if 'Tính Lương' in feat_name:
        ws_rd.cell(row=r, column=2).value = "Tháng 8/2026"
        ws_rd.cell(row=r, column=4).value = "Tính Lương"
        ws_rd.cell(row=r, column=5).value = (
            "• Công cụ tính lương Gross sang Net đã hoàn thiện và đang hoạt động thực tế trên Kênh Web (Pure Client-side JS, không phụ thuộc API ngoài).\n"
            "• Đã nhúng cơ chế gợi ý thông minh (Smart Suggestion) sau khi tính lương: Gợi ý trích 20% lương gửi Tiết Kiệm TKO | Trích 10% tham gia Thực Tập Sinh Đầu Tư | Kiểm tra Điểm Tín Dụng CIC."
        )
        ws_rd.cell(row=r, column=6).value = (
            "Bộ 3 nút CTA trực tiếp sau khi tính lương: 'Mở Sổ Tiết Kiệm Bản Việt / VPBank' | 'Tích Sản Quỹ Mở SIP Chỉ Từ 10k' | 'Kiểm Tra Điểm Tín Dụng CIC'."
        )
        c_st = ws_rd.cell(row=r, column=7)
        c_st.value = "Done"
        c_st.fill = status_done_fill
        c_st.font = status_done_font
        ws_rd.row_dimensions[r].height = 55.0

wb.save('05_HUBS/financial-hub-roadmap.xlsx')
print("Successfully marked Tính Lương as Done (Tháng 8/2026) in Roadmap!")
