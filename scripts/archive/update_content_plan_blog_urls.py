import openpyxl

wb = openpyxl.load_workbook('05_HUBS/financial-hub-roadmap.xlsx')
ws_cp = wb['Content Plan']

updated_count = 0
for r in range(3, ws_cp.max_row + 1):
    url = str(ws_cp.cell(row=r, column=7).value or '').strip()
    fmt = str(ws_cp.cell(row=r, column=6).value or '').strip()
    
    # Check if this is a blog article under /diem-tin-dung
    if url.startswith('/diem-tin-dung/') and not url.startswith('/diem-tin-dung/blog/'):
        slug = url.replace('/diem-tin-dung/', '')
        new_url = f"/diem-tin-dung/blog/{slug}"
        ws_cp.cell(row=r, column=7).value = new_url
        updated_count += 1
        print(f"Row {r}: {url} -> {new_url}")
    
    # Check if this is a blog article under /tai-chinh (exclude programmatic subpages like ty-gia/usd-vnd or ngan-hang)
    elif url.startswith('/tai-chinh/') and not url.startswith('/tai-chinh/blog/') and not '/ty-gia/' in url and url != '/tai-chinh/ngan-hang':
        slug = url.replace('/tai-chinh/', '')
        new_url = f"/tai-chinh/blog/{slug}"
        ws_cp.cell(row=r, column=7).value = new_url
        updated_count += 1
        print(f"Row {r}: {url} -> {new_url}")

wb.save('05_HUBS/financial-hub-roadmap.xlsx')
print(f"Successfully updated {updated_count} blog URLs to include '/blog/' segment in financial-hub-roadmap.xlsx!")
