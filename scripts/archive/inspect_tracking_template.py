import openpyxl

template_path = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/07_REPORTS/sheets/web-event-tracking-template.xlsx'
wb = openpyxl.load_workbook(template_path, data_only=True)
print("Sheet names:", wb.sheetnames)

for s in wb.sheetnames:
    ws = wb[s]
    print(f"\n=== Sheet: {s} (rows: {ws.max_row}, cols: {ws.max_column}) ===")
    for r in range(1, min(15, ws.max_row + 1)):
        row_vals = [ws.cell(r, c).value for c in range(1, min(12, ws.max_column + 1))]
        if any(row_vals):
            print(f"Row {r}:", row_vals)

