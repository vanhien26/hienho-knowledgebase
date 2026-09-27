import openpyxl

wb = openpyxl.load_workbook('05_HUBS/financial-hub-roadmap.xlsx')
ws_rd = wb['Roadmap']

# Define exact Driven-Traffic & Suggestion Mechanisms for all features
updates = {
    'Financial Hub': (
        "• Master Gateway quy hoạch 3 Hero Product Banners nổi bật nhất tại Hero Zone: Điểm Tín Dụng CIC, Gửi Tiết Kiệm TKO, Thực Tập Sinh Đầu Tư.\n"
        "• Thanh Smart Financial Recommendation tự động đề xuất 1 trong 3 sản phẩm theo hành vi người dùng.\n"
        "• Module Dashboard hiển thị song song: Lãi suất Tiết kiệm cao nhất vs Tỷ suất lợi nhuận Quỹ Mở.",
        "Universal OneLink dẫn trực diện vào 3 phễu chính: Kiểm Tra CIC Miễn Phí / Mở Sổ Tiết Kiệm Bản Việt & VPBank / Mở Tài Khoản Thực Tập Sinh Đầu Tư."
    ),
    'CIC & Tiết Kiệm': (
        "• Widget CIC Simulator đánh giá nhóm nợ 1-5 ➔ Gợi ý mở Sổ Tiết Kiệm để làm đẹp hồ sơ tín dụng.\n"
        "• Hero Widget Tính Lãi Tiết Kiệm TKO (Bản Việt / VPBank) ➔ Tích hợp tab so sánh 'Gửi Tiết Kiệm vs Thực Tập Sinh Đầu Tư'.",
        "• CIC: Nút 'Kiểm Tra Điểm Tín Dụng MoMo Miễn Phí' mở báo cáo in-app.\n• Tiết Kiệm: Nút 'Mở Sổ Tiết Kiệm Online' nhận thêm +0.2% lãi suất."
    ),
    'Blog & Cẩm Nang Tài Chính': (
        "• 100% bài viết cẩm nang nhúng Sticky Recommendation Bar ở chân trang dẫn về 1 trong 3 sản phẩm (CIC / Tiết Kiệm / Thực Tập Sinh Đầu Tư) theo ngữ cảnh chủ đề bài viết.\n"
        "• Cơ chế tự động inject In-text Banner & Calculator Widget liên quan.",
        "Dynamic In-text CTA & Sticky Bottom Dock dẫn thẳng vào luồng đăng ký 3 sản phẩm ưu tiên."
    ),
    'Giá Vàng': (
        "• Bảng giá vàng realtime đa thương hiệu SJC, PNJ, DOJI, 9999 kèm biểu đồ lịch sử.\n"
        "• Widget So Sánh Hiệu Suất Tích Sản: Đặt cạnh nhau 3 kênh: 'Vàng vs Gửi Tiết Kiệm vs Thực Tập Sinh Đầu Tư' ➔ Điều hướng người dùng sang Tiết Kiệm và Quỹ Mở.",
        "CTA 'Đa Dạng Hóa Tích Sản: Mở Sổ Tiết Kiệm 6.5%/năm Hoặc Đầu Tư Quỹ Mở Thực Tập Sinh Chỉ Từ 10k'."
    ),
    'Tính Lương & Thuế TNCN': (
        "• Công cụ tính lương Gross sang Net ➔ Sau khi ra số tiền Net thực nhận, tự động hiển thị Hộp Đề Xuất 3 Bước:\n"
        "  1. Trích 20% gửi Tiết Kiệm TKO sinh lời an toàn.\n"
        "  2. Trích 10% tham gia Thực Tập Sinh Đầu Tư tích sản dài hạn.\n"
        "  3. Kiểm tra Điểm Tín Dụng CIC dựa trên mức thu nhập.",
        "Bộ 3 Nút Chuyển Đổi Trực Tiếp: 'Gửi Tiết Kiệm Ngay' | 'Bắt Đầu Tích Sản SIP 10k' | 'Kiểm Tra CIC Miễn Phí'."
    ),
    'Phân Bổ Lương': (
        "• Công cụ phân bổ 50/30/20 tự động mapping:\n"
        "  • Hũ Tích Lũy 20% ➔ Suggest trực diện mở Sổ Tiết Kiệm Bản Việt / VPBank.\n"
        "  • Hũ Đầu Tư 10% ➔ Suggest tham gia Thực Tập Sinh Đầu Tư Quỹ Mở.\n"
        "  • Đo lường chỉ số an toàn đòn bẩy ➔ Suggest kiểm tra CIC.",
        "CTA 1-Click: 'Tự Động Trích Tiền Sang Tiết Kiệm TKO' và 'Kích Hoạt Gói Thực Tập Sinh Đầu Tư'."
    ),
    'Tỷ Giá & Ngoại Tệ': (
        "• Cổng tỷ giá master và programmatic subpages theo cặp tiền (/tai-chinh/ty-gia/[pair-slug]).\n"
        "• Widget Phân Tích Lợi Suất: So sánh 'Giữ Ngoại Tệ vs Gửi Tiết Kiệm VND Lãi Cao' ➔ Kéo traffic về sản phẩm Tiết Kiệm Online.\n"
        "• Đề xuất kênh đầu tư đón sóng tỷ giá ➔ Dẫn về Thực Tập Sinh Đầu Tư.",
        "CTA 'Tối Ưu Dòng Tiền: Đổi Sang VND Gửi Tiết Kiệm Lãi Suất Đến 6.5%' hoặc 'Đầu Tư Đón Sóng Quốc Tế'."
    ),
    'Chứng Khoán & Quỹ Mở': (
        "• Cổng thông tin chính thức của Chương Trình 'Thực Tập Sinh Đầu Tư' (Chứng khoán CVX & Quỹ mở SIP Dragon Capital, VinaCapital, SSIAM).\n"
        "• Bộ giả lập Lãi Kép Tích Lũy Định Kỳ từ 10k/ngày.\n"
        "• Tab so sánh cân bằng danh mục: Kết hợp Thực Tập Sinh Đầu Tư với Gửi Tiết Kiệm.",
        "CTA Trọng Tâm: 'Đăng Ký Tham Gia Thực Tập Sinh Đầu Tư - Nhận Quà Khởi Nghiệp' & 'Mở Sổ Tiết Kiệm Cân Bằng'."
    ),
    'Quản Lý Dòng Tiền & Tiêu Dùng': (
        "• Bộ tính Dòng Tiền Tự Do (Free Cash Flow) & Hạn mức tiêu dùng ➔ Tự động phân luồng dòng tiền dư thừa:\n"
        "  • Tiền dự phòng ➔ Đề xuất gửi Tiết Kiệm Online kỳ hạn linh hoạt.\n"
        "  • Tiền nhàn rỗi dài hạn ➔ Đề xuất gói Thực Tập Sinh Đầu Tư SIP.\n"
        "  • Cảnh báo nợ quá hạn ➔ Đề xuất kiểm tra CIC.",
        "CTA 'Phân Bổ Dòng Tiền Thặng Dư: Gửi Tiết Kiệm Online / Tích Sản SIP Đầu Tư / Kiểm Tra Sức Khỏe CIC'."
    ),
    'Đầu Tư & Tích Sản FIRE': (
        "• Bộ Giả Lập Tự Do Tài Chính & Nghỉ Hưu Sớm (FIRE Simulator).\n"
        "• Cơ chế phân bổ vốn mục tiêu theo nguyên tắc 2 Trụ Cột: Vốn Phòng Thủ (Gửi Tiết Kiệm 30 Bank) + Vốn Tăng Trưởng (Thực Tập Sinh Đầu Tư Quỹ Mở).\n"
        "• Đo lường điểm sức khỏe tín dụng CIC để tối ưu chi phí sử dụng vốn.",
        "CTA 'Xây Dựng Danh Mục FIRE: 50% Gửi Tiết Kiệm An Toàn + 50% Thực Tập Sinh Đầu Tư Tăng Trưởng'."
    )
}

for r in range(3, ws_rd.max_row + 1):
    feat_name = str(ws_rd.cell(row=r, column=4).value or '').strip()
    if feat_name in updates:
        scope_text, hook_text = updates[feat_name]
        ws_rd.cell(row=r, column=5).value = scope_text
        ws_rd.cell(row=r, column=6).value = hook_text
        num_lines = max(scope_text.count('\n') + 1, hook_text.count('\n') + 1)
        ws_rd.row_dimensions[r].height = max(60.0, num_lines * 19.0)

wb.save('05_HUBS/financial-hub-roadmap.xlsx')
print("Successfully embedded Driven-Traffic & Suggestion Matrix to CIC, Tiết Kiệm, and Thực Tập Sinh Đầu Tư across all Roadmap features!")
