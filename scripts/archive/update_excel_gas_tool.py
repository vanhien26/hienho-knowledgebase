import openpyxl

file_path = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/vehicle-hub-roadmap.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb['Utilities & Formula Specs']

for row in range(2, ws.max_row + 1):
    val = ws.cell(row=row, column=2).value
    if val and ('Đầy Bình' in str(val) or 'Giá Xăng' in str(val)):
        ws.cell(row=row, column=2, value='Bộ Công Cụ Tính Toán Giá Xăng Đa Năng (3 Chế Độ)')
        ws.cell(row=row, column=3, value='Interactive Multi-tab Calculator (Lít / Loại Xe / Lộ Trình)')
        ws.cell(row=row, column=4, value='Dự toán chính xác chi phí nhiên liệu theo số lít, theo dòng xe thực tế (đầy bình) hoặc theo lộ trình di chuyển (Km).')
        ws.cell(row=row, column=5, value='• Chế độ 1: Số lít cần đổ (L)\n• Chế độ 2: Loại xe (Dung tích V_tank) & % xăng còn lại\n• Chế độ 3: Quãng đường di chuyển (Km) & Định mức tiêu hao (L/100km)\n• Chọn loại nhiên liệu (RON 95, E5, DO)')
        ws.cell(row=row, column=6, value='1. Theo Lít: Cost = Lít * Giá_hiện_tại.\n2. Theo Loại Xe: Lít_cần = V_tank * (1 - %_còn_lại) -> Cost = Lít_cần * Giá_hiện_tại.\n3. Theo Lộ Trình: Lít_tiêu_thụ = (Km / 100) * FC -> Cost_trip = Lít_tiêu_thụ * Giá_hiện_tại.\n4. Chênh lệch: Delta = Lít * (Giá_hiện_tại - Giá_kỳ_trước).')
        ws.cell(row=row, column=7, value='• Tổng chi phí nhiên liệu (VND).\n• Lượng xăng tiêu thụ (Lít).\n• Mức chênh lệch tăng/giảm so với kỳ điều hành trước.\n• Chi phí nhiên liệu bình quân/km.')
        ws.cell(row=row, column=8, value="Nút 'Tìm Cây Xăng Gần Nhất' & 'Thu Thập Voucher Giảm 20K Đổ Xăng MoMo'.")
        break

wb.save(file_path)
print("Updated Excel Tool 4 successfully.")
