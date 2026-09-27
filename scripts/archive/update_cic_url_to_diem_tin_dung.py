import openpyxl

wb = openpyxl.load_workbook('05_HUBS/financial-hub-roadmap.xlsx')

count = 0
for s in wb.sheetnames:
    ws = wb[s]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(row=r, column=c).value or '')
            if '/tra-cuu-cic' in val:
                new_val = val.replace('/tra-cuu-cic', '/diem-tin-dung')
                ws.cell(row=r, column=c).value = new_val
                count += 1
                print(f"Updated {s} Row {r} Col {c}: {new_val}")

wb.save('05_HUBS/financial-hub-roadmap.xlsx')
print(f"Successfully updated {count} URL paths to '/diem-tin-dung' in financial-hub-roadmap.xlsx!")
