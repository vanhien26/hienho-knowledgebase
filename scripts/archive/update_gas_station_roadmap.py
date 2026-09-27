import openpyxl

file_path = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/vehicle-hub-roadmap.xlsx'
wb = openpyxl.load_workbook(file_path)

# 1. Update Roadmap Sheet
ws_road = wb['Roadmap']
for r in range(3, ws_road.max_row + 1):
    val = str(ws_road.cell(r, 4).value or '')
    if 'Cây Xăng' in val:
        new_scope = (
            "• Trang chính (/cay-xang): Tập trung vào tính năng 'Cây xăng gần đây' & 'Cây xăng mở cửa 24/24'.\n"
            "• Xây dựng 10 trang con trọng điểm có Volume lớn nhất (Local SEO & pSEO):\n"
            "  - 4 Trang Tỉnh/Thành: TP.HCM, Hà Nội, Đà Nẵng, Bình Dương.\n"
            "  - 6 Trang Quận/Đường lớn: Q.1, Q.7, TP.Thủ Đức (TP.HCM); Cầu Giấy, Đống Đa (Hà Nội); Tuyến Quốc Lộ 1A.\n"
            "(Search Volume toàn cụm: 550.000 search/tháng)"
        )
        ws_road.cell(r, 5, value=new_scope)
        break

# 2. Update Content Structure Sheet
if 'Content Structure' in wb.sheetnames:
    ws_cs = wb['Content Structure']
    for r in range(2, ws_cs.max_row + 1):
        page_val = str(ws_cs.cell(r, 2).value or '')
        if 'cay-xang' in page_val or 'Cây Xăng' in page_val:
            ws_cs.cell(r, 2, value="Bản Đồ Cây Xăng Gần Đây\n(/cay-xang & 10 trang con)")
            cs_structure = (
                "1. Hero Section: H1 ('Cây Xăng Gần Đây Mở Cửa 24/24') + Nút GPS 'Tìm Cây Xăng Gần Tôi Nhất'\n"
                "2. Horizontal Filter Bar: Lọc Thương hiệu (Petrolimex, PVOil) | Trạng thái (Đang mở, 24/7) | MoMo QR\n"
                "3. Dual-View Component: Bản đồ GPS tương tác (Pin có logo trạm) + Danh sách trạm xếp theo khoảng cách\n"
                "4. 1-Tap Action: Nút 'Chỉ đường' mở Google Maps + Nút 'Quét mã MoMo' mở app\n"
                "5. Programmatic Cluster: Cụm 10 trang con có Vol to (TP.HCM, Hà Nội, Đà Nẵng, Q.1, Cầu Giấy, QL1A...)\n"
                "6. Voucher & Loyalty Hook: Banner nhận voucher 20k đổ xăng MoMo\n"
                "7. FAQ & Long Content SEO Local"
            )
            ws_cs.cell(r, 3, value=cs_structure)
            break

wb.save(file_path)
print("Updated vehicle-hub-roadmap.xlsx successfully!")
