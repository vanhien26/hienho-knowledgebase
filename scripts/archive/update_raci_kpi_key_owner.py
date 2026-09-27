import openpyxl
from openpyxl.styles import Font

wb = openpyxl.load_workbook('05_HUBS/financial-hub-roadmap.xlsx')
ws = wb['RACI & Resource Allocation']

font_role_key_owner = Font(name='Arial', size=9.5, bold=True, color='B91C1C') # Red bold
font_role_co_owner = Font(name='Arial', size=9.5, bold=True, color='15803D')  # Green bold

# Find row 1.1 and update Web Platform to Co-Owner (A/R) and FinHub BU to Key Owner (A/R)
for r in range(4, ws.max_row + 1):
    val1 = str(ws.cell(row=r, column=1).value or '')
    if '1.1 Thiết Lập Mục Tiêu KPI' in val1:
        ws.cell(row=r, column=2).value = "Co-Owner (A/R)"
        ws.cell(row=r, column=2).font = font_role_co_owner
        
        ws.cell(row=r, column=3).value = "Key Owner (A/R)"
        ws.cell(row=r, column=3).font = font_role_key_owner
        
        ws.cell(row=r, column=5).value = (
            "FinHub BU (Cell Team) giữ vai trò Key Owner chịu trách nhiệm chính về mục tiêu tăng trưởng kinh doanh tổng thể; "
            "Web Platform đóng vai trò Co-Owner cùng cam kết và tối ưu hóa các chỉ số phễu: Lượng truy cập Web (1M MPV), "
            "Người dùng tương tác tiện ích (500K MEU), Tỷ lệ nhấp Web-to-App (15-25%) và Giao dịch tài chính đầu tiên (+25% MoM)."
        )

wb.save('05_HUBS/financial-hub-roadmap.xlsx')
print("Successfully updated Row 1.1: FinHub BU (Cell Team) is Key Owner (A/R) and Web Platform is Co-Owner (A/R)!")
