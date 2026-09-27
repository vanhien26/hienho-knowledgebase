import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load financial-hub-roadmap.xlsx
wb = openpyxl.load_workbook('05_HUBS/financial-hub-roadmap.xlsx')

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
section_font = Font(name='Arial', size=10.5, bold=True, color='1F4E78')
thin_side = Side(style='thin', color='64748B')
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

p1_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid") # Green
p2_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Yellow
p3_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Orange

status_done_fill = PatternFill(start_color="1E7E34", end_color="1E7E34", fill_type="solid") # Dark Green
status_done_font = Font(name='Arial', size=9.5, bold=True, color='FFFFFF')

status_prog_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Orange/Yellow
status_prog_font = Font(name='Arial', size=9.5, bold=True, color='B9770E')

status_plan_fill = PatternFill(start_color="1B4F72", end_color="1B4F72", fill_type="solid") # Dark Blue
status_plan_font = Font(name='Arial', size=9.5, bold=True, color='FFFFFF')

# Exact Real Execution Roadmap matching the User's Real Master Sheet
real_roadmap_data = [
    # Phase, Time, Proj_Name, Spec_Desc, Status, Route, Target_Vol
    ("Phase 1", "Tháng 8/2026", "Dự Án 1: Master Financial Hub",
     "Xây dựng trang cổng trung tâm, thiết kế Mega Menu 5 nhóm nhu cầu tài chính, thanh Trợ thủ đồng hành và các thẻ điều hướng sản phẩm.",
     "Done", "momo.vn/tai-chinh", "1,000,000+ PV/tháng"),

    ("Phase 1", "Tháng 8/2026", "Dự Án 2: CIC, Tiết Kiệm",
     "Công cụ CIC Simulator và Nợ Xấu (tra cứu 339k volume). Công cụ tính lãi suất tiết kiệm cho TKO (Bản Việt, VPBank).",
     "Done", "momo.vn/tra-cuu-cic & /tiet-kiem-online", "569,110 search/tháng"),

    ("Phase 1", "Tháng 8/2026", "Dự Án 3: Content Strategy & Blog Tài Chính",
     "Thiết lập quy trình tự động hóa xuất bản cẩm nang tài chính bằng AI độc lập trên MoSpark.",
     "In Progress", "momo.vn/tai-chinh/blog", "500,000+ search/tháng"),

    ("Phase 1", "Tháng 8/2026", "Dự Án 4: Tra Cứu Giá Vàng Realtime",
     "Xây dựng bảng giá vàng realtime (SJC, PNJ, DOJI, 24K, 9999, Nhẫn trơn), biểu đồ lịch sử 7-30 ngày.",
     "Done", "momo.vn/gia-vang", "78,879,270 search/tháng"),

    ("Phase 1", "Tháng 9/2026", "Dự Án 5: Tính Lương & Thuế TNCN 2026",
     "Bộ công cụ tính lương Gross - Net (áp dụng luật thuế mới 2026), Tool quyết toán thuế TNCN tự động và thanh trượt phân bổ 50/30/20.",
     "In Progress", "momo.vn/tinh-luong & /thue-tncn", "1,350,140 search/tháng"),

    ("Phase 1", "Tháng 9/2026", "Dự Án 6: Bộ Quy Đổi Ngoại Tệ & Tỷ Giá",
     "Bộ công cụ quy đổi tiền tệ tức thì giữa 20+ ngoại tệ (USD, JPY, EUR...) và bảng tỷ giá so sánh giữa các ngân hàng thương mại lớn.",
     "In Progress", "momo.vn/ty-gia & /ngoai-te", "9,476,250 search/tháng"),

    ("Phase 1", "Tháng 9/2026", "Dự Án 7: Thực Tập Sinh Đầu Tư (Chứng Khoán)",
     "Công cụ tra cứu, hướng dẫn và giả lập Đầu Tư Chứng Khoán dành cho người mới bắt đầu (F0), mua Quỹ mở SIP từ 10.000đ.",
     "Planned", "momo.vn/chung-khoan & /chung-chi-quy", "4,183,460 search/tháng"),

    ("Phase 1", "Tháng 9/2026", "Dự Án 8: Programmatic Bank Hub (34 Ngân Hàng)",
     "Triển khai Programmatic template tự động cho 34 ngân hàng đối tác và Cổng chuyển tiền Napas 247 kèm quà tân thủ 500k.",
     "In Progress", "momo.vn/ngan-hang/[slug] & /napas", "7,015,550 search/tháng"),

    ("Phase 2", "Tháng 10/2026", "Dự Án 9: Bổ Sung Các Công Cụ Tính Toán Tài Chính",
     "Máy tính trả góp 0% tổng quát (Ví Trả Sau 20 triệu), Công cụ tính lãi vay mua nhà/vay tín chấp dư nợ giảm dần và Ma trận so sánh thẻ tín dụng hoàn tiền.",
     "Planned", "momo.vn/tra-gop & /the-tin-dung", "1,288,500 search/tháng"),

    ("Phase 2", "Tháng 10/2026", "Dự Án 10: Trung Tâm Đầu Tư & Tích Sản FIRE",
     "Công cụ 'Có Tiền Đầu Tư Gì?' phân bổ tài sản Lump Sum (An toàn, Tăng trưởng, Mạo hiểm), Bộ giả lập Lãi kép tích lũy định kỳ SIP và FIRE.",
     "Planned", "momo.vn/dau-tu-fire", "600,000+ search/tháng"),

    ("Phase 2", "Tháng 11/2026", "Dự Án 11: Bảng So Sánh Lãi Suất 30+ Ngân Hàng",
     "Bảng so sánh đa chiều biểu lãi suất tiền gửi tiết kiệm theo kỳ hạn 1-36 tháng của 30+ ngân hàng, tích hợp bộ lọc lãi suất cao nhất.",
     "Planned", "momo.vn/lai-suat-ngan-hang", "1,246,620 search/tháng"),

    ("Phase 3", "Tháng 12/2026", "Dự Án 12: AI Financial Pulse & News Stream",
     "Hệ thống tự động hóa ingest tin tức tài chính, lọc nhiễu AI, tóm tắt 2 gạch đầu dòng và alert biến động số liệu 24/7 (mô hình Finpath AI).",
     "Planned", "momo.vn/tai-chinh (Pulse Module)", "Daily Active Engagement"),

    ("Phase 3", "Tháng 01/2027", "Dự Án 13: Đóng Gói Toàn Diện Embeddable Widgets SDK",
     "Đóng gói 7 bộ công cụ tiện ích thành React Web Components nhúng linh hoạt vào mọi điểm chạm Web MoMo và hệ thống đối tác ngoài.",
     "Planned", "Embed Component SDK", "Cross-sell Toàn Sàn"),

    ("Phase 3", "Q1/2027", "Dự Án 14: Cá Nhân Hóa Thời Gian Thực & Phễu U18/F0",
     "Triển khai Real-time Personalization dựa trên dữ liệu người dùng tự khai báo trên Web, tối ưu hóa hành trình chuyển đổi Web-to-App.",
     "Planned", "Toàn bộ hệ thống Hub", "Nâng CVR W2A +30%")
]

# Write to Roadmap sheet in financial-hub-roadmap.xlsx
if 'Roadmap' in wb.sheetnames:
    del wb['Roadmap']

ws_rm = wb.create_sheet('Roadmap', index=1)
ws_rm.append(["LỘ TRÌNH TRIỂN KHAI TOÀN DIỆN CÁC DỰ ÁN FINANCIAL HUB (THỰC TẾ TRIỂN KHAI 2026 - 2027)"])
ws_rm.merge_cells("A1:G1")
ws_rm.cell(row=1, column=1).fill = section_fill
ws_rm.cell(row=1, column=1).font = section_font
ws_rm.cell(row=1, column=1).alignment = Alignment(horizontal='left', vertical='center')

headers = [
    "Giai Đoạn (Phase)",
    "Thời Gian",
    "Tên Dự Án / Module Trọng Tâm",
    "Chi Tiết Triển Khai & Nhiệm Vụ Kỹ Thuật",
    "Trạng Thái",
    "URL Route Sản Phẩm",
    "Dung Lượng Search / Tác Động"
]
ws_rm.append(headers)
for c_i in range(1, len(headers) + 1):
    c = ws_rm.cell(row=2, column=c_i)
    c.fill = header_fill; c.font = header_font; c.border = thin_border
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws_rm.row_dimensions[2].height = 26.0

for item in real_roadmap_data:
    phase, time_val, name, desc, status_val, route, vol = item
    ws_rm.append([phase, time_val, name, desc, status_val, route, vol])
    r = ws_rm.max_row
    ws_rm.row_dimensions[r].height = 24.0
    
    ws_rm.cell(row=r, column=1).alignment = Alignment(horizontal='center', vertical='center')
    ws_rm.cell(row=r, column=2).alignment = Alignment(horizontal='center', vertical='center')
    ws_rm.cell(row=r, column=3).alignment = Alignment(horizontal='left', vertical='center'); ws_rm.cell(row=r, column=3).font = Font(name='Arial', size=9.5, bold=True)
    ws_rm.cell(row=r, column=4).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws_rm.cell(row=r, column=5).alignment = Alignment(horizontal='center', vertical='center')
    ws_rm.cell(row=r, column=6).alignment = Alignment(horizontal='left', vertical='center')
    ws_rm.cell(row=r, column=7).alignment = Alignment(horizontal='center', vertical='center'); ws_rm.cell(row=r, column=7).font = Font(name='Arial', size=9.5, bold=True, color='002060')

    if "Phase 1" in phase: ws_rm.cell(row=r, column=1).fill = p1_fill
    elif "Phase 2" in phase: ws_rm.cell(row=r, column=1).fill = p2_fill
    elif "Phase 3" in phase: ws_rm.cell(row=r, column=1).fill = p3_fill

    # Status Pill Styling
    c_st = ws_rm.cell(row=r, column=5)
    if status_val == "Done":
        c_st.fill = status_done_fill; c_st.font = status_done_font
    elif status_val == "In Progress":
        c_st.fill = status_prog_fill; c_st.font = status_prog_font
    elif status_val == "Planned":
        c_st.fill = status_plan_fill; c_st.font = status_plan_font

    for c_i in range(1, 8): ws_rm.cell(row=r, column=c_i).border = thin_border

widths = [14, 16, 36, 65, 16, 35, 25]
for idx, w in enumerate(widths, start=1):
    col_letter = get_column_letter(idx)
    ws_rm.column_dimensions[col_letter].width = w

wb.save('05_HUBS/financial-hub-roadmap.xlsx')
print("Successfully synchronized financial-hub-roadmap.xlsx with exact real execution roadmap!")

# 2. Update financial-hub-brd.md Section 10
with open('05_HUBS/financial-hub-brd.md', 'r', encoding='utf-8') as f:
    brd = f.read()

new_s10_content = """## 10. LỘ TRÌNH TRIỂN KHAI THỰC TẾ 3 GIAI ĐOẠN (ROADMAP & STATUS)

<table style="width:100%; border-collapse:collapse; border:1.5px solid #64748b; margin:1em 0; font-size:0.85em;">
  <thead>
    <tr style="background-color:#f1f5f9;">
      <th style="border:1.5px solid #64748b; padding:6px; text-align:center; font-weight:700;">Giai Đoạn</th>
      <th style="border:1.5px solid #64748b; padding:6px; text-align:center; font-weight:700;">Thời Gian</th>
      <th style="border:1.5px solid #64748b; padding:6px; text-align:left; font-weight:700;">Tên Dự Án / Module Trọng Tâm</th>
      <th style="border:1.5px solid #64748b; padding:6px; text-align:left; font-weight:700;">Chi Tiết Triển Khai & Nhiệm Vụ Kỹ Thuật</th>
      <th style="border:1.5px solid #64748b; padding:6px; text-align:center; font-weight:700;">Trạng Thái Thực Tế</th>
      <th style="border:1.5px solid #64748b; padding:6px; text-align:left; font-weight:700;">URL Route Sản Phẩm</th>
    </tr>
  </thead>
  <tbody>
    <!-- PHASE 1 - THÁNG 8 -->
    <tr style="background-color:#f0fdf4;">
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700;" rowspan="8"><strong>Phase 1</strong></td>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center;">Tháng 8/2026</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><strong>Dự Án 1: Master Financial Hub</strong></td>
      <td style="border:1px solid #94a3b8; padding:6px;">Xây dựng trang cổng trung tâm, thiết kế Mega Menu 5 nhóm nhu cầu tài chính, thanh Trợ thủ đồng hành và các thẻ điều hướng sản phẩm.</td>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700; color:#15803d; background-color:#dcfce7;">Done</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">momo.vn/tai-chinh</code></td>
    </tr>
    <tr style="background-color:#f0fdf4;">
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center;">Tháng 8/2026</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><strong>Dự Án 2: CIC, Tiết Kiệm</strong></td>
      <td style="border:1px solid #94a3b8; padding:6px;">Công cụ CIC Simulator và Nợ Xấu (tra cứu 339k volume). Công cụ tính lãi suất tiết kiệm cho TKO.</td>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700; color:#15803d; background-color:#dcfce7;">Done</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">/tra-cuu-cic</code> & <code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">/tiet-kiem-online</code></td>
    </tr>
    <tr style="background-color:#fefce8;">
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center;">Tháng 8/2026</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><strong>Dự Án 3: Content Strategy & Blog Tài Chính</strong></td>
      <td style="border:1px solid #94a3b8; padding:6px;">Thiết lập quy trình tự động hóa xuất bản cẩm nang tài chính bằng AI độc lập trên MoSpark.</td>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700; color:#b45309; background-color:#fef9c3;">In Progress</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">/tai-chinh/blog</code></td>
    </tr>
    <tr style="background-color:#f0fdf4;">
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center;">Tháng 8/2026</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><strong>Dự Án 4: Tra Cứu Giá Vàng Realtime</strong></td>
      <td style="border:1px solid #94a3b8; padding:6px;">Xây dựng bảng giá vàng realtime (SJC, PNJ, DOJI, 24K, 9999, Nhẫn trơn), biểu đồ lịch sử 7-30 ngày.</td>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700; color:#15803d; background-color:#dcfce7;">Done</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">/gia-vang</code></td>
    </tr>
    <!-- PHASE 1 - THÁNG 9 -->
    <tr style="background-color:#fefce8;">
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center;">Tháng 9/2026</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><strong>Dự Án 5: Tính Lương & Thuế TNCN 2026</strong></td>
      <td style="border:1px solid #94a3b8; padding:6px;">Bộ công cụ tính lương Gross - Net (áp dụng luật thuế mới 2026), Tool quyết toán thuế TNCN tự động và thanh trượt phân bổ 50/30/20.</td>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700; color:#b45309; background-color:#fef9c3;">In Progress</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">/tinh-luong</code> & <code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">/thue-tncn</code></td>
    </tr>
    <tr style="background-color:#fefce8;">
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center;">Tháng 9/2026</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><strong>Dự Án 6: Bộ Quy Đổi Ngoại Tệ & Tỷ Giá</strong></td>
      <td style="border:1px solid #94a3b8; padding:6px;">Bộ công cụ quy đổi tiền tệ tức thì giữa 20+ ngoại tệ (USD, JPY, EUR...) và bảng tỷ giá so sánh giữa các ngân hàng thương mại lớn.</td>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700; color:#b45309; background-color:#fef9c3;">In Progress</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">/ty-gia</code> & <code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">/ngoai-te</code></td>
    </tr>
    <tr style="background-color:#eff6ff;">
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center;">Tháng 9/2026</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><strong>Dự Án 7: Thực Tập Sinh Đầu Tư (Chứng Khoán)</strong></td>
      <td style="border:1px solid #94a3b8; padding:6px;">Công cụ tra cứu, hướng dẫn và giả lập Đầu Tư Chứng Khoán dành cho người mới bắt đầu (F0), mua Quỹ mở SIP từ 10.000đ.</td>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700; color:#1d4ed8; background-color:#dbeafe;">Planned</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">/chung-khoan</code> & <code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">/chung-chi-quy</code></td>
    </tr>
    <tr style="background-color:#fefce8;">
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center;">Tháng 9/2026</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><strong>Dự Án 8: Programmatic Bank Hub (34 Ngân Hàng)</strong></td>
      <td style="border:1px solid #94a3b8; padding:6px;">Triển khai Programmatic template tự động cho 34 ngân hàng đối tác và Cổng chuyển tiền Napas 247 kèm quà tân thủ 500k.</td>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700; color:#b45309; background-color:#fef9c3;">In Progress</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">/ngan-hang/[slug]</code> & <code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">/napas</code></td>
    </tr>

    <!-- PHASE 2 - THÁNG 10 - 11 -->
    <tr style="background-color:#f8fafc;">
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700;" rowspan="3"><strong>Phase 2</strong></td>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center;">Tháng 10/2026</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><strong>Dự Án 9: Bổ Sung Các Công Cụ Tính Toán Tài Chính</strong></td>
      <td style="border:1px solid #94a3b8; padding:6px;">Máy tính trả góp 0% tổng quát (Ví Trả Sau 20 triệu), Công cụ tính lãi vay mua nhà/vay tín chấp dư nợ giảm dần và Ma trận so sánh thẻ tín dụng.</td>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700; color:#1d4ed8; background-color:#dbeafe;">Planned</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">/tra-gop</code> & <code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">/the-tin-dung</code></td>
    </tr>
    <tr style="background-color:#f8fafc;">
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center;">Tháng 10/2026</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><strong>Dự Án 10: Trung Tâm Đầu Tư & Tích Sản FIRE</strong></td>
      <td style="border:1px solid #94a3b8; padding:6px;">Công cụ 'Có Tiền Đầu Tư Gì?' phân bổ tài sản Lump Sum (An toàn, Tăng trưởng, Mạo hiểm), Bộ giả lập Lãi kép tích lũy định kỳ SIP và FIRE.</td>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700; color:#1d4ed8; background-color:#dbeafe;">Planned</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">/dau-tu-fire</code></td>
    </tr>
    <tr style="background-color:#f8fafc;">
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center;">Tháng 11/2026</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><strong>Dự Án 11: Bảng So Sánh Lãi Suất 30+ Ngân Hàng</strong></td>
      <td style="border:1px solid #94a3b8; padding:6px;">Bảng so sánh đa chiều biểu lãi suất tiền gửi tiết kiệm theo kỳ hạn 1-36 tháng của 30+ ngân hàng, tích hợp bộ lọc lãi suất cao nhất.</td>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700; color:#1d4ed8; background-color:#dbeafe;">Planned</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">/lai-suat-ngan-hang</code></td>
    </tr>

    <!-- PHASE 3 - THÁNG 12/2026 - Q1/2027 -->
    <tr>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700;" rowspan="3"><strong>Phase 3</strong></td>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center;">Tháng 12/2026</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><strong>Dự Án 12: AI Financial Pulse & News Stream</strong></td>
      <td style="border:1px solid #94a3b8; padding:6px;">Hệ thống tự động hóa ingest tin tức tài chính, lọc nhiễu AI, tóm tắt 2 gạch đầu dòng và alert biến động số liệu 24/7 (mô hình Finpath AI).</td>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700; color:#1d4ed8; background-color:#dbeafe;">Planned</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><code style="background:#f1f5f9;padding:2px 4px;border-radius:4px;font-family:monospace;">momo.vn/tai-chinh</code></td>
    </tr>
    <tr>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center;">Tháng 01/2027</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><strong>Dự Án 13: Đóng Gói Toàn Diện Embeddable Widgets SDK</strong></td>
      <td style="border:1px solid #94a3b8; padding:6px;">Đóng gói 7 bộ công cụ tiện ích thành React Web Components nhúng linh hoạt vào mọi điểm chạm Web MoMo và hệ thống đối tác ngoài.</td>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700; color:#1d4ed8; background-color:#dbeafe;">Planned</td>
      <td style="border:1px solid #94a3b8; padding:6px;">Embed Component SDK</td>
    </tr>
    <tr>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center;">Q1/2027</td>
      <td style="border:1px solid #94a3b8; padding:6px;"><strong>Dự Án 14: Cá Nhân Hóa Thời Gian Thực & Phễu U18/F0</strong></td>
      <td style="border:1px solid #94a3b8; padding:6px;">Triển khai Real-time Personalization dựa trên dữ liệu người dùng tự khai báo trên Web, tối ưu hóa hành trình chuyển đổi Web-to-App.</td>
      <td style="border:1px solid #94a3b8; padding:6px; text-align:center; font-weight:700; color:#1d4ed8; background-color:#dbeafe;">Planned</td>
      <td style="border:1px solid #94a3b8; padding:6px;">Toàn bộ hệ thống Hub</td>
    </tr>
  </tbody>
</table>
"""

s10_start = brd.find("## 10. LỘ TRÌNH")
if s10_start != -1:
    brd = brd[:s10_start] + new_s10_content
    with open('05_HUBS/financial-hub-brd.md', 'w', encoding='utf-8') as f:
        f.write(brd)
    print("Successfully updated BRD Section 10 with exact real execution roadmap!")
