import openpyxl

file_cs = '/Users/hienhv/Downloads/[WP] Vehicle Hub - Content Strategy.xlsx'
wb = openpyxl.load_workbook(file_cs, data_only=True)
ws_kw = wb['Keyword']

tram_kws = []
for r in range(2, ws_kw.max_row + 1):
    kw = str(ws_kw.cell(r, 1).value or ws_kw.cell(r, 2).value or '').strip().lower()
    if 'trạm thu phí' in kw or 'tram thu phi' in kw or 'thu phí không dừng' in kw or 'phí cao tốc' in kw:
        vol = ws_kw.cell(r, 2).value or ws_kw.cell(r, 3).value or 0
        try:
            vol = float(vol)
        except:
            vol = 0
        tram_kws.append((kw, vol))

tram_kws.sort(key=lambda x: x[1], reverse=True)
print(f"Total 'trạm thu phí' keywords: {len(tram_kws)}, Sum Vol: {sum(x[1] for x in tram_kws):,.0f}")
for k, v in tram_kws[:25]:
    print(f"  • {k}: {v:,.0f}")

