import openpyxl

file_cs = '/Users/hienhv/Downloads/[WP] Vehicle Hub - Content Strategy.xlsx'
wb_cs = openpyxl.load_workbook(file_cs, data_only=True)

if 'Đăng kiểm' in wb_cs.sheetnames:
    ws = wb_cs['Đăng kiểm']
    print(f"Sheet 'Đăng kiểm' has {ws.max_row} rows and {ws.max_column} cols")
    print("Header:", [ws.cell(1, c).value for c in range(1, ws.max_column + 1)])
    
    dk_rows = []
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        # Check which column is volume
        # Typically Col 1: Theme, Col 2: Pillar, Col 3: Cluster, Col 4: Keyword, Col 5: Volume
        kw = str(vals[3] if len(vals) > 3 else '')
        vol = vals[4] if len(vals) > 4 else 0
        try:
            vol = float(vol)
        except:
            vol = 0
        dk_rows.append((kw, vol, vals))
        
    dk_rows.sort(key=lambda x: x[1], reverse=True)
    total_vol = sum(x[1] for x in dk_rows)
    print(f"Total rows: {len(dk_rows)}, Total Volume in sheet: {total_vol:,.0f}")
    print("\nTop 25 keywords by volume in sheet 'Đăng kiểm':")
    for i, (k, v, raw) in enumerate(dk_rows[:25], start=1):
        print(f"  {i}. {k}: {v:,.0f} | Cluster: {raw[2] if len(raw)>2 else ''}")
        
    # Check location keywords
    loc_kws = [x for x in dk_rows if any(c in x[0].lower() for c in ['hà nội', 'hcm', 'sài gòn', 'đà nẵng', 'bình dương', 'quận'])]
    print(f"\nLocation keywords in sheet 'Đăng kiểm': {len(loc_kws)}")
    for k, v, raw in loc_kws[:15]:
        print(f"  - {k}: {v:,.0f}")
