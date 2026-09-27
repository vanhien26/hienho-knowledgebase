import openpyxl

file_cs = '/Users/hienhv/Downloads/[WP] Vehicle Hub - Content Strategy.xlsx'
wb = openpyxl.load_workbook(file_cs, data_only=True)
ws_kw = wb['Keyword']

bot_kws = []
for r in range(2, ws_kw.max_row + 1):
    kw = str(ws_kw.cell(r, 1).value or ws_kw.cell(r, 2).value or '').strip().lower()
    if 'bot' in kw or 'trạm thu' in kw or 'vé cầu đường' in kw or 'phí qua trạm' in kw:
        vol = ws_kw.cell(r, 2).value or ws_kw.cell(r, 3).value or 0
        try:
            vol = float(vol)
        except:
            vol = 0
        bot_kws.append((kw, vol))

bot_kws.sort(key=lambda x: x[1], reverse=True)
print(f"Total 'bot / tram thu' keywords: {len(bot_kws)}, Sum Vol: {sum(x[1] for x in bot_kws):,.0f}")
for k, v in bot_kws[:20]:
    print(f"  • {k}: {v:,.0f}")

