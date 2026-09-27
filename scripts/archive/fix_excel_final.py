import openpyxl
from openpyxl.styles import Alignment, Border, Side, PatternFill, Font

file_path = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/vehicle-hub-roadmap.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb['Roadmap']

# Properly clear all merged cells ranges
ws.merged_cells.ranges.clear()

# Delete rows from 4 downwards
ws.delete_rows(4, ws.max_row)

align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

fill_p1 = PatternFill("solid", fgColor="D9EAD3")
fill_p2 = PatternFill("solid", fgColor="FFF2CC")
font_bold = Font(bold=True)

data = [
    # PHASE 1 (Tháng 9)
    ['Phase 1: Tiện ích Bản đồ, Content & Phạt Nguội (Tháng 9/2026)', 'Master Hub', 'Trang Chủ Master Hub (/tien-ich-giao-thong)', 'Mặt tiền trung tâm hệ sinh thái tiện ích giao thông.', 'Tháng 9/2026'],
    ['None', 'Magnet + Transaction', 'Mở rộng Phạt Nguội & Location Page (/phat-nguoi, /phat-nguoi/[tinh])', 'Mở rộng trang Phạt nguội theo tỉnh/thành. Đẩy mạnh traffic dẫn về luồng Nộp Phạt In-App và Master Hub.', 'Tháng 9/2026'],
    ['None', 'Location & Maps', 'Mạng lưới Trạm, Bãi đỗ & Cây xăng (/gia-xang, /cay-xang, /tram-sac, /bai-do-xe, /bao-duong)', 'Xây dựng bản đồ định vị: Giá xăng, Cây xăng, Garage/Rửa xe, Trạm sạc EV, Bãi đỗ xe.', 'Tháng 9/2026'],
    ['None', 'Content & SEO', 'Blog & Content Production (/tra-cuu-muc-phat, /bien-bao-giao-thong, /kinh-nghiem-lai-xe)', 'Sản xuất nội dung cẩm nang, tra cứu mức phạt NĐ 100/123, biển báo giao thông.', 'Tháng 9/2026'],
    ['None', 'Utilities', 'Tra Cứu Hạn Đăng Kiểm (/dang-kiem)', 'Xây dựng trang tra cứu hạn đăng kiểm & đặt lịch kiểm định.', 'Tháng 9/2026'],

    # PHASE 2 (Tháng 10 trở đi)
    ['Phase 2: Giao dịch, Tài chính & Khẩn Cấp (Tháng 10/2026)', 'Transaction (Mở rộng)', 'Thu Phí BOT (/phi-khong-dung, /tram-thu-phi)', 'Phát triển & mở rộng luồng nạp tiền thu phí không dừng ePass/VETC đã build.', 'Tháng 10/2026'],
    ['None', 'Utilities', 'Cứu Hộ Ô Tô (/cuu-ho)', 'Danh bạ bản đồ gọi cứu hộ ô tô khẩn cấp 24/7.', 'Tháng 10/2026'],
    ['None', 'Monetize', 'Bảo Hiểm Phương Tiện (/bao-hiem-o-to, /bao-hiem-xe-may)', 'Luồng cross-sell bảo hiểm, prefill thông tin xe.', 'Tháng 10/2026'],
    ['None', 'Calculator', 'Công cụ Tài chính Xe (/lan-banh, /chi-phi-nuoi-xe, /so-sanh-xe)', 'Tính phí lăn bánh, chi phí nuôi xe, so sánh mẫu xe.', 'Tháng 10/2026'],
    ['None', 'Smart City', 'Camera & Lộ trình (/camera-giao-thong, /chuyen-di)', 'Live stream camera điểm kẹt xe, nhật ký chuyến đi.', 'Tháng 10/2026']
]

for idx, row_data in enumerate(data, start=4):
    for col, val in enumerate(row_data, start=1):
        cell = ws.cell(row=idx, column=col, value=val if val != 'None' else None)
        cell.alignment = align_left if col in [3, 4] else align_center
        cell.border = border
        if col == 1 and val != 'None':
            cell.font = font_bold
            
        if idx <= 8:
            cell.fill = fill_p1
        else:
            cell.fill = fill_p2

# Merge columns
ws.merge_cells('A4:A8')
ws.merge_cells('A9:A13')

wb.save(file_path)
print("Excel Roadmap cleanly updated with Phat Nguoi expanded and Cuu Ho pushed to Phase 2.")
