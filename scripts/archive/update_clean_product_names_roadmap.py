import openpyxl

wb = openpyxl.load_workbook('05_HUBS/financial-hub-roadmap.xlsx')
ws = wb['Roadmap']

# Name cleanup mapping
clean_names = {
    'Dự Án 1: Master Financial Hub': 'Financial Hub',
    'Dự Án 2: CIC, Tiết Kiệm': 'CIC & Tiết Kiệm',
    'Dự Án 3: Content Strategy & Blog Tài Chính': 'Blog & Cẩm Nang Tài Chính',
    'Dự Án 4: Tra Cứu Giá Vàng Realtime': 'Giá Vàng',
    'Dự Án 5: Tính Lương & Thuế TNCN 2026': 'Tính Lương & Thuế TNCN',
    'Dự Án 7: Bộ Quy Đổi Tỷ Giá & Hub Các Cặp Tiền Tệ': 'Tỷ Giá & Ngoại Tệ',
    'Dự Án 8: Thực Tập Sinh Đầu Tư (Chứng Khoán)': 'Chứng Khoán & Quỹ Mở',
    'Dự Án 9: Bổ Sung Các Công Cụ Tính Toán Tài Chính': 'Công Cụ Tài Chính Cá Nhân',
    'Dự Án 9: Trung Tâm Đầu Tư & Tích Sản FIRE': 'Đầu Tư & Tích Sản FIRE'
}

for r in range(3, ws.max_row + 1):
    current_val = str(ws.cell(row=r, column=4).value or '').strip()
    for old_prefix, clean_name in clean_names.items():
        if old_prefix in current_val or current_val in old_prefix:
            ws.cell(row=r, column=4).value = clean_name
            break

wb.save('05_HUBS/financial-hub-roadmap.xlsx')
print("Successfully cleaned product names in Roadmap: removed 'Dự Án' and verbose suffixes!")
