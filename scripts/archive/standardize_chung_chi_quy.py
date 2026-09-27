import openpyxl

wb = openpyxl.load_workbook('05_HUBS/financial-hub-roadmap.xlsx')

# 1. UPDATE UTILITIES SHEET
ws_u = wb['Utilities']
for r in range(4, ws_u.max_row + 1):
    for c in range(1, ws_u.max_column + 1):
        val = str(ws_u.cell(row=r, column=c).value or '')
        if 'Quỹ Mở SIP' in val:
            val = val.replace('Quỹ Mở SIP (Icon Dragon/SSIAM', 'Chứng Chỉ Quỹ (Icon Chứng Chỉ Quỹ - Dragon Capital/SSIAM')
            val = val.replace('Quỹ Mở SIP', 'Chứng Chỉ Quỹ (Tích sản định kỳ SIP)')
            ws_u.cell(row=r, column=c).value = val
        if 'Icon Quỹ Mở' in val:
            val = val.replace('Icon Quỹ Mở', 'Icon Chứng Chỉ Quỹ')
            ws_u.cell(row=r, column=c).value = val
        if 'đầu tư Quỹ Mở SIP' in val:
            val = val.replace('đầu tư Quỹ Mở SIP', 'đầu tư Chứng Chỉ Quỹ định kỳ (SIP)')
            ws_u.cell(row=r, column=c).value = val

# 2. UPDATE ROADMAP SHEET
ws_rd = wb['Roadmap']
for r in range(3, ws_rd.max_row + 1):
    # Track & Product Name in Row 11 (Chứng Khoán & Quỹ Mở)
    p_name = str(ws_rd.cell(row=r, column=4).value or '')
    if 'Chứng Khoán & Quỹ Mở' in p_name:
        ws_rd.cell(row=r, column=3).value = "Đầu Tư & Chứng Chỉ Quỹ"
        ws_rd.cell(row=r, column=4).value = "Chứng Khoán & Chứng Chỉ Quỹ"
    
    for c in range(1, ws_rd.max_column + 1):
        val = str(ws_rd.cell(row=r, column=c).value or '')
        if 'Quỹ mở SIP' in val:
            val = val.replace('Quỹ mở SIP', 'Chứng Chỉ Quỹ (Tích sản SIP)')
            ws_rd.cell(row=r, column=c).value = val
        if 'Quỹ Mở SIP' in val:
            val = val.replace('Quỹ Mở SIP', 'Chứng Chỉ Quỹ (Tích sản SIP)')
            ws_rd.cell(row=r, column=c).value = val

wb.save('05_HUBS/financial-hub-roadmap.xlsx')
print("Successfully standardized 'Chứng Chỉ Quỹ' terminology across financial-hub-roadmap.xlsx!")
