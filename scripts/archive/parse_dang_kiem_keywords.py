import openpyxl
from collections import defaultdict

file_cs = '/Users/hienhv/Downloads/[WP] Vehicle Hub - Content Strategy.xlsx'
wb = openpyxl.load_workbook(file_cs, data_only=True)
ws = wb['Đăng kiểm']

keywords_by_group = defaultdict(list)

for r in range(1, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    pillar = str(vals[1] or '')
    cluster = str(vals[2] or '')
    kw = str(vals[3] or '').strip().lower()
    vol = vals[4] or 0
    try:
        vol = float(vol)
    except:
        vol = 0
    
    if not kw or kw == 'none':
        continue

    # Categorize into user search intents:
    if any(w in kw for w in ['phí', 'giá', 'bao nhiêu tiền', 'chi phí', 'lệ phí']):
        keywords_by_group['1. Chi Phí & Lệ Phí Đăng Kiểm'].append((kw, vol))
    elif any(w in kw for w in ['chu kỳ', 'hạn', 'thời gian', 'mấy năm', 'bao lâu', 'mấy tháng']):
        keywords_by_group['2. Chu Kỳ & Thời Hạn Đăng Kiểm'].append((kw, vol))
    elif any(w in kw for w in ['phạt', 'quá hạn', 'chậm', 'hết hạn']):
        keywords_by_group['3. Mức Phạt & Quá Hạn Đăng Kiểm'].append((kw, vol))
    elif any(w in kw for w in ['thủ tục', 'hồ sơ', 'giấy tờ', 'cần gì', 'quy trình', 'mang gì']):
        keywords_by_group['4. Hồ Sơ, Giấy Tờ & Thủ Tục'].append((kw, vol))
    elif any(w in kw for w in ['đặt lịch', 'online', 'app', 'trực tuyến']):
        keywords_by_group['5. Đặt Lịch Hẹn Đăng Kiểm Online'].append((kw, vol))
    elif any(w in kw for w in ['trung tâm', 'trạm', 'ở đâu', 'địa chỉ']):
        keywords_by_group['6. Trung Tâm & Trạm Đăng Kiểm'].append((kw, vol))
    elif any(w in kw for w in ['xe điện', 'vinfast', 'ô tô điện']):
        keywords_by_group['7. Đăng Kiểm Xe Điện EV'].append((kw, vol))
    elif any(w in kw for w in ['tra cứu', 'biển số', 'check']):
        keywords_by_group['8. Tra Cứu Biển Số & Thông Tin Xe'].append((kw, vol))
    else:
        keywords_by_group['9. Từ Khóa Tổng Quan & Phổ Thông'].append((kw, vol))

for group, kws in sorted(keywords_by_group.items()):
    kws.sort(key=lambda x: x[1], reverse=True)
    total_vol = sum(x[1] for x in kws)
    print(f"\n=== {group} (Tổng Volume: {total_vol:,.0f} search/tháng, {len(kws)} từ khóa) ===")
    for k, v in kws[:8]:
        if v > 0:
            print(f"  • {k}: {v:,.0f}")

