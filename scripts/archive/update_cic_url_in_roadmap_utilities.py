import openpyxl

wb = openpyxl.load_workbook('05_HUBS/financial-hub-roadmap.xlsx')

# Update Roadmap Row 4 (CIC & Tiết Kiệm)
ws_rd = wb['Roadmap']
for r in range(3, ws_rd.max_row + 1):
    feat = str(ws_rd.cell(row=r, column=4).value or '')
    if 'CIC' in feat:
        ws_rd.cell(row=r, column=5).value = (
            "• Chuyên trang Điểm Tín Dụng & Tra Cứu CIC chính thức: /diem-tin-dung\n"
            "• Widget CIC Simulator đánh giá nhóm nợ 1-5 ➔ Gợi ý mở Sổ Tiết Kiệm để làm đẹp hồ sơ tín dụng.\n"
            "• Hero Widget Tính Lãi Tiết Kiệm TKO (Bản Việt / VPBank) ➔ Tích hợp tab so sánh 'Gửi Tiết Kiệm vs Thực Tập Sinh Đầu Tư'.\n"
            "• Chuẩn UX bắt buộc: Tích hợp Khối Chốt Hạ Chuyển Đổi gồm Hộp MoMo Gợi Ý (AI) phản hồi động theo mức nhập + Lưới Icon Dịch Vụ MoMo (chi tiết xem sheet Utilities)."
        )
        ws_rd.row_dimensions[r].height = 70.0

# Update Utilities Row 8
ws_u = wb['Utilities']
for r in range(4, ws_u.max_row + 1):
    tool_name = str(ws_u.cell(row=r, column=2).value or '')
    if 'CIC' in tool_name:
        ws_u.cell(row=r, column=2).value = "CIC Simulator & Điểm Tín Dụng (/diem-tin-dung)"

wb.save('05_HUBS/financial-hub-roadmap.xlsx')
print("Successfully synced /diem-tin-dung URL across Roadmap and Utilities sheets!")
