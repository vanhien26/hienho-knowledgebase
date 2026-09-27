import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


from scripts._lib import load_workbook_safe, save_with_backup, standard_argparser

def main(args=None):
    wb = load_workbook_safe('05_HUBS/student-hub-roadmap.xlsx')

    # Common styles
    font_disclaimer = Font(name='Calibri', size=11, bold=True, color='990000', italic=True)
    font_title = Font(name='Calibri', size=13, bold=True, color='1E293B')
    font_section_hdr = Font(name='Calibri', size=11, bold=True, color='0F172A')
    font_bold = Font(name='Calibri', size=11, bold=True, color='000000')
    font_regular = Font(name='Calibri', size=11, color='000000')
    font_link = Font(name='Calibri', size=11, color='1D4ED8', underline='single')

    fill_disclaimer = PatternFill('solid', fgColor='FCE4D6')
    fill_section = PatternFill('solid', fgColor='F1F5F9')
    fill_header_sub = PatternFill('solid', fgColor='E2E8F0')

    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    border_thin = Side(style='thin', color='CBD5E1')
    border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    if 'Readme' in wb.sheetnames:
        ws_readme = wb['Readme']
        ws_readme.delete_rows(1, ws_readme.max_row)
        ws_readme.views.sheetView[0].showGridLines = True

        # Disclaimer row
        ws_readme.cell(row=1, column=1, value="Note: Do not share access to all materials without project manager's permission").font = font_disclaimer
        ws_readme.cell(row=1, column=1).fill = fill_disclaimer
        ws_readme.merge_cells('A1:B1')

        readme_content = [
            # Section Playbook
            (3, "Quy Trình Hợp Tác Chiến Lược Mới (Web Platform x Youth & Student Segment BU)", None, "SECTION"),
            (4, "Tôn Chỉ Hợp Tác", "Phát triển Student Hub (momo.vn/sinh-vien) & Cổng Việc Làm (/sinh-vien/viec-lam) làm Cổng Khám Phá & Định Danh Sinh Viên Toàn Diện, dẫn dắt chuyển đổi sang App MoMo", "PLAYBOOK"),
            (5, "1. Đồng Sở Hữu KPI (KPI Co-ownership)", "Web Product đồng sở hữu KPI với Youth Segment BU (Target: 500.000 MPV/Q4, 70.000 New Verified Students, 1.5M Student Pass MEU, W2A CTR 10-15%)", "PLAYBOOK"),
            (6, "2. Quản Lý Dự Án Bài Bản (Single SSOT)", "Sử dụng file này làm công cụ SSOT quản lý lộ trình phát triển Web Product (Timeline 3 Phase, 100 Blog Content Plan, 12 Trường ĐH Trọng Điểm & Đặc tả Tiện ích Web)", "PLAYBOOK"),
            (7, "3. Vận Hành Sprint 1 Tháng (Cadence)", "Họp lập kế hoạch (Sprint Planning) đầu tháng; Rà soát tiến độ (Check-in) định kỳ 2 tuần/lần (giữa và cuối tháng)", "PLAYBOOK"),
            (8, "4. Trách Nhiệm Giải Trình (Accountability)", "Báo cáo tiến độ minh bạch hàng tháng cho IVP Leadership. Nếu BU trễ 1-2 tháng, Web Platform sẽ tái phân bổ nguồn lực sang dự án ưu tiên khác", "PLAYBOOK"),

            # Docs
            (10, "Project Management Docs", None, "SECTION"),
            (11, "BRD (Master Specification)", "[Web Platform] Student Hub - Master BRD (SSOT - 05_HUBS/student-hub-brd.md)", "LINK"),
            (12, "Roadmap & Monthly Sprint Plan", "[Web Platform] Student Hub - Roadmap 2026 (05_HUBS/student-hub-roadmap.xlsx)", "LINK"),
            (13, "Content Strategy & Keyword Plan", "[WP] Student Hub - Content Strategy & MoSpark AI Pipeline (100 Articles & 12 Priority Universities)", "LINK"),
            (14, "Tracking & Onelink Plan (W2A)", "Web Platform x Youth Segment BU - Appsflyer Onelink & Student Pass W2A Tracking Spec", "LINK"),
            (15, "Market Sizing & Search Demand", "Student Hub - Master Inventory Sizing (Search Volume 12 Trường ĐH & Cổng Việc Làm /viec-lam)", "LINK"),
            (16, "Budget & Resource Tracker", "2026_WebPlatform_StudentHub_Tracker", "LINK"),

            # Design
            (18, "Design & UX/UI (Figma)", None, "SECTION"),
            (19, "Master Design Board (Figma)", "https://www.figma.com/design/momo-student-hub-2026", "LINK"),
            (20, "Design System & UI Kit", "@momo-webplatform/mobase (Gen Z UI Theme & School Pass Custom Card Kit)", "LINK"),
            (21, "Responsive Prototype (Web & Mobile)", "https://www.figma.com/proto/momo-student-hub-prototype", "LINK"),
            (22, "Asset Library (Brand Logos, Badges)", "MoMo Student Pass & 12 Priority Universities Asset Library", "LINK"),

            # Environments
            (24, "Environments & Testing (UAT / Staging / Live)", None, "SECTION"),
            (25, "Môi Trường UAT / Staging", "https://staging-web.momo.vn/sinh-vien", "LINK"),
            (26, "Môi Trường Production (Live)", "https://www.momo.vn/sinh-vien", "LINK"),
            (27, "Cổng Việc Làm Sinh Viên Master", "https://www.momo.vn/sinh-vien/viec-lam", "LINK"),
            (28, "Cổng Quản Trị Nội Dung (CMS)", "https://mospark.mservice.io/student-hub (MoSpark CMS / MS Park)", "LINK"),
            (29, "Hạ Tầng Data Connectors & APIs", "Google Transit API (Xe buýt), Merchant POI DB (Điểm thanh toán MoMo), Student Pass eKYC Verification API", "LINK"),

            # Materials
            (31, "Materials & Marketing Deliverables", None, "SECTION"),
            (32, "Quick snapshot of users' Insights", "Gen Z Freshmen & Student Financial & Career Pain Points 2026", "LINK"),
            (33, "Value Exchange Concept (Student Pass)", "No-Login Value First: Trải nghiệm tiện ích tra cứu trước ➔ Mở App eKYC Thẻ Sinh Viên Số nhận ưu đãi", "LINK"),
            (34, "QR Code dẫn về trang Hub in-app", "Dynamic QR Code Generator for Desktop Web (Universal OneLink)", "LINK"),
            (35, "RefId / Tracking QR / Onelink", "Appsflyer Onelink Config (pid=web_hub, c=student_pass_2026, af_dp=momo://student_pass)", "LINK"),

            # Legal
            (37, "Legal & Compliance", None, "SECTION"),
            (38, "File đưa legal check", "Student Hub - Data Privacy, Decree 13 Compliance & Student Identity Disclaimers", "LINK"),
            (39, "Quy tắc bảo mật thông tin (0-PII)", "Ẩn 100% thông tin cá nhân & PII sinh viên trên public web; eKYC thực hiện hoàn toàn trong App MoMo", "LINK"),

            # Alignment
            (41, "Working Board & Alignment", None, "SECTION"),
            (42, "Monthly Sprint Planning & Check-in", "Lập kế hoạch đầu tháng | Review tiến độ ngày 15 & 30 hàng tháng cùng Youth Segment BU", "LINK"),
            (43, "Báo Cáo Tiến Độ (IVP Leadership)", "Báo cáo minh bạch hàng tháng gửi Ban Lãnh Đạo IVP (Theo chuẩn Key Highlights & Business Impact)", "LINK"),

            # Framework Architecture
            (45, "Khung Quy Hoạch 3 Phase Chiến Lược Trên Web Platform Student Hub", None, "SECTION"),
            (46, "Phase 1: Foundation Launch (Tháng 8/2026 - Completed)", "Hoàn thiện hạ tầng Gateway (/sinh-vien) & 5 trường Pilot (UEH, FTU, TDTU, HCMUT, VLU). Tích hợp 3 tính năng cốt lõi: Học phí tín chỉ & điểm chuẩn 3 năm, Tìm điểm thanh toán MoMo quanh campus (F&B/SME), Tra cứu tuyến xe buýt.", "FRAMEWORK"),
            (47, "Phase 2: Scale, Student Verification Drive & Gamification (Tháng 9/2026 - In Progress)", "Mở rộng 7 trường mới (HUST, NEU, HANU, HAU, USSH, HUIT, UEL -> 12 trường). Promote & Drive 100% sinh viên đến eKYC Verify Student Pass In-App nhận gói voucher 500K. Tích hợp Gamification: Web Photobooth Real-time Event 26-27/09, Nhiệm vụ Đại sứ (/ambassador), Minigame tích điểm In-App.", "FRAMEWORK"),
            (48, "Phase 3: Career Master, Programmatic Scale & Ecosystem (Quý 4/2026 - Planned)", "Ra mắt Cổng Việc Làm Sinh Viên Master (/sinh-vien/viec-lam) kết hợp AI Resume Builder (ATS Score), Bản đồ việc làm 5km quanh campus từ đối tác F&B/bán lẻ MoMo; 1-Tap Apply via Student Pass mở Ví Trả Sau 0% HSSV. Programmatic pSEO Scale-out 100+ trang trường ĐH/CĐ toàn quốc; Sub-pages (/nha-tro, /review, /workshop).", "FRAMEWORK")
        ]

        for row_info in readme_content:
            r_idx, col1_val, col2_val, item_type = row_info
        
            c1 = ws_readme.cell(row=r_idx, column=1, value=col1_val)
            c2 = ws_readme.cell(row=r_idx, column=2, value=col2_val)
        
            c1.border = border
            c2.border = border
        
            if item_type == "SECTION":
                c1.font = font_section_hdr
                c1.fill = fill_section
                c2.fill = fill_section
                ws_readme.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=2)
            elif item_type == "PLAYBOOK":
                c1.font = font_bold
                c2.font = font_regular
                c1.alignment = align_left
                c2.alignment = align_left
            elif item_type == "FRAMEWORK":
                c1.font = font_bold
                c2.font = font_regular
                c1.alignment = align_left
                c2.alignment = align_left
                c1.fill = fill_header_sub
            else: # LINK
                c1.font = font_bold
                c2.font = font_link if (col2_val and col2_val.startswith('http')) else font_regular
                c1.alignment = align_left
                c2.alignment = align_left

        ws_readme.column_dimensions['A'].width = 45
        ws_readme.column_dimensions['B'].width = 110

    save_with_backup(wb, '05_HUBS/student-hub-roadmap.xlsx', dry_run=args.dry_run, make_backup=not args.no_backup)
    print('Successfully updated sheet Readme in student-hub-roadmap.xlsx to match Vehicle Hub & Financial Hub standards exactly!')

if __name__ == "__main__":
    _ap = standard_argparser("Fix runner - safe with --dry-run")
    main(_ap.parse_args())

