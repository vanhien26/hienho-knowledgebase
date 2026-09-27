import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


from scripts._lib import load_workbook_safe, save_with_backup, standard_argparser

def main(args=None):
    wb = load_workbook_safe('05_HUBS/financial-hub-roadmap.xlsx')
    ws_old = wb['Sitemap']

    old_rows = []
    for r in range(4, ws_old.max_row + 1):
        ptype = ws_old.cell(row=r, column=2).value
        pname = ws_old.cell(row=r, column=3).value
        purl = ws_old.cell(row=r, column=4).value
        pkw = ws_old.cell(row=r, column=5).value
        old_rows.append((ptype, pname, purl, pkw))

    # Define new entries
    lai_suat_entry = (
        "Công Cụ Tiện Ích Under /tai-chinh",
        "Bảng Lãi Suất Ngân Hàng & Lãi Suất Vay Mới Nhất",
        "/tai-chinh/lai-suat",
        "lãi suất ngân hàng, bảng lãi suất vay, lãi suất điều hành nhnn (1237000 search/tháng - Lãi suất thị trường & lãi suất vay, không liên quan gửi tiết kiệm)"
    )

    tu_do_tai_chinh_entry = (
        "Công Cụ Tiện Ích Under /tai-chinh",
        "Bộ Giả Lập Tự Do Tài Chính & Tích Sản FIRE",
        "/tai-chinh/tu-do-tai-chinh",
        "tự do tài chính, fire simulator, tính lãi kép tích sản, nghỉ hưu sớm (120000 search/tháng)"
    )

    luong_huu_entry = (
        "Công Cụ Tiện Ích Under /tai-chinh",
        "Dự Toán Lương Hưu & Tuổi Nghỉ Hưu Mới Nhất",
        "/tai-chinh/luong-huu",
        "cách tính lương hưu, dự toán lương hưu, tuổi nghỉ hưu mới nhất (180000 search/tháng)"
    )

    # Re-build list in logical order
    refined_rows = []
    for item in old_rows:
        refined_rows.append(item)
        # Insert /tai-chinh/lai-suat right after /tai-chinh/gui-tiet-kiem
        if item[2] == '/tai-chinh/gui-tiet-kiem':
            refined_rows.append(lai_suat_entry)
        # Insert /tai-chinh/tu-do-tai-chinh and /tai-chinh/luong-huu right after /tai-chinh/phan-bo-luong
        if item[2] == '/tai-chinh/phan-bo-luong':
            refined_rows.append(tu_do_tai_chinh_entry)
            refined_rows.append(luong_huu_entry)

    del wb['Sitemap']
    ws_s = wb.create_sheet('Sitemap', index=3)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    section_font = Font(name='Arial', size=10.5, bold=True, color='1F4E78')
    thin_side = Side(style='thin', color='64748B')
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    font_bold = Font(name='Arial', size=9.5, bold=True)
    font_regular = Font(name='Arial', size=9.5)
    font_url = Font(name='Arial', size=9.5, color='0563C1', underline='single')

    ws_s.append(["DANH BẠ KIẾN TRÚC THÔNG TIN & SITEMAP CÁC TRANG CẦN TẠO (FINANCIAL HUB)"])
    ws_s.merge_cells("A1:E1")
    ws_s.cell(row=1, column=1).fill = section_fill
    ws_s.cell(row=1, column=1).font = section_font
    ws_s.row_dimensions[1].height = 25.0

    ws_s.append(["Bản đồ định tuyến tinh gọn toàn bộ các trang cần tạo thuộc Financial Hub: Cổng trung tâm, Cổng tiện ích /tai-chinh/cong-cu, các công cụ tiện ích trực thuộc /tai-chinh/*, URL tượng trưng cho cặp tiền & ngân hàng, chuyên trang độc lập và blog hub."])
    ws_s.merge_cells("A2:E2")
    ws_s.cell(row=2, column=1).font = Font(name='Arial', size=9.5, italic=True, color='475569')
    ws_s.row_dimensions[2].height = 20.0

    s_headers = [
        "STT",
        "Phân Cấp / Loại Trang",
        "Tên Trang Cần Tạo (H1 / Title)",
        "Định Tuyến URL Chuẩn Hóa (Canonical URL)",
        "Mục Tiêu Từ Khóa & Mục Đích Tạo Trang"
    ]
    ws_s.append(s_headers)
    for c_i in range(1, len(s_headers) + 1):
        c = ws_s.cell(row=3, column=c_i)
        c.fill = header_fill; c.font = header_font; c.border = thin_border
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws_s.row_dimensions[3].height = 30.0

    for idx, row_item in enumerate(refined_rows, start=1):
        full_row = [idx] + list(row_item)
        ws_s.append(full_row)
        r = ws_s.max_row
        num_lines = str(row_item[3] or '').count('\n') + 1
        ws_s.row_dimensions[r].height = max(34.0, num_lines * 18.0)

        ws_s.cell(row=r, column=1).alignment = Alignment(horizontal='center', vertical='center')
        ws_s.cell(row=r, column=2).alignment = Alignment(horizontal='center', vertical='center'); ws_s.cell(row=r, column=2).font = font_bold
        ws_s.cell(row=r, column=3).alignment = Alignment(horizontal='left', vertical='center'); ws_s.cell(row=r, column=3).font = font_bold
        ws_s.cell(row=r, column=4).alignment = Alignment(horizontal='left', vertical='center'); ws_s.cell(row=r, column=4).font = font_url
        ws_s.cell(row=r, column=5).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

        for c_i in range(1, 6):
            ws_s.cell(row=r, column=c_i).border = thin_border

    col_widths_s = [6, 35, 45, 38, 60]
    for col_idx, w in enumerate(col_widths_s, start=1):
        ws_s.column_dimensions[get_column_letter(col_idx)].width = w

    save_with_backup(wb, '05_HUBS/financial-hub-roadmap.xlsx', dry_run=args.dry_run, make_backup=not args.no_backup)
    print(f"Successfully added 3 pages! Sitemap now has {len(refined_rows)} rows (1 to {len(refined_rows)}).")

if __name__ == "__main__":
    _ap = standard_argparser("Fix runner - safe with --dry-run")
    main(_ap.parse_args())

