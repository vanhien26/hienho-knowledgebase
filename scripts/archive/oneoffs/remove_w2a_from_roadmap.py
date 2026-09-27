import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


from scripts._lib import load_workbook_safe, save_with_backup, standard_argparser

def main(args=None):
    wb = load_workbook_safe('05_HUBS/financial-hub-roadmap.xlsx')

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    section_font = Font(name='Arial', size=10.5, bold=True, color='1F4E78')
    thin_side = Side(style='thin', color='64748B')
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    p0_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    p1_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
    p2_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    status_done_fill = PatternFill(start_color="1E7E34", end_color="1E7E34", fill_type="solid")
    status_done_font = Font(name='Arial', size=9.5, bold=True, color='FFFFFF')
    status_prog_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    status_prog_font = Font(name='Arial', size=9.5, bold=True, color='B9770E')
    status_plan_fill = PatternFill(start_color="1B4F72", end_color="1B4F72", fill_type="solid")
    status_plan_font = Font(name='Arial', size=9.5, bold=True, color='FFFFFF')

    font_bold = Font(name='Arial', size=9.5, bold=True)
    font_regular = Font(name='Arial', size=9.5)

    ws_old = wb['Roadmap']
    old_rows = []
    for r in range(3, ws_old.max_row + 1):
        phase = ws_old.cell(row=r, column=1).value
        time_str = ws_old.cell(row=r, column=2).value
        track = ws_old.cell(row=r, column=3).value
        prod = ws_old.cell(row=r, column=4).value
        scope = ws_old.cell(row=r, column=5).value
        # Col 6 was W2A -> REMOVE!
        status = ws_old.cell(row=r, column=7).value
        old_rows.append((phase, time_str, track, prod, scope, status))

    del wb['Roadmap']
    ws_rd = wb.create_sheet('Roadmap', index=1)

    ws_rd.append(["LỘ TRÌNH TRIỂN KHAI TOÀN DIỆN CÁC TÍNH NĂNG FINANCIAL MASTER HUB (2026 - 2027)"])
    ws_rd.merge_cells("A1:F1")
    ws_rd.cell(row=1, column=1).fill = section_fill
    ws_rd.cell(row=1, column=1).font = section_font
    ws_rd.row_dimensions[1].height = 25.0

    rd_headers = [
        "Giai Đoạn (Phase)",
        "Thời Gian",
        "Cụm Sản Phẩm (Track)",
        "Tên Tính Năng / Sản Phẩm",
        "Chi Tiết Triển Khai Kỹ Thuật & UX Scope",
        "Trạng Thái"
    ]
    ws_rd.append(rd_headers)
    for c_i in range(1, len(rd_headers) + 1):
        c = ws_rd.cell(row=2, column=c_i)
        c.fill = header_fill; c.font = header_font; c.border = thin_border
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws_rd.row_dimensions[2].height = 28.0

    for row_item in old_rows:
        ws_rd.append(list(row_item))
        r = ws_rd.max_row
        num_lines = str(row_item[4] or '').count('\n') + 1
        ws_rd.row_dimensions[r].height = max(50.0, num_lines * 17.5)

        ws_rd.cell(row=r, column=1).alignment = Alignment(horizontal='center', vertical='center')
        ws_rd.cell(row=r, column=2).alignment = Alignment(horizontal='center', vertical='center')
        ws_rd.cell(row=r, column=3).alignment = Alignment(horizontal='center', vertical='center'); ws_rd.cell(row=r, column=3).font = font_bold
        ws_rd.cell(row=r, column=4).alignment = Alignment(horizontal='left', vertical='center'); ws_rd.cell(row=r, column=4).font = font_bold
        ws_rd.cell(row=r, column=5).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws_rd.cell(row=r, column=6).alignment = Alignment(horizontal='center', vertical='center')

        phase = str(row_item[0] or '')
        time_str = str(row_item[1] or '')
        if "Tháng 7" in time_str: ws_rd.cell(row=r, column=1).fill = p0_fill
        elif "Tháng 8" in time_str or "Tháng 9" in time_str: ws_rd.cell(row=r, column=1).fill = p1_fill
        else: ws_rd.cell(row=r, column=1).fill = p2_fill

        status_val = str(row_item[5] or '')
        c_st = ws_rd.cell(row=r, column=6)
        if status_val == "Done": c_st.fill = status_done_fill; c_st.font = status_done_font
        elif status_val == "In Progress": c_st.fill = status_prog_fill; c_st.font = status_prog_font
        elif status_val == "Planned": c_st.fill = status_plan_fill; c_st.font = status_plan_font

        for c_i in range(1, 7): ws_rd.cell(row=r, column=c_i).border = thin_border

    col_widths = [14, 16, 26, 28, 85, 16]
    for idx, w in enumerate(col_widths, start=1):
        ws_rd.column_dimensions[get_column_letter(idx)].width = w

    save_with_backup(wb, '05_HUBS/financial-hub-roadmap.xlsx', dry_run=args.dry_run, make_backup=not args.no_backup)
    print("Successfully removed W2A column from Roadmap sheet! Total 6 clean columns remaining.")

if __name__ == "__main__":
    _ap = standard_argparser("Fix runner - safe with --dry-run")
    main(_ap.parse_args())

