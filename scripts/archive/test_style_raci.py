import openpyxl
from openpyxl.styles import Alignment, Border, Side, PatternFill, Font

file_fin = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/financial-hub-roadmap.xlsx'
file_veh = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/vehicle-hub-roadmap.xlsx'

wb_fin = openpyxl.load_workbook(file_fin)
ws_fin = wb_fin['RACI']

wb_veh = openpyxl.load_workbook(file_veh)

# Remove outdated RACI sheet if exists, we will standardize into Resources
if 'RACI & Resource Allocation' in wb_veh.sheetnames:
    del wb_veh['RACI & Resource Allocation']

if 'Resources' in wb_veh.sheetnames:
    del wb_veh['Resources']

ws = wb_veh.create_sheet('Resources', index=5)

# Styling from Financial Hub RACI
fill_title = PatternFill(fill_type='solid', start_color='D9E1F2', end_color='D9E1F2')
font_title = Font(name='Calibri', size=13, bold=True, color='1F4E78')

fill_header = PatternFill(fill_type='solid', start_color='1F4E78', end_color='1F4E78')
font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')

fill_group = PatternFill(fill_type='solid', start_color='D9E1F2', end_color='D9E1F2')
font_group = Font(name='Calibri', size=11, bold=True, color='1F4E78')

font_sub = Font(name='Calibri', size=10, italic=False, color='1F4E78')

font_key_owner = Font(name='Calibri', size=11, bold=True, color='B91C1C') # Red
font_co_owner = Font(name='Calibri', size=11, bold=True, color='15803D')  # Green
font_consulted = Font(name='Calibri', size=11, bold=False, color='D97706') # Amber
font_task = Font(name='Calibri', size=11, bold=True, color='000000')
font_desc = Font(name='Calibri', size=10, bold=False, color='000000')

align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
align_desc = Alignment(horizontal='left', vertical='center', wrap_text=True)

border_thin = Border(left=Side(style='thin', color='D9D9D9'),
                     right=Side(style='thin', color='D9D9D9'),
                     top=Side(style='thin', color='D9D9D9'),
                     bottom=Side(style='thin', color='D9D9D9'))

border_header = Border(left=Side(style='thin', color='1F4E78'),
                       right=Side(style='thin', color='1F4E78'),
                       top=Side(style='thin', color='1F4E78'),
                       bottom=Side(style='thin', color='1F4E78'))

# Set Column Widths (matching Financial Hub)
ws.column_dimensions['A'].width = 46.0
ws.column_dimensions['B'].width = 24.0
ws.column_dimensions['C'].width = 24.0
ws.column_dimensions['D'].width = 86.0

# Row 1: Main Title
ws.append(['MA TRẬN PHÂN ĐỊNH TRÁCH NHIỆM RACI: WEB PLATFORM x CELL TEAM (INSURTECH)', None, None, None])
ws.merge_cells('A1:D1')
ws.row_dimensions[1].height = 32
for col in range(1, 5):
    c = ws.cell(1, col)
    c.fill = fill_title
    c.font = font_title
    c.alignment = align_center

# Row 2: Principle Summary
ws.append(['• Web Platform: Xây dựng Sản Phẩm UI/UX, Lập Content Plan, Sản Xuất Nội Dung (Content Production), Track Luồng Web | • InsurTech (Cell Team): QC Content, QC Sản Phẩm, Cung Cấp API/Database, Track Luồng In-App, Growth Plan Tổng Thể', None, None, None])
ws.merge_cells('A2:D2')
ws.row_dimensions[2].height = 24
for col in range(1, 5):
    c = ws.cell(2, col)
    c.font = font_sub
    c.alignment = align_center

# Row 3: Header
headers = ['Hạng Mục / Đầu Việc Triển Khai', 'Web Platform Team', 'Cell Team (InsurTech BU)', 'Mô Tả Trách Nhiệm Chi Tiết & Tiêu Chí Nghiệm Thu']
ws.append(headers)
ws.row_dimensions[3].height = 28
for col in range(1, 5):
    c = ws.cell(3, col)
    c.fill = fill_header
    c.font = font_header
    c.alignment = align_center
    c.border = border_header

# Content Data Structure
content_data = [
    # Group I
    ('GROUP', 'NHÓM I: CHIẾN LƯỢC TĂNG TRƯỞNG & VẬN HÀNH SPRINT', None, None, None),
    ('TASK', '1.1 Hoạch Định Kế Hoạch Tăng Trưởng Tổng Thể (Growth Plan)', 'Consulted (C)', 'Key Owner (A/R)', 
     'InsurTech giữ vai trò Key Owner chịu trách nhiệm toàn diện về Growth Plan tổng thể, cấp ngân sách Marketing, chạy Paid Ads và cam kết chỉ tiêu kinh doanh (Giao dịch bảo hiểm, Nộp phạt in-app); Web Platform tham vấn phễu SEO.'),
    ('TASK', '1.2 Lập Kế Hoạch Sprint Web 1 Tháng & Nhịp Check-in 2 Tuần/Lần', 'Key Owner (A/R)', 'Co-Owner (A/R)', 
     'Web Platform chủ trì tổ chức Sprint Planning đầu tháng để chốt backlog tính năng Web; InsurTech đồng chủ trì check-in ngày 15 & 30 hàng tháng để rà soát tiến độ và tháo gỡ điểm nghẽn.'),
    
    # Group II
    ('GROUP', 'NHÓM II: THIẾT KẾ & XÂY DỰNG SẢN PHẨM WEB (PRODUCT UI/UX)', None, None, None),
    ('TASK', '2.1 Thiết Kế Giao Diện UI/UX & Phễu Chuyển Đổi (Product UI/UX)', 'Key Owner (A/R)', 'Consulted (C)', 
     'Web Platform chịu trách nhiệm thiết kế toàn bộ luồng trải nghiệm UI/UX cho 22 URLs thuộc hệ sinh thái (Trang chủ Master Hub, Cây xăng, Giá xăng, Đăng kiểm...), tối ưu hóa giao diện Mobile-first và Onelink W2A.'),
    ('TASK', '2.2 Lập Trình Frontend Web & Bộ Máy Tính Tiện Ích (Utilities & Calculators)', 'Key Owner (A/R)', 'Consulted (C)', 
     'Web Platform trực tiếp lập trình giao diện Web trên Next.js/MoSpark, phát triển bộ công cụ tương tác (Máy tính giá xăng 3 chế độ, Tính chi phí lăn bánh, Phong thủy biển số xe...); InsurTech tư vấn nghiệp vụ.'),
    ('TASK', '2.3 Kiểm Thử & Nghiệm Thu Chất Lượng Sản Phẩm Trước Go-Live (Product QC)', 'Contributor (R)', 'QC Owner (A/R)', 
     'InsurTech giữ vai trò Key Owner chịu trách nhiệm QC sản phẩm cuối cùng (Product QC), kiểm thử toàn diện luồng tương tác, tính ổn định và độ chính xác của giao diện trước khi bấm nút Go-live.'),
    
    # Group III
    ('GROUP', 'NHÓM III: KẾ HOẠCH NỘI DUNG, SẢN XUẤT & KIỂM DUYỆT (CONTENT PLAN & QC)', None, None, None),
    ('TASK', '3.1 Xây Dựng Kế Hoạch Nội Dung & Từ Khóa SEO (Content Plan)', 'Key Owner (A/R)', 'Consulted (C)', 
     'Web Platform chịu trách nhiệm lập kế hoạch nội dung toàn diện (Content Plan), nghiên cứu từ khóa theo 20 chủ đề giao thông và thiết lập lịch biên soạn cẩm nang định kỳ.'),
    ('TASK', '3.2 Trực Tiếp Sản Xuất Bài Viết Cẩm Nang & SEO (Content Production)', 'Key Owner (A/R)', 'Consulted (C)', 
     'Web Platform trực tiếp triển khai sản xuất bài viết (Content Production) qua MoSpark GenAI Pipeline, tối ưu hóa tiêu chuẩn E-E-A-T và cấu trúc dữ liệu Schema (FAQ, HowTo, LocalBusiness).'),
    ('TASK', '3.3 Kiểm Duyệt Chuyên Môn & Đảm Bảo Độ Chính Xác Nội Dung (Content QC)', 'Contributor (R)', 'QC Owner (A/R)', 
     'InsurTech giữ vai trò Key Owner kiểm duyệt chất lượng nội dung (Content QC), thẩm định tính chuẩn xác của các quy định pháp luật giao thông, mức phạt Nghị định 100/123 và nghiệp vụ bảo hiểm xe.'),
    
    # Group IV
    ('GROUP', 'NHÓM IV: HẠ TẦNG DỮ LIỆU ĐỊA ĐIỂM & KẾT NỐI HỆ THỐNG (DATA & API)', None, None, None),
    ('TASK', '4.1 Cung Cấp Hạ Tầng API Dữ Liệu & Nghiệp Vụ Giao Thông (Provide APIs)', 'Consulted (C)', 'Key Owner (A/R)', 
     'InsurTech chịu trách nhiệm làm việc với các đối tác (VTTI, Cơ quan đăng kiểm, ePass/VETC) để cung cấp API kết nối (nếu có), hỗ trợ cấp quyền truy cập và xử lý các lỗi kỹ thuật phía máy chủ.'),
    ('TASK', '4.2 Cung Cấp Cơ Sở Dữ Liệu Địa Điểm & Trạm Nhiên Liệu (Provide Database)', 'Consulted (C)', 'Key Owner (A/R)', 
     'InsurTech chịu trách nhiệm cung cấp và đồng bộ bộ dữ liệu chuẩn hóa (Database): danh mục 17.000 cây xăng, trạm sạc EV VinFast/V-Green, mạng lưới garage cứu hộ và biểu phí BOT để Web nạp vào CDN.'),
    ('TASK', '4.3 Tích Hợp API & Hiển Thị Bản Đồ Địa Điểm Trên Web (Data Integration)', 'Key Owner (A/R)', 'Consulted (C)', 
     'Web Platform chịu trách nhiệm tiếp nhận dữ liệu, xử lý thuật toán lọc tọa độ GPS (Haversine) và render mượt mà lên bản đồ tương tác Web (Mapbox / Google Maps).'),
    
    # Group V
    ('GROUP', 'NHÓM V: ĐO LƯỜNG ĐA KÊNH & THEO DÕI HÀNH TRÌNH CHUYỂN ĐỔI (DATA TRACKING)', None, None, None),
    ('TASK', '5.1 Cấu Hình Đo Lường & Tracking Toàn Bộ Luồng Web (Web Tracking)', 'Key Owner (A/R)', 'Consulted (C)', 
     'Web Platform chịu trách nhiệm toàn diện về hạ tầng đo lường Kênh Web (Web Tracking): cấu hình DataLayer, sự kiện GA4, GTM, UTM parameters, Appsflyer OneLink và Dynamic Desktop QR Code.'),
    ('TASK', '5.2 Kiểm Tra & Theo Dõi Luồng Người Dùng Phía Ứng Dụng (In-App Tracking)', 'Consulted (C)', 'Key Owner (A/R)', 
     'InsurTech chịu trách nhiệm 100% về In-App Tracking: tiếp nhận schema điều hướng từ Web sang App, theo dõi tỷ lệ rớt phễu thanh toán (Nộp phạt, Mua bảo hiểm, Nạp ePass) và tối ưu hóa CR (+10% Uplift).'),
    
    # Group VI
    ('GROUP', 'NHÓM VI: BÁO CÁO GIẢI TRÌNH & ĐIỀU PHỐI NGUỒN LỰC', None, None, None),
    ('TASK', '6.1 Báo Cáo Hiệu Quả Phễu Web-to-App Hàng Tháng Lên Ban Lãnh Đạo IVP', 'Co-Owner (A/R)', 'Key Owner (A/R)', 
     'InsurTech chủ trì báo cáo kết quả kinh doanh và tỷ lệ chuyển đổi In-App; Web Platform đồng chủ trì báo cáo lưu lượng truy cập (Traffic Web), thứ hạng SEO và tỷ lệ nhấp nút (CTR) lên Ban Lãnh Đạo IVP.')
]

current_row = 4
for item in content_data:
    if item[0] == 'GROUP':
        ws.append([item[1], None, None, None])
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
        ws.row_dimensions[current_row].height = 24
        for col in range(1, 5):
            c = ws.cell(current_row, col)
            c.fill = fill_group
            c.font = font_group
            c.alignment = align_left
            c.border = border_thin
    else:
        ws.append([item[1], item[2], item[3], item[4]])
        ws.row_dimensions[current_row].height = 36
        c1 = ws.cell(current_row, 1)
        c2 = ws.cell(current_row, 2)
        c3 = ws.cell(current_row, 3)
        c4 = ws.cell(current_row, 4)
        
        c1.font = font_task
        c1.alignment = align_left
        c1.border = border_thin
        
        # Color role badges
        for cell_role in [c2, c3]:
            val = cell_role.value
            cell_role.alignment = align_center
            cell_role.border = border_thin
            if 'Key Owner' in val or 'QC Owner' in val:
                cell_role.font = font_key_owner
            elif 'Co-Owner' in val or 'Contributor' in val:
                cell_role.font = font_co_owner
            elif 'Consulted' in val:
                cell_role.font = font_consulted
            else:
                cell_role.font = font_desc
                
        c4.font = font_desc
        c4.alignment = align_desc
        c4.border = border_thin
        
    current_row += 1

# Add spacing and Legend Table (exact like Financial Hub)
current_row += 1
ws.cell(current_row, 1, 'BẢNG CHÚ GIẢI KÝ HIỆU & NGUYÊN TẮC PHÂN CÔNG TRÁCH NHIỆM:')
ws.cell(current_row, 1).font = Font(name='Calibri', size=11, bold=True, color='1F4E78')
ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
current_row += 1

legend_headers = ['Thuật Ngữ RACI', 'Tên Tiếng Anh', 'Vai Trò Thực Tế Trong Dự Án', 'Định Nghĩa Phân Nhiệm']
ws.append(legend_headers)
ws.row_dimensions[current_row].height = 24
for col in range(1, 5):
    c = ws.cell(current_row, col)
    c.fill = fill_group
    c.font = font_group
    c.alignment = align_center
    c.border = border_thin
current_row += 1

legends = [
    ('Key Owner (A/R)', 'Accountable & Responsible', 'Chủ Trì & Chịu Trách Nhiệm Chính', 'Chịu trách nhiệm trực tiếp lập kế hoạch, thực thi và cam kết kết quả đầu ra cuối cùng.'),
    ('QC Owner (A/R)', 'Quality Control Owner', 'Chủ Trì Kiểm Duyệt & Nghiệm Thu', 'Có quyền phủ quyết (veto), chịu trách nhiệm nghiệm thu chất lượng trước khi Go-live.'),
    ('Co-Owner (A/R)', 'Co-Accountable & Responsible', 'Đồng Sở Hữu & Đồng Thực Thi', 'Cùng chịu trách nhiệm cam kết chỉ số và phối hợp xử lý điểm nghẽn liên phòng ban.'),
    ('Contributor (R)', 'Responsible Contributor', 'Đội Ngũ Thực Thi Hỗ Trợ', 'Trực tiếp thực hiện các đầu việc chuyên môn theo sự điều phối của Key Owner.'),
    ('Consulted (C)', 'Consulted', 'Tham Vấn Chuyên Môn 2 Chiều', 'Được lấy ý kiến chuyên môn bắt buộc; phản hồi đóng góp để hoàn thiện giải pháp.'),
    ('Informed (I)', 'Informed', 'Nhận Thông Tin 1 Chiều', 'Được cập nhật tiến độ định kỳ, không can thiệp trực tiếp vào quá trình thực thi.')
]

for leg in legends:
    ws.append([leg[0], leg[1], leg[2], leg[3]])
    ws.row_dimensions[current_row].height = 22
    c1 = ws.cell(current_row, 1)
    c2 = ws.cell(current_row, 2)
    c3 = ws.cell(current_row, 3)
    c4 = ws.cell(current_row, 4)
    
    if 'Key Owner' in leg[0] or 'QC Owner' in leg[0]:
        c1.font = font_key_owner
    elif 'Co-Owner' in leg[0] or 'Contributor' in leg[0]:
        c1.font = font_co_owner
    elif 'Consulted' in leg[0]:
        c1.font = font_consulted
    else:
        c1.font = font_desc
        
    c1.alignment = align_center
    c1.border = border_thin
    
    for c in [c2, c3]:
        c.font = font_desc
        c.alignment = align_center
        c.border = border_thin
        
    c4.font = font_desc
    c4.alignment = align_left
    c4.border = border_thin
    current_row += 1

wb_veh.save(file_veh)
print("Standardized Resources RACI sheet in vehicle-hub-roadmap.xlsx successfully!")
