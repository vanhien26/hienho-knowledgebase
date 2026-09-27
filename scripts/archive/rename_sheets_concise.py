import openpyxl

wb = openpyxl.load_workbook('05_HUBS/financial-hub-roadmap.xlsx')
print("Current sheet names:", wb.sheetnames)

# Mapping old names to short, concise names
name_mapping = {
    'Readme': 'Readme',
    'Roadmap': 'Roadmap',
    'Partner Banks Directory': 'Banks',
    'Consolidated Financial Markets': 'Markets',
    'RACI & Resource Allocation': 'RACI',
    'Market Difficulty & Quick Wins': 'Quick Wins'
}

for old_name, new_name in name_mapping.items():
    if old_name in wb.sheetnames:
        wb[old_name].title = new_name

print("Updated sheet names:", wb.sheetnames)
wb.save('05_HUBS/financial-hub-roadmap.xlsx')
print("Successfully renamed all sheet names to short and concise titles!")
