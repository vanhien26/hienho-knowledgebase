import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.load_workbook('05_HUBS/financial-hub-roadmap.xlsx')

# 1. UPDATE ROADMAP SHEET
ws_rd = wb['Roadmap']
for r in range(3, ws_rd.max_row + 1):
    feat_name = str(ws_rd.cell(row=r, column=4).value or '').strip()
    if 'Công Cụ Tài Chính Cá Nhân' in feat_name:
        ws_rd.cell(row=r, column=5).value = (
            "• Bộ Dự Toán Thưởng Tết & Lương Tháng 13 (tính chính xác số tiền Net thực nhận sau thuế TNCN lũy tiến vào dịp cuối năm).\n"
            "• Bộ Dự Toán Quyết Toán Thuế TNCN & Tiền Hoàn Thuế (ước tính số thuế nộp thừa được cơ quan thuế hoàn trả để chuyển sang tích sản sinh lời).\n"
            "• Bộ Quy Đổi Đơn Vị Vàng & Tính Lợi Nhuận Tích Sản Vàng (quy đổi đa chiều Ounce, Lượng, Cây, Chỉ, Gram và tính tiền lời/lỗ theo giá realtime)."
        )
        ws_rd.cell(row=r, column=6).value = (
            "Dòng tiền thưởng Tết & hoàn thuế ➔ Dẫn mở Túi Thần Tài / Gửi Tiết Kiệm Bản Việt & VPBank sinh lời 6.5%/năm hoặc Mua Vàng Tài Lộc MoMo."
        )
        # Adjust row height
        ws_rd.row_dimensions[r].height = 65.0

# 2. UPDATE MARKETS SHEET (Remove Bảo Hiểm Xã Hội if present)
ws_m = wb['Markets']
bhxh_row = None
for r in range(3, ws_m.max_row):
    mkt = str(ws_m.cell(row=r, column=2).value or '').strip()
    if 'Bảo Hiểm Xã Hội' in mkt:
        bhxh_row = r
        break

if bhxh_row:
    ws_m.delete_rows(bhxh_row)
    print(f"Deleted Bảo Hiểm Xã Hội row {bhxh_row} from Markets sheet.")

# Recalculate total volume and percentages in Markets
total_vol = 0
for r in range(3, ws_m.max_row):
    v = ws_m.cell(row=r, column=4).value
    if isinstance(v, int):
        total_vol += v

# Update STT, percentages, and total row
for idx, r in enumerate(range(3, ws_m.max_row), start=1):
    ws_m.cell(row=r, column=1).value = idx
    v = ws_m.cell(row=r, column=4).value
    if isinstance(v, int) and total_vol > 0:
        ws_m.cell(row=r, column=5).value = f"{(v / total_vol) * 100:.2f}%"

# Total row
r_tot = ws_m.max_row
ws_m.cell(row=r_tot, column=1).value = "TỔNG CỘNG FINANCIAL MASTER HUB"
ws_m.cell(row=r_tot, column=4).value = total_vol
ws_m.cell(row=r_tot, column=4).number_format = '0'
ws_m.cell(row=r_tot, column=5).value = "100.00%"

wb.save('05_HUBS/financial-hub-roadmap.xlsx')
print(f"Successfully updated Wealth Management scope: pure tools in Roadmap, new Markets total: {total_vol}!")
