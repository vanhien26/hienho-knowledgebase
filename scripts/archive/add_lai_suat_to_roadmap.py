import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook('05_HUBS/financial-hub-roadmap.xlsx')
ws_rd = wb['Roadmap']

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
section_font = Font(name='Arial', size=10.5, bold=True, color='1F4E78')
thin_side = Side(style='thin', color='64748B')
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

p1_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid") # Green
p2_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid") # Light Blue

status_done_fill = PatternFill(start_color="1E7E34", end_color="1E7E34", fill_type="solid")
status_done_font = Font(name='Arial', size=9.5, bold=True, color='FFFFFF')
status_prog_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
status_prog_font = Font(name='Arial', size=9.5, bold=True, color='B9770E')
status_plan_fill = PatternFill(start_color="1B4F72", end_color="1B4F72", fill_type="solid")
status_plan_font = Font(name='Arial', size=9.5, bold=True, color='FFFFFF')

font_bold = Font(name='Arial', size=9.5, bold=True)
font_regular = Font(name='Arial', size=9.5)

# Complete refreshed 11 roadmap items with Tính Lương marked as Done, and Lãi Suất added for Tháng 9/2026
roadmap_items = [
    # 0: Phase, 1: Timeline, 2: Track, 3: Product Name, 4: Tech & UX Scope, 5: W2A Hook, 6: Status
    ("Phase 1", "Tháng 8/2026", "Master Gateway", "Financial Hub",
     "• Master Gateway quy hoạch 3 Hero Product Banners nổi bật nhất tại Hero Zone: Điểm Tín Dụng CIC, Gửi Tiết Kiệm TKO, Thực Tập Sinh Đầu Tư.\n• Thanh Smart Financial Recommendation tự động đề xuất 1 trong 3 sản phẩm theo hành vi người dùng.\n• Module Dashboard hiển thị song song: Lãi suất Tiết kiệm cao nhất vs Tỷ suất lợi nhuận Quỹ Mở.\n• Cấu hình 301 Redirect toàn diện từ /trung-tam-tai-chinh sang /tai-chinh.",
     "Universal OneLink dẫn trực diện vào 3 phễu chính: Kiểm Tra CIC Miễn Phí / Mở Sổ Tiết Kiệm Bản Việt & VPBank / Mở Tài Khoản Thực Tập Sinh Đầu Tư.",
     "Done"),

    ("Phase 1", "Tháng 8/2026", "Sức Khỏe Tín Dụng & Tiết Kiệm", "CIC & Tiết Kiệm",
     "• Widget CIC Simulator đánh giá nhóm nợ 1-5 ➔ Gợi ý mở Sổ Tiết Kiệm để làm đẹp hồ sơ tín dụng.\n• Hero Widget Tính Lãi Tiết Kiệm TKO (Bản Việt / VPBank) ➔ Tích hợp tab so sánh 'Gửi Tiết Kiệm vs Thực Tập Sinh Đầu Tư'.",
     "• CIC: Nút 'Kiểm Tra Điểm Tín Dụng MoMo Miễn Phí' mở báo cáo in-app.\n• Tiết Kiệm: Nút 'Mở Sổ Tiết Kiệm Online' nhận thêm +0.2% lãi suất.",
     "Done"),

    ("Phase 1", "Tháng 8/2026", "Content & AI Pipeline", "Blog & Cẩm Nang Tài Chính",
     "• 100% bài viết cẩm nang nhúng Sticky Recommendation Bar ở chân trang dẫn về 1 trong 3 sản phẩm (CIC / Tiết Kiệm / Thực Tập Sinh Đầu Tư) theo ngữ cảnh chủ đề bài viết.\n• Cơ chế tự động inject In-text Banner & Calculator Widget liên quan.",
     "Dynamic In-text CTA & Sticky Bottom Dock dẫn thẳng vào luồng đăng ký 3 sản phẩm ưu tiên.",
     "In Progress"),

    ("Phase 1", "Tháng 8/2026", "Đầu Tư & Vàng", "Giá Vàng",
     "• Bảng giá vàng realtime đa thương hiệu SJC, PNJ, DOJI, 9999 kèm biểu đồ lịch sử.\n• Widget So Sánh Hiệu Suất Tích Sản: Đặt cạnh nhau 3 kênh: 'Vàng vs Gửi Tiết Kiệm vs Thực Tập Sinh Đầu Tư' ➔ Điều hướng người dùng sang Tiết Kiệm và Quỹ Mở.",
     "CTA 'Đa Dạng Hóa Tích Sản: Mở Sổ Tiết Kiệm 6.5%/năm Hoặc Đầu Tư Quỹ Mở Thực Tập Sinh Chỉ Từ 10k'.",
     "Done"),

    ("Phase 1", "Tháng 8/2026", "Thu Nhập & Thuế", "Tính Lương",
     "• Công cụ tính lương Gross sang Net đã hoàn thiện và đang hoạt động trực tiếp trên Kênh Web (Pure Client-side JS, không phụ thuộc API ngoài).\n• Đã nhúng cơ chế gợi ý thông minh (Smart Suggestion) sau khi tính lương: Gợi ý trích 20% lương gửi Tiết Kiệm TKO | Trích 10% tham gia Thực Tập Sinh Đầu Tư | Kiểm tra Điểm Tín Dụng CIC.",
     "Bộ 3 nút CTA trực tiếp sau khi tính lương: 'Mở Sổ Tiết Kiệm Bản Việt / VPBank' | 'Tích Sản Quỹ Mở SIP Chỉ Từ 10k' | 'Kiểm Tra Điểm Tín Dụng CIC'.",
     "Done"),

    ("Phase 1", "Tháng 9/2026", "Tiết Kiệm & Lãi Suất", "Lãi Suất Ngân Hàng",
     "• Xây dựng Bảng So Sánh Lãi Suất Tiền Gửi 30+ Ngân Hàng (/lai-suat) theo các kỳ hạn 1 - 36 tháng (Online TKO vs Tại quầy).\n• Ghim sản phẩm Tiết Kiệm Online MoMo (Bản Việt & VPBank) ở vị trí Top 1 kèm nhãn ưu đãi '+0.2% Lãi Suất Khi Mở Trên MoMo'.\n• Tab đối chiếu: 'So sánh Gửi Tiết Kiệm (5%/năm) vs Thực Tập Sinh Đầu Tư Quỹ Mở (10-15%/năm)' để kéo user trẻ tích sản.",
     "Nút CTA 'Mở Sổ Tiết Kiệm Online Bản Việt / VPBank Lãi Suất Ưu Đãi' & 'Tích Sản SIP Cùng Thực Tập Sinh Đầu Tư'.",
     "In Progress"),

    ("Phase 1", "Tháng 9/2026", "Thu Nhập & Quản Lý Chi Tiêu", "Phân Bổ Lương",
     "• Công cụ phân bổ 50/30/20 tự động mapping:\n  • Hũ Tích Lũy 20% ➔ Suggest trực diện mở Sổ Tiết Kiệm Bản Việt / VPBank.\n  • Hũ Đầu Tư 10% ➔ Suggest tham gia Thực Tập Sinh Đầu Tư Quỹ Mở.\n  • Đo lường chỉ số an toàn đòn bẩy ➔ Suggest kiểm tra CIC.",
     "CTA 1-Click: 'Tự Động Trích Tiền Sang Tiết Kiệm TKO' và 'Kích Hoạt Gói Thực Tập Sinh Đầu Tư'.",
     "In Progress"),

    ("Phase 1", "Tháng 9/2026", "Tiền Tệ & Ngoại Hối", "Tỷ Giá & Ngoại Tệ",
     "• Cổng tỷ giá master (/tai-chinh/ty-gia) và programmatic subpages theo từng cặp tiền (/tai-chinh/ty-gia/[pair-slug]) cho 20+ ngoại tệ (USD/VND, JPY/VND, EUR/VND...).\n• Widget Phân Tích Lợi Suất: So sánh 'Giữ Ngoại Tệ vs Gửi Tiết Kiệm VND Lãi Cao' ➔ Kéo traffic về sản phẩm Tiết Kiệm Online.",
     "CTA 'Tối Ưu Dòng Tiền: Đổi Sang VND Gửi Tiết Kiệm Lãi Suất Đến 6.5%' hoặc 'Mở Thẻ Quốc Tế Visa/Mastercard MoMo'.",
     "In Progress"),

    ("Phase 1", "Tháng 9/2026", "Đầu Tư & Quỹ Mở F0", "Chứng Khoán & Quỹ Mở",
     "• Cổng thông tin chính thức của Chương Trình 'Thực Tập Sinh Đầu Tư' (Chứng khoán CVX & Quỹ mở SIP Dragon Capital, VinaCapital, SSIAM).\n• Bộ giả lập Lãi Kép Tích Lũy Định Kỳ từ 10k/ngày.\n• Tab so sánh cân bằng danh mục: Kết hợp Thực Tập Sinh Đầu Tư với Gửi Tiết Kiệm.",
     "CTA Trọng Tâm: 'Đăng Ký Tham Gia Thực Tập Sinh Đầu Tư - Nhận Quà Khởi Nghiệp' & 'Mở Sổ Tiết Kiệm Cân Bằng'.",
     "Planned"),

    ("Phase 2", "Tháng 10/2026", "Dòng Tiền & Tiêu Dùng", "Quản Lý Dòng Tiền & Tiêu Dùng",
     "• Bộ tính Dòng Tiền Tự Do (Free Cash Flow) & Hạn mức tiêu dùng ➔ Tự động phân luồng dòng tiền dư thừa:\n  • Tiền dự phòng ➔ Đề xuất gửi Tiết Kiệm Online kỳ hạn linh hoạt.\n  • Tiền nhàn rỗi dài hạn ➔ Đề xuất gói Thực Tập Sinh Đầu Tư SIP.\n  • Cảnh báo nợ quá hạn ➔ Đề xuất kiểm tra CIC.",
     "CTA 'Phân Bổ Dòng Tiền Thặng Dư: Gửi Tiết Kiệm Online / Tích Sản SIP Đầu Tư / Kiểm Tra Sức Khỏe CIC'.",
     "Planned"),

    ("Phase 2", "Tháng 10/2026", "Tích Sản & Hưu Trí", "Đầu Tư & Tích Sản FIRE",
     "• Bộ Giả Lập Tự Do Tài Chính & Nghỉ Hưu Sớm (FIRE Simulator).\n• Cơ chế phân bổ vốn mục tiêu theo nguyên tắc 2 Trụ Cột: Vốn Phòng Thủ (Gửi Tiết Kiệm 30 Bank) + Vốn Tăng Trưởng (Thực Tập Sinh Đầu Tư Quỹ Mở).\n• Đo lường điểm sức khỏe tín dụng CIC để tối ưu chi phí sử dụng vốn.",
     "CTA 'Xây Dựng Danh Mục FIRE: 50% Gửi Tiết Kiệm An Toàn + 50% Thực Tập Sinh Đầu Tư Tăng Trưởng'.",
     "Planned")
]

if 'Roadmap' in wb.sheetnames:
    del wb['Roadmap']

ws_rd = wb.create_sheet('Roadmap', index=1)
ws_rd.append(["LỘ TRÌNH TRIỂN KHAI TOÀN DIỆN CÁC TÍNH NĂNG FINANCIAL MASTER HUB (2026 - 2027)"])
ws_rd.merge_cells("A1:G1")
ws_rd.cell(row=1, column=1).fill = section_fill
ws_rd.cell(row=1, column=1).font = section_font
ws_rd.row_dimensions[1].height = 24.0

rd_headers = [
    "Giai Đoạn (Phase)",
    "Thời Gian",
    "Cụm Sản Phẩm (Track)",
    "Tên Tính Năng / Sản Phẩm",
    "Chi Tiết Triển Khai Kỹ Thuật & UX Scope",
    "Cơ Chế Chuyển Đổi Web-to-App (W2A Hook)",
    "Trạng Thái"
]
ws_rd.append(rd_headers)
for c_i in range(1, len(rd_headers) + 1):
    c = ws_rd.cell(row=2, column=c_i)
    c.fill = header_fill; c.font = header_font; c.border = thin_border
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws_rd.row_dimensions[2].height = 28.0

for row_item in roadmap_items:
    ws_rd.append(list(row_item))
    r = ws_rd.max_row
    
    num_lines = max(row_item[4].count('\n') + 1, row_item[5].count('\n') + 1)
    ws_rd.row_dimensions[r].height = max(55.0, num_lines * 18.0)
    
    ws_rd.cell(row=r, column=1).alignment = Alignment(horizontal='center', vertical='center')
    ws_rd.cell(row=r, column=2).alignment = Alignment(horizontal='center', vertical='center')
    ws_rd.cell(row=r, column=3).alignment = Alignment(horizontal='center', vertical='center'); ws_rd.cell(row=r, column=3).font = font_bold
    ws_rd.cell(row=r, column=4).alignment = Alignment(horizontal='left', vertical='center'); ws_rd.cell(row=r, column=4).font = font_bold
    ws_rd.cell(row=r, column=5).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws_rd.cell(row=r, column=6).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws_rd.cell(row=r, column=7).alignment = Alignment(horizontal='center', vertical='center')

    phase = row_item[0]
    if "Phase 1" in phase: ws_rd.cell(row=r, column=1).fill = p1_fill
    elif "Phase 2" in phase: ws_rd.cell(row=r, column=1).fill = p2_fill

    status_val = row_item[6]
    c_st = ws_rd.cell(row=r, column=7)
    if status_val == "Done": c_st.fill = status_done_fill; c_st.font = status_done_font
    elif status_val == "In Progress": c_st.fill = status_prog_fill; c_st.font = status_prog_font
    elif status_val == "Planned": c_st.fill = status_plan_fill; c_st.font = status_plan_font

    for c_i in range(1, 8): ws_rd.cell(row=r, column=c_i).border = thin_border

col_widths = [14, 16, 28, 28, 75, 48, 16]
for idx, w in enumerate(col_widths, start=1):
    ws_rd.column_dimensions[get_column_letter(idx)].width = w

wb.save('05_HUBS/financial-hub-roadmap.xlsx')
print("Successfully updated Roadmap: Tính Lương marked Done, Lãi Suất added for Sept 2026!")
