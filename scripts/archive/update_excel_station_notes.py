import openpyxl

file_path = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/vehicle-hub-roadmap.xlsx'
wb = openpyxl.load_workbook(file_path)

ws_road = wb['Roadmap']
for r in range(3, ws_road.max_row + 1):
    val = str(ws_road.cell(r, 4).value or '')
    if 'Cây Xăng' in val:
        new_scope = (
            "• Trang chính (/cay-xang) & cụm 10 trang con trọng điểm (Local SEO & pSEO):\n"
            "  - 4 Tỉnh/Thành: TP.HCM, Hà Nội, Đà Nẵng, Bình Dương.\n"
            "  - 6 Quận/Đường lớn: Q.1, Q.7, Thủ Đức, Cầu Giấy, Đống Đa, Quốc Lộ 1A.\n"
            "• 2 YÊU CẦU BẮT BUỘC:\n"
            "  1. Component Filter Engine: Lọc đa thuộc tính (Khu vực, Brand, Mở 24/24, Thanh toán MoMo, Xăng Euro 5).\n"
            "  2. Content SEO/GEO: Tự động render Schema LocalBusiness, FAQPage, Dynamic Metadata theo khu vực."
        )
        ws_road.cell(r, 5, value=new_scope)
        break

wb.save(file_path)
print("Updated Excel Roadmap with 2 critical requirements successfully!")
