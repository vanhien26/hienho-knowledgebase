import openpyxl

wb = openpyxl.load_workbook('/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/vehicle-hub-roadmap.xlsx')
ws = wb['Roadmap']

for col in range(1, ws.max_column + 1):
    col_letter = openpyxl.utils.get_column_letter(col)
    print(f"Col {col_letter} width: {ws.column_dimensions[col_letter].width}")

c = ws.cell(2, 1)
print("Header fill:", c.fill.start_color.rgb, "Font:", c.font.name, c.font.size, c.font.color.rgb if c.font.color else None)

c_row3 = ws.cell(3, 1)
print("Data fill:", c_row3.fill.start_color.rgb if c_row3.fill else None, "Font:", c_row3.font.name, c_row3.font.size)
