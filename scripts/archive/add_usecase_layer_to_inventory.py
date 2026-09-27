import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.load_workbook('/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/inventory.xlsx')
ws1 = wb['Inventory']
ws2 = wb['Quick Win Action Playbook']
ws3 = wb['Competitor Intelligence']
ws4 = wb['Division Rollup']

lead_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
strike_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
gap_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
qw_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
bb_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
tc_fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
dp_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
sub_header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")

thin_side = Side(style='thin', color='CBD5E1')
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
font_bold = Font(name='Arial', size=10, bold=True, color='000000')
font_vol_bold_blue = Font(name='Arial', size=10, bold=True, color='002060')
font_header_white = Font(name='Arial', size=10, bold=True, color='FFFFFF')

use_case_mapping = {
    # PS
    'Xổ Số Kiến Thiết': 'Lottery',
    'Vietlott (Xổ số điện toán)': 'Lottery',
    'Nạp Game': 'Digital Entertainment',
    'Nạp xu Tiktok': 'Digital Entertainment',
    'Dịch vụ công': 'Public Services',
    'Phạt nguội': 'Phạt Nguội',
    'Đặt lịch khám bệnh & Bệnh viện': 'Billpay',
    'Điện': 'Billpay',
    'Nước': 'Billpay',
    'Internet': 'Billpay',
    'Truyền hình': 'Billpay',
    'Đánh giá năng lực': 'Billpay',
    'Sim Chính Chủ (Đăng ký sim & Sim số đẹp)': 'Telco',
    'Nạp Data (Gói cước 3G/4G/5G)': 'Telco',
    'Sim Du Lịch (eSIM & Sim Quốc Tế)': 'Telco',
    'Nạp tiền điện thoại': 'Telco',
    'Chuyển tiền': 'P2P',
    'Soundbox': 'Offline Payment',
    'Thanh toán quốc tế': 'Cross Border',
    
    # CreditTech
    'Giá vàng': 'Gold & Commodity',
    'Tỷ giá': 'Forex',
    'Ngoại Tệ': 'Forex',
    'Ngân hàng': 'Banking',
    'Lãi suất': 'Banking',
    'Gửi tiết kiệm': 'Banking',
    'Thẻ Visa': 'Banking',
    'Thẻ ghi nợ (ATM)': 'Banking',
    'Napas': 'Banking',
    'Mastercard': 'Banking',
    'Mã số thuế': 'Tax',
    'Thuế (TNCN)': 'Tax',
    'Chứng khoán': 'Stock & Investment',
    'Cổ phiếu': 'Stock & Investment',
    'Trái phiếu': 'Stock & Investment',
    'Chứng chỉ quỹ': 'Stock & Investment',
    'Thẻ tín dụng': 'Loan & Credit',
    'Vay tín chấp (Vay nhanh)': 'Loan & Credit',
    'Vay thế chấp (Mua nhà)': 'Loan & Credit',
    'Trả góp': 'Loan & Credit',
    'Điểm tín dụng CIC': 'Loan & Credit',
    'Nợ xấu': 'Loan & Credit',
    'Tính lương': 'Payroll',
    'Crypto': 'Digital Assets',
    'Tiền số & Tiền ảo': 'Digital Assets',
    
    # InsurTech
    'Giá xăng dầu (Vehicle Hub)': 'Vehicle Hub',
    'Trạm sạc xe điện EV (Vehicle Hub)': 'Vehicle Hub',
    'Biển số xe & Đấu giá (Vehicle Hub)': 'Vehicle Hub',
    'Sửa xe & Cứu hộ giao thông (Vehicle Hub)': 'Vehicle Hub',
    'Ôn thi bằng lái xe (Vehicle Hub)': 'Vehicle Hub',
    'Tìm cây xăng (Vehicle Hub)': 'Vehicle Hub',
    'Gara ô tô & Bảo dưỡng (Vehicle Hub)': 'Vehicle Hub',
    'Rửa xe & Chăm sóc xe (Vehicle Hub)': 'Vehicle Hub',
    'Bãi đỗ xe & Giữ xe (Vehicle Hub)': 'Vehicle Hub',
    'Thu phí không dừng (ePass / VETC)': 'Vehicle Hub',
    'Định giá xe (Vehicle Hub)': 'Vehicle Hub',
    'Bảo Hiểm Y Tế': 'Social Insurance',
    'Bảo hiểm xã hội': 'Social Insurance',
    'Bảo hiểm Ô tô': 'Commercial Insurance',
    'Bảo hiểm xe máy': 'Commercial Insurance',
    'Bảo hiểm sức khỏe +': 'Commercial Insurance',
    'Bảo hiểm Du lịch': 'Commercial Insurance',
    'Thanh toán phí bảo hiểm': 'Commercial Insurance',
    
    # MDS
    'Cinema': 'Cinema',
    'Vé xe khách': 'OTA',
    'Vé máy bay': 'OTA',
    'Du lịch - Đi lại': 'OTA',
    'Vé tàu hỏa': 'OTA',
    'Thuê xe tự lái': 'OTA',
    'Đặt khách sạn': 'OTA',
    
    # GPD
    'Student Hub (Trường Đại Học & Cẩm nang sinh viên)': 'Student Hub',
    
    # Other
    'Lịch vạn niên & Lịch âm': 'Lifestyle & Almanac',
    'Mã giảm giá & Voucher': 'Promotion'
}

# 1. Update Sheet 1: Inventory (Add Use Case Column)
s1_rows = []
for r in range(2, ws1.max_row + 1):
    div = ws1.cell(row=r, column=2).value
    name = str(ws1.cell(row=r, column=3).value or '')
    vol = ws1.cell(row=r, column=4).value
    impact = ws1.cell(row=r, column=5).value
    win = ws1.cell(row=r, column=6).value
    posture = ws1.cell(row=r, column=7).value
    momo_pos = ws1.cell(row=r, column=8).value
    top5 = ws1.cell(row=r, column=9).value
    season = ws1.cell(row=r, column=10).value
    uc = use_case_mapping.get(name, 'General')
    s1_rows.append([r-1, div, uc, name, vol, impact, win, posture, momo_pos, top5, season])

while ws1.max_row > 0:
    ws1.delete_rows(1)

ws1_headers = ['STT', 'Division', 'Use Case', 'Thị Trường (Dịch Vụ Cụ Thể)', 'Search Volume / Tháng', 'Mức Độ Tác Động', 'Khả Năng Thắng', 'Định Vị Thực Thi', 'Vị Thế MoMo', 'Top 5 Đối Thủ Google SERP', 'Mùa Vụ & Yếu Tố Thúc Đẩy']
ws1.append(ws1_headers)
ws1.row_dimensions[1].height = 28.0
for c_i, h in enumerate(ws1_headers, start=1):
    cell = ws1.cell(row=1, column=c_i)
    cell.fill = header_fill
    cell.font = font_header_white
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border

for idx, r_data in enumerate(s1_rows, start=2):
    r_data[0] = f"#{idx - 1}"
    for c_i, val in enumerate(r_data, start=1):
        c = ws1.cell(row=idx, column=c_i, value=val)
        c.border = thin_border
        if c_i in [1, 2, 3, 6, 7, 8, 9]: c.alignment = Alignment(horizontal='center', vertical='center')
        elif c_i == 5: c.alignment = Alignment(horizontal='right', vertical='center'); c.number_format = '#,##0'; c.font = font_vol_bold_blue
        else: c.alignment = Alignment(horizontal='left', vertical='center')
        if c_i in [1, 2, 3, 4, 8, 9]: c.font = font_bold
        
        if c_i == 8:
            if "Đánh Chiếm Ngay" in str(val): c.fill = qw_fill
            elif "Nuôi Dài Hạn" in str(val): c.fill = bb_fill
            elif "Bổ Trợ" in str(val): c.fill = tc_fill
            elif "Tạm Chưa" in str(val): c.fill = dp_fill
        if c_i == 9:
            if "Dẫn Đầu" in str(val): c.fill = lead_fill
            elif "Bám Đuổi" in str(val): c.fill = strike_fill
            else: c.fill = gap_fill

ws1.column_dimensions['A'].width = 8
ws1.column_dimensions['B'].width = 14
ws1.column_dimensions['C'].width = 22
ws1.column_dimensions['D'].width = 44
ws1.column_dimensions['E'].width = 22
ws1.column_dimensions['F'].width = 26
ws1.column_dimensions['G'].width = 26
ws1.column_dimensions['H'].width = 26
ws1.column_dimensions['I'].width = 28
ws1.column_dimensions['J'].width = 60
ws1.column_dimensions['K'].width = 45

# 2. Update Sheet 2: Quick Win Action Playbook (Add Use Case Column)
s2_rows = []
for r in range(3, ws2.max_row + 1):
    div = ws2.cell(row=r, column=2).value
    name = str(ws2.cell(row=r, column=3).value or '')
    vol = ws2.cell(row=r, column=4).value
    win = ws2.cell(row=r, column=5).value
    season = ws2.cell(row=r, column=6).value
    url = ws2.cell(row=r, column=7).value
    tool = ws2.cell(row=r, column=8).value
    gap = ws2.cell(row=r, column=9).value
    uc = use_case_mapping.get(name, 'General')
    s2_rows.append([r-2, div, uc, name, vol, win, season, url, tool, gap])

while ws2.max_row > 0:
    ws2.delete_rows(1)

ws2.append([f"KẾ HOẠCH HÀNH ĐỘNG {len(s2_rows)} THỊ TRƯỜNG ĐÁNH CHIẾM NGAY (QUICK WINS - XẾP HẠNG #1 ĐẾN #{len(s2_rows)})"])
ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
ws2.cell(row=1, column=1).fill = header_fill
ws2.cell(row=1, column=1).font = font_header_white
ws2.cell(row=1, column=1).alignment = Alignment(horizontal='center', vertical='center')
ws2.row_dimensions[1].height = 32.0

ws2_headers = ['Thứ Hạng', 'Division', 'Use Case', 'Thị Trường Đánh Chiếm', 'Search Volume / Tháng', 'Mức Độ Dễ Thắng', 'Mùa Cao Điểm', 'URL Route Dự Kiến', 'Công Cụ / Trải Nghiệm Tương Tác', 'Khoảng Trống Cạnh Tranh (Moat)']
ws2.append(ws2_headers)
ws2.row_dimensions[2].height = 28.0
for c_i, h in enumerate(ws2_headers, start=1):
    cell = ws2.cell(row=2, column=c_i)
    cell.fill = sub_header_fill
    cell.font = font_header_white
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border

for idx, r_data in enumerate(s2_rows, start=3):
    r_data[0] = f"#{idx - 2}"
    for c_i, val in enumerate(r_data, start=1):
        c = ws2.cell(row=idx, column=c_i, value=val)
        c.border = thin_border
        if c_i in [1, 2, 3, 6]: c.alignment = Alignment(horizontal='center', vertical='center')
        elif c_i == 5: c.alignment = Alignment(horizontal='right', vertical='center'); c.number_format = '#,##0'; c.font = font_vol_bold_blue
        else: c.alignment = Alignment(horizontal='left', vertical='center')
        if c_i in [1, 2, 3, 4]: c.font = font_bold
        if c_i == 6: c.fill = qw_fill
        if c_i in [9, 10]: c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

ws2.column_dimensions['A'].width = 12
ws2.column_dimensions['B'].width = 14
ws2.column_dimensions['C'].width = 22
ws2.column_dimensions['D'].width = 44
ws2.column_dimensions['E'].width = 22
ws2.column_dimensions['F'].width = 25
ws2.column_dimensions['G'].width = 30
ws2.column_dimensions['H'].width = 32
ws2.column_dimensions['I'].width = 45
ws2.column_dimensions['J'].width = 50

# 3. Create / Rebuild Sheet 5: Use Case Rollup
if 'Use Case Rollup' in wb.sheetnames:
    del wb['Use Case Rollup']
ws5 = wb.create_sheet(title='Use Case Rollup')

ws5.append(['TỔNG HỢP THEO USE CASE TOÀN BỘ HỆ SINH THÁI MOMO 2026'])
ws5.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
ws5.cell(row=1, column=1).fill = header_fill
ws5.cell(row=1, column=1).font = font_header_white
ws5.cell(row=1, column=1).alignment = Alignment(horizontal='center', vertical='center')
ws5.row_dimensions[1].height = 32.0

ws5_headers = ['STT', 'Division', 'Use Case', 'Tổng Search Volume / Tháng', 'Số Lượng Thị Trường', 'Danh Sách Thị Trường Thành Viên']
ws5.append(ws5_headers)
ws5.row_dimensions[2].height = 28.0
for c_i, h in enumerate(ws5_headers, start=1):
    cell = ws5.cell(row=2, column=c_i)
    cell.fill = sub_header_fill
    cell.font = font_header_white
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border

uc_map = {}
for r_data in s1_rows:
    div = r_data[1]
    uc = r_data[2]
    name = r_data[3]
    vol = r_data[4]
    key = (div, uc)
    if key not in uc_map:
        uc_map[key] = {'vol': 0, 'count': 0, 'markets': []}
    uc_map[key]['vol'] += vol
    uc_map[key]['count'] += 1
    uc_map[key]['markets'].append(name)

sorted_ucs = sorted(uc_map.items(), key=lambda x: -x[1]['vol'])
for idx, ((d_name, uc_name), val) in enumerate(sorted_ucs, start=3):
    names_str = ', '.join(val['markets'])
    ws5.append([idx - 2, d_name, uc_name, val['vol'], val['count'], names_str])
    r = ws5.max_row
    ws5.row_dimensions[r].height = 24.0
    ws5.cell(row=r, column=1).alignment = Alignment(horizontal='center', vertical='center'); ws5.cell(row=r, column=1).number_format = '0'
    ws5.cell(row=r, column=2).alignment = Alignment(horizontal='center', vertical='center'); ws5.cell(row=r, column=2).font = font_bold
    ws5.cell(row=r, column=3).alignment = Alignment(horizontal='left', vertical='center'); ws5.cell(row=r, column=3).font = font_bold
    ws5.cell(row=r, column=4).alignment = Alignment(horizontal='right', vertical='center'); ws5.cell(row=r, column=4).number_format = '#,##0'; ws5.cell(row=r, column=4).font = font_vol_bold_blue
    ws5.cell(row=r, column=5).alignment = Alignment(horizontal='center', vertical='center'); ws5.cell(row=r, column=5).font = font_bold
    ws5.cell(row=r, column=6).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    for c_i in range(1, 7): ws5.cell(row=r, column=c_i).border = thin_border

ws5.column_dimensions['A'].width = 8
ws5.column_dimensions['B'].width = 14
ws5.column_dimensions['C'].width = 24
ws5.column_dimensions['E'].width = 20
ws5.column_dimensions['D'].width = 24
ws5.column_dimensions['F'].width = 75

wb.save('/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/inventory.xlsx')
print("Successfully added Use Case layer and created Sheet 5 'Use Case Rollup' in inventory.xlsx!")
