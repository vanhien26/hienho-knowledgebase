import openpyxl

file_cs = '/Users/hienhv/Downloads/[WP] Vehicle Hub - Content Strategy.xlsx'
wb_cs = openpyxl.load_workbook(file_cs, data_only=True)
ws_kw = wb_cs['Keyword']

rescue_kws = []
for r in range(2, ws_kw.max_row + 1):
    kw = str(ws_kw.cell(r, 1).value or ws_kw.cell(r, 2).value or '').strip().lower()
    if 'cứu hộ' in kw or 'cuu ho' in kw or 'kéo xe' in kw or 'kích bình' in kw or 'vá vỏ' in kw:
        vol = ws_kw.cell(r, 2).value or ws_kw.cell(r, 3).value or 0
        try:
            vol = float(vol)
        except:
            vol = 0
        rescue_kws.append((kw, vol))

rescue_kws.sort(key=lambda x: x[1], reverse=True)
print(f"Total rescue keywords: {len(rescue_kws)}, Sum Vol: {sum(x[1] for x in rescue_kws):,.0f}")
for k, v in rescue_kws[:15]:
    print(f"  • {k}: {v:,.0f}")

