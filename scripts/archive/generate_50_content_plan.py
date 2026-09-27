import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook('05_HUBS/financial-hub-roadmap.xlsx')

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
section_font = Font(name='Arial', size=10.5, bold=True, color='1F4E78')
thin_side = Side(style='thin', color='64748B')
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

font_bold = Font(name='Arial', size=9.5, bold=True)
font_regular = Font(name='Arial', size=9.5)
font_vol = Font(name='Arial', size=9.5, bold=True, color='002060')

status_plan_fill = PatternFill(start_color="1B4F72", end_color="1B4F72", fill_type="solid")
status_plan_font = Font(name='Arial', size=9.5, bold=True, color='FFFFFF')

# 50 Detailed Financial Articles across 5 Pillars
content_articles = [
    # --- KHỐI I: TIẾT KIỆM & LÃI SUẤT (10 Bài) ---
    (1, "Tiết Kiệm & Lãi Suất", "Bảng Lãi Suất Ngân Hàng Nào Cao Nhất Tháng 9/2026? So Sánh 30+ Bank", "lãi suất ngân hàng", 135000, "Cẩm Nang & Bảng So Sánh", "/tai-chinh/lai-suat-ngan-hang-cao-nhat", "Mở Sổ Tiết Kiệm Bản Việt / VPBank", "Tuần 1", "Planned"),
    (2, "Tiết Kiệm & Lãi Suất", "Gửi Tiết Kiệm Online Có An Toàn Không? 5 Tiêu Chí Chọn Ngân Hàng", "gửi tiết kiệm online", 22000, "Hướng Dẫn & FAQ", "/tai-chinh/gui-tiet-kiem-online-an-toan", "Mở Sổ Tiết Kiệm Online MoMo", "Tuần 1", "Planned"),
    (3, "Tiết Kiệm & Lãi Suất", "Cách Tính Tiền Lãi Tiết Kiệm Không Kỳ Hạn & Có Kỳ Hạn Chính Xác", "cách tính lãi suất tiết kiệm", 18100, "Công Thức & Tool Nhúng", "/tai-chinh/cach-tinh-lai-suat-tiet-kiem", "Hero Widget Tính Lãi TKO", "Tuần 1", "Planned"),
    (4, "Tiết Kiệm & Lãi Suất", "So Sánh Lãi Suất Tiền Gửi Bản Việt (BVBank) và VPBank Trên MoMo", "lãi suất bản việt vpbank", 14800, "So Sánh Chuyên Sâu", "/tai-chinh/so-sanh-tiet-kiem-ban-viet-vpbank", "Mở Sổ Nhận Thêm +0.2% Lãi", "Tuần 2", "Planned"),
    (5, "Tiết Kiệm & Lãi Suất", "Nên Gửi Tiết Kiệm Kỳ Hạn 6 Tháng Hay 12 Tháng? Phân Tích Lợi Suất", "gửi tiết kiệm 6 tháng hay 12 tháng", 12100, "Tư Vấn Tài Chính", "/tai-chinh/nen-gui-tiet-kiem-ky-han-nao", "Widget So Sánh Kỳ Hạn", "Tuần 2", "Planned"),
    (6, "Tiết Kiệm & Lãi Suất", "Lãi Suất Tiết Kiệm Tích Lũy Hàng Tháng Là Gì? Bí Quyết Cho Người Trẻ", "tiết kiệm tích lũy hàng tháng", 9900, "Cẩm Nang F0", "/tai-chinh/tiet-kiem-tich-luy-hang-thang", "Mở Túi Thần Tài Sinh Lời", "Tuần 2", "Planned"),
    (7, "Tiết Kiệm & Lãi Suất", "Rút Tiền Tiết Kiệm Trước Hạn Có Bị Mất Lãi Không? Quy Định Mới NHNN", "rút tiết kiệm trước hạn", 8100, "Pháp Lý & FAQ", "/tai-chinh/rut-tiet-kiem-truoc-han", "Mở Sổ Rút Gốc Linh Hoạt", "Tuần 3", "Planned"),
    (8, "Tiết Kiệm & Lãi Suất", "Có 50 Triệu Nên Gửi Tiết Kiệm Hay Mở Túi Thần Tài? So Sánh Dòng Tiền", "có 50 triệu nên gửi tiết kiệm", 6600, "Bài Toán Thực Tế", "/tai-chinh/co-50-trieu-nen-gui-tiet-kiem-khong", "Mở Túi Thần Tài / Tiết Kiệm", "Tuần 3", "Planned"),
    (9, "Tiết Kiệm & Lãi Suất", "Sổ Tiết Kiệm Điện Tử Là Gì? Thủ Tục Mở & Tất Toán 100% Online", "sổ tiết kiệm điện tử", 5400, "Hướng Dẫn Thao Tác", "/tai-chinh/so-tiet-kiem-dien-tu-la-gi", "Mở Sổ Tiết Kiệm In-App", "Tuần 4", "Planned"),
    (10, "Tiết Kiệm & Lãi Suất", "Lãi Suất Bậc Thang Là Gì? Cách Tối Đa Tiền Lời Cho Khoản Tiền Lớn", "lãi suất bậc thang", 4400, "Chiến Lược Tiết Kiệm", "/tai-chinh/lai-suat-bac-thang", "Bảng So Sánh Lãi Suất 30 Bank", "Tuần 4", "Planned"),

    # --- KHỐI II: ĐẦU TƯ & TÍCH SẢN (12 Bài) ---
    (11, "Đầu Tư & Tích Sản", "Phân Biệt Vàng 24K, 9999, Vàng Tây Và Vàng Nhẫn Trơn Khi Đầu Tư", "vàng 24k và 9999", 74000, "Kiến Thức Vàng", "/tai-chinh/phan-biet-cac-loai-vang-dau-tu", "Mua Vàng Nhẫn Tài Lộc", "Tuần 1", "Planned"),
    (12, "Đầu Tư & Tích Sản", "Nên Mua Vàng Miếng SJC Hay Vàng Nhẫn Trơn Tích Trữ Lúc Này?", "nên mua vàng sjc hay vàng nhẫn", 49500, "Tư Vấn Thị Trường", "/tai-chinh/nen-mua-vang-sjc-hay-vang-nhan", "Bảng Giá Vàng Realtime", "Tuần 1", "Planned"),
    (13, "Đầu Tư & Tích Sản", "1 Lượng Vàng Bằng Bao Nhiêu Cây, Chỉ, Gram? Bảng Quy Đổi Chuẩn", "1 lượng vàng bao nhiêu chỉ", 40500, "Tra Cứu Nhanh", "/tai-chinh/quy-doi-don-vi-vang", "Bộ Quy Đổi Đơn Vị Vàng", "Tuần 1", "Planned"),
    (14, "Đầu Tư & Tích Sản", "Hướng Dẫn Mở Tài Khoản Chứng Khoán Cho Người Mới Bắt Đầu (F0)", "mở tài khoản chứng khoán", 33100, "Cẩm Nang F0", "/tai-chinh/huong-dan-mo-tai-khoan-chung-khoan", "Mở Tài Khoản CVX Trên MoMo", "Tuần 2", "Planned"),
    (15, "Đầu Tư & Tích Sản", "Cách Đọc Bảng Giá Chứng Khoán HOSE, HNX, UPCOM Chuẩn Xác 100%", "cách đọc bảng giá chứng khoán", 27100, "Hướng Dẫn Kỹ Thuật", "/tai-chinh/cach-doc-bang-gia-chung-khoan", "Cổng Tra Cứu 957 Tickers", "Tuần 2", "Planned"),
    (16, "Đầu Tư & Tích Sản", "Chứng Chỉ Quỹ Là Gì? Có Nên Đầu Tư Chứng Chỉ Quỹ Thay Gửi Tiết Kiệm?", "chứng chỉ quỹ là gì", 14800, "Cẩm Nang Đầu Tư", "/tai-chinh/chung-chi-quy-la-gi", "Đầu Tư Quỹ Mở Chỉ Từ 10k", "Tuần 2", "Planned"),
    (17, "Đầu Tư & Tích Sản", "Chiến Lược Tích Sản SIP Là Gì? Sức Mạnh Lãi Kép Sau 10 Năm", "tích sản chứng chỉ quỹ sip", 9900, "Chiến Lược Dài Hạn", "/tai-chinh/chien-luoc-tich-san-sip", "Bộ Giả Lập Lãi Kép SIP", "Tuần 3", "Planned"),
    (18, "Đầu Tư & Tích Sản", "Top 5 Quỹ Mở Cổ Phiếu Tăng Trưởng Tốt Nhất Thị Trường Việt Nam", "top quỹ mở uy tín", 8100, "Đánh Giá Xếp Hạng", "/tai-chinh/top-quy-mo-co-phieu-tot-nhat", "Mua Quỹ Dragon / SSIAM", "Tuần 3", "Planned"),
    (19, "Đầu Tư & Tích Sản", "Trái Phiếu Doanh Nghiệp Là Gì? Rủi Ro & Cách Chọn Trái Phiếu An Toàn", "trái phiếu doanh nghiệp là gì", 8100, "Phân Tích Rủi Ro", "/tai-chinh/trai-phieu-doanh-nghiep", "Tư Vấn Danh Mục Tích Sản", "Tuần 3", "Planned"),
    (20, "Đầu Tư & Tích Sản", "Cách Mua Vàng Online Và Gửi Tiết Kiệm Vàng An Toàn Trên Điện Thoại", "mua vàng online", 6600, "Hướng Dẫn Thao Tác", "/tai-chinh/huong-dan-mua-vang-online", "Mở Hũ Tích Lũy Vàng MoMo", "Tuần 4", "Planned"),
    (21, "Đầu Tư & Tích Sản", "Cổ Tức Là Gì? Nên Nhận Cổ Tức Bằng Tiền Mặt Hay Bằng Cổ Phiếu?", "cổ tức là gì", 5400, "Kiến Thức Cổ Phiếu", "/tai-chinh/co-tuc-la-gi", "Mở Tài Khoản Đầu Tư CVX", "Tuần 4", "Planned"),
    (22, "Đầu Tư & Tích Sản", "Tiền Điện Tử & Crypto Là Gì? Khung Pháp Lý & Những Cảnh Báo Lừa Đảo", "tiền điện tử crypto là gì", 4400, "Cảnh Báo & Pháp Lý", "/tai-chinh/tien-dien-tu-crypto-canh-bao", "Khuyến Nghị Kênh Đầu Tư Chuẩn", "Tuần 4", "Planned"),

    # --- KHỐI III: THU NHẬP, THUẾ & QUẢN LÝ CHI TIÊU (10 Bài) ---
    (23, "Thu Nhập & Thuế", "Lương Gross Và Lương Net Là Gì? Công Thức Quy Đổi Chuẩn 2026", "lương gross net là gì", 27100, "Cẩm Nang Tiền Lương", "/tai-chinh/phan-biet-luong-gross-va-net", "Công Cụ Tính Lương Gross-Net", "Tuần 1", "Planned"),
    (24, "Thu Nhập & Thuế", "Cách Tra Cứu Mã Số Thuế Cá Nhân Online Bằng CCCD Gắn Chip Nhanh Nhất", "tra cứu mã số thuế cá nhân", 90500, "Hướng Dẫn Dịch Vụ Công", "/tai-chinh/tra-cuu-ma-so-thue-ca-nhan", "Cổng Tra Cứu Mã Số Thuế", "Tuần 1", "Planned"),
    (25, "Thu Nhập & Thuế", "Biểu Thuế Thu Nhập Cá Nhân 7 Bậc Mới Nhất & Cách Tự Tính Tiền Thuế", "thuế thu nhập cá nhân", 49500, "Công Thức Thuế", "/tai-chinh/bieu-thue-thu-nhap-ca-nhan", "Bộ Dự Toán Thuế TNCN", "Tuần 2", "Planned"),
    (26, "Thu Nhập & Thuế", "Mức Giảm Trừ Gia Cảnh 2026 Là Bao Nhiêu? Thủ Tục Đăng Ký Người Phụ Thuộc", "mức giảm trừ gia cảnh", 22000, "Quy Định Pháp Lý", "/tai-chinh/muc-giam-tru-gia-canh-moi-nhat", "Dự Toán Giảm Trừ Thuế", "Tuần 2", "Planned"),
    (27, "Thu Nhập & Thuế", "Quy Tắc Quản Lý Tài Chính 50/30/20: Bí Quyết Tiết Kiệm 20% Lương Mỗi Tháng", "quy tắc 50 30 20", 18100, "Phương Pháp Chi Tiêu", "/tai-chinh/quy-tac-quan-ly-tai-chinh-50-30-20", "Công Cụ Phân Bổ Lương MoMo", "Tuần 2", "Planned"),
    (28, "Thu Nhập & Thuế", "Phương Pháp 6 Chiếc Hũ Tài Chính (JARS) Cho Người Mới Đi Làm", "6 chiếc hũ tài chính", 14800, "Kỹ Năng Tài Chính", "/tai-chinh/6-chiec-hu-tai-chinh-jars", "Mở Hũ Chi Tiêu Linh Hoạt", "Tuần 3", "Planned"),
    (29, "Thu Nhập & Thuế", "Hướng Dẫn Tự Quyết Toán Thuế TNCN Online Qua App Thuế Điện Tử", "tự quyết toán thuế tncn", 12100, "Hướng Dẫn Thủ Tục", "/tai-chinh/tu-quyet-toan-thue-tncn-online", "Nhận Hoàn Thuế Về Túi Thần Tài", "Tuần 3", "Planned"),
    (30, "Thu Nhập & Thuế", "Thưởng Tết Có Phải Đóng Thuế TNCN Không? Cách Tính Tiền Thưởng Thực Nhận", "thưởng tết có đóng thuế không", 9900, "Hỏi Đáp Mùa Tết", "/tai-chinh/thuong-tet-co-dong-thue-tncn-khong", "Bộ Dự Toán Thưởng Tết", "Tuần 3", "Planned"),
    (31, "Thu Nhập & Thuế", "Tra Cứu Mã Số Thuế Người Phụ Thuộc Ở Đâu? Hướng Dẫn Chi Tiết", "tra cứu mã số thuế người phụ thuộc", 8100, "Tra Cứu Thuế", "/tai-chinh/tra-cuu-ma-so-thue-nguoi-phu-thuoc", "Cổng Tra Cứu MST MoMo", "Tuần 4", "Planned"),
    (32, "Thu Nhập & Thuế", "Cách Lập Kế Hoạch Tài Chính Cá Nhân 1 Năm Đơn Giản Cho Người Trẻ", "kế hoạch tài chính cá nhân", 6600, "Template Kế Hoạch", "/tai-chinh/lap-ke-hoach-tai-chinh-ca-nhan", "Bộ Giả Lập FIRE Hưu Trí", "Tuần 4", "Planned"),

    # --- KHỐI IV: NGÂN HÀNG, TỶ GIÁ & THẺ (10 Bài) ---
    (33, "Ngân Hàng & Thẻ", "Tỷ Giá USD/VND Hôm Nay: Ngân Hàng Nào Mua Bán Đô La Giá Tốt Nhất?", "tỷ giá usd hôm nay", 110000, "Bảng So Sánh Tỷ Giá", "/tai-chinh/ty-gia/usd-vnd", "Mở Thẻ Visa Quốc Tế", "Tuần 1", "Planned"),
    (34, "Ngân Hàng & Thẻ", "Đổi Tiền Yên Nhật (JPY) Ở Đâu Rẻ & An Toàn Khi Đi Du Lịch / Du Học?", "đổi tiền yên nhật", 27100, "Cẩm Nang Đổi Ngoại Tệ", "/tai-chinh/ty-gia/jpy-vnd", "Nhận Kiều Hối Nhật Bản", "Tuần 1", "Planned"),
    (35, "Ngân Hàng & Thẻ", "Thẻ Tín Dụng Là Gì? 7 Nguyên Tắc Dùng Thẻ Thông Minh Không Lo Nợ Nần", "thẻ tín dụng là gì", 22000, "Cẩm Nang Thẻ", "/tai-chinh/the-tin-dung-la-gi", "Mở Thẻ Tín Dụng Online", "Tuần 2", "Planned"),
    (36, "Ngân Hàng & Thẻ", "Phân Biệt Thẻ Visa Và Mastercard: Điểm Giống, Khác Nhau & Nên Dùng Loại Nào?", "so sánh visa và mastercard", 18100, "So Sánh Thẻ", "/tai-chinh/so-sanh-the-visa-va-mastercard", "Mở Thẻ Quốc Tế MoMo", "Tuần 2", "Planned"),
    (37, "Ngân Hàng & Thẻ", "Thẻ Ghi Nợ (Debit Card) Là Gì? Khác Gì Với Thẻ Tín Dụng & Thẻ ATM?", "thẻ ghi nợ là gì", 14800, "Kiến Thức Thẻ", "/tai-chinh/the-ghi-no-debit-card-la-gi", "Liên Kết Thẻ ATM Miễn Phí", "Tuần 2", "Planned"),
    (38, "Ngân Hàng & Thẻ", "Cách Mở Thẻ Tín Dụng Online Không Cần Chứng Minh Thu Nhập", "mở thẻ tín dụng online", 12100, "Hướng Dẫn Mở Thẻ", "/tai-chinh/mo-the-tin-dung-online-khong-chung-minh-thu-nhap", "Mở Thẻ Tín Dụng Đối Tác", "Tuần 3", "Planned"),
    (39, "Ngân Hàng & Thẻ", "Tỷ Giá Won Hàn Quốc (KRW) Hôm Nay & Kinh Nghiệm Đổi Tiền Đi Hàn", "tỷ giá won hàn quốc", 9900, "Cẩm Nang Ngoại Tệ", "/tai-chinh/ty-gia/krw-vnd", "Thanh Toán Quốc Tế MoMo", "Tuần 3", "Planned"),
    (40, "Ngân Hàng & Thẻ", "Phí Chuyển Đổi Ngoại Tệ Thẻ Tín Dụng Của Các Ngân Hàng: Bảng So Sánh", "phí chuyển đổi ngoại tệ", 8100, "Bảng So Sánh Phí", "/tai-chinh/phi-chuyen-doi-ngoai-te", "Mở Thẻ Phí Thấp 1.5%", "Tuần 3", "Planned"),
    (41, "Ngân Hàng & Thẻ", "Danh Bạ 34 Ngân Hàng Đối Tác MoMo: Biểu Phí & Gói Quà Liên Kết Ví 500k", "ngân hàng liên kết momo", 6600, "Danh Bạ Đối Tác", "/tai-chinh/ngan-hang", "Nhận Quà Tân Thủ 500k", "Tuần 4", "Planned"),
    (42, "Ngân Hàng & Thẻ", "Napas Là Gì? Hệ Thống Chuyển Tiền Nhanh 24/7 Hoạt Động Ra Sao?", "napas là gì", 5400, "Kiến Thức Thanh Toán", "/tai-chinh/napas-la-gi", "Chuyển Tiền Miễn Phí MoMo", "Tuần 4", "Planned"),

    # --- KHỐI V: SỨC KHỎE TÍN DỤNG & CIC (8 Bài) ---
    (43, "Sức Khỏe Tín Dụng & CIC", "Hướng Dẫn Kiểm Tra Điểm Tín Dụng CIC Cá Nhân Miễn Phí Nhanh Nhất", "kiểm tra cic miễn phí", 49500, "Cẩm Nang CIC", "/tra-cuu-cic/huong-dan-kiem-tra-cic-mien-phi", "Check Báo Cáo Tín Dụng In-App", "Tuần 1", "Planned"),
    (44, "Sức Khỏe Tín Dụng & CIC", "Nợ Xấu Là Gì? Phân Loại 5 Nhóm Nợ Xấu Ngân Hàng Mới Nhất", "nợ xấu là gì", 33100, "Kiến Thức Tín Dụng", "/tra-cuu-cic/no-xau-la-gi-5-nhom-no", "Widget CIC Simulator", "Tuần 1", "Planned"),
    (45, "Sức Khỏe Tín Dụng & CIC", "Điểm Tín Dụng Bao Nhiêu Là Tốt Để Được Duyệt Hạn Mức Vay / Mở Thẻ?", "điểm tín dụng bao nhiêu là tốt", 22000, "Tư Vấn Điểm Số", "/tra-cuu-cic/diem-tin-dung-bao-nhieu-la-tot", "Kiểm Tra Điểm MoMo", "Tuần 2", "Planned"),
    (46, "Sức Khỏe Tín Dụng & CIC", "Nợ Xấu Nhóm 2, Nhóm 3 Sau Bao Lâu Được Xóa Lịch Sử Trên Hệ Thống CIC?", "xóa nợ xấu mất bao lâu", 18100, "Quy Định Xóa Nợ", "/tra-cuu-cic/xoa-no-xau-mat-bao-lau", "Tư Vấn Phục Hồi Tín Dụng", "Tuần 2", "Planned"),
    (47, "Sức Khỏe Tín Dụng & CIC", "5 Cách Tăng Điểm Tín Dụng Nhanh Chóng Trong Vòng 3 Đến 6 Tháng", "cách tăng điểm tín dụng", 12100, "Chiến Lược Tín Dụng", "/tra-cuu-cic/cach-tang-diem-tin-dung-nhanh", "Bật Thanh Toán Tự Động Hóa Đơn", "Tuần 3", "Planned"),
    (48, "Sức Khỏe Tín Dụng & CIC", "Bị Kẻ Gian Lấy Thông Tin Mở Thẻ / Vay Tiền Dính Nợ Xấu Phải Làm Sao?", "bị lừa nợ xấu phải làm sao", 9900, "Cảnh Báo & Xử Lý", "/tra-cuu-cic/bi-lay-thong-tin-vay-no-xau", "Tra Cứu CIC Bảo Vệ Danh Tính", "Tuần 3", "Planned"),
    (49, "Sức Khỏe Tín Dụng & CIC", "Điểm Tín Dụng MoMo Là Gì? Quyền Lợi Khi Đạt Điểm Tín Dụng Cao", "điểm tín dụng momo", 8100, "Sản Phẩm Độc Quyền", "/tra-cuu-cic/diem-tin-dung-momo", "Mở Điểm Tín Dụng MoMo In-App", "Tuần 4", "Planned"),
    (50, "Sức Khỏe Tín Dụng & CIC", "Dính Nợ Xấu Có Mở Được Thẻ Tín Dụng Hoặc Mở Tài Khoản Ngân Hàng Không?", "nợ xấu có làm thẻ ngân hàng được không", 6600, "Hỏi Đáp Chuyên Gia", "/tra-cuu-cic/no-xau-co-mo-the-ngan-hang-duoc-khong", "Khám Phá Gói Tín Dụng Phù Hợp", "Tuần 4", "Planned")
]

if 'Content Plan' in wb.sheetnames:
    del wb['Content Plan']

ws_cp = wb.create_sheet('Content Plan', index=5)
ws_cp.append(["KẾ HOẠCH NỘI DUNG TÀI CHÍNH THÁNG 9/2026 (50 BÀI VIẾT CHUẨN E-E-A-T THEO 5 TRỤ CỘT)"])
ws_cp.merge_cells("A1:J1")
ws_cp.cell(row=1, column=1).fill = section_fill
ws_cp.cell(row=1, column=1).font = section_font
ws_cp.row_dimensions[1].height = 24.0

cp_headers = [
    "STT",
    "Trụ Cột Nghiệp Vụ",
    "Tiêu Đề Bài Viết Chuẩn SEO (H1)",
    "Từ Khóa Mục Tiêu (Primary KW)",
    "Search Volume Ước Tính",
    "Định Dạng Nội Dung",
    "URL Slug Dự Kiến",
    "Điểm Chạm Chuyển Đổi (W2A Hook)",
    "Tuần Phát Hành",
    "Trạng Thái"
]
ws_cp.append(cp_headers)
for c_i in range(1, len(cp_headers) + 1):
    c = ws_cp.cell(row=2, column=c_i)
    c.fill = header_fill; c.font = header_font; c.border = thin_border
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws_cp.row_dimensions[2].height = 28.0

for row_item in content_articles:
    ws_cp.append(list(row_item))
    r = ws_cp.max_row
    ws_cp.row_dimensions[r].height = 22.0
    
    ws_cp.cell(row=r, column=1).alignment = Alignment(horizontal='center', vertical='center')
    ws_cp.cell(row=r, column=2).alignment = Alignment(horizontal='left', vertical='center'); ws_cp.cell(row=r, column=2).font = font_bold
    ws_cp.cell(row=r, column=3).alignment = Alignment(horizontal='left', vertical='center')
    ws_cp.cell(row=r, column=4).alignment = Alignment(horizontal='left', vertical='center')
    ws_cp.cell(row=r, column=5).alignment = Alignment(horizontal='right', vertical='center'); ws_cp.cell(row=r, column=5).font = font_vol; ws_cp.cell(row=r, column=5).number_format = '#,##0'
    ws_cp.cell(row=r, column=6).alignment = Alignment(horizontal='center', vertical='center')
    ws_cp.cell(row=r, column=7).alignment = Alignment(horizontal='left', vertical='center')
    ws_cp.cell(row=r, column=8).alignment = Alignment(horizontal='left', vertical='center')
    ws_cp.cell(row=r, column=9).alignment = Alignment(horizontal='center', vertical='center')
    ws_cp.cell(row=r, column=10).alignment = Alignment(horizontal='center', vertical='center')
    
    c_st = ws_cp.cell(row=r, column=10)
    c_st.fill = status_plan_fill; c_st.font = status_plan_font

    for c_i in range(1, 11): ws_cp.cell(row=r, column=c_i).border = thin_border

col_widths = [6, 25, 45, 28, 20, 24, 38, 32, 16, 14]
for idx, w in enumerate(col_widths, start=1):
    ws_cp.column_dimensions[get_column_letter(idx)].width = w

wb.save('05_HUBS/financial-hub-roadmap.xlsx')
print("Successfully generated Content Plan sheet with 50 financial articles in 05_HUBS/financial-hub-roadmap.xlsx!")
