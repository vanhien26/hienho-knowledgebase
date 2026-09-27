import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook('05_HUBS/financial-hub-roadmap.xlsx')

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
section_font = Font(name='Arial', size=10.5, bold=True, color='1F4E78')
thin_side = Side(style='thin', color='64748B')
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

font_bold = Font(name='Arial', size=9.5, bold=True)
font_regular = Font(name='Arial', size=9.5)
font_role_key_owner = Font(name='Arial', size=9.5, bold=True, color='B91C1C') # Red bold
font_role_co_owner = Font(name='Arial', size=9.5, bold=True, color='15803D')  # Green bold
font_role_contributor = Font(name='Arial', size=9.5, bold=True, color='1D4ED8') # Blue bold
font_role_consulted = Font(name='Arial', size=9.5, bold=False, color='D97706') # Amber
font_role_informed = Font(name='Arial', size=9.5, bold=False, color='64748B')  # Slate

if 'RACI & Resource Allocation' in wb.sheetnames:
    del wb['RACI & Resource Allocation']

ws_raci = wb.create_sheet('RACI & Resource Allocation', index=4)

ws_raci.append(["MA TRẬN PHÂN ĐỊNH TRÁCH NHIỆM RACI: WEB PLATFORM x CELL TEAM (FINHUB)"])
ws_raci.merge_cells("A1:E1")
ws_raci.cell(row=1, column=1).fill = section_fill
ws_raci.cell(row=1, column=1).font = section_font
ws_raci.row_dimensions[1].height = 24.0

ws_raci.append(["• Web Platform: Xây dựng Sản Phẩm Web (Utilities), Lập Content Plan, Tracking Toàn Diện Luồng Web | • FinHub (Cell Team): Tăng Trưởng (Growth), Content QC, Legal QC, Check Luồng In-App"])
ws_raci.merge_cells("A2:E2")
ws_raci.cell(row=2, column=1).font = Font(name='Arial', size=9.5, italic=True, color='1F4E78')
ws_raci.row_dimensions[2].height = 20.0

raci_headers = [
    "Hạng Mục / Đầu Việc Triển Khai",
    "Web Platform Team",
    "Cell Team (FinHub BU)",
    "Hạ Tầng / Backend Team",
    "Mô Tả Trách Nhiệm Chi Tiết & Tiêu Chí Nghiệm Thu"
]
ws_raci.append(raci_headers)
for c_i in range(1, len(raci_headers) + 1):
    c = ws_raci.cell(row=3, column=c_i)
    c.fill = header_fill; c.font = header_font; c.border = thin_border
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws_raci.row_dimensions[3].height = 28.0

raci_matrix_data = [
    ("GROUP", "NHÓM I: CHIẾN LƯỢC TĂNG TRƯỞNG & VẬN HÀNH SPRINT (GROWTH & SPRINT CADENCE)", "", "", ""),
    ("ROW", "1.1 Quản Trị Mục Tiêu Tăng Trưởng Kinh Doanh (Business Growth & KPIs)",
     "Co-Owner (A/R)", "Key Owner (A/R)", "Informed (I)",
     "Cell Team (FinHub) chịu trách nhiệm cao nhất về kết quả tăng trưởng kinh doanh (Revenue, MAU, Giao dịch tài chính đầu tiên +25% MoM); Web Platform đồng sở hữu và tối ưu hóa phễu Web Traffic (1M MPV, 500K MEU)."),
    
    ("ROW", "1.2 Lập Kế Hoạch Sprint Web 1 Tháng & Nhịp Check-in 2 Tuần/Lần",
     "Key Owner (A/R)", "Co-Owner (A/R)", "Contributor (R)",
     "Web Platform chủ trì tổ chức Sprint Planning đầu tháng để chốt backlog tính năng Web; Cell Team đồng chủ trì check-in ngày 15 & 30 hàng tháng để rà soát tiến độ và tháo gỡ điểm nghẽn."),

    ("GROUP", "NHÓM II: XÂY DỰNG SẢN PHẨM WEB & BỘ TIỆN ÍCH (WEB PRODUCT & UTILITIES)", "", "", ""),
    ("ROW", "2.1 Lập Trình Master Gateway & Giao Diện Cổng Tài Chính (momo.vn/tai-chinh)",
     "Key Owner (A/R)", "Consulted (C)", "Contributor (R)",
     "Web Platform chịu trách nhiệm 100% về thiết kế UI/UX, lập trình mặt tiền Cổng Khám Phá Tài Chính, Mega Navigation Bar và Dashboard thị trường tổng hợp."),

    ("ROW", "2.2 Phát Triển Bộ Tiện Ích & Công Cụ Tính Toán Web (Utilities & Calculators)",
     "Key Owner (A/R)", "Consulted (C)", "Informed (I)",
     "Web Platform lập trình các công cụ tiện ích tương tác không cần đăng nhập (Tính Lương, Thuế, Lãi Tiết Kiệm, FIRE Simulator...); Cell Team tư vấn và chuẩn hóa công thức tài chính."),

    ("ROW", "2.3 Đóng Gói Bộ Công Cụ Tiện Ích Dùng Chung (Embeddable Widgets SDK)",
     "Key Owner (A/R)", "Consulted (C)", "Informed (I)",
     "Web Platform đóng gói các tiện ích thành Web Components nhúng linh hoạt vào mọi điểm chạm Web MoMo và hệ thống đối tác ngoài."),

    ("GROUP", "NHÓM III: KẾ HOẠCH NỘI DUNG, KIỂM DUYỆT CHẤT LƯỢNG & PHÁP LÝ (CONTENT & LEGAL QC)", "", "", ""),
    ("ROW", "3.1 Xây Dựng Kế Hoạch Nội Dung & Từ Khóa SEO (Content Plan & MoSpark AI)",
     "Key Owner (A/R)", "Consulted (C)", "Contributor (R)",
     "Web Platform chịu trách nhiệm lập kế hoạch nội dung (Content Plan), nghiên cứu từ khóa theo 24 thị trường và vận hành pipeline GenAI trên MoSpark CMS."),

    ("ROW", "3.2 Kiểm Duyệt & Đảm Bảo Chất Lượng Nội Dung Tài Chính (Content QC)",
     "Contributor (R)", "Key Owner (A/R)", "Informed (I)",
     "Cell Team (FinHub) giữ vai trò Key Owner thẩm định và phê duyệt chất lượng chuyên môn (Content QC), đảm bảo tính chuẩn xác của các gói tài chính, lãi suất và thuật ngữ nghiệp vụ trước khi xuất bản."),

    ("ROW", "3.3 Thẩm Định Pháp Lý & Tuân Thủ Quy Định YMYL / NHNN (Legal QC)",
     "Contributor (R)", "Key Owner (A/R)", "Informed (I)",
     "Cell Team (FinHub) chịu trách nhiệm làm việc với Phòng Pháp Chế (Legal) và đối tác tài chính để kiểm duyệt tính tuân thủ pháp lý (Legal QC), disclaimers và quy định của Ngân hàng Nhà nước."),

    ("GROUP", "NHÓM IV: TRACKING LUỒNG WEB & KIỂM TRA LUỒNG IN-APP (TRACKING & IN-APP FLOW QC)", "", "", ""),
    ("ROW", "4.1 Cấu Hình & Tracking Toàn Diện Luồng Web (Web Tracking, GA4, GTM, OneLink, QR)",
     "Key Owner (A/R)", "Consulted (C)", "Contributor (R)",
     "Web Platform chịu trách nhiệm toàn bộ hạ tầng đo lường Kênh Web: GTM, GA4 events, UTM parameters, Appsflyer OneLink và Dynamic Desktop QR Code."),

    ("ROW", "4.2 Kiểm Tra & Tối Ưu Hóa Hành Trình Người Dùng In-App (Check Luồng In-App, Deeplink, eKYC)",
     "Co-Owner (A/R)", "Key Owner (A/R)", "Contributor (R)",
     "Cell Team (FinHub) chịu trách nhiệm kiểm tra và tối ưu hóa luồng In-App: đảm bảo khi người dùng click CTA Web sẽ mở đúng màn hình App tương ứng, luồng eKYC và kích hoạt giao dịch đầu tiên mượt mà."),

    ("GROUP", "NHÓM V: BÁO CÁO GIẢI TRÌNH & ĐIỀU PHỐI TÀI NGUYÊN (ACCOUNTABILITY & DEPRIORITIZATION)", "", "", ""),
    ("ROW", "5.1 Báo Cáo Hiệu Quả Phễu Web-to-App Hàng Tháng Lên Ban Lãnh Đạo IVP",
     "Co-Owner (A/R)", "Key Owner (A/R)", "Contributor (R)",
     "Cell Team (FinHub) chủ trì báo cáo kết quả tăng trưởng kinh doanh; Web Platform đồng chủ trì báo cáo hiệu quả chuyển đổi Kênh Web lên Ban Lãnh Đạo IVP hàng tháng."),

    ("ROW", "5.2 Cơ Chế Rút Lui & Điều Phối Tài Nguyên Kỹ Thuật (Deprioritization Rule)",
     "Key Owner (A)", "Co-Owner (A/R)", "Informed (I)",
     "Nếu sau 1-2 tháng phối hợp mà Cell Team không tập trung (chậm trễ Content/Legal QC, thiếu deeplink, trễ cam kết), Web Platform sẽ chủ động bàn giao lại dự án và chuyển tài nguyên kỹ thuật sang BU khác.")
]

for item in raci_matrix_data:
    item_type = item[0]
    if item_type == "GROUP":
        ws_raci.append([item[1], "", "", "", ""])
        r = ws_raci.max_row
        ws_raci.merge_cells(f"A{r}:E{r}")
        ws_raci.cell(row=r, column=1).fill = section_fill
        ws_raci.cell(row=r, column=1).font = section_font
        ws_raci.row_dimensions[r].height = 22.0
    elif item_type == "ROW":
        ws_raci.append(list(item[1:]))
        r = ws_raci.max_row
        ws_raci.row_dimensions[r].height = 36.0
        
        c1 = ws_raci.cell(row=r, column=1); c1.font = font_bold; c1.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        c2 = ws_raci.cell(row=r, column=2); c2.alignment = Alignment(horizontal='center', vertical='center')
        c3 = ws_raci.cell(row=r, column=3); c3.alignment = Alignment(horizontal='center', vertical='center')
        c4 = ws_raci.cell(row=r, column=4); c4.alignment = Alignment(horizontal='center', vertical='center')
        c5 = ws_raci.cell(row=r, column=5); c5.font = font_regular; c5.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

        for cell, role_txt in [(c2, item[2]), (c3, item[3]), (c4, item[4])]:
            if "Key Owner" in role_txt: cell.font = font_role_key_owner
            elif "Co-Owner" in role_txt: cell.font = font_role_co_owner
            elif "Contributor" in role_txt: cell.font = font_role_contributor
            elif "Consulted" in role_txt: cell.font = font_role_consulted
            elif "Informed" in role_txt: cell.font = font_role_informed
            else: cell.font = font_bold

        for c_i in range(1, 6): ws_raci.cell(row=r, column=c_i).border = thin_border

# Definitions at bottom
ws_raci.append([""])
ws_raci.append(["BẢNG CHÚ GIẢI KÝ HIỆU & NGUYÊN TẮC PHÂN CÔNG TRÁCH NHIỆM:"])
r_def = ws_raci.max_row
ws_raci.cell(row=r_def, column=1).font = Font(name='Arial', size=10, bold=True, color='002060')

raci_glossary = [
    ("Key Owner (A/R)", "Accountable & Responsible", "Chủ Trì & Chịu Trách Nhiệm Chính", "Đội ngũ chịu trách nhiệm giải trình cao nhất (Accountability) trước Ban Lãnh Đạo IVP và trực tiếp thực thi/nghiệm thu hạng mục."),
    ("Co-Owner (A/R)", "Co-Accountable & Responsible", "Đồng Sở Hữu & Đồng Thực Thi", "Đội ngũ cùng cam kết kết quả đầu ra, chịu trách nhiệm cung cấp chính sách, deeplink, dữ liệu nghiệp vụ và đồng sở hữu KPI."),
    ("Contributor (R)", "Responsible Contributor", "Đội Ngũ Thực Thi Hỗ Trợ", "Đội ngũ trực tiếp tham gia hỗ trợ kỹ thuật, đấu nối hệ thống hoặc cung cấp hạ tầng dữ liệu theo đúng cam kết thời hạn."),
    ("Consulted (C)", "Consulted", "Tham Vấn Chuyên Môn 2 Chiều", "Đội ngũ được tham khảo ý kiến chuyên môn (công thức lãi suất, kiểm duyệt luật tài chính YMYL, góp ý luồng UX) trước khi ban hành."),
    ("Informed (I)", "Informed", "Nhận Thông Tin 1 Chiều", "Đội ngũ nhận báo cáo cập nhật tiến độ định kỳ để nắm bối cảnh liên phòng ban, không bắt buộc tham gia vận hành chi tiết.")
]

ws_raci.append(["Thuật Ngữ RACI", "Tên Tiếng Anh", "Vai Trò Thực Tế", "Mô Tả Trách Nhiệm Chi Tiết"])
r_h = ws_raci.max_row
for c_i in range(1, 5):
    c = ws_raci.cell(row=r_h, column=c_i)
    c.fill = header_fill; c.font = header_font; c.border = thin_border
    c.alignment = Alignment(horizontal='center', vertical='center')
ws_raci.row_dimensions[r_h].height = 24.0

for g in raci_glossary:
    ws_raci.append(list(g))
    r = ws_raci.max_row
    ws_raci.row_dimensions[r].height = 22.0
    ws_raci.cell(row=r, column=1).alignment = Alignment(horizontal='center', vertical='center'); ws_raci.cell(row=r, column=1).font = font_bold
    ws_raci.cell(row=r, column=2).alignment = Alignment(horizontal='left', vertical='center'); ws_raci.cell(row=r, column=2).font = font_bold
    ws_raci.cell(row=r, column=3).alignment = Alignment(horizontal='left', vertical='center')
    ws_raci.cell(row=r, column=4).alignment = Alignment(horizontal='left', vertical='center')
    for c_i in range(1, 5): ws_raci.cell(row=r, column=c_i).border = thin_border

raci_widths = [44, 22, 24, 22, 78]
for idx, w in enumerate(raci_widths, start=1):
    ws_raci.column_dimensions[get_column_letter(idx)].width = w

wb.save('05_HUBS/financial-hub-roadmap.xlsx')
print("Successfully updated RACI matrix reflecting exact strategic scopes: Web Platform (Utilities, Content Plan, Tracking) & FinHub (Growth, Content QC, Legal QC, In-App Flow)!")
