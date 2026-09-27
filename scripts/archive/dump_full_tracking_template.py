import openpyxl

template_path = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/07_REPORTS/sheets/web-event-tracking-template.xlsx'
wb = openpyxl.load_workbook(template_path, data_only=True)

for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n==========================================")
    print(f"SHEET: {name}")
    print(f"==========================================")
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(vals):
            print(f"R{r}:", [v for v in vals if v is not None])

