import openpyxl

file_excel = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/vehicle-hub-roadmap.xlsx'
wb = openpyxl.load_workbook(file_excel)

ws_road = wb['Roadmap']
for r in range(3, ws_road.max_row + 1):
    val_name = str(ws_road.cell(r, 4).value or '')
    
    # 1. Update Giá Xăng row
    if 'Giá Xăng' in val_name:
        scope = (
            "• Trang duy nhất (/gia-xang) - TUYỆT ĐỐI KHÔNG tạo trang con theo tỉnh/thành (xăng dầu quy định theo Vùng 1 & Vùng 2 toàn quốc, tránh phạt Duplicate Content).\n"
            "• Bảng giá realtime 7 loại nhiên liệu (RON 95-V Euro 5, DO 0.001S-V Euro 5...) kèm nút gạt [Vùng 1] vs [Vùng 2].\n"
            "• Bộ máy tính giá xăng 1-chạm (Đầy bình theo xe, Nút tiền tròn/lít, Dự toán lộ trình Km).\n"
            "(Search Volume: 10.523.210/tháng)"
        )
        ws_road.cell(r, 5, value=scope)
        
    # 2. Update Đăng Kiểm row
    elif 'Đăng Kiểm' in val_name:
        scope = (
            "• Trang duy nhất (/dang-kiem) - KHÔNG tạo trang con tỉnh/thành (do volume search địa phương = 0, tập trung 100% SEO vào 1 URL duy nhất).\n"
            "• Tra cứu chu kỳ & ngày hết hạn tự động theo Thông tư 02 & 08/2023.\n"
            "• Dự toán phí đăng kiểm (340k) + phí bảo trì đường bộ (1.560k/năm).\n"
            "• Bản đồ gần 300 trạm kiểm định xe toàn quốc tích hợp bộ lọc dropdown.\n"
            "• Phễu bán chéo Bảo hiểm TNDS bắt buộc (điều kiện tiên quyết tại trạm)."
        )
        ws_road.cell(r, 5, value=scope)

wb.save(file_excel)
print("Updated Excel Roadmap with anti-duplicate rules successfully!")
