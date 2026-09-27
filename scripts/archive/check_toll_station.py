import openpyxl

file_inv = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/inventory.xlsx'
file_cs = '/Users/hienhv/Downloads/[WP] Vehicle Hub - Content Strategy.xlsx'
file_road = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/vehicle-hub-roadmap.xlsx'

print("=== 1. INVENTORY.XLSX ===")
wb_inv = openpyxl.load_workbook(file_inv, data_only=True)
ws_inv = wb_inv.active
for r in range(1, ws_inv.max_row + 1):
    vals = [str(ws_inv.cell(r, c).value) for c in range(1, ws_inv.max_column + 1)]
    if any(k in v.lower() for v in vals for k in ['thu phí', 'epass', 'vetc', 'bot', 'trạm thu phí']):
        print(" ", vals[:6])

print("\n=== 2. ROADMAP EXCEL ===")
wb_road = openpyxl.load_workbook(file_road, data_only=True)
ws_road = wb_road['Roadmap']
for r in range(2, ws_road.max_row + 1):
    vals = [str(ws_road.cell(r, c).value) for c in range(1, ws_road.max_column + 1)]
    if any(k in v.lower() for v in vals for k in ['thu phí', 'epass', 'vetc', 'bot', 'trạm thu phí']):
        print(f"  Row {r}:", vals[0], "|", vals[1], "|", vals[3], "|", vals[6])

print("\n=== 3. CONTENT STRATEGY KEYWORDS ===")
wb_cs = openpyxl.load_workbook(file_cs, data_only=True)
print("Available sheets in Content Strategy:", wb_cs.sheetnames)
ws_kw = wb_cs['Keyword']
toll_kws = []
for r in range(2, ws_kw.max_row + 1):
    kw = str(ws_kw.cell(r, 1).value or ws_kw.cell(r, 2).value or '').strip().lower()
    if any(k in kw for k in ['thu phí', 'epass', 'vetc', 'bot', 'trạm thu phí']):
        vol = ws_kw.cell(r, 2).value or ws_kw.cell(r, 3).value or 0
        try:
            vol = float(vol)
        except:
            vol = 0
        toll_kws.append((kw, vol))

toll_kws.sort(key=lambda x: x[1], reverse=True)
print(f"Total toll keywords in Keyword sheet: {len(toll_kws)}, Sum Vol: {sum(x[1] for x in toll_kws):,.0f}")
for k, v in toll_kws[:15]:
    print(f"  • {k}: {v:,.0f}")

