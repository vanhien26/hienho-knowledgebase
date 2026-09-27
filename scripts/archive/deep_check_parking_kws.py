import openpyxl
from collections import defaultdict

file_cs = '/Users/hienhv/Downloads/[WP] Vehicle Hub - Content Strategy.xlsx'
wb = openpyxl.load_workbook(file_cs, data_only=True)
ws_kw = wb['Keyword']

parking_kws = []
for r in range(2, ws_kw.max_row + 1):
    kw = str(ws_kw.cell(r, 1).value or ws_kw.cell(r, 2).value or '').strip().lower()
    if any(k in kw for k in ['bãi đỗ', 'bãi gửi', 'bãi giữ', 'đỗ xe', 'gửi xe', 'parking']):
        vol = ws_kw.cell(r, 2).value or ws_kw.cell(r, 3).value or 0
        try:
            vol = float(vol)
        except:
            vol = 0
        parking_kws.append((kw, vol))

parking_kws.sort(key=lambda x: x[1], reverse=True)
print(f"Total parking keywords found: {len(parking_kws)}, Total volume: {sum(x[1] for x in parking_kws):,.0f}")

# Group into categories:
cats = defaultdict(list)
for k, v in parking_kws:
    if any(loc in k for loc in ['hà nội', 'hcm', 'sài gòn', 'đà nẵng', 'quận', 'hoàn kiếm', 'tân sơn nhất', 'nội bài']):
        cats['Location / Landmark'].append((k, v))
    elif any(t in k for t in ['qua đêm', '24/24', '24 24', '24/7', 'theo giờ', 'theo tháng']):
        cats['Time / 24/7 / Overnight'].append((k, v))
    elif any(veh in k for veh in ['ô tô', 'oto', 'xe hơi']):
        cats['Car Specific'].append((k, v))
    elif any(veh in k for veh in ['xe máy', 'xe may']):
        cats['Motorbike Specific'].append((k, v))
    elif any(fee in k for fee in ['giá', 'phí', 'bảng giá', 'bao nhiêu']):
        cats['Pricing / Fee'].append((k, v))
    else:
        cats['General / Head terms'].append((k, v))

for cat_name, items in sorted(cats.items()):
    items.sort(key=lambda x: x[1], reverse=True)
    print(f"\n=== {cat_name} ({len(items)} keywords, Sum Vol: {sum(x[1] for x in items):,.0f}) ===")
    for k, v in items[:10]:
        if v > 0:
            print(f"  • {k}: {v:,.0f}")

