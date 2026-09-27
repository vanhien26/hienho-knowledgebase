import openpyxl
from openpyxl.styles import Alignment, Border, Side, PatternFill, Font

file_path = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/vehicle-hub-roadmap.xlsx'
wb = openpyxl.load_workbook(file_path)

sheet_name = 'Resources'
if sheet_name in wb.sheetnames:
    del wb[sheet_name]
ws = wb.create_sheet(sheet_name)

# Define Styles
font_header = Font(bold=True, color="FFFFFF")
fill_header = PatternFill("solid", fgColor="4C1130") # Dark MoMo Pink

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

headers = ['Hạng Mục Công Việc Khối (Task Area)', 'Web Platform', 'InsurTech']
ws.append(headers)

for col_num in range(1, 4):
    cell = ws.cell(row=1, column=col_num)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = border

data = [
    ['1. Phát triển Sản phẩm (Product UI/UX)', 'A/R (Chịu trách nhiệm chính)\n- Build toàn bộ giao diện Web, luồng UI/UX.\n- Tối ưu phễu Web-to-App.', 'QC (Nghiệm thu)\n- Chịu trách nhiệm QC nghiệm thu sản phẩm cuối trước khi Go-live.'],
    ['2. Sản xuất Nội dung (Content Pipeline)', 'A/R (Chịu trách nhiệm chính)\n- Lập kế hoạch nội dung (Content Plan).\n- Sản xuất bài viết, cẩm nang SEO (Content Production).', 'QC (Nghiệm thu)\n- QC nội dung, kiểm duyệt tính chính xác của luật giao thông & nghiệp vụ tài chính.'],
    ['3. Hạ tầng Dữ liệu & Hệ thống (Data & API)', 'C (Phối hợp tích hợp)\n- Gọi và render API hiển thị lên giao diện Web.', 'A/R (Chịu trách nhiệm chính)\n- Cung cấp API (nếu có).\n- Cung cấp Database liên quan (Giá xăng, Trạm sạc, Garage, Phạt nguội).'],
    ['4. Đo lường & Phân tích (Data Tracking)', 'A/R (Chịu trách nhiệm Web)\n- Thiết lập đo lường, tracking DataLayer toàn bộ luồng Web.', 'A/R (Chịu trách nhiệm In-App)\n- Thiết lập tracking luồng In-App.\n- Theo dõi tỷ lệ rớt phễu thanh toán.'],
    ['5. Truyền thông & Tăng trưởng (Growth Plan)', 'C (Hỗ trợ kéo Free Traffic)\n- Đẩy mạnh SEO để hứng Organic Search Volume.', 'A/R (Chịu trách nhiệm chính)\n- Lên kế hoạch Growth Plan tổng thể.\n- Bơm ngân sách Paid Ads, In-App Marketing.']
]

for idx, row_data in enumerate(data, start=2):
    for col, val in enumerate(row_data, start=1):
        cell = ws.cell(row=idx, column=col, value=val)
        cell.alignment = align_left if col > 1 else align_center
        cell.border = border
        if col == 1:
            cell.font = Font(bold=True)

# Adjust column widths
ws.column_dimensions['A'].width = 35
ws.column_dimensions['B'].width = 55
ws.column_dimensions['C'].width = 55

wb.save(file_path)
print("Resources (RACI) sheet created successfully.")
