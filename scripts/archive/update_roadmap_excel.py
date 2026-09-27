import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

file_path = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/vehicle-hub-roadmap.xlsx'
wb = openpyxl.load_workbook(file_path)

if 'Roadmap' in wb.sheetnames:
    ws = wb['Roadmap']
    # Clear existing rows from 4 downwards
    ws.delete_rows(4, ws.max_row)
    
    # Styles
    font_bold = Font(bold=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    fill_phase1 = PatternFill("solid", fgColor="D9EAD3")
    fill_phase2 = PatternFill("solid", fgColor="FFF2CC")
    fill_blocked = PatternFill("solid", fgColor="F4CCCC")

    data = [
        # Phase 1
        ['Phase 1: Build-Up & Foundation (Nền tảng hút Traffic & Tiện ích Ưu tiên P1)', 'Master Hub Gateway', 'Trang Chủ Master Hub (/tien-ich-giao-thong)', '1. Mặt tiền trung tâm với Hero Widget tra cứu 3-in-1 (Biển số, ePass).\n2. Gắn điểm chạm Onelink/QR Code điều hướng In-App.\nVai trò: Trang đích hợp nhất hệ sinh thái Giao thông Web.', 'Go-live Sớm'],
        ['None', 'Magnet + Transaction (P1)', 'Cổng Tra Cứu Phạt Nguội (/phat-nguoi)', '1. Trả kết quả Inline 2 nhánh (Có vi phạm / Không vi phạm).\n2. Block Cross-sell gợi ý nộp phạt In-App & Mua Bảo hiểm.\nVai trò: Ưu tiên P1 (3.6M search, 10.1 điểm), phễu thu thập Lead lớn nhất.', 'Check-in Giữa Sprint'],
        ['None', 'Magnet + Transaction (P1)', 'Thu Phí Không Dừng ePass/VETC (/phi-khong-dung)', '1. Hướng dẫn & luồng liên kết ePass/VETC theo biển số.\n2. Nút W2A nạp ví thu phí không dừng.\nVai trò: Ưu tiên P1 (8.7 điểm), monetization nạp tiền định kỳ.', 'Check-in Giữa Sprint'],
        ['None', 'Giữ Chân & Acquisition (P2)', 'Bảng Giá Xăng & Bản Đồ Cây Xăng (/gia-xang, /cay-xang)', '1. Giá xăng tự động update kỳ điều hành & Bản đồ định vị GPS.\nVai trò: Volume tìm kiếm cực lớn (10M+), phễu Acquisition tiếp cận tệp chủ xe.', 'Check-in Đầu Sprint'],
        ['None', 'Giữ Chân & Acquisition (P2)', 'Bản Đồ Trạm Sạc EV (/tram-sac)', '1. Định vị trạm sạc VinFast, V-Green theo tỉnh/thành phố.\nVai trò: Đặt nền tảng cho tệp chủ xe điện (EV) tăng trưởng dài hạn.', 'Check-in Đầu Sprint'],
        
        # Phase 2
        ['Phase 2: Convert & Monetize (Tối ưu Chuyển Đổi & Bán Chéo)', 'Conversion', 'Tối Ưu Chuyển Đổi Bảo Hiểm & W2A Handoff', '1. Tối ưu Onelink giữ context biển số khi quét QR mở App.\n2. Prefill Vehicle Profile (Auto-fill) khi sang luồng mua BH.\n3. A/B Testing block cross-sell chuẩn hóa theo loại xe (Ô tô/Xe máy).', 'Backlog / Sprint Kế'],
        ['None', 'Calculator & Guide', 'Tính Lăn Bánh & Thuế Trước Bạ (/lan-banh)', '1. Công cụ tính phí trước bạ ô tô (10/12%).\nVai trò: Nguồn Lead chất lượng để bán chéo Bảo hiểm VCX/TNDS xe mới.', 'Backlog / Sprint Kế'],
        ['None', 'Giữ Chân & Khẩn Cấp', 'Danh Bạ Cứu Hộ & Bãi Đỗ Xe (/cuu-ho, /bai-do-xe)', '1. Bản đồ VETC Parking và danh bạ cứu hộ 24/7 theo định vị.\nVai trò: Xử lý Intent khẩn cấp & Tiện ích giữ chân.', 'Backlog / Sprint Kế'],
        
        # Blocked / Pending
        ['Blocked / Pending (Vướng Phụ Thuộc Đối Tác)', 'Magnet + Transaction (P1)', 'Tra Cứu Hạn Đăng Kiểm (/dang-kiem)', '1. Tra cứu hạn đăng kiểm theo biển số & đặt lịch kiểm định.\nTRẠNG THÁI: Blocked. API tra cứu đang bị đóng/bảo trì từ phía TTĐK. Chờ mở lại.', 'Blocked'],
        ['None', 'Technical & Compliance', 'Tuân Thủ Dữ Liệu & Retargeting', '1. Xây UI Consent lưu biển số theo Nghị định 13.\n2. Kênh Retargeting Zalo OA thông báo Phạt nguội.', 'Pending']
    ]

    for idx, row_data in enumerate(data, start=4):
        for col, val in enumerate(row_data, start=1):
            cell = ws.cell(row=idx, column=col, value=val if val != 'None' else None)
            cell.alignment = align_left if col in [3, 4] else align_center
            cell.border = border
            
            # Row merges and coloring
            if "Phase 1" in row_data[0]:
                cell.fill = fill_phase1
            elif "Phase 2" in row_data[0]:
                cell.fill = fill_phase2
            elif "Blocked" in row_data[0]:
                cell.fill = fill_blocked
            elif idx <= 8: # belong to phase 1
                cell.fill = fill_phase1
            elif idx <= 11:
                cell.fill = fill_phase2
            else:
                cell.fill = fill_blocked

    # Merge Phase cells vertically
    ws.merge_cells('A4:A8')
    ws.merge_cells('A9:A11')
    ws.merge_cells('A12:A13')
    
    # Also adjust column width just in case
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 35
    ws.column_dimensions['D'].width = 50
    ws.column_dimensions['E'].width = 20

wb.save(file_path)
print("Roadmap updated successfully based on Cell Team input.")
