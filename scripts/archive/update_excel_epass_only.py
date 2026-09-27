import openpyxl
from openpyxl.styles import Alignment, Border, Side, PatternFill, Font

file_excel = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/vehicle-hub-roadmap.xlsx'
wb = openpyxl.load_workbook(file_excel)

border_thin = Border(left=Side(style='thin', color='D9D9D9'),
                     right=Side(style='thin', color='D9D9D9'),
                     top=Side(style='thin', color='D9D9D9'),
                     bottom=Side(style='thin', color='D9D9D9'))
align_left = Alignment(horizontal='left', vertical='top', wrap_text=True)
align_center_top = Alignment(horizontal='center', vertical='top', wrap_text=True)
font_data = Font(name='Calibri', size=10)
font_data_bold = Font(name='Calibri', size=10, bold=True)
font_url = Font(name='Consolas', size=9, color='1F4E78')

# 1. Update Roadmap Sheet
ws_road = wb['Roadmap']
for r in range(3, ws_road.max_row + 1):
    val_name = str(ws_road.cell(r, 4).value or '')
    if 'Thu Phí' in val_name:
        ws_road.cell(r, 4, value="Thu Phí BOT ePass (/phi-khong-dung)")
        ws_road.cell(r, 5, value="• 1 Trang duy nhất (/phi-khong-dung) - Đối tác chiến lược độc quyền: ePass (VDTC - Viettel).\n• Tra cứu số dư tài khoản giao thông, nạp tiền ePass miễn phí, đăng ký dán thẻ tại nhà.\n(Search Volume: 108.210 search/tháng)")
        ws_road.cell(r, 6, value="Nút CTA liên kết tài khoản ePass & Bật tính năng 'Tự động bù tiền BOT trên MoMo'.")
        break

# Insert a new row in Roadmap for Trạm Thu Phí Tuyến A > B if not present
has_tram_thu_phi = any('tram-thu-phi' in str(ws_road.cell(r, 4).value or '') for r in range(3, ws_road.max_row + 1))
if not has_tram_thu_phi:
    new_row = [
        "Phase 2", "Tháng 10/2026", "Location & Trip", "Tra Cứu Trạm Thu Phí Tuyến A ➔ B (/tram-thu-phi)",
        "• 1 Trang duy nhất (/tram-thu-phi): Trợ thủ tính tổng phí cầu đường khi di chuyển từ A đến B.\n• Chọn tuyến cao tốc mẫu hoặc nhập Điểm đi - Điểm đến.\n• Báo tổng số trạm BOT, vị trí Km, tổng số tiền vé cầu đường cần chuẩn bị.\n• Bản đồ và bảng kê gần 100 trạm BOT toàn quốc.",
        "Nút 1-chạm: 'Nạp đúng [X] nghìn đồng vào ePass qua MoMo để khởi hành'.", "Planned"
    ]
    ws_road.append(new_row)
    r_idx = ws_road.max_row
    ws_road.row_dimensions[r_idx].height = 80
    for c in range(1, 8):
        cell = ws_road.cell(r_idx, c)
        cell.border = border_thin
        cell.font = font_data
        if c in [1, 2, 7]:
            cell.alignment = align_center_top
        elif c in [3, 4]:
            cell.alignment = align_center_top
            cell.font = font_data_bold
        else:
            cell.alignment = align_left

# 2. Update Content Structure Sheet
if 'Content Structure' in wb.sheetnames:
    ws_cs = wb['Content Structure']
    for r in range(2, ws_cs.max_row + 1):
        for c in range(1, ws_cs.max_column + 1):
            val = str(ws_cs.cell(r, c).value or '')
            if 'VETC' in val or 'vetc' in val:
                ws_cs.cell(r, c, value=val.replace('ePass/VETC', 'ePass').replace('ePass / VETC', 'ePass').replace('VETC', 'ePass'))

# 3. Update Subpages Directory Sheet
if 'Subpages Directory' in wb.sheetnames:
    ws_sub = wb['Subpages Directory']
    # Add /tram-thu-phi
    sub_row = [
        str(ws_sub.max_row - 2), "Trạm Thu Phí", "Trang Độc Lập", "Tra Cứu Trạm Thu Phí & Dự Toán Phí Tuyến A ➔ B",
        "/tien-ich-giao-thong/tram-thu-phi", "trạm thu phí, phí cao tốc, giá vé bot",
        "Bộ tính phí cầu đường tuyến A > B, Bảng kê chi tiết từng trạm, Nạp tiền ePass chuẩn xác",
        "Phase 2 (Planned)"
    ]
    ws_sub.append(sub_row)
    idx = ws_sub.max_row
    ws_sub.row_dimensions[idx].height = 25
    for c in range(1, 9):
        cell = ws_sub.cell(idx, c)
        cell.border = border_thin
        cell.font = font_data
        if c in [1, 3, 8]:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif c == 2:
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.font = font_data_bold
        elif c == 5:
            cell.alignment = Alignment(horizontal='left', vertical='center')
            cell.font = font_url
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center')

# 4. Clean VETC across Resources & Utilities
for sheet_name in ['Resources', 'Utilities & Formula Specs', 'Market Sizing & Opportunities']:
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                val = str(ws.cell(r, c).value or '')
                if 'VETC' in val or 'vetc' in val:
                    ws.cell(r, c, value=val.replace('ePass/VETC', 'ePass').replace('ePass / VETC', 'ePass').replace('VETC Parking', 'Bãi Đỗ Xe Thông Minh').replace('VETC', 'ePass'))

wb.save(file_excel)
print("Updated Excel workbook with ePass exclusivity successfully!")
