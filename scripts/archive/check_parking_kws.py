import openpyxl

file_inv = '/Users/hienhv/HienHv/Klaus/Hienho_MoMo-Base/05_HUBS/inventory.xlsx'
file_cs = '/Users/hienhv/Downloads/[WP] Vehicle Hub - Content Strategy.xlsx'

wb_cs = openpyxl.load_workbook(file_cs, data_only=True)
ws_kw = wb_cs['Keyword']

parking_kws = []
for r in range(2, ws_kw.max_row + 1):
    kw = str(ws_kw.cell(r, 1).value or ws_kw.cell(r, 2).value or '').strip().lower()
    if any(k in kw for k in ['bãi đỗ', 'bãi gửi', 'bãi giữ', 'đỗ xe', 'gửi xe']):
        vol = ws_kw.cell(r, 2).value or ws_kw.cell(r, 3).value or 0
        try:
            vol = float(vol)
        except:
            vol = 0
        parking_kws.append((kw, vol))

parking_kws.sort(key=lambda x: x[1], reverse=True)
print(f"Total parking keywords in Keyword sheet: {len(parking_kws)}, Sum Vol: {sum(x[1] for x in parking_kws):,.0f}")
for k, v in parking_kws[:20]:
    print(f"  • {k}: {v:,.0f}")

