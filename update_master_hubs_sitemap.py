import openpyxl


from scripts._lib import load_workbook_safe, save_with_backup, standard_argparser

def main(args=None):
    wb = load_workbook_safe('05_HUBS/student-hub-roadmap.xlsx')

    # 1. Update Sitemap sheet to master hubs architecture
    if 'Sitemap' in wb.sheetnames:
        ws_sm = wb['Sitemap']
        ws_sm.delete_rows(5, ws_sm.max_row)
    
        sitemap_clean = [
            [1, 'Master Gateway', 'Trang Chủ Student Hub', '/sinh-vien', 'sinh viên momo, student pass momo, ưu đãi sinh viên', 'Tổng quan Thẻ Sinh Viên Số, Banner ưu đãi hot, Widget Bạn đồng hành MoMo, Top Review Trường.', 'Phase 1.0 (Completed)'],
            [2, 'Trường ĐH Pilot', 'Trang Chi Tiết Trường ĐH (12 Trường)', '/sinh-vien/[truong-slug]', 'review ueh, học phí ueh, điểm chuẩn ueh, tiện ích ueh', 'Thông tin tuyển sinh, điểm chuẩn 3 năm, học phí tín chỉ, xe buýt, nhúng Inline Widget Nhà trọ & Workshop.', 'Phase 1.1 & Phase 2 (Completed/In Progress)'],
            [3, 'Cổng Việc Làm', 'Cổng Việc Làm Master & AI Resume', '/sinh-vien/viec-lam', 'việc làm sinh viên, tìm việc part time, cv thực tập sinh', 'Bản đồ việc làm part-time 5km quanh campus, AI Resume Builder (ATS Score), 1-Tap Apply qua Student Pass.', 'Phase 3.1 (Planned)'],
            [4, 'Nhà Trọ Sinh Viên', 'Cổng Master Nhà Trọ & KTX Sinh Viên', '/sinh-vien/nha-tro', 'thuê nhà trọ sinh viên, phòng trọ gần trường đại học', 'Danh sách phòng trọ an toàn quanh 12+ trường ĐH (dùng bộ lọc ?truong=[slug]), tool AI Bill Splitter chia tiền trọ.', 'Phase 3.2 (Planned)'],
            [5, 'Sự Kiện & Workshop', 'Cổng Master Event & Workshop Sinh Viên', '/sinh-vien/workshop', 'workshop sinh viên, webinar hướng nghiệp, ai talk sinh viên', 'Cổng đăng ký Webinar Career & AI Talk, sự kiện CLB trường ĐH, lịch workshop kỹ năng sinh viên.', 'Phase 3.3 (Planned)'],
            [6, 'Đại Sứ Sinh Viên', 'Trang Showcase Đại Sứ Sinh Viên', '/sinh-vien/ambassador', 'đại sứ sinh viên momo, momo campus ambassador', 'Showcase project đại sứ sinh viên, bảng vinh danh leaderboard, form tuyển dụng Đại sứ.', 'Phase 2.1 (In Progress)']
        ]
    
        for r_offset, r_data in enumerate(sitemap_clean, start=5):
            for c_idx, val in enumerate(r_data, start=1):
                cell = ws_sm.cell(row=r_offset, column=c_idx, value=val)

    # 2. Update Roadmap sheet rows to master hubs
    if 'Roadmap' in wb.sheetnames:
        ws_rm = wb['Roadmap']
        for r in range(5, ws_rm.max_row+1):
            item_val = str(ws_rm.cell(row=r, column=4).value or '')
            if '/sinh-vien/[truong-slug]/nha-tro' in item_val:
                ws_rm.cell(row=r, column=4, value='Cổng Master Nhà Trọ & KTX Sinh Viên (/sinh-vien/nha-tro)')
                ws_rm.cell(row=r, column=5, value='• Cổng Master Nhà Trọ (/sinh-vien/nha-tro) gom toàn bộ dữ liệu phòng trọ an toàn.\n• Sử dụng bộ lọc ?truong=[truong-slug] hiển thị nhà trọ quanh campus 5km.\n• Tích hợp tool AI Bill Splitter (chia tiền trọ).')
            elif '/sinh-vien/[truong-slug]/workshop' in item_val:
                ws_rm.cell(row=r, column=4, value='Cổng Master Webinar & Event Workshop (/sinh-vien/workshop)')
                ws_rm.cell(row=r, column=5, value='• Cổng Master Workshop (/sinh-vien/workshop) gom toàn bộ sự kiện/webinar kỹ năng & AI talk.\n• Sử dụng bộ lọc ?truong=[truong-slug] hiển thị workshop theo trường.')

    save_with_backup(wb, '05_HUBS/student-hub-roadmap.xlsx', dry_run=args.dry_run, make_backup=not args.no_backup)
    print('Successfully updated sitemap and roadmap in student-hub-roadmap.xlsx to Master Hubs architecture!')

if __name__ == "__main__":
    _ap = standard_argparser("Fix runner - safe with --dry-run")
    main(_ap.parse_args())

