import openpyxl

# 1. Update financial-hub-roadmap.xlsx
wb = openpyxl.load_workbook('05_HUBS/financial-hub-roadmap.xlsx')

# Update Sheet Roadmap
if 'Roadmap' in wb.sheetnames:
    ws = wb['Roadmap']
    for r in range(3, ws.max_row + 1):
        v4 = str(ws.cell(row=r, column=4).value or '')
        v5 = str(ws.cell(row=r, column=5).value or '')
        if 'Tỷ Giá' in v4:
            ws.cell(row=r, column=5).value = (
                "• Xây dựng Trang Cổng Tỷ Giá Master (/tai-chinh/ty-gia) tích hợp bảng so sánh realtime 30+ ngân hàng.\n"
                "• Triển khai Hệ thống Programmatic Subpages theo từng Cặp Tiền Tệ (/tai-chinh/ty-gia/[pair-slug]) cho 20+ ngoại tệ chủ lực (USD/VND, JPY/VND, EUR/VND, KRW/VND, CNY/VND, GBP/VND, AUD/VND, SGD/VND, THB/VND...).\n"
                "• Mỗi trang con cặp tiền tích hợp: Bộ quy đổi 2 chiều tức thì, Biểu đồ biến động lịch sử 7-30 ngày của riêng cặp tiền, Bảng so sánh tỷ giá Mua/Bán giữa các ngân hàng thương mại và Cẩm nang đổi tiền/du lịch.\n"
                "• Biểu đồ theo dõi biến động tỷ giá ngân hàng trung ương và thị trường tự do."
            )

wb.save('05_HUBS/financial-hub-roadmap.xlsx')
print("Successfully updated /tai-chinh/ty-gia routing in 05_HUBS/financial-hub-roadmap.xlsx!")

# 2. Update financial-hub-brd.md
with open('05_HUBS/financial-hub-brd.md', 'r', encoding='utf-8') as f:
    brd = f.read()

brd = brd.replace('/ty-gia/usd-vnd', '/tai-chinh/ty-gia/usd-vnd')
brd = brd.replace('/ty-gia/jpy-vnd', '/tai-chinh/ty-gia/jpy-vnd')
brd = brd.replace('/ty-gia/eur-vnd', '/tai-chinh/ty-gia/eur-vnd')
brd = brd.replace('/ty-gia/krw-vnd', '/tai-chinh/ty-gia/krw-vnd')
brd = brd.replace('/ty-gia/cny-vnd', '/tai-chinh/ty-gia/cny-vnd')
brd = brd.replace('/ty-gia/gbp-vnd', '/tai-chinh/ty-gia/gbp-vnd')
brd = brd.replace('/ty-gia/aud-vnd', '/tai-chinh/ty-gia/aud-vnd')
brd = brd.replace('/ty-gia/sgd-vnd', '/tai-chinh/ty-gia/sgd-vnd')
brd = brd.replace('/ty-gia/thb-vnd', '/tai-chinh/ty-gia/thb-vnd')
brd = brd.replace('/ty-gia/[pair-slug]', '/tai-chinh/ty-gia/[pair-slug]')
brd = brd.replace('`momo.vn/ty-gia`', '`momo.vn/tai-chinh/ty-gia`')
brd = brd.replace('`/ty-gia`', '`/tai-chinh/ty-gia`')

with open('05_HUBS/financial-hub-brd.md', 'w', encoding='utf-8') as f:
    f.write(brd)

print("Successfully updated /tai-chinh/ty-gia routing in 05_HUBS/financial-hub-brd.md!")
