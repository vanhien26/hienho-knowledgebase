import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


from scripts._lib import load_workbook_safe, save_with_backup, standard_argparser

def main(args=None):
    wb = load_workbook_safe('05_HUBS/student-hub-roadmap.xlsx')

    font_bold = Font(name='Calibri', size=11, bold=True, color='000000')
    font_regular = Font(name='Calibri', size=11, color='000000')

    fill_inprogress = PatternFill('solid', fgColor='FEF08A')
    fill_planned = PatternFill('solid', fgColor='F3F4F6')

    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    border_thin = Side(style='thin', color='CBD5E1')
    border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    # 1. Update Roadmap Sheet
    if 'Roadmap' in wb.sheetnames:
        ws_rm = wb['Roadmap']
        for r in range(5, ws_rm.max_row+1):
            item_val = str(ws_rm.cell(row=r, column=4).value or '')
            phase_val = str(ws_rm.cell(row=r, column=1).value or '')
            if '/viec-lam' in item_val or 'Career Master' in str(ws_rm.cell(row=r, column=3).value or ''):
                ws_rm.cell(row=r, column=1, value='Phase 2')
                ws_rm.cell(row=r, column=2, value='Tháng 9/2026')
                ws_rm.cell(row=r, column=7, value='Phase 2 (In Progress / Fast-Tracked)')
                for c in range(1, 8):
                    cell = ws_rm.cell(row=r, column=c)
                    cell.fill = fill_inprogress
        print('Updated Roadmap sheet: Moved /viec-lam to Phase 2 (Tháng 9/2026).')

    # 2. Update Sitemap Sheet
    if 'Sitemap' in wb.sheetnames:
        ws_sm = wb['Sitemap']
        for r in range(5, ws_sm.max_row+1):
            url_val = str(ws_sm.cell(row=r, column=4).value or '')
            if '/viec-lam' in url_val:
                ws_sm.cell(row=r, column=7, value='Phase 2 (In Progress)')
        print('Updated Sitemap sheet: Set /viec-lam status to Phase 2.')

    # 3. Update Readme Sheet
    if 'Readme' in wb.sheetnames:
        ws_rd = wb['Readme']
        for r in range(1, ws_rd.max_row+1):
            c1_val = str(ws_rd.cell(row=r, column=1).value or '')
            c2_val = str(ws_rd.cell(row=r, column=2).value or '')
            if 'Phase 2:' in c1_val:
                ws_rd.cell(row=r, column=2, value='Mở rộng 7 trường mới (HUST, NEU, HANU, HAU, USSH, HUIT, UEL -> 12 trường). Fast-track ra mắt Cổng Việc Làm Master (/sinh-vien/viec-lam) kết hợp AI Resume Builder (ATS Score), Bản đồ việc làm 5km quanh campus từ đối tác F&B/bán lẻ MoMo; 1-Tap Apply via Student Pass mở Ví Trả Sau 0% HSSV. Promote & Drive 100% sinh viên đến eKYC Verify Student Pass In-App nhận gói voucher 500K. Tích hợp Gamification: Web Photobooth Real-time Event 26-27/09, Nhiệm vụ Đại sứ (/ambassador), Minigame tích điểm In-App.')
            elif 'Phase 3:' in c1_val:
                ws_rd.cell(row=r, column=2, value='Programmatic pSEO Scale-out 100+ trang trường ĐH/CĐ toàn quốc theo Tier từ khóa. Ra mắt Cổng Master Nhà Trọ (/nha-tro ?truong=[slug]) & Cổng Master Workshop (/workshop).')
        print('Updated Readme sheet: Updated Phase 2 and Phase 3 descriptions.')

    save_with_backup(wb, '05_HUBS/student-hub-roadmap.xlsx', dry_run=args.dry_run, make_backup=not args.no_backup)
    print('Successfully saved student-hub-roadmap.xlsx with /viec-lam moved to Phase 2!')

if __name__ == "__main__":
    _ap = standard_argparser("Fix runner - safe with --dry-run")
    main(_ap.parse_args())

