import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


from scripts._lib import load_workbook_safe, save_with_backup, standard_argparser

def main(args=None):
    wb = load_workbook_safe('05_HUBS/student-hub-roadmap.xlsx')

    # Common styles
    font_title = Font(name='Calibri', size=14, bold=True, color='1E293B')
    font_subtitle = Font(name='Calibri', size=11, italic=True, color='64748B')
    font_header = Font(name='Calibri', size=11, bold=True, color='1E293B')
    font_section = Font(name='Calibri', size=11, bold=True, color='0F172A')
    font_bold = Font(name='Calibri', size=11, bold=True, color='000000')
    font_regular = Font(name='Calibri', size=11, color='000000')

    fill_header = PatternFill('solid', fgColor='F1F5F9')
    fill_completed = PatternFill('solid', fgColor='DCFCE7')
    fill_inprogress = PatternFill('solid', fgColor='FEF08A')
    fill_planned = PatternFill('solid', fgColor='F3F4F6')

    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center')

    border_thin = Side(style='thin', color='CBD5E1')
    border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    if 'Roadmap' in wb.sheetnames:
        ws_roadmap = wb['Roadmap']
        ws_roadmap.delete_rows(4, ws_roadmap.max_row)
    
        ws_roadmap.cell(row=1, column=1, value='LỘ TRÌNH PHÁT TRIỂN SẢN PHẨM STUDENT HUB (STUDENT HUB ROADMAP 2026)').font = font_title
        ws_roadmap.cell(row=2, column=1, value='Quy hoạch 3 Phase chiến lược chuẩn theo định hướng: Phase 1 (T8 Foundation & 5 trường), Phase 2 (T9 Scale 7 trường, Drive Verify Student Pass & Gamification), Phase 3 (Q4 Career Master & pSEO Scale 100+).').font = font_subtitle

        headers_roadmap = ['Phase', 'Thời Gian', 'Giai Đoạn & Nhóm Dịch Vụ', 'Tên Trang / Utility / Công Việc', 'Mô Tả Sản Phẩm & Quy Hoạch Kỹ Thuật (UX/UI & Specs)', 'Dữ Liệu, API & Phối Hợp Kỹ Thuật', 'Phễu Chuyển Đổi Kinh Doanh (Web-to-App)']
        for c_idx, h in enumerate(headers_roadmap, start=1):
            cell = ws_roadmap.cell(row=4, column=c_idx, value=h)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = border

        exact_roadmap_data = [
            ['Phase 1 (Completed)', 'Tháng 8/2026', 'Foundation Launch | Core Features', 'Master Gateway (/sinh-vien) & 5 Trang Trường Pilot (UEH, FTU, TDTU, HCMUT, VLU)', '• Dựng khung HTML/CSS Master Gateway (/sinh-vien) và 5 trang trường Pilot trọng điểm.\n• Hoàn thiện 3 tính năng cốt lõi: (1) Thông tin trường (Học phí, điểm chuẩn 3 năm), (2) Tìm điểm thanh toán chấp nhận MoMo quanh campus (F&B/SME), (3) Tra cứu tuyến xe buýt đến trường.\n• ĐÃ HOÀN TẤT & GO-LIVE.', '- Master Data 5 Trường Pilot\n- Google Transit API (Xe buýt)\n- Merchant POI Database (Điểm thanh toán MoMo)', 'Anchor links #hoc-phi, #diem-chuan & OneLink giới thiệu Student Pass.'],
            ['Phase 2 (In Progress)', 'Tháng 9/2026', 'Scale-out, Verification Drive & Gamification', 'Mở Rộng 7 Trường Mới + Phễu Verify Student Pass + Module Gamification & Event Offline', '• Mở rộng 7 trang trường mới (HUST, NEU, HANU, HAU, USSH, HUIT, UEL).\n• Tập trung Promote & Drive người dùng đến phễu Xác Thực Sinh Viên (Verify Student Pass eKYC / Thẻ Sinh Viên Số Custom Card UI).\n• Tích hợp hệ thống Gamification: Module Web Photobooth Real-time sự kiện Offline Campus 26-27/09, Nhiệm vụ Đại sứ Sinh viên (Ambassador Quests /sinh-vien/ambassador) và Minigame tích điểm đổi quà In-App.', '- API Xác thực Thẻ Sinh Viên Số In-App\n- Realtime Photobooth Image Engine\n- Ambassador Leaderboard System', 'Focus 100% phễu Web-to-App CTR kéo sinh viên hoàn tất eKYC Verify Student Pass In-App nhận gói voucher 50K-500K & Đặc quyền sinh viên.'],
            ['Phase 3 (Planned)', 'Quý 4/2026', 'Career Master, Programmatic pSEO 100+ & Ecosystem', 'Cổng Việc Làm Master (/sinh-vien/viec-lam) + AI Resume Builder + Programmatic pSEO 100+ Trường', '• Ra mắt Cổng Việc Làm Sinh Viên Master (/sinh-vien/viec-lam) kết hợp công cụ AI Resume Builder (ATS Score), Bản đồ việc làm 5km quanh campus từ đối tác F&B/bán lẻ MoMo; 1-Tap Apply via Student Pass mở Ví Trả Sau 0% HSSV.\n• Programmatic pSEO Engine tự động mở rộng 100+ trang trường ĐH/CĐ toàn quốc.\n• Mở rộng Sub-pages chuyên biệt (/nha-tro, /review, /workshop).', '- Merchant Jobs Database\n- MoMo GenAI ATS Resume Engine\n- Programmatic Content Engine 100+ Trường', 'Phễu 1-Tap Apply ứng tuyển -> Verify Student Pass -> Mở Ví Trả Sau 0% HSSV & Túi Thần Tài.']
        ]

        for r_offset, r_data in enumerate(exact_roadmap_data, start=5):
            for c_idx, val in enumerate(r_data, start=1):
                cell = ws_roadmap.cell(row=r_offset, column=c_idx, value=val)
                cell.font = font_bold if c_idx in [1, 2, 4] else font_regular
                cell.border = border
                cell.alignment = align_center if c_idx in [1, 2] else align_left
                if 'Phase 1' in r_data[0]:
                    cell.fill = fill_completed
                elif 'Phase 2' in r_data[0]:
                    cell.fill = fill_inprogress
                else:
                    cell.fill = fill_planned

        ws_roadmap.column_dimensions['A'].width = 22
        ws_roadmap.column_dimensions['B'].width = 16
        ws_roadmap.column_dimensions['C'].width = 34
        ws_roadmap.column_dimensions['D'].width = 45
        ws_roadmap.column_dimensions['E'].width = 68
        ws_roadmap.column_dimensions['F'].width = 50
        ws_roadmap.column_dimensions['G'].width = 50

    save_with_backup(wb, '05_HUBS/student-hub-roadmap.xlsx', dry_run=args.dry_run, make_backup=not args.no_backup)
    print('Successfully updated sheet Roadmap in student-hub-roadmap.xlsx with exact Phase definition!')

if __name__ == "__main__":
    _ap = standard_argparser("Fix runner - safe with --dry-run")
    main(_ap.parse_args())

